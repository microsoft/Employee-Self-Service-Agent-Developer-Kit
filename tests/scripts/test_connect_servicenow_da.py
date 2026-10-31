# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.

from __future__ import annotations

import copy
import csv
import importlib.util
import json
import re
import subprocess
import sys
import uuid
import zipfile

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Protection

import connect_servicenow_da as snow


ENVIRONMENT_ID = "00000000-0000-4000-8000-000000001111"
TENANT_ID = "00000000-0000-4000-8000-000000009999"
AGENT_ID = "00000000-0000-4000-8000-000000002222"
OTHER_AGENT_ID = "00000000-0000-4000-8000-000000002223"
CONNECTION_ID = "00000000-0000-4000-8000-000000003333"
APP_CLIENT_ID = "00000000-0000-4000-8000-000000006666"
AGENT_SLUG = "employee-self-service-hr"
_RECOVERY_HEADINGS = (
    "**Problem**",
    "**Category**",
    "**Owner**",
    "**Where**",
    "**Values to use**",
    "**Action**",
    "**Success**",
    "**Recheck and resume**",
)


def _assert_recovery_handoff(value: dict) -> None:
    assert set(value) == {
        "problem",
        "category",
        "owner",
        "where",
        "values",
        "action",
        "success",
        "recheck",
        "message",
    }
    assert all(value[key] for key in value)
    assert all(heading in value["message"] for heading in _RECOVERY_HEADINGS)
    assert "{" not in value["message"]
_CONTRACT_PATH = (
    Path(__file__).resolve().parents[2]
    / "solutions"
    / "ess-maker-skills"
    / "src"
    / "skills"
    / "connect"
    / snow.PROVIDER_KEY
    / "contract.json"
)
_SCRIPT_PATH = (
    Path(__file__).resolve().parents[2]
    / "solutions"
    / "ess-maker-skills"
    / "scripts"
    / "connect_servicenow_da.py"
)


@pytest.fixture(autouse=True)
def _disable_real_servicenow_telemetry(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("ESS_ADK_TELEMETRY", "off")
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, _event: None,
    )


def _context() -> dict:
    return {
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.HR_SCHEMA_NAME,
            "workspace_slug": AGENT_SLUG,
        },
        "active": {"slug": AGENT_SLUG},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
    }


def _lifecycle_path(root: Path) -> Path:
    return (
        root
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )


def _context_at(root: Path) -> dict:
    context = _context()
    context["root"] = root
    return context


def _fill_entra_workbook(
    path: Path,
    candidates: tuple[str, ...],
) -> None:
    workbook = load_workbook(path)
    if snow.IDENTITY_ENTRA_HANDOFF_SHEET in workbook.sheetnames:
        if len(candidates) != 1:
            workbook.close()
            raise AssertionError(
                "The v2 Entra handoff requires exactly one selected claim."
            )
        workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["C2"] = candidates[0]
        workbook.save(path)
        workbook.close()
        return
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    for offset, candidate in enumerate(candidates):
        sheet.cell(
            snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW + offset,
            1,
        ).value = candidate
    workbook.save(path)
    workbook.close()


def _fill_servicenow_workbook(
    path: Path,
    *,
    selected_claim: str,
    user_field: str,
) -> None:
    workbook = load_workbook(path)
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    for row in range(
        snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW,
        sheet.max_row + 1,
    ):
        if sheet.cell(row, 1).value == selected_claim:
            sheet.cell(row, 2).value = user_field
            break
    workbook.save(path)
    workbook.close()


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        return list(reader.fieldnames or []), list(reader)


def _write_csv(
    path: Path,
    headers: list[str],
    rows: list[dict[str, str]],
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=headers,
            lineterminator="\r\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def _fill_entra_csv(
    path: Path,
    *,
    client_id: str,
    user_claim: str,
) -> None:
    headers, rows = _read_csv(path)
    assert len(rows) == 1
    rows[0]["Client ID"] = client_id
    rows[0]["User Claim"] = user_claim
    _write_csv(path, headers, rows)


def _fill_servicenow_csv(
    path: Path,
    *,
    user_field: str,
) -> None:
    headers, rows = _read_csv(path)
    assert [row["Record"] for row in rows] == ["Entra handoff", "Steps"]
    rows[1]["ServiceNow User Field"] = user_field
    _write_csv(path, headers, rows)


def _confirm_identity_requirement(
    context: dict,
    *,
    choice: str = "email",
) -> dict:
    return snow.record_identity_requirement(
        context,
        choice=choice,
    )


def _author_universal_entra_template(path: Path) -> Path:
    template = (
        Path(__file__).resolve().parents[2]
        / "solutions"
        / "ess-maker-skills"
        / snow.IDENTITY_WORKBOOK_TEMPLATE_ENTRA
    )
    path.write_bytes(template.read_bytes())
    return path


def _author_servicenow_template(path: Path) -> Path:
    template = (
        Path(__file__).resolve().parents[2]
        / "solutions"
        / "ess-maker-skills"
        / snow.IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW
    )
    path.write_bytes(template.read_bytes())
    return path

def _validate_dynamic_visible_question(question: dict) -> dict:
    validator_path = (
        Path(__file__).resolve().parents[1]
        / "setup"
        / "test_visible_handoff_payloads.py"
    )
    spec = importlib.util.spec_from_file_location(
        "_visible_handoff_validator",
        validator_path,
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module._validate_payload(
        [question],
        required_phrases=(
            "Choose the type of employee identifier",
            "Do not provide an employee value",
        ),
        minimum_length=100,
        rendered=True,
    )


def _visible_prose_budget(text: str) -> tuple[int, int]:
    without_links = re.sub(
        r"\[([^\]]+)\]\([^)]+\)",
        r"\1",
        text,
    )
    without_literals = re.sub(r"`[^`]+`", "", without_links)
    words = re.findall(r"[A-Za-z]+(?:'[A-Za-z]+)?", without_literals)
    lines = [line for line in text.splitlines() if line.strip()]
    return len(words), len(lines)


def _record_test_identity_mapping(
    root: Path,
    *,
    candidates: tuple[str, ...] = ("employee_id",),
    selected_claim: str = "employee_id",
    user_field: str = "employee_number",
) -> dict:
    context = _context_at(root)
    state_path = _lifecycle_path(root)
    if not state_path.exists():
        state = snow._state_base(context, _components())
        state["adminSetup"]["preflight"]["instanceName"] = "dev123"
        snow._write_json_atomic(state_path, state)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    if (
        "identityRequirement" not in state.get("evidence", {})
        and "identityClaimCandidates" not in state.get("evidence", {})
        and not snow._identity_workbook_path(
            context,
            stage="entra",
        ).exists()
    ):
        _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, candidates)
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    snow.record_admin_phase(
        context,
        phase="entra-registration",
        status="completed",
        client_id=APP_CLIENT_ID,
    )
    service_now = snow.identity_mapping_handoff(
        context,
        stage="servicenow",
    )
    service_now_path = Path(service_now["workbookPath"])
    _fill_servicenow_workbook(
        service_now_path,
        selected_claim=selected_claim,
        user_field=user_field,
    )
    result = snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=service_now_path,
        client_id=None,
        confirmed=True,
    )
    return result


def _phase_contract(phase_id: str) -> dict:
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    return next(
        phase for phase in contract["phases"] if phase["id"] == phase_id
    )


def _seed_passing_test(
    root: Path,
    components: dict,
    *,
    publish: dict | None = None,
) -> dict:
    state = snow._state_base(_context(), components)
    normalized_connection_id = CONNECTION_ID.replace("-", "")
    state["evidence"]["credential"] = {
        "connectionId": normalized_connection_id,
    }
    state["evidence"]["agentConnection"] = {
        "kind": "maker-attestation",
        "status": "completed",
        "connectionId": normalized_connection_id,
        "physicalStatus": "Connected",
        "makerAttested": True,
        "recordedAt": "2026-10-01T00:00:00Z",
        "binding": {
            "draftSemanticHash": snow._draft_semantic_hash(components),
        },
    }
    connection_binding_hash = snow._connection_binding_hash(
        _context(),
        state,
        components,
    )
    state["connectionBindingHash"] = connection_binding_hash
    state["evidence"]["test"] = {
        "kind": "maker-attestation",
        "status": "completed",
        "promptCategory": "list-my-open-hr-cases",
        "result": "pass",
        "failureCategory": None,
        "recordedAt": "2026-10-01T00:01:00Z",
        "binding": {
            "provider": snow.PROVIDER_KEY,
            "profile": snow.PROFILE_KEY,
            "agentSlug": AGENT_SLUG,
            "agentId": AGENT_ID,
            "environmentId": ENVIRONMENT_ID,
            "connectionId": normalized_connection_id,
            "draftSemanticHash": snow._draft_semantic_hash(components),
            "connectionBindingHash": connection_binding_hash,
        },
    }
    if publish is not None:
        state["evidence"]["publish"] = publish
    state_path = _lifecycle_path(root)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps(state), encoding="utf-8")
    return state


def _seed_complete_dev_state(
    root: Path,
    *,
    profile: str,
    legacy_publish_status: str | None = None,
) -> tuple[dict, dict]:
    product = snow.PRODUCT_PROFILES[profile]
    slug = (
        "employee-self-service-hr"
        if profile == "hrsd"
        else "employee-self-service-it"
    )
    agent_id = AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID
    context = _profile_context(
        profile,
        agent_id=agent_id,
        agent_slug=slug,
    )
    context["root"] = root
    components = _components(CONNECTION_ID)
    contract_path = (
        root
        / "src"
        / "skills"
        / "connect"
        / product.provider_key
        / "contract.json"
    )
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    source_contract = (
        Path(__file__).resolve().parents[2]
        / "solutions"
        / "ess-maker-skills"
        / "src"
        / "skills"
        / "connect"
        / product.provider_key
        / "contract.json"
    )
    contract = json.loads(source_contract.read_text(encoding="utf-8"))
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    with snow.product_profile_context(profile):
        state = snow._state_base(context, components)
        state["attested"] = True
        state["attestedAt"] = "2026-10-23T00:00:00Z"
        state["acceptedContractRevision"] = contract["contractRevision"]
        state["evidence"]["credential"] = {
            "connectionId": CONNECTION_ID.replace("-", ""),
            "instanceName": "dev123",
            "resourceUri": APP_CLIENT_ID,
            "authMode": snow.AUTH_MODE,
            "physicalStatus": "Connected",
        }
        state["evidence"]["agentConnection"] = {
            "kind": "maker-attestation",
            "status": "completed",
            "recordedAt": "2026-10-23T00:00:30Z",
            "makerAttested": True,
            "physicalStatus": "Connected",
            "authMode": snow.AUTH_MODE,
            "connectionId": CONNECTION_ID.replace("-", ""),
            "binding": {
                "provider": product.provider_key,
                "profile": profile,
                "agentSlug": slug,
                "agentId": agent_id,
                "environmentId": ENVIRONMENT_ID,
                "draftSemanticHash": state["draftSemanticHash"],
            },
        }
        state["connectionBindingHash"] = snow._connection_binding_hash(
            context,
            state,
            components,
        )
        state["phases"] = {}
        for phase in contract["phases"]:
            checkpoint_results = {}
            for target in phase["checkpoints"]:
                checkpoints = (
                    [
                        f"{target[:-1]}{suffix}-001"
                        for suffix in (
                            "APP",
                            "CLAIMS",
                            "SCOPE",
                            "PREAUTH",
                            "PERMISSIONS",
                            "CONSENT",
                        )
                    ]
                    if target.endswith("*")
                    else [target]
                )
                for checkpoint in checkpoints:
                    checkpoint_results[checkpoint] = (
                        "Passed"
                        if "Passed" in phase["completionStatuses"]
                        else phase["completionStatuses"][0]
                    )
            state["phases"][phase["id"]] = {
                "status": "done",
                "checkpointResults": checkpoint_results,
                "checkpointAcknowledgements": {},
            }
        state["evidence"]["test"] = {
            "kind": "maker-attestation",
            "status": "completed",
            "promptCategory": product.test_prompt_category,
            "result": "pass",
            "failureCategory": None,
            "recordedAt": "2026-10-23T00:01:00Z",
            "binding": {
                "provider": product.provider_key,
                "profile": profile,
                "agentSlug": slug,
                "agentId": agent_id,
                "environmentId": ENVIRONMENT_ID,
                "connectionId": CONNECTION_ID.replace("-", ""),
                "draftSemanticHash": state["draftSemanticHash"],
                "connectionBindingHash": state["connectionBindingHash"],
            },
        }
        if profile == "hrsd":
            state["adminSetup"]["phaseHandoffs"][
                "plugin-prerequisites"
            ] = {
                "status": "completed",
                "evidence": {
                    "kind": "structured-admin-attestation",
                    "recordedAt": "2026-10-23T00:00:05Z",
                },
            }
        state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] = {
            "status": "completed",
            "evidence": {
                "kind": "structured-admin-attestation",
                "recordedAt": "2026-10-23T00:00:10Z",
                "oidcCapabilityConfirmed": True,
                "runbookCompleted": True,
                "mappingDetailsCollected": True,
                "claimMappingConfirmed": True,
            },
        }
        if legacy_publish_status is not None:
            state["phases"]["publish"] = {
                "status": legacy_publish_status,
                "checkpointResults": {},
            }
            state["evidence"]["publish"] = {
                "status": "completed",
                "componentHash": state["componentHash"],
            }
        state_path = snow._lifecycle_state_path(context)
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(json.dumps(state), encoding="utf-8")
    return context, state


def _components(connection_id: str | None = None) -> dict:
    return {
        "changeToken": "token",
        "botComponentChanges": [
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000005555",
                    "version": 7,
                    "displayName": "ServiceNow HRSD Common Orchestrator",
                    "schemaName": (
                        f"{snow.HR_SCHEMA_NAME}.topic."
                        "ServiceNowHRSDSystemCommonOrchestrator"
                    ),
                    "state": "Active",
                    "status": "Active",
                    "dialog": {
                        "$kind": "TaskDialog",
                        "action": {
                            "$kind": "InvokeConnectorAction",
                            "connectionReference": "servicenow-ref",
                        },
                    },
                },
            },
                {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "schemaName": f"{snow.HR_SCHEMA_NAME}.topic.WorkdayRuntime",
                    "state": "Active",
                    "dialog": {
                        "$kind": "TaskDialog",
                        "action": {
                            "$kind": "InvokeFlowAction",
                            "flowId": "flow-1",
                        },
                    },
                },
            },
        ],
        "connectionReferenceChanges": [
            {
                "$kind": "ConnectionReferenceInsert",
                "connectionReference": {
                    "$kind": "ConnectionReference",
                    "version": 4,
                    "id": "00000000-0000-4000-8000-000000004444",
                    "connectionId": connection_id,
                    "connectorId": snow.CONNECTOR_ID,
                    "connectionReferenceLogicalName": "servicenow-ref",
                    "displayName": "ServiceNow",
                    "sharedConnectionParameters": json.dumps(
                        {
                            "name": "entraIDUserLogin",
                            "values": {
                                "token:ResourceUri": {"value": APP_CLIENT_ID},
                                "token:InstanceName": {"value": "dev123"},
                            },
                        }
                    ),
                },
            }
        ],
        "connectorDefinitionChanges": [
            {"$kind": "ConnectorDefinitionInsert"}
        ],
    }


def _portal_components() -> dict:
    components = _components()
    components["botComponentChanges"].insert(
        1,
        {
            "$kind": "BotComponentInsert",
            "component": {
                "$kind": "DialogComponent",
                "id": "00000000-0000-4000-8000-000000005556",
                "version": 3,
                "displayName": "ServiceNow HRSD Setup Configurations",
                "schemaName": (
                    f"{snow.HR_SCHEMA_NAME}.topic."
                    "ServiceNowHRSDSetupConfigurations"
                ),
                "state": "Active",
                "status": "Active",
                "dialog": {
                    "$kind": "TaskDialog",
                    "beginDialog": {
                        "actions": [
                            {
                                "$kind": "SetVariable",
                                "displayName": (
                                    "Set ServiceNow Portal BaseURI"
                                ),
                                "variable": (
                                    "Global.ServiceNowHRSDPortalBaseURI"
                                ),
                                "value": {
                                    "$kind": "Literal",
                                    "literalValue": (
                                        "https://dev123.service-now.com/"
                                        "employee"
                                    ),
                                },
                            }
                        ]
                    },
                },
            },
        },
    )
    return components


def _connection_record(
    connection_id: str = CONNECTION_ID,
    *,
    display_name: str = "ServiceNow",
    status: str = "Connected",
    auth_mode: str = snow.AUTH_MODE,
    instance_name: str = "dev123",
    resource_uri: str = APP_CLIENT_ID,
) -> dict:
    return {
        "name": connection_id.replace("-", ""),
        "properties": {
            "displayName": display_name,
            "apiId": snow.CONNECTOR_ID,
            "statuses": [{"target": "token", "status": status}],
            "connectionParametersSet": {
                "name": auth_mode,
                "values": {
                    "token:InstanceName": {"value": instance_name},
                    "token:ResourceUri": {"value": resource_uri},
                },
            },
        },
    }


def _profile_context(
    profile: str,
    *,
    agent_id: str,
    agent_slug: str,
) -> dict:
    schema = (
        snow.HR_SCHEMA_NAME if profile == "hrsd" else snow.IT_SCHEMA_NAME
    )
    return {
        "root": Path.cwd(),
        "agent": {
            "id": agent_id,
            "schema_name": schema,
            "workspace_slug": agent_slug,
        },
        "active": {"slug": agent_slug},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
    }


def _write_shared_source_state(
    root: Path,
    *,
    profile: str,
    agent_id: str,
    agent_slug: str,
    environment_id: str = ENVIRONMENT_ID,
    instance_name: str = "dev123",
    include_oidc: bool = True,
    include_credential: bool = True,
) -> Path:
    source_profile = snow.PRODUCT_PROFILES[profile]
    source_path = (
        root
        / ".local"
        / "connect"
        / source_profile.provider_key
        / "agents"
        / agent_slug
        / "lifecycle.json"
    )
    binding = {
        "tenantId": TENANT_ID,
        "environmentId": environment_id,
        "instanceName": instance_name,
        "applicationClientId": APP_CLIENT_ID,
        "provider": source_profile.provider_key,
        "profile": profile,
        "agentSlug": agent_slug,
        "agentId": agent_id,
    }
    candidates = {
        "kind": "administrator-confirmed-claim-candidates",
        "status": "entra-confirmed",
        "applicationClientId": APP_CLIENT_ID,
        "candidates": ["employee_id"],
        "entraAdministratorConfirmed": True,
        "tokenClaimEmissionConfirmed": True,
        "binding": binding,
        "importedWorkbookSha256": "entra-workbook-hash",
        "recordedAt": "2026-10-01T00:01:00Z",
    }
    candidates["candidateSetSha256"] = snow._identity_candidate_set_hash(
        application_client_id=APP_CLIENT_ID,
        candidates=candidates["candidates"],
        binding=binding,
    )
    mapping = {
        "kind": "administrator-confirmed-identity-mapping",
        "status": "servicenow-confirmed",
        "applicationClientId": APP_CLIENT_ID,
        "tokenClaimName": "employee_id",
        "serviceNowUserField": "employee_number",
        "candidateSetSha256": candidates["candidateSetSha256"],
        "entraAdministratorConfirmed": True,
        "tokenClaimEmissionConfirmed": True,
        "serviceNowAdministratorConfirmed": True,
        "binding": binding,
        "importedWorkbookSha256": "servicenow-workbook-hash",
        "recordedAt": "2026-10-01T00:01:00Z",
    }
    mapping["relationshipSha256"] = snow._identity_mapping_relationship_hash(
        mapping
    )
    phase_handoffs = {
        "preflight": {
            "status": "completed",
            "evidence": {
                "kind": "read-only-resource-discovery",
                "instanceName": instance_name,
                "recordedAt": "2026-10-01T00:00:00Z",
            },
        },
        "entra-registration": {
            "status": "completed",
            "evidence": {
                "kind": "structured-admin-attestation",
                "clientId": APP_CLIENT_ID,
                "tokenClaimCandidates": candidates["candidates"],
                "tokenClaimEmissionConfirmed": True,
                "candidateSetSha256": candidates["candidateSetSha256"],
                "recordedAt": "2026-10-01T00:00:00Z",
            },
        },
        "servicenow-oidc": (
            {
                "status": "completed",
                "evidence": {
                    "kind": "structured-admin-attestation",
                    "oidcCapabilityConfirmed": True,
                    "runbookCompleted": True,
                    "mappingDetailsCollected": True,
                    "claimMappingConfirmed": True,
                    "tokenClaimName": "employee_id",
                    "serviceNowUserField": "employee_number",
                    "candidateSetSha256": candidates[
                        "candidateSetSha256"
                    ],
                    "mappingRelationshipSha256": mapping[
                        "relationshipSha256"
                    ],
                    "recordedAt": "2026-10-01T00:01:00Z",
                },
            }
            if include_oidc
            else {"status": "pending"}
        ),
        "credential": (
            {
                "status": "completed",
                "evidence": {
                    "kind": "maker-connection-return",
                    "connectionId": CONNECTION_ID.replace("-", ""),
                    "recordedAt": "2026-10-01T00:02:00Z",
                },
            }
            if include_credential
            else {"status": "pending"}
        ),
    }
    state = {
        "schemaVersion": snow.LIFECYCLE_SCHEMA_VERSION,
        "provider": source_profile.provider_key,
        "profile": profile,
        "agentSlug": agent_slug,
        "agentId": agent_id,
        "agentSchemaName": source_profile.schema_name,
        "tenantId": TENANT_ID,
        "environmentId": environment_id,
        "ring": "test",
        "attested": True,
        "phases": {
            "preflight": {"status": "done"},
            "entra-registration": {"status": "done"},
            "servicenow-oidc": {
                "status": "done" if include_oidc else "pending"
            },
            "credential": {
                "status": "done" if include_credential else "pending"
            },
            "topics": {"status": "pending"},
            "test": {"status": "pending"},
            "publish": {"status": "pending"},
        },
        "adminSetup": {
            "schemaVersion": snow.ADMIN_SETUP_SCHEMA_VERSION,
            "scope": profile,
            "authMode": snow.AUTH_MODE,
            "preflight": {
                "instanceName": instance_name,
                "instanceOrigin": (
                    f"https://{instance_name}.service-now.com"
                ),
                "discovery": {},
            },
            "phaseHandoffs": phase_handoffs,
        },
        "evidence": {
            "identityClaimCandidates": candidates,
            **({"identityMapping": mapping} if include_oidc else {}),
            **(
                {
                    "credential": {
                        "connectionId": CONNECTION_ID.replace("-", ""),
                        "displayName": "Shared ServiceNow",
                        "connectorId": snow.CONNECTOR_ID,
                        "physicalStatus": "Connected",
                        "authMode": snow.AUTH_MODE,
                        "environmentId": environment_id,
                        "instanceName": instance_name,
                        "resourceUri": APP_CLIENT_ID,
                    },
                    "topics": {"customerChoice": "source-only"},
                    "test": {"status": "completed", "result": "pass"},
                    "publish": {"status": "completed"},
                }
                if include_credential
                else {}
            ),
        },
        "transactions": {"topics": {}},
        "migration": {},
    }
    source_path.parent.mkdir(parents=True, exist_ok=True)
    source_path.write_text(json.dumps(state), encoding="utf-8")
    return source_path


def _write_shared_target_state(
    root: Path,
    *,
    profile: str,
    agent_id: str,
    agent_slug: str,
    instance_name: str | None = None,
) -> tuple[dict, Path]:
    context = _profile_context(
        profile,
        agent_id=agent_id,
        agent_slug=agent_slug,
    )
    target_profile = snow.PRODUCT_PROFILES[profile]
    path = (
        root
        / ".local"
        / "connect"
        / target_profile.provider_key
        / "agents"
        / agent_slug
        / "lifecycle.json"
    )
    with snow.product_profile_context(profile):
        state = snow._state_base(context, _components())
        state["attested"] = True
        state["acceptedContractRevision"] = 1
        state["phases"] = {
            phase: {"status": "pending", "checkpointResults": {}}
            for phase in (
                "preflight",
                "entra-registration",
                "servicenow-oidc",
                "credential",
                "topics",
                "portal-url",
                "agent-connection",
                "test",
            )
        }
        if instance_name:
            state["adminSetup"]["preflight"]["instanceName"] = instance_name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state), encoding="utf-8")
    return context, path


def _write_portal_value_source_state(
    root: Path,
    *,
    profile: str,
    agent_id: str,
    agent_slug: str,
    portal_url: str,
    instance_name: str = "dev123",
    environment_id: str = ENVIRONMENT_ID,
    tenant_id: str = TENANT_ID,
    valid_proof: bool = True,
) -> Path:
    path = _write_shared_source_state(
        root,
        profile=profile,
        agent_id=agent_id,
        agent_slug=agent_slug,
        environment_id=environment_id,
        instance_name=instance_name,
        include_oidc=False,
        include_credential=False,
    )
    state = json.loads(path.read_text(encoding="utf-8"))
    state["tenantId"] = tenant_id
    state["adminSetup"]["phaseHandoffs"]["preflight"] = {
        "status": "completed",
        "evidence": {
            "kind": "maker-confirmed-target",
            "instanceName": instance_name,
            "recordedAt": "2026-10-01T00:00:00Z",
        },
    }
    binding = {
        "provider": state["provider"],
        "profile": state["profile"],
        "agentSlug": agent_slug,
        "agentId": agent_id,
        "environmentId": environment_id,
        "tenantId": tenant_id,
        "instanceName": instance_name,
        "topicId": f"topic-{agent_slug}",
        "topicSchemaName": f"{state['agentSchemaName']}.topic.SetupConfigurations",
    }
    state.setdefault("evidence", {})["portalConfirmation"] = {
        "kind": "maker-confirmed-portal-url",
        "status": "confirmed",
        "source": "committed",
        "portalUrl": portal_url,
        "recordedAt": "2026-10-01T00:03:00Z",
        "readbackAt": "2026-10-01T00:03:00Z",
        "binding": binding,
        "readbackEvidence": {
            "topicVersion": 1 if valid_proof else None,
            "topicContentHashExcludingVersion": (
                "a" * 64 if valid_proof else ""
            ),
        },
    }
    path.write_text(json.dumps(state), encoding="utf-8")
    return path


def _seed_admin_credential_requirements(root: Path) -> None:
    state = snow._state_base(_context(), _components())
    setup = snow._admin_setup(state)
    setup["preflight"]["instanceName"] = "dev123"
    setup["phaseHandoffs"]["entra-registration"] = {
        "status": "completed",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
        },
    }
    snow._write_json_atomic(_lifecycle_path(root), state)


def _write_foundation_context(
    root: Path,
    *,
    components: dict | None = None,
    canonical_account: str | None = None,
    operational_account: str | None = None,
) -> Path:
    setup_path = root / snow.SETUP_STATE
    config_path = root / snow.ACTIVE_CONFIG
    snapshot = Path(
        "workspace/agents/employee-self-service-hr/"
        ".agentbuilder/components.json"
    )
    setup_path.parent.mkdir(parents=True, exist_ok=True)
    setup_path.write_text(
        json.dumps(
            {
                "schema_version": 4,
                "intent": "DA foundation setup",
                "environment": {
                    "id": ENVIRONMENT_ID,
                    "tenant_id": "00000000-0000-4000-8000-000000009999",
                    "ring": "test",
                    "api_version": "2024-10-01",
                    "power_platform_api_endpoint": (
                        "https://example.environment.api.test.powerplatform.com"
                    ),
                    **(
                        {"account": canonical_account}
                        if canonical_account
                        else {}
                    ),
                },
                "agents": {
                    AGENT_ID: {
                        "authoring_ready": True,
                        "connect_ready": False,
                        "agent": {
                            "id": AGENT_ID,
                            "schema_name": snow.HR_SCHEMA_NAME,
                            "realm": "dev",
                            "workspace_slug": AGENT_SLUG,
                        },
                        "workspace": {
                            "folder": f"workspace/agents/{AGENT_SLUG}",
                            "agent_path": "agent.mcs.yml",
                        },
                        "steps": {
                            "SETUP-01": {"state": "done"},
                            "SETUP-02.1": {"state": "done"},
                            "SETUP-02.2": {"state": "blocked"},
                            "SETUP-03": {"state": "done"},
                            "SETUP-04": {"state": "done"},
                            "SETUP-05": {"state": "blocked"},
                            "SETUP-06": {"state": "done"},
                            "SETUP-07": {"state": "done"},
                        },
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    config_path.write_text(
        json.dumps(
            {
                "activeAgent": AGENT_SLUG,
                **(
                    {"account": operational_account}
                    if operational_account
                    else {}
                ),
                "agents": [
                    {
                        "slug": AGENT_SLUG,
                        "botId": AGENT_ID,
                        "agentBuilderChangeSetPath": snapshot.as_posix(),
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    if components is not None:
        snapshot_path = root / snapshot
        snapshot_path.parent.mkdir(parents=True, exist_ok=True)
        snapshot_path.write_text(json.dumps(components), encoding="utf-8")
    return snapshot


def test_summarize_components_separates_servicenow_from_workday_flows() -> None:
    result = snow.summarize_components(_components(CONNECTION_ID))

    assert result["serviceNowTopicCount"] == 1
    assert result["activeServiceNowTopicCount"] == 1
    assert result["serviceNowTopics"][0]["displayName"] == (
        "ServiceNow HRSD Common Orchestrator"
    )
    assert result["serviceNowInvokeConnectorActionCount"] == 1
    assert result["serviceNowInvokeFlowActionCount"] == 0
    assert result["invokeFlowActionCount"] == 1
    assert result["cloudFlowDefinitionCount"] == 0
    assert result["reference"]["instanceName"] == "dev123"
    assert result["reference"]["resourceUri"] == APP_CLIENT_ID


def test_backend_import_defers_openpyxl_until_workbook_use() -> None:
    scripts = snow.SOLUTION_ROOT / "scripts"
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            (
                f"import sys; sys.path.insert(0, {str(scripts)!r}); "
                "import connect_servicenow_da; "
                "raise SystemExit(1 if 'openpyxl' in sys.modules else 0)"
            ),
        ],
        cwd=snow.SOLUTION_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr


def test_summarize_components_requires_active_state_and_status() -> None:
    components = _components(CONNECTION_ID)
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Active"
    topic["status"] = "Inactive"

    result = snow.summarize_components(components)

    assert result["serviceNowTopicCount"] == 1
    assert result["activeServiceNowTopicCount"] == 0


def test_draft_semantic_hash_ignores_only_known_publish_metadata() -> None:
    components = _components(CONNECTION_ID)
    baseline = snow._draft_semantic_hash(components)
    metadata_only = copy.deepcopy(components)
    metadata_only["changeToken"] = "server-revision"
    topic = metadata_only["botComponentChanges"][0]["component"]
    topic["version"] = 99
    topic["auditInfo"] = {"modifiedAt": "2026-10-01T00:00:00Z"}

    assert snow._draft_semantic_hash(metadata_only) == baseline

    authored_edit = copy.deepcopy(metadata_only)
    authored_edit["botComponentChanges"][0]["component"]["displayName"] = (
        "Edited authored topic"
    )
    assert snow._draft_semantic_hash(authored_edit) != baseline

    connection_edit = copy.deepcopy(metadata_only)
    reference = connection_edit["connectionReferenceChanges"][0][
        "connectionReference"
    ]
    parameters = json.loads(reference["sharedConnectionParameters"])
    parameters["values"]["token:InstanceName"]["value"] = "dev456"
    reference["sharedConnectionParameters"] = json.dumps(parameters)
    assert snow._draft_semantic_hash(connection_edit) != baseline


def test_draft_semantic_hash_fails_closed_on_malformed_relevant_shape() -> None:
    components = _components(CONNECTION_ID)
    components["botComponentChanges"] = {"unexpected": True}

    with pytest.raises(snow.ServiceNowConnectError, match="topic collection"):
        snow._draft_semantic_hash(components)


def test_portal_url_payload_updates_only_exact_literal_node() -> None:
    components = _portal_components()
    before = copy.deepcopy(components)

    payload, preimage, postimage = snow.build_portal_url_update_payload(
        components,
        "https://dev123.service-now.com/hrportal",
        expected_instance_name="dev123",
    )

    assert payload["changeToken"] == "token"
    assert len(payload["botComponentChanges"]) == 1
    update = payload["botComponentChanges"][0]
    assert update["$kind"] == "BotComponentUpdate"
    assert update["component"] == postimage
    assert snow._portal_literal_value(preimage) == (
        "https://dev123.service-now.com/employee"
    )
    assert snow._portal_literal_value(postimage) == (
        "https://dev123.service-now.com/hrportal"
    )
    assert components == before
    restored = copy.deepcopy(postimage)
    snow._portal_value_node(restored)["value"]["literalValue"] = (
        snow._portal_literal_value(preimage)
    )
    assert snow._component_content_hash(restored) == (
        snow._component_content_hash(preimage)
    )


def _portal_expression_components() -> dict:
    components = _portal_components()
    topic = components["botComponentChanges"][1]["component"]
    node = snow._portal_value_node(topic)
    node["value"] = {
        "$kind": "ValueExpression",
        "expressionText": (
            '"https://dev123.service-now.com/employee"'
        ),
    }
    return components


def _portal_literal_value_expression_components(
    literal: str = "https://<instance>.service-now.com/sp",
) -> dict:
    components = _portal_components()
    topic = components["botComponentChanges"][1]["component"]
    node = snow._portal_value_node(topic)
    node["value"] = {
        "$kind": "ValueExpression",
        "literalValue": literal,
    }
    return components


def test_portal_value_expression_literal_wrapper_preserves_shape_and_preimage(
) -> None:
    components = _portal_literal_value_expression_components()
    before = copy.deepcopy(components)

    summary = snow.portal_configuration_summary(
        components,
        expected_instance_name="dev123",
    )
    payload, preimage, postimage = snow.build_portal_url_update_payload(
        components,
        "https://dev123.service-now.com/hrportal",
        expected_instance_name="dev123",
    )

    assert summary["configured"] is False
    assert summary["valid"] is False
    assert summary["reason"] == "placeholder"
    assert snow._portal_value_node(preimage)["value"] == {
        "$kind": "ValueExpression",
        "literalValue": "https://<instance>.service-now.com/sp",
    }
    assert snow._portal_value_node(postimage)["value"] == {
        "$kind": "ValueExpression",
        "literalValue": "https://dev123.service-now.com/hrportal",
    }
    assert payload["botComponentChanges"][0]["component"] == postimage
    restored = copy.deepcopy(postimage)
    snow._write_portal_value(
        snow._portal_value_node(restored)["value"],
        snow._portal_literal_value(preimage),
    )
    assert restored == preimage
    assert components == before


def test_portal_value_expression_reads_actual_package_shape() -> None:
    components = _portal_expression_components()

    summary = snow.portal_configuration_summary(
        components,
        expected_instance_name="dev123",
    )

    assert summary["valid"] is True
    assert summary["path"] == "/employee"
    topic = components["botComponentChanges"][1]["component"]
    assert snow._portal_literal_value(topic) == (
        "https://dev123.service-now.com/employee"
    )


def test_portal_value_expression_update_preserves_representation() -> None:
    components = _portal_expression_components()
    before = copy.deepcopy(components)

    payload, preimage, postimage = snow.build_portal_url_update_payload(
        components,
        "https://dev123.service-now.com/hrportal",
        expected_instance_name="dev123",
    )

    pre_value = snow._portal_value_node(preimage)["value"]
    post_value = snow._portal_value_node(postimage)["value"]
    assert pre_value == {
        "$kind": "ValueExpression",
        "expressionText": (
            '"https://dev123.service-now.com/employee"'
        ),
    }
    assert post_value == {
        "$kind": "ValueExpression",
        "expressionText": (
            '"https://dev123.service-now.com/hrportal"'
        ),
    }
    assert payload["botComponentChanges"][0]["component"] == postimage
    assert components == before


@pytest.mark.parametrize(
    "value",
    [
        {
            "$kind": "ValueExpression",
            "expressionText": (
                'Concatenate("https://dev123.service-now.com", "/sp")'
            ),
        },
        {
            "$kind": "ValueExpression",
            "expressionText": "Global.ServiceNowPortalBaseURI",
        },
        {
            "$kind": "ValueExpression",
            "expressionText": " true ",
        },
        {
            "$kind": "ValueExpression",
            "expressionText": (
                '"https://dev123.service-now.com/sp"'
            ),
            "extra": True,
        },
        {
            "$kind": "UnknownExpression",
            "expressionText": (
                '"https://dev123.service-now.com/sp"'
            ),
        },
        {
            "$kind": "ValueExpression",
            "literalValue": 123,
        },
        {
            "$kind": "ValueExpression",
            "literalValue": "https://dev123.service-now.com/sp",
            "expressionText": (
                '"https://dev123.service-now.com/sp"'
            ),
        },
    ],
)
def test_portal_expression_rejects_computed_or_unknown_shapes(
    value: dict,
) -> None:
    components = _portal_expression_components()
    topic = components["botComponentChanges"][1]["component"]
    snow._portal_value_node(topic)["value"] = value

    with pytest.raises(snow.ServiceNowConnectError):
        snow.portal_configuration_summary(
            components,
            expected_instance_name="dev123",
        )


@pytest.mark.parametrize(
    ("value", "message"),
    [
        ("https://other.service-now.com/sp", "confirmed ServiceNow instance"),
        ("http://dev123.service-now.com/sp", "plain HTTPS"),
        ("https://dev123.service-now.com/sp?x=1", "plain HTTPS"),
        ("https://dev123.service-now.com:8443/sp", "plain HTTPS"),
        ("https://dev123.service-now.com/hr portal", "plain HTTPS"),
    ],
)
def test_portal_url_validation_rejects_unsafe_or_inferred_values(
    value: str,
    message: str,
) -> None:
    with pytest.raises(snow.ServiceNowConnectError, match=message):
        snow.normalize_portal_url(
            value,
            expected_instance_name="dev123",
        )


def test_portal_payload_rejects_zero_or_multiple_target_nodes() -> None:
    components = _portal_components()
    setup_topic = components["botComponentChanges"][1]["component"]
    setup_topic["dialog"]["beginDialog"]["actions"] = []
    with pytest.raises(snow.ServiceNowConnectError, match="exactly one Set"):
        snow.build_portal_url_update_payload(
            components,
            "https://dev123.service-now.com/hrportal",
            expected_instance_name="dev123",
        )

    duplicate = copy.deepcopy(components)
    duplicate["botComponentChanges"][1]["component"]["dialog"][
        "beginDialog"
    ]["actions"] = [
        {
            "$kind": "SetVariable",
            "displayName": snow.PORTAL_NODE_DISPLAY_NAME,
            "variable": snow.PORTAL_VARIABLE_NAME,
            "value": {"$kind": "Literal", "literalValue": ""},
        },
        {
            "$kind": "SetVariable",
            "displayName": snow.PORTAL_NODE_DISPLAY_NAME,
            "variable": snow.PORTAL_VARIABLE_NAME,
            "value": {"$kind": "Literal", "literalValue": ""},
        },
    ]
    with pytest.raises(snow.ServiceNowConnectError, match="found 2"):
        snow.build_portal_url_update_payload(
            duplicate,
            "https://dev123.service-now.com/hrportal",
            expected_instance_name="dev123",
        )


@pytest.mark.parametrize(
    "value",
    [
        "https://dev123.service-now.com",
        "https://dev123.service-now.com/",
    ],
)
def test_portal_url_validation_accepts_explicit_root(value: str) -> None:
    assert snow.normalize_portal_url(
        value,
        expected_instance_name="dev123",
    ) == "https://dev123.service-now.com"


def test_portal_summary_never_echoes_url_credentials() -> None:
    components = _portal_components()
    setup_topic = components["botComponentChanges"][1]["component"]
    snow._portal_value_node(setup_topic)["value"]["literalValue"] = (
        "https://user:password@dev123.service-now.com/employee"
    )

    summary = snow.portal_configuration_summary(
        components,
        expected_instance_name="dev123",
    )

    assert summary["valid"] is False
    assert summary["reason"] == "invalid-url"
    assert summary["origin"] == "https://dev123.service-now.com"
    assert "user" not in json.dumps(summary)
    assert "password" not in json.dumps(summary)


def test_portal_payload_rejects_zero_or_multiple_setup_topics() -> None:
    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Setup Configurations topic; found 0",
    ):
        snow.build_portal_url_update_payload(
            _components(),
            "https://dev123.service-now.com/hrportal",
            expected_instance_name="dev123",
        )

    components = _portal_components()
    duplicate = copy.deepcopy(
        components["botComponentChanges"][1]
    )
    duplicate["component"]["id"] = (
        "00000000-0000-4000-8000-000000005557"
    )
    components["botComponentChanges"].append(duplicate)
    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Setup Configurations topic; found 2",
    ):
        snow.build_portal_url_update_payload(
            components,
            "https://dev123.service-now.com/hrportal",
            expected_instance_name="dev123",
        )


def test_topic_state_update_replays_full_dialog_component() -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]

    payload = snow.build_topic_state_update_payload(
        components,
        topic["id"],
        "Inactive",
    )

    assert set(payload) == {"changeToken", "botComponentChanges"}
    assert payload["changeToken"] == "token"
    change = payload["botComponentChanges"][0]
    assert change["$kind"] == "BotComponentUpdate"
    updated = change["component"]
    assert updated["id"] == topic["id"]
    assert updated["version"] == 7
    assert updated["schemaName"] == topic["schemaName"]
    assert updated["dialog"] == topic["dialog"]
    assert updated["state"] == "Inactive"
    assert updated["status"] == "Inactive"
    assert topic["state"] == "Active"
    assert topic["status"] == "Active"


def test_topic_state_update_requires_change_token() -> None:
    components = _components()
    del components["changeToken"]
    topic_id = components["botComponentChanges"][0]["component"]["id"]

    with pytest.raises(snow.ServiceNowConnectError, match="change token"):
        snow.build_topic_state_update_payload(
            components,
            topic_id,
            "Inactive",
        )


def test_topic_state_mutation_requires_confirmation() -> None:
    with pytest.raises(snow.ServiceNowConnectError, match="confirmation"):
        snow.set_topic_state(
            {},
            "00000000-0000-4000-8000-000000005555",
            "Inactive",
            confirmed=False,
        )


def test_enable_all_topics_payload_contains_only_inactive_servicenow_topics() -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"

    payload = snow.build_enable_all_topics_payload(components)

    assert payload["changeToken"] == "token"
    assert len(payload["botComponentChanges"]) == 1
    change = payload["botComponentChanges"][0]
    assert change["$kind"] == "BotComponentUpdate"
    assert change["component"]["id"] == topic["id"]
    assert change["component"]["state"] == "Active"
    assert change["component"]["status"] == "Active"


def test_enable_all_topics_refetches_and_verifies(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(
            self,
            _agent_id: str,
            payload: dict,
        ) -> dict:
            update = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][0]["component"] = update
            components["changeToken"] = "token-2"
            return {}

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.chdir(tmp_path)

    result = snow.enable_all_servicenow_topics(
        {
            "agent": {
                "id": AGENT_ID,
                "schema_name": snow.HR_SCHEMA_NAME,
            },
            "environment": {
                "id": ENVIRONMENT_ID,
                "ring": "test",
            },
        },
        confirmed=True,
    )

    assert result["customerChoice"] == "enable-all"
    assert result["before"] == {"total": 1, "active": 0, "inactive": 1}
    assert result["after"] == {"total": 1, "active": 1, "inactive": 0}
    assert result["changedTopics"][0]["id"] == topic["id"]
    assert result["published"] is False


def test_enable_all_topics_requires_confirmation() -> None:
    with pytest.raises(snow.ServiceNowConnectError, match="confirmation"):
        snow.enable_all_servicenow_topics({}, confirmed=False)


def test_enable_all_topics_is_noop_when_all_are_active(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            raise AssertionError(f"Unexpected update: {payload}")

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.chdir(tmp_path)

    result = snow.enable_all_servicenow_topics(
        {
            "agent": {
                "id": AGENT_ID,
                "schema_name": snow.HR_SCHEMA_NAME,
            },
            "environment": {
                "id": ENVIRONMENT_ID,
                "ring": "test",
            },
        },
        confirmed=True,
    )

    assert result["status"] == "already-active"
    assert result["changedTopics"] == []
    assert result["published"] is False


def test_keep_current_topic_choice_records_without_mutation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            raise AssertionError(f"Unexpected update: {payload}")

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.record_keep_current_topic_choice(_context())

    assert result["customerChoice"] == "keep-current"
    assert result["changedTopics"] == []
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert state["evidence"]["topics"]["customerChoice"] == "keep-current"
    assert state["evidence"]["topics"]["kind"] == "maker-attestation"
    assert state["evidence"]["topics"]["status"] == "recorded"
    assert state["provider"] == snow.PROVIDER_KEY


def test_enable_all_topics_fails_when_refetch_is_still_inactive(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="before any topic mutation",
    ):
        snow.enable_all_servicenow_topics(_context(), confirmed=True)


def test_connection_summary_prefers_token_status() -> None:
    result = snow.connection_summary(
        {
            "name": CONNECTION_ID.replace("-", ""),
            "properties": {
                "displayName": "ServiceNow",
                "statuses": [
                    {"target": "connector", "status": "Ready"},
                    {"target": "token", "status": "Connected"},
                ],
                "connectionParametersSet": {
                    "name": "entraIDUserLogin",
                    "values": {
                        "token:InstanceName": {"value": "dev123"},
                        "token:ResourceUri": {"value": APP_CLIENT_ID},
                        "password": {"value": "must-not-be-returned"},
                    },
                },
            },
        }
    )

    assert result["status"] == "Connected"
    assert result["statusTarget"] == "token"
    assert result["parameterValues"]["token:InstanceName"] == "dev123"
    assert result["parameterValues"]["token:ResourceUri"] == APP_CLIENT_ID
    assert "password" not in result["parameterValues"]


def test_load_context_requires_matching_schema_v4_hr_agent(
    tmp_path: Path,
) -> None:
    snapshot = _write_foundation_context(tmp_path)

    result = snow.load_context(tmp_path)

    assert result["agent"]["id"] == AGENT_ID
    assert result["snapshotPath"] == tmp_path / snapshot


def test_load_context_preserves_matching_setup_account(
    tmp_path: Path,
) -> None:
    account = "maker@example.test"
    _write_foundation_context(
        tmp_path,
        canonical_account=account,
        operational_account=account,
    )

    result = snow.load_context(tmp_path)

    assert result["account"] == account


def test_load_context_rejects_conflicting_setup_accounts(
    tmp_path: Path,
) -> None:
    _write_foundation_context(
        tmp_path,
        canonical_account="expected@example.test",
        operational_account="other@example.test",
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="different Microsoft accounts",
    ):
        snow.load_context(tmp_path)


@pytest.mark.parametrize(
    ("config_key", "value"),
    [
        ("environmentId", "00000000-0000-4000-8000-000000008888"),
        ("tenantId", "00000000-0000-4000-8000-000000007777"),
        (
            "powerPlatformApiEndpoint",
            "https://other.environment.api.test.powerplatform.com",
        ),
        ("ring", "prod"),
        ("agentBuilderApiVersion", "2099-01-01"),
    ],
)
def test_load_context_rejects_operational_environment_contradiction(
    tmp_path: Path,
    config_key: str,
    value: str,
) -> None:
    _write_foundation_context(tmp_path)
    config_path = tmp_path / snow.ACTIVE_CONFIG
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config.update(
        {
            "environmentId": ENVIRONMENT_ID,
            "tenantId": "00000000-0000-4000-8000-000000009999",
            "powerPlatformApiEndpoint": (
                "https://example.environment.api.test.powerplatform.com"
            ),
            "ring": "test",
            "agentBuilderApiVersion": "2024-10-01",
        }
    )
    config[config_key] = value
    config_path.write_text(json.dumps(config), encoding="utf-8")

    with pytest.raises(
        snow.ServiceNowConnectError,
        match=f"different environment {config_key} values",
    ):
        snow.load_context(tmp_path)


def test_agentbuilder_client_uses_setup_account_hint(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    observed: dict[str, object] = {}
    context = {
        "environment": {
            "tenant_id": "00000000-0000-4000-8000-000000009999",
            "ring": "test",
            "power_platform_api_endpoint": (
                "https://example.environment.api.test.powerplatform.com"
            ),
            "api_version": "2024-10-01",
        },
        "account": "maker@example.test",
    }

    def fake_authenticate(*_args, **kwargs):
        observed.update(kwargs)
        return "token"

    monkeypatch.setattr(snow, "authenticate", fake_authenticate)

    client = snow._agentbuilder_client(context)

    assert client is not None
    assert observed["account_hint"] == "maker@example.test"


def test_load_context_rejects_incomplete_foundation_step(
    tmp_path: Path,
) -> None:
    test_load_context_requires_matching_schema_v4_hr_agent(tmp_path)
    setup_path = tmp_path / snow.SETUP_STATE
    setup = json.loads(setup_path.read_text(encoding="utf-8"))
    setup["agents"][AGENT_ID]["steps"]["SETUP-03"]["state"] = "blocked"
    setup["agents"][AGENT_ID]["authoring_ready"] = False
    setup_path.write_text(json.dumps(setup), encoding="utf-8")

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="not ready for authoring",
    ):
        snow.load_context(tmp_path)


def test_record_agent_connection_validates_health_and_persists_attestation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / "servicenow"
        / "agents"
        / AGENT_ID
        / "state.json"
    )
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 1,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
            }
        ),
        encoding="utf-8",
    )

    class FakeConnectivity:
        def get_connection(self, connection_id: str) -> dict:
            assert connection_id == CONNECTION_ID
            return {
                "name": CONNECTION_ID.replace("-", ""),
                "properties": {
                    "displayName": "ServiceNow",
                    "statuses": [{"target": "token", "status": "Connected"}],
                    "connectionParametersSet": {
                        "name": "entraIDUserLogin",
                    },
                },
            }

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {"fetch_components": lambda self, _agent_id: _components()},
        )(),
    )

    result = snow.record_agent_connection_attestation(
        _context(),
        CONNECTION_ID,
    )

    assert result["connectionId"] == CONNECTION_ID.replace("-", "")
    assert result["status"] == "completed"
    assert result["binding"]["connectionId"] == result["connectionId"]
    assert result["physicalStatus"] == "Connected"
    assert result["makerAttested"] is True
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert state["evidence"]["agentConnection"]["makerAttested"] is True
    assert (
        state["evidence"]["agentConnection"]["binding"]["connectionId"]
        == state["evidence"]["agentConnection"]["connectionId"]
    )
    assert state["migration"]["legacySourcePath"].endswith("state.json")


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_publish_is_disabled_before_any_remote_client_or_write(
    monkeypatch: pytest.MonkeyPatch,
    profile: str,
) -> None:
    client_calls = 0

    def forbidden_client(_context: dict) -> None:
        nonlocal client_calls
        client_calls += 1
        pytest.fail("Disabled publish must not construct a remote client.")

    monkeypatch.setattr(snow, "_agentbuilder_client", forbidden_client)
    context = _profile_context(
        profile,
        agent_id=AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID,
        agent_slug=(
            "employee-self-service-hr"
            if profile == "hrsd"
            else "employee-self-service-it"
        ),
    )
    with snow.product_profile_context(profile):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="Automatic publishing is disabled",
        ):
            snow.publish(context, confirmed=True)

    assert client_calls == 0


def test_publish_cli_returns_explicit_disabled_error_without_remote_client(
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    client_calls = 0

    def forbidden_client(_context: dict) -> None:
        nonlocal client_calls
        client_calls += 1
        pytest.fail("Disabled publish CLI must not construct a client.")

    monkeypatch.setattr(snow, "load_context", lambda **_kwargs: _context())
    monkeypatch.setattr(snow, "_agentbuilder_client", forbidden_client)

    assert snow.main(
        [
            "--provider",
            "servicenow-da-hrsd",
            "publish",
            "--yes",
        ]
    ) == 1

    error = json.loads(capsys.readouterr().out)
    assert error["status"] == "error"
    assert "Automatic publishing is disabled" in error["message"]
    assert client_calls == 0


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_validation_realm_handoff_requires_exact_complete_dev_record(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    product = snow.PRODUCT_PROFILES[profile]
    context = _profile_context(
        profile,
        agent_id=AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID,
        agent_slug=(
            "employee-self-service-hr"
            if profile == "hrsd"
            else "employee-self-service-it"
        ),
    )
    context["root"] = tmp_path
    with snow.product_profile_context(profile):
        handoff = snow.validation_realm_handoff(context)

    assert handoff["status"] == "dev-connect-required"
    assert handoff["question"] is None
    assert handoff["devCompletion"]["reason"] == "no-dev-lifecycle"
    assert handoff["provider"] == product.provider_key
    _assert_recovery_handoff(handoff["recoveryHandoff"])
    assert "Dev prerequisite incomplete" in handoff["recoveryHandoff"]["category"]


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_complete_dev_record_unlocks_realm_chooser_without_state_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile=profile)
    with snow.product_profile_context(profile):
        path = snow._lifecycle_state_path(context)
        before = path.read_bytes()
        handoff = snow.validation_realm_handoff(context)

    assert handoff["status"] == "choice-required"
    assert handoff["devCompletion"]["reason"] == "complete"
    assert handoff["question"]["options"] == [
        {"label": "Dev"},
        {"label": "Test"},
        {"label": "Prod"},
    ]
    assert path.read_bytes() == before


@pytest.mark.parametrize(
    ("profile", "realm", "prompt"),
    (
        ("hrsd", "test", "List my open HR cases"),
        ("hrsd", "prod", "List my open HR cases"),
        ("itsm", "test", "List my open IT tickets"),
        ("itsm", "prod", "List my open IT tickets"),
    ),
)
def test_realm_functional_question_uses_profile_prompt(
    profile: str,
    realm: str,
    prompt: str,
) -> None:
    with snow.product_profile_context(profile):
        question = snow._realm_functional_test_question(
            realm=realm,
            test_pane_url="https://target.example/test-pane",
        )

    assert question["header"] == (
        f"Validate the {realm.title()} ServiceNow target"
    )
    assert f"`{prompt}`" in question["question"]
    assert "https://target.example/test-pane" in question["question"]
    assert "Do not add an employee name" in question["question"]
    assert "Do not paste the response" in question["question"]
    assert "directs you back to the Dev agent" in question["question"]
    assert "failure category" not in question["question"]
    assert [option["label"] for option in question["options"]] == [
        "Pass",
        "Not pass",
    ]


def test_manual_realm_status_is_not_presented_as_manually_finished() -> None:
    rendered = snow._customer_facing_realm_check(
        {
            "checkpoint": "SN-REALM-OIDC-MANUAL",
            "status": "Manual",
            "automationType": "manual",
            "result": "No supported read API proves the target OIDC mapping.",
        }
    )

    assert rendered["status"] == "Manual"
    assert rendered["customerStatus"] == "Not automatically verified"
    assert "does not mean" in rendered["customerExplanation"]
    assert "already finished or confirmed" in rendered["customerExplanation"]


@pytest.mark.parametrize("realm", ("test", "prod"))
def test_missing_realm_deployment_handoff_is_resumable_and_ordered(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    realm: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(tmp_path, profile="itsm")
    with snow.product_profile_context("itsm"):
        path = snow._lifecycle_state_path(context)
    phases_before = copy.deepcopy(state["phases"])
    dev_test_before = copy.deepcopy(state["evidence"]["test"])

    with snow.product_profile_context("itsm"):
        recorded = snow.record_realm_deployment(
            context,
            realm=realm,
            status="not-deployed",
        )

    assert recorded["status"] == "deployment-required"
    assert recorded["remoteMutationCount"] == 0
    question = recorded["question"]
    body = question["question"]
    realm_label = realm.title()
    assert question["header"] == f"Deploy this agent to {realm_label}"
    assert [option["label"] for option in question["options"]] == [
        "Completed",
        "Not yet",
    ]
    assert question["allowFreeformInput"] is False
    deploy_index = body.index(f"Deploy to {realm_label}")
    save_index = body.index("**Save**", deploy_index)
    second_deploy_index = body.index("**Deploy**", save_index)
    open_index = body.index("**Open Agent**", second_deploy_index)
    assert deploy_index < save_index < second_deploy_index < open_index
    assert "upper-right" in body
    assert "lower-left" in body
    assert "will not deploy, promote, import, publish" in body
    assert "does not prove validation passed" in body
    assert "actual ALM owner" in body
    assert "do not deploy it again" in body
    assert "employee-self-service-it" in body
    assert context["environment"]["id"] in body

    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved["phases"] == phases_before
    assert saved["evidence"]["test"] == dev_test_before
    assert saved["evidence"]["realmDeployments"][realm]["status"] == (
        "not-deployed"
    )
    assert "realmValidations" not in saved["evidence"]

    monkeypatch.setattr(
        snow,
        "_agentbuilder_token",
        lambda _context: pytest.fail(
            "Deployment-required resume must not call target APIs."
        ),
    )
    with snow.product_profile_context("itsm"):
        resumed = snow.validate_realm(context, realm=realm)

    assert resumed["status"] == "deployment-required"
    assert resumed["question"] == question
    assert resumed["remoteMutationCount"] == 0


def test_deployment_completion_requires_exact_read_only_validation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(tmp_path, profile="itsm")
    with snow.product_profile_context("itsm"):
        path = snow._lifecycle_state_path(context)
    phases_before = copy.deepcopy(state["phases"])
    with snow.product_profile_context("itsm"):
        snow.record_realm_deployment(
            context,
            realm="test",
            status="not-deployed",
        )
        completed = snow.record_realm_deployment(
            context,
            realm="test",
            status="deployment-completed",
        )

    assert completed["status"] == "validation-required"
    assert completed["question"] is None
    assert completed["remoteMutationCount"] == 0
    assert completed["next"]["action"] == (
        "rerun-read-only-realm-validation"
    )
    assert "not validation evidence" in completed["next"]["instruction"]

    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved["phases"] == phases_before
    assert saved["evidence"]["realmDeployments"]["test"]["status"] == (
        "deployment-completed"
    )
    assert "realmValidations" not in saved["evidence"]

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: (_ for _ in ()).throw(
            snow.ServiceNowConnectError("Exact target is still unproved.")
        ),
    )
    with snow.product_profile_context("itsm"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="Exact target is still unproved",
        ):
            snow.validate_realm(context, realm="test")

    saved_after = json.loads(path.read_text(encoding="utf-8"))
    assert saved_after["phases"] == phases_before
    assert "realmValidations" not in saved_after["evidence"]


@pytest.mark.parametrize("result", ("fail", None))
def test_failed_or_missing_current_test_does_not_unlock_realm_chooser(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    result: str | None,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    if result is None:
        state["evidence"].pop("test")
    else:
        state["evidence"]["test"]["result"] = result
        state["evidence"]["test"]["failureCategory"] = "connector"
    snow._lifecycle_state_path(context).write_text(
        json.dumps(state),
        encoding="utf-8",
    )

    with snow.product_profile_context("hrsd"):
        handoff = snow.validation_realm_handoff(context)

    assert handoff["status"] == "dev-connect-required"
    assert handoff["devCompletion"]["reason"] == "required-phases-incomplete"
    assert "test" in handoff["devCompletion"]["phases"]
    _assert_recovery_handoff(handoff["recoveryHandoff"])


def test_missing_entra_family_member_does_not_unlock_realm_chooser(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    state["phases"]["entra-registration"]["checkpointResults"].pop(
        "SN-DA-HRSD-ENTRA-PREAUTH-001"
    )
    snow._lifecycle_state_path(context).write_text(
        json.dumps(state),
        encoding="utf-8",
    )

    with snow.product_profile_context("hrsd"):
        handoff = snow.validation_realm_handoff(context)

    assert handoff["status"] == "dev-connect-required"
    assert "entra-registration" in handoff["devCompletion"]["phases"]


def test_missing_manual_ack_does_not_unlock_realm_chooser(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] = {
        "status": "pending"
    }
    snow._lifecycle_state_path(context).write_text(
        json.dumps(state),
        encoding="utf-8",
    )

    with snow.product_profile_context("hrsd"):
        handoff = snow.validation_realm_handoff(context)

    assert handoff["status"] == "dev-connect-required"
    assert "servicenow-oidc" in handoff["devCompletion"]["phases"]


@pytest.mark.parametrize("legacy_status", ("pending", "done"))
def test_schema_v7_retires_legacy_publish_without_resetting_test(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    legacy_status: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(
        tmp_path,
        profile="hrsd",
        legacy_publish_status=legacy_status,
    )
    state["schemaVersion"] = 6
    state["acceptedContractRevision"] = 8
    path = snow._lifecycle_state_path(context)
    path.write_text(json.dumps(state), encoding="utf-8")
    test_before = copy.deepcopy(state["evidence"]["test"])

    with snow.product_profile_context("hrsd"):
        migrated = snow._load_lifecycle_state(context)

    assert migrated["schemaVersion"] == 8
    assert "publish" not in migrated["phases"]
    assert migrated["retiredPhases"]["publish"]["status"] == legacy_status
    assert migrated["evidence"]["publish"]["status"] == "completed"
    assert migrated["evidence"]["test"] == test_before
    assert migrated["migration"]["schemaV7"][
        "retiredAutomaticPublishPhase"
    ] is True
    assert migrated["migration"]["schemaV8"][
        "retiredEntraCheckpointIds"
    ] == [
        "SN-DA-HRSD-ENTRA-CONSENT-001",
        "SN-DA-HRSD-ENTRA-PERMISSIONS-001",
    ]


def test_schema_v8_retires_graph_gates_without_fabricating_completion(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["schemaVersion"] = 7
    state["phases"] = {
        "entra-registration": {
            "status": "in-progress",
            "checkpointResults": {
                "SN-DA-HRSD-ENTRA-APP-001": "Passed",
                "SN-DA-HRSD-ENTRA-CLAIMS-001": "Manual",
                "SN-DA-HRSD-ENTRA-SCOPE-001": "Passed",
                "SN-DA-HRSD-ENTRA-PREAUTH-001": "Passed",
                "SN-DA-HRSD-ENTRA-PERMISSIONS-001": "NotConfigured",
                "SN-DA-HRSD-ENTRA-CONSENT-001": "Failed",
            },
            "checkpointAcknowledgements": {
                "SN-DA-HRSD-ENTRA-PERMISSIONS-001": {
                    "status": "NotConfigured",
                },
                "SN-DA-HRSD-ENTRA-CONSENT-001": {
                    "status": "Failed",
                },
            },
        }
    }
    path = snow._lifecycle_state_path(context)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state), encoding="utf-8")
    emitted: list[dict] = []
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, event: emitted.append(copy.deepcopy(event)),
    )

    migrated = snow._load_lifecycle_state(context)
    snow._write_lifecycle_state(context, migrated)

    phase = migrated["phases"]["entra-registration"]
    assert phase["status"] == "in-progress"
    assert set(phase["checkpointResults"]) == {
        "SN-DA-HRSD-ENTRA-APP-001",
        "SN-DA-HRSD-ENTRA-CLAIMS-001",
        "SN-DA-HRSD-ENTRA-SCOPE-001",
        "SN-DA-HRSD-ENTRA-PREAUTH-001",
    }
    assert phase["checkpointAcknowledgements"] == {}
    retired = migrated["retiredCheckpointResults"]["entra-registration"]
    assert retired["SN-DA-HRSD-ENTRA-PERMISSIONS-001"]["status"] == (
        "NotConfigured"
    )
    assert retired["SN-DA-HRSD-ENTRA-CONSENT-001"]["status"] == "Failed"
    assert retired["SN-DA-HRSD-ENTRA-CONSENT-001"][
        "acknowledgement"
    ] == {"status": "Failed"}
    assert migrated["migration"]["schemaV8"][
        "retiredEntraCheckpointIds"
    ] == [
        "SN-DA-HRSD-ENTRA-CONSENT-001",
        "SN-DA-HRSD-ENTRA-PERMISSIONS-001",
    ]
    assert emitted == []


def test_validate_dev_records_choice_without_changing_active_agent(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    config_path = tmp_path / snow.ACTIVE_CONFIG
    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text(
        json.dumps({"activeAgent": "other-agent"}),
        encoding="utf-8",
    )

    with snow.product_profile_context("hrsd"):
        result = snow.validate_realm(context, realm="dev")

    assert result["status"] == "dev-resume"
    assert result["remoteMutationCount"] == 0
    assert json.loads(config_path.read_text(encoding="utf-8"))[
        "activeAgent"
    ] == "other-agent"
    with snow.product_profile_context("hrsd"):
        saved = json.loads(
            snow._lifecycle_state_path(context).read_text(encoding="utf-8")
        )
    assert saved["evidence"]["validationRealmSelection"]["realm"] == "dev"
    assert saved["evidence"]["test"]["result"] == "pass"


@pytest.mark.parametrize("realm,realm_value", (("test", 1), ("prod", 2)))
def test_validate_target_realm_is_read_only_and_uses_exact_target_context(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    realm: str,
    realm_value: int,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="itsm")
    target_agent_id = "00000000-0000-4000-8000-000000007777"
    target_environment_id = "00000000-0000-4000-8000-000000008888"
    inventory_environment_id = (
        f"Default-{target_environment_id}"
        if realm == "test"
        else target_environment_id
    )
    target_instance = "target-snow"
    target_client_id = "00000000-0000-4000-8000-000000006666"

    class FakeClient:
        def __init__(self, target: bool) -> None:
            self.target = target
            self.mutation_count = 0

        def get_realms(self, _agent_id: str) -> dict:
            return {
                "routeRealm": 0,
                "siblingRealms": [
                    {"realm": realm_value, "botId": target_agent_id}
                ],
            }

        def get_realm_configuration(
            self,
            _agent_id: str,
            requested_realm: int,
        ) -> dict:
            assert requested_realm == realm_value
            return {
                "realm": realm.title(),
                "cdsBotId": target_agent_id,
                "schemaName": f"{snow.IT_SCHEMA_NAME}_{realm}",
                "grsRepositoryId": "family-1",
                "commitSha": "abc123",
            }

        def get_agent(self, _agent_id: str) -> dict:
            assert self.target
            return {
                "botId": target_agent_id,
                "realm": realm,
                "schemaName": f"{snow.IT_SCHEMA_NAME}_{realm}",
            }

        def publish_agent(self, _agent_id: str) -> dict:
            self.mutation_count += 1
            pytest.fail("Realm validation must not publish.")

        def update_components(self, *_args: object) -> dict:
            self.mutation_count += 1
            pytest.fail("Realm validation must not update components.")

    dev_client = FakeClient(target=False)
    target_client = FakeClient(target=True)

    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return [
                {
                    "name": CONNECTION_ID.replace("-", ""),
                    "properties": {
                        "displayName": "Target ServiceNow",
                        "statuses": [
                            {"target": "token", "status": "Connected"}
                        ],
                        "connectionParametersSet": {
                            "name": snow.AUTH_MODE,
                            "values": {
                                "token:InstanceName": {
                                    "value": target_instance
                                },
                                "token:ResourceUri": {
                                    "value": target_client_id
                                },
                            },
                        },
                    },
                }
            ]

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: dev_client)
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [
            {"id": inventory_environment_id, "displayName": "Target"}
        ],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        lambda *_args, **_kwargs: target_client,
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client_for_environment",
        lambda *_args, **_kwargs: FakeConnectivity(),
    )

    with snow.product_profile_context("itsm"):
        result = snow.validate_realm(
            context,
            realm=realm,
            target_instance=target_instance,
            target_client_id=target_client_id,
        )

    assert result["status"] == "ready-for-test-pane"
    assert result["deploymentEvidence"] is None
    assert result["readinessSummary"] == {
        "automaticPrechecks": "Complete",
        "functionalTest": "Pending",
        "endToEndValidation": "Not yet verified",
    }
    question = result["question"]
    assert question["header"] == (
        f"Validate the {realm.title()} ServiceNow target"
    )
    assert target_environment_id in question["question"]
    assert target_agent_id in question["question"]
    assert "`List my open IT tickets`" in question["question"]
    assert [option["label"] for option in question["options"]] == [
        "Pass",
        "Not pass",
    ]
    assert "empty list is not automatically Not pass" in question["question"]
    assert "generic conversational answer is not automatically Pass" in (
        question["question"]
    )
    assert result["allAutomatedChecksPassed"] is True
    assert result["target"]["environmentId"] == target_environment_id
    assert result["target"]["inventoryEnvironmentId"] == (
        inventory_environment_id
    )
    assert result["target"]["agentId"] == target_agent_id
    assert result["target"]["schemaName"] == f"{snow.IT_SCHEMA_NAME}_{realm}"
    assert result["automatedChecks"][0]["checkpoint"] == "ESS-SOLN-001"
    assert result["automatedChecks"][0]["status"] == "Passed"
    assert result["automatedChecks"][1]["checkpoint"] == (
        "SN-REALM-CREDENTIAL-001"
    )
    assert result["automatedChecks"][1]["status"] == "Passed"
    assert {
        row["checkpoint"] for row in result["manualChecks"]
    } == {
        "SN-REALM-OIDC-MANUAL",
        "SN-REALM-AGENT-BINDING-MANUAL",
    }
    assert all(
        row["status"] == "Manual"
        and row["customerStatus"] == "Not automatically verified"
        and "does not mean" in row["customerExplanation"]
        for row in result["manualChecks"]
    )
    with snow.product_profile_context("itsm"):
        saved = json.loads(
            snow._lifecycle_state_path(context).read_text(encoding="utf-8")
        )
    realm_record = saved["evidence"]["realmValidations"][realm]
    assert realm_record["targetEnvironmentId"] == target_environment_id
    assert realm_record["targetAgentId"] == target_agent_id
    assert realm_record["sourceDevTestBinding"]["provider"] == (
        snow.PRODUCT_PROFILES["itsm"].provider_key
    )
    assert len(realm_record["targetConfigurationSha256"]) == 64
    assert realm_record["targetCredentialExpectation"] == {
        "source": "explicit-target-realm-confirmation",
        "instanceName": target_instance,
        "applicationClientId": target_client_id,
    }
    assert result["remoteMutationCount"] == 0
    assert target_client.mutation_count == 0
    assert result["next"]["url"].startswith(
        "https://copilotstudio.test.microsoft.com/"
    )
    assert target_environment_id in result["next"]["url"]
    assert target_agent_id in result["next"]["url"]


def test_validate_target_realm_blocks_when_exact_environment_is_unproven(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": "00000000-0000-4000-8000-000000007777",
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
                "commitSha": "test-commit",
            }

    class MissingClient:
        def get_agent(self, _agent_id: str) -> dict:
            raise snow.AgentBuilderHTTPError("Direct agent", 404)

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: DevClient())
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [
            {
                "id": "00000000-0000-4000-8000-000000008888",
                "displayName": "Unknown",
            }
        ],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        lambda *_args, **_kwargs: MissingClient(),
    )

    with snow.product_profile_context("hrsd"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="Could not prove one exact Test environment",
        ) as raised:
            snow.validate_realm(context, realm="test")

    assert raised.value.details["matchCount"] == 0
    _assert_recovery_handoff(raised.value.details["recoveryHandoff"])
    recovery = raised.value.details["recoveryHandoff"]
    assert recovery["category"] == "Target agent not found"
    assert "does not by itself prove missing permission" in recovery["action"]
    discovery_error = raised.value.details["discoveryErrors"][0]
    assert discovery_error["statusCode"] == 404
    assert discovery_error["requestPath"].endswith(
        "/00000000-0000-4000-8000-000000007777"
    )


def test_target_discovery_normalizes_default_ids_and_preserves_unknown_target(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="itsm")
    target_agent_id = "63e6cd4a-1c3b-4a18-bbc9-28bf67ac4d42"
    default_guid = "935884d7-bdee-469b-a461-fcc530a3ac83"
    trial_guid = "b0a07253-5000-eb91-abfc-757c90091b0f"
    source_guid = context["environment"]["id"]

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": f"{snow.IT_SCHEMA_NAME}_test",
                "grsRepositoryId": "family-1",
                "commitSha": "test-commit",
            }

    class CandidateClient:
        def __init__(self, outcome: str, environment_id: str) -> None:
            self.outcome = outcome
            self.host = f"https://{environment_id}.environment.example"

        def get_agent(self, _agent_id: str) -> dict:
            if self.outcome == "unavailable":
                raise OSError("candidate endpoint unavailable")
            if self.outcome == "not-found":
                raise snow.AgentBuilderHTTPError("Direct agent", 404)
            return {
                "botId": target_agent_id,
                "realm": "dev",
            }

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: DevClient(),
    )
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [
            {"id": f"Default-{default_guid}", "displayName": "Default"},
            {"id": trial_guid, "displayName": "Trial"},
            {"id": source_guid, "displayName": "Source"},
        ],
    )

    def candidate_client(
        _context: dict,
        *,
        environment_id: str,
        token: str,
        require_resolved_host: bool,
    ) -> CandidateClient:
        assert token == "token"
        assert require_resolved_host is True
        return CandidateClient(
            "unavailable"
            if environment_id == default_guid
            else "not-found"
            if environment_id == trial_guid
            else "source",
            environment_id,
        )

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        candidate_client,
    )

    with snow.product_profile_context("itsm"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="Could not prove one exact Test environment",
        ) as raised:
            snow.validate_realm(context, realm="test")

    errors = raised.value.details["discoveryErrors"]
    assert errors == [
        {
            "environmentId": default_guid,
            "inventoryEnvironmentId": f"Default-{default_guid}",
            "error": "OSError",
            "requestPath": (
                f"/copilotstudio/minimalBots/api/{target_agent_id}"
            ),
            "requestHost": (
                f"https://{default_guid}.environment.example"
            ),
            "environmentIdSource": "default-prefix-tenant-guid",
            "transportCategory": "os-transport",
            "errorTypeChain": ["OSError"],
            "errno": None,
            "detail": "candidate endpoint unavailable",
        },
        {
            "environmentId": trial_guid,
            "inventoryEnvironmentId": trial_guid,
            "error": "AgentBuilderHTTPError",
            "statusCode": 404,
            "requestPath": (
                f"/copilotstudio/minimalBots/api/{target_agent_id}"
            ),
        },
        {
            "environmentId": source_guid,
            "inventoryEnvironmentId": source_guid,
            "error": "TargetRealmMismatch",
            "expectedRealm": "test",
            "observedRealm": "dev",
            "targetAgentId": target_agent_id,
        },
    ]
    recovery = raised.value.details["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["category"] == "Candidate environment read unavailable"
    assert f"Default-{default_guid}" in recovery["values"]
    assert "Target presence in these environments remains unknown" in (
        recovery["values"]
    )
    assert "Do not repair ALM binding" in recovery["action"]
    assert "Deploy to Test" not in recovery["message"]


def test_target_discovery_accepts_exact_native_match_with_unreadable_bystander(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="itsm")
    target_agent_id = "63e6cd4a-1c3b-4a18-bbc9-28bf67ac4d42"
    unreadable_guid = "935884d7-bdee-469b-a461-fcc530a3ac83"
    matched_guid = "b0a07253-5000-eb91-abfc-757c90091b0f"

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": f"{snow.IT_SCHEMA_NAME}_test",
                "grsRepositoryId": "family-1",
                "commitSha": "test-commit",
            }

    class CandidateClient:
        def __init__(self, readable: bool, environment_id: str) -> None:
            self.readable = readable
            self.host = f"https://{environment_id}.environment.example"

        def get_agent(self, _agent_id: str) -> dict:
            if not self.readable:
                raise OSError("candidate endpoint unavailable")
            return {
                "botId": target_agent_id,
                "realm": "test",
                "schemaName": f"{snow.IT_SCHEMA_NAME}_test",
            }

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: DevClient(),
    )

    def list_candidates(*_args: object, **kwargs: object) -> list[dict]:
        session = kwargs["session"]
        assert session.adapters["https://"].max_retries.total == 0
        return [
            {"id": f"Default-{unreadable_guid}"},
            {"id": matched_guid},
        ]

    monkeypatch.setattr(snow, "list_environments", list_candidates)
    def candidate_client(
        _context: dict,
        *,
        environment_id: str,
        token: str,
        require_resolved_host: bool,
    ) -> CandidateClient:
        assert token == "token"
        assert require_resolved_host is True
        return CandidateClient(
            environment_id == matched_guid,
            environment_id,
        )

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        candidate_client,
    )

    with snow.product_profile_context("itsm"):
        target = snow._resolve_target_realm(context, realm="test")

    assert target["environmentId"] == matched_guid
    assert target["agentId"] == target_agent_id
    assert target["schemaName"] == f"{snow.IT_SCHEMA_NAME}_test"
    assert target["almFamilyId"] == "family-1"
    assert target["discoveryWarnings"] == [
        {
            "environmentId": unreadable_guid,
            "inventoryEnvironmentId": f"Default-{unreadable_guid}",
            "error": "OSError",
            "requestPath": (
                f"/copilotstudio/minimalBots/api/{target_agent_id}"
            ),
            "requestHost": (
                f"https://{unreadable_guid}.environment.example"
            ),
            "environmentIdSource": "default-prefix-tenant-guid",
            "transportCategory": "os-transport",
            "errorTypeChain": ["OSError"],
            "errno": None,
            "detail": "candidate endpoint unavailable",
        }
    ]


def test_target_discovery_disables_client_automatic_retries() -> None:
    client = SimpleNamespace(session=snow.requests.Session())

    assert snow._without_automatic_retries(client) is client
    retries = client.session.adapters["https://"].max_retries
    assert retries.total == 0
    assert retries.connect == 0
    assert retries.read == 0
    assert retries.status == 0


def test_target_discovery_requires_a_resolved_environment_host(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    attempted: list[str] = []

    def unresolved(hostname: str, _port: int) -> None:
        attempted.append(hostname)
        raise snow.socket.gaierror("not found")

    monkeypatch.setattr(snow.socket, "getaddrinfo", unresolved)

    with pytest.raises(
        OSError,
        match="No supported AgentBuilder environment host resolved",
    ) as raised:
        snow._derive_resolved_environment_host(
            "935884d7-bdee-469b-a461-fcc530a3ac83",
            "prod",
        )

    assert len(attempted) == 2
    assert all(host in str(raised.value) for host in attempted)


@pytest.mark.parametrize(
    ("error", "category"),
    (
        (
            snow.requests.exceptions.ConnectionError(
                "NameResolutionError: getaddrinfo failed"
            ),
            "dns",
        ),
        (
            snow.requests.exceptions.ConnectTimeout("connect timed out"),
            "connect-timeout",
        ),
        (
            snow.requests.exceptions.SSLError("TLS handshake failed"),
            "tls",
        ),
        (OSError(10061, "connection refused"), "connection-refused"),
    ),
)
def test_target_discovery_transport_error_categories(
    error: BaseException,
    category: str,
) -> None:
    evidence = snow._transport_error_evidence(error)

    assert evidence["transportCategory"] == category
    assert evidence["errorTypeChain"]
    assert evidence["detail"]


def test_validate_target_realm_schema_mismatch_reports_expected_and_actual(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="itsm")
    target_agent_id = "00000000-0000-4000-8000-000000007777"

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {
                "routeRealm": 0,
                "siblingRealms": [
                    {"realm": 1, "botId": target_agent_id}
                ],
            }

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
                "commitSha": "test-commit",
            }

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: DevClient(),
    )
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: pytest.fail(
            "Schema mismatch must stop before environment discovery."
        ),
    )

    with snow.product_profile_context("itsm"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="different agent schema",
        ) as raised:
            snow.validate_realm(context, realm="test")

    assert raised.value.details["expectedSchemaName"] == snow.IT_SCHEMA_NAME
    assert raised.value.details["expectedSchemaNames"] == sorted(
        {
            snow.IT_SCHEMA_NAME.casefold(),
            f"{snow.IT_SCHEMA_NAME}_test".casefold(),
        }
    )
    assert raised.value.details["configuredSchemaName"] == snow.HR_SCHEMA_NAME
    assert raised.value.details["targetAgentId"] == target_agent_id
    assert raised.value.details["targetEnvironmentId"] is None
    assert raised.value.details["almFamilyId"] == "family-1"
    assert raised.value.details["commitSha"] == "test-commit"
    assert raised.value.details["observedTargetAgentId"] is None
    assert raised.value.details["observedSchemaName"] is None
    recovery = raised.value.details["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert snow.IT_SCHEMA_NAME in recovery["values"]
    assert snow.HR_SCHEMA_NAME in recovery["values"]
    assert target_agent_id in recovery["values"]
    assert "not reached and unknown" in recovery["values"]
    assert "employee-self-service-it" in recovery["values"]
    assert recovery["owner"] == (
        "Copilot Studio / Power Platform ALM owner for this agent family"
    )
    assert "do not select a different agent automatically" in recovery["action"]


@pytest.mark.parametrize(
    ("observed_schema", "expected_error"),
    (
        (None, "TargetSchemaUnknown"),
        (snow.HR_SCHEMA_NAME, "TargetSchemaMismatch"),
    ),
)
def test_validate_target_realm_blocks_unknown_or_mismatched_observed_schema(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    observed_schema: str | None,
    expected_error: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="itsm")
    target_agent_id = "00000000-0000-4000-8000-000000007777"
    target_environment_id = "00000000-0000-4000-8000-000000008888"

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": "",
                "grsRepositoryId": "family-1",
            }

    class TargetClient:
        def get_agent(self, _agent_id: str) -> dict:
            return {
                "botId": target_agent_id,
                "realm": "test",
                **(
                    {"schemaName": observed_schema}
                    if observed_schema is not None
                    else {}
                ),
            }

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: DevClient(),
    )
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [{"id": target_environment_id}],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        lambda *_args, **_kwargs: TargetClient(),
    )

    with snow.product_profile_context("itsm"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="Could not prove one exact Test environment",
        ) as raised:
            snow.validate_realm(context, realm="test")

    schema_error = raised.value.details["discoveryErrors"][0]
    assert schema_error["error"] == expected_error
    assert schema_error["expectedSchemaName"] == snow.IT_SCHEMA_NAME
    assert schema_error["configuredSchemaName"] is None
    assert schema_error["observedSchemaName"] == observed_schema
    recovery = raised.value.details["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert snow.IT_SCHEMA_NAME in recovery["values"]
    assert f"`{observed_schema or 'unknown'}`" in recovery["values"]


@pytest.mark.parametrize("http_status", (403, 429))
def test_validate_target_realm_surfaces_environment_inventory_failure(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    http_status: int,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": "00000000-0000-4000-8000-000000007777",
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
            }

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: DevClient())

    def fail_inventory(*_args: object, **_kwargs: object) -> list[dict]:
        raise snow.AgentBuilderHTTPError("Environment listing", http_status)

    monkeypatch.setattr(snow, "list_environments", fail_inventory)

    with snow.product_profile_context("hrsd"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="Could not list the environments",
        ) as raised:
            snow.validate_realm(context, realm="test")

    assert raised.value.details["httpStatus"] == http_status
    recovery = raised.value.details["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["category"] == (
        "Read throttled" if http_status == 429 else "Read permission"
    )


def test_target_realm_throttling_stops_before_later_environment_reads(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    target_agent_id = "00000000-0000-4000-8000-000000007777"
    environment_ids = [
        "00000000-0000-4000-8000-000000008888",
        "00000000-0000-4000-8000-000000009999",
    ]
    client_count = 0

    class DevClient:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
            }

    class ThrottledClient:
        def get_agent(self, _agent_id: str) -> dict:
            raise snow.AgentBuilderHTTPError(
                "Direct agent",
                429,
                request_id="request-1",
                retry_after="30",
            )

    def environment_client(*_args: object, **_kwargs: object) -> object:
        nonlocal client_count
        client_count += 1
        if client_count > 1:
            pytest.fail("Throttling must stop later environment reads.")
        return ThrottledClient()

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: DevClient())
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [
            {"id": environment_id} for environment_id in environment_ids
        ],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        environment_client,
    )

    with snow.product_profile_context("hrsd"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="cannot continue",
        ) as raised:
            snow.validate_realm(context, realm="test")

    assert client_count == 1
    assert raised.value.details["statusCode"] == 429
    assert raised.value.details["retryAfter"] == "30"
    assert raised.value.details["requestId"] == "request-1"


def test_nonpassed_target_check_requires_configuration_not_test_pane(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    target_agent_id = "00000000-0000-4000-8000-000000007777"
    target_environment_id = "00000000-0000-4000-8000-000000008888"
    target_instance = "target-snow"
    target_client_id = "00000000-0000-4000-8000-000000006666"

    class Client:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
                # Deliberately no commitSha: ESS-SOLN-001 must fail.
            }

        def get_agent(self, _agent_id: str) -> dict:
            return {"botId": target_agent_id, "realm": "test"}

    client = Client()

    class Connectivity:
        def list_connections(self) -> list[dict]:
            return [
                {
                    "name": CONNECTION_ID.replace("-", ""),
                    "properties": {
                        "statuses": [
                            {"target": "token", "status": "Connected"}
                        ],
                        "connectionParametersSet": {
                            "name": snow.AUTH_MODE,
                            "values": {
                                "token:InstanceName": {
                                    "value": target_instance
                                },
                                "token:ResourceUri": {
                                    "value": target_client_id
                                },
                            },
                        },
                    },
                }
            ]

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: client)
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [{"id": target_environment_id}],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        lambda *_args, **_kwargs: client,
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client_for_environment",
        lambda *_args, **_kwargs: Connectivity(),
    )

    with snow.product_profile_context("hrsd"):
        result = snow.validate_realm(
            context,
            realm="test",
            target_instance=target_instance,
            target_client_id=target_client_id,
        )

    assert result["status"] == "needs-configuration"
    assert result["allAutomatedChecksPassed"] is False
    assert result["automatedChecks"][0]["status"] == "Failed"
    _assert_recovery_handoff(
        result["automatedChecks"][0]["recoveryHandoff"]
    )
    assert result["automatedChecks"][0]["remediation"] == (
        result["automatedChecks"][0]["recoveryHandoff"]["message"]
    )
    assert result["next"]["action"] == "configure-and-rerun"
    assert "url" not in result["next"]


def test_missing_target_credential_prevents_ready_for_test_pane(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    target_agent_id = "00000000-0000-4000-8000-000000007777"
    target_environment_id = "00000000-0000-4000-8000-000000008888"
    target_instance = "target-snow"
    target_client_id = "00000000-0000-4000-8000-000000006666"

    class Client:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
                "commitSha": "abc123",
            }

        def get_agent(self, _agent_id: str) -> dict:
            return {"botId": target_agent_id, "realm": "test"}

    class EmptyConnectivity:
        def list_connections(self) -> list[dict]:
            return []

    client = Client()
    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: client)
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [{"id": target_environment_id}],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        lambda *_args, **_kwargs: client,
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client_for_environment",
        lambda *_args, **_kwargs: EmptyConnectivity(),
    )

    with snow.product_profile_context("hrsd"):
        result = snow.validate_realm(
            context,
            realm="test",
            target_instance=target_instance,
            target_client_id=target_client_id,
        )

    credential = next(
        row
        for row in result["automatedChecks"]
        if row["checkpoint"] == "SN-REALM-CREDENTIAL-001"
    )
    assert credential["status"] == "NotConfigured"
    _assert_recovery_handoff(credential["recoveryHandoff"])
    assert target_instance in credential["recoveryHandoff"]["values"]
    assert target_client_id in credential["recoveryHandoff"]["values"]
    assert result["status"] == "needs-configuration"
    assert result["question"] is None
    assert result["readinessSummary"]["automaticPrechecks"] == (
        "Needs attention"
    )
    assert result["readinessSummary"]["functionalTest"] == "Pending"
    assert result["next"]["action"] == "configure-and-rerun"
    assert "url" not in result["next"]


def test_unconfirmed_target_identity_is_manual_and_does_not_copy_dev_values(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, _state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    target_agent_id = "00000000-0000-4000-8000-000000007777"
    target_environment_id = "00000000-0000-4000-8000-000000008888"

    class Client:
        def get_realms(self, _agent_id: str) -> dict:
            return {"routeRealm": 0, "siblingRealms": []}

        def get_realm_configuration(
            self,
            _agent_id: str,
            _realm: int,
        ) -> dict:
            return {
                "realm": "Test",
                "cdsBotId": target_agent_id,
                "schemaName": snow.HR_SCHEMA_NAME,
                "grsRepositoryId": "family-1",
                "commitSha": "abc123",
            }

        def get_agent(self, _agent_id: str) -> dict:
            return {"botId": target_agent_id, "realm": "test"}

    client = Client()
    connectivity_calls = 0

    def forbidden_connectivity(*_args: object, **_kwargs: object) -> None:
        nonlocal connectivity_calls
        connectivity_calls += 1
        pytest.fail("Missing target identity must not query credential inventory.")

    monkeypatch.setattr(snow, "_agentbuilder_token", lambda _context: "token")
    monkeypatch.setattr(snow, "_agentbuilder_client", lambda _context: client)
    monkeypatch.setattr(
        snow,
        "list_environments",
        lambda *_args, **_kwargs: [{"id": target_environment_id}],
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client_for_environment",
        lambda *_args, **_kwargs: client,
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client_for_environment",
        forbidden_connectivity,
    )

    with snow.product_profile_context("hrsd"):
        result = snow.validate_realm(context, realm="test")

    credential = next(
        row
        for row in result["manualChecks"]
        if row["checkpoint"] == "SN-REALM-CREDENTIAL-001"
    )
    assert credential["status"] == "Manual"
    assert "do not copy Dev" in credential["remediation"]
    _assert_recovery_handoff(credential["recoveryHandoff"])
    assert result["status"] == "ready-for-test-pane"
    assert result["blockingAutomatedChecks"] == []
    assert result["next"]["action"] == "customer-test-pane-validation"
    assert connectivity_calls == 0


@pytest.mark.parametrize(
    ("result", "failure_category"),
    (
        ("pass", None),
        ("fail", None),
        ("fail", "identity-mapping"),
        ("fail", "connection"),
    ),
)
def test_record_realm_test_keeps_manual_rows_and_dev_evidence_isolated(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    result: str,
    failure_category: str | None,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, state = _seed_complete_dev_state(tmp_path, profile="hrsd")
    manual_checks = [
        {
            "checkpoint": "SN-REALM-OIDC-MANUAL",
            "status": "Manual",
        },
        {
            "checkpoint": "SN-REALM-AGENT-BINDING-MANUAL",
            "status": "Manual",
        },
    ]
    state["evidence"]["realmValidations"] = {
        "test": {
            "targetEnvironmentId": (
                "00000000-0000-4000-8000-000000008888"
            ),
            "targetAgentId": "00000000-0000-4000-8000-000000007777",
            "targetConfigurationSha256": "a" * 64,
            "manualChecks": manual_checks,
        }
    }
    path = snow._lifecycle_state_path(context)
    path.write_text(json.dumps(state), encoding="utf-8")
    config_path = tmp_path / snow.ACTIVE_CONFIG
    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text(
        json.dumps({"activeAgent": "do-not-change"}),
        encoding="utf-8",
    )
    dev_test_before = copy.deepcopy(state["evidence"]["test"])
    phases_before = copy.deepcopy(state["phases"])

    with snow.product_profile_context("hrsd"):
        recorded = snow.record_realm_test(
            context,
            realm="test",
            result=result,
            failure_category=failure_category,
        )

    assert recorded["targetedManualCheckpoints"] == []
    assert recorded["targetedRecoveryHandoffs"] == []
    assert recorded["question"] is None
    assert recorded["manualChecks"] == manual_checks
    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved["evidence"]["test"] == dev_test_before
    assert saved["phases"] == phases_before
    realm_test = saved["evidence"]["realmValidations"]["test"][
        "functionalTest"
    ]
    assert realm_test["result"] == result
    assert realm_test["failureCategory"] == failure_category
    assert realm_test["binding"]["targetEnvironmentId"].endswith("8888")
    assert all(row["status"] == "Manual" for row in manual_checks)
    assert json.loads(config_path.read_text(encoding="utf-8")) == {
        "activeAgent": "do-not-change"
    }
    if result == "fail":
        assert recorded["next"] == {
            "action": "return-to-dev-connect",
            "url": (
                "https://copilotstudio.test.microsoft.com/environments/"
                f"{ENVIRONMENT_ID}/copilots/{AGENT_ID}/"
                "details?agentBackend=cosmos"
            ),
            "instruction": (
                "Not pass recorded; return to the Dev agent and rerun "
                "`/connect servicenow` to recheck the connection setup "
                "step by step."
            ),
        }
    else:
        assert recorded["next"]["action"] == (
            "realm-functional-validation-complete"
        )


def test_reconcile_legacy_publish_receipt_requires_stable_double_read(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    component_hash = snow._component_hash(components)
    last_published_at = "2026-09-30T20:22:44.5041453Z"
    _seed_passing_test(
        tmp_path,
        components,
        publish={
            "requestedAt": "2026-09-30T20:22:44.806391Z",
            "completedAt": "2026-09-30T20:22:54.575668Z",
            "status": "confirmation-required",
            "componentHash": "a" * 64,
            "response": {"validationPending": None},
        },
    )

    class FakeAgentBuilder:
        def get_agent(self, _agent_id: str) -> dict:
            return {"lastPublishedAt": last_published_at}

        def fetch_components(self, _agent_id: str) -> dict:
            return components

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    inspected = snow.inspect_publish_state(_context())
    reconciled = snow.reconcile_publish_receipt(
        _context(),
        expected_component_hash=inspected["componentHash"],
        expected_last_published_at=inspected["serverLastPublishedAt"],
        confirmed=True,
    )

    assert inspected["reconciliationEligible"] is True
    assert reconciled["status"] == "confirmation-required"
    assert reconciled["publishedComponentHash"] == component_hash
    assert reconciled["componentHash"] == component_hash
    assert reconciled["verifiedBy"] == "maker-attested-current-revision"
    assert reconciled["recordedAt"] == reconciled["reconciledAt"]
    assert reconciled["completedAt"] == "2026-09-30T20:22:54.575668Z"
    assert reconciled["legacyReceipt"]["componentHash"] == "a" * 64

    test_record = snow.record_test_attestation(
        _context(),
        prompt_category="list-my-open-hr-cases",
        result="pass",
        failure_category=None,
    )
    assert test_record["binding"]["draftSemanticHash"] == (
        reconciled["publishedSemanticHash"]
    )
    assert test_record["binding"]["connectionBindingHash"] == (
        reconciled["connectionBindingHash"]
    )
    acknowledgement = _phase_contract("test")[
        "manualAcknowledgementEvidence"
    ]
    assert test_record["status"] in acknowledgement[
        "acceptedRecordStatuses"
    ]
    assert test_record["kind"] == acknowledgement["requiredEvidenceKind"]
    assert test_record["result"] == acknowledgement["requiredValues"]["result"]


def test_reconcile_legacy_publish_receipt_uses_server_time_when_completion_missing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    component_hash = snow._component_hash(components)
    last_published_at = "2026-09-30T20:22:44.5041453Z"
    _seed_passing_test(
        tmp_path,
        components,
        publish={
            "requestedAt": "2026-09-30T20:22:44.806391Z",
            "status": "needs_remediation",
            "componentHash": "a" * 64,
        },
    )

    class FakeAgentBuilder:
        def get_agent(self, _agent_id: str) -> dict:
            return {"lastPublishedAt": last_published_at}

        def fetch_components(self, _agent_id: str) -> dict:
            return components

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    reconciled = snow.reconcile_publish_receipt(
        _context(),
        expected_component_hash=component_hash,
        expected_last_published_at=last_published_at,
        confirmed=True,
    )

    assert reconciled["completedAt"] == last_published_at
    assert reconciled["recordedAt"]


def test_reconcile_publish_receipt_rejects_revision_drift(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    state_path = _lifecycle_path(tmp_path)
    _seed_passing_test(
        tmp_path,
        components,
        publish={
            "requestedAt": "2026-09-30T20:22:44.806391Z",
            "status": "confirmation-required",
            "componentHash": "a" * 64,
        },
    )

    class FakeAgentBuilder:
        def get_agent(self, _agent_id: str) -> dict:
            return {"lastPublishedAt": "2026-09-30T20:22:44.5041453Z"}

        def fetch_components(self, _agent_id: str) -> dict:
            changed = copy.deepcopy(components)
            changed["changeToken"] = "changed-after-confirmation"
            return changed

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="component revision changed",
    ):
        snow.reconcile_publish_receipt(
            _context(),
            expected_component_hash=snow._component_hash(components),
            expected_last_published_at="2026-09-30T20:22:44.5041453Z",
            confirmed=True,
        )

    unchanged = json.loads(state_path.read_text(encoding="utf-8"))
    assert unchanged["evidence"]["publish"]["componentHash"] == "a" * 64


def test_record_agent_connection_rejects_unsupported_auth_mode(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeConnectivity:
        def get_connection(self, _connection_id: str) -> dict:
            return {
                "name": CONNECTION_ID.replace("-", ""),
                "properties": {
                    "displayName": "ServiceNow",
                    "statuses": [{"target": "token", "status": "Connected"}],
                    "connectionParametersSet": {"name": "oauth2"},
                },
            }

    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Microsoft Entra ID User Login",
    ):
        snow.record_agent_connection_attestation(
            {
                "agent": {"id": AGENT_ID},
                "environment": {"id": ENVIRONMENT_ID},
            },
            CONNECTION_ID,
        )


def test_inspect_preserves_progress_and_reports_completed_steps(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / "servicenow"
        / "agents"
        / AGENT_ID
        / "state.json"
    )
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 1,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "agentConnection": {
                    "connectionId": CONNECTION_ID.replace("-", ""),
                    "makerAttested": True,
                },
                "parameterSharing": {"status": "not-exposed"},
                "publish": {"completedAt": "2026-09-24T00:00:00Z"},
                "test": {"result": "pass"},
                "steps": {
                    "topics": "done",
                    "agentConnection": "done",
                    "parameterSharing": "done",
                    "publish": "done",
                    "test": "done",
                },
            }
        ),
        encoding="utf-8",
    )

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return _components(CONNECTION_ID)

    class FakeConnectivity:
        def get_connector(self) -> dict:
            raise AssertionError("inspect must not read connector definition")

        def list_connections(self) -> list[dict]:
            return [
                {
                    "name": CONNECTION_ID.replace("-", ""),
                    "properties": {
                        "displayName": "maker@example.com",
                        "statuses": [
                            {"target": "token", "status": "Connected"}
                        ],
                        "connectionParametersSet": {
                            "name": "entraIDUserLogin",
                        },
                    },
                }
            ]

    context = _context()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    result = snow.inspect(context)

    assert "connector" not in result["connectivity"]
    assert result["connectivity"]["connections"][0]["connectionId"] == (
        CONNECTION_ID.replace("-", "")
    )
    assert result["connectivity"]["userConnectionBinding"]["mode"] == (
        "maker-ui"
    )
    assert result["progress"]["topics"]["status"] == "done"
    assert result["progress"]["credential"]["status"] == "done"
    assert result["progress"]["legacyPublish"]["status"] == "historical"
    assert (
        result["progress"]["agentConnection"]["status"]
        == "confirmation-required"
    )
    assert "parameterSharing" not in result["progress"]
    assert result["progress"]["test"]["status"] == "pending"
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert state["evidence"]["agentConnection"]["reconfirmRequired"] is True
    assert state["evidence"]["parameterSharing"]["reconfirmRequired"] is True
    assert state["evidence"]["test"]["reconfirmRequired"] is True


def test_inspect_reopens_topics_after_later_deactivation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / "servicenow"
        / "agents"
        / AGENT_ID
        / "state.json"
    )
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 1,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "topicEnablement": {"customerChoice": "enable-all"},
                "steps": {"topics": "done"},
            }
        ),
        encoding="utf-8",
    )
    components = _components(CONNECTION_ID)
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

    class FakeConnectivity:
        def get_connector(self) -> dict:
            raise AssertionError("inspect must not read connector definition")

        def list_connections(self) -> list[dict]:
            return []

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    result = snow.inspect(
        {
            "agent": {
                "id": AGENT_ID,
                "schema_name": snow.HR_SCHEMA_NAME,
            },
            "environment": {
                "id": ENVIRONMENT_ID,
                "tenant_id": TENANT_ID,
                "ring": "test",
            },
        }
    )

    assert "connector" not in result["connectivity"]
    assert result["connectivity"]["connections"] == []
    assert result["progress"]["topics"]["status"] == "pending"


def test_inspect_preserves_explicit_keep_current_topic_choice(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / "servicenow"
        / "agents"
        / AGENT_ID
        / "state.json"
    )
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 1,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "topicEnablement": {"customerChoice": "keep-current"},
                "steps": {"topics": "done"},
            }
        ),
        encoding="utf-8",
    )
    components = _components(CONNECTION_ID)
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

    class FakeConnectivity:
        def get_connector(self) -> dict:
            raise AssertionError("inspect must not read connector definition")

        def list_connections(self) -> list[dict]:
            return []

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    result = snow.inspect(
        {
            "agent": {
                "id": AGENT_ID,
                "schema_name": snow.HR_SCHEMA_NAME,
            },
            "environment": {
                "id": ENVIRONMENT_ID,
                "tenant_id": TENANT_ID,
                "ring": "test",
            },
        }
    )

    assert "connector" not in result["connectivity"]
    assert result["connectivity"]["connections"] == []
    assert result["progress"]["topics"]["status"] == "pending"


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_inspect_live_uses_components_and_connection_list_only(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    components = (
        _components(CONNECTION_ID)
        if profile == "hrsd"
        else _itsm_components()
    )
    context = _profile_context(
        profile,
        agent_id=AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID,
        agent_slug=(
            "employee-self-service-hr"
            if profile == "hrsd"
            else "employee-self-service-it"
        ),
    )
    context["root"] = tmp_path
    calls: list[tuple[str, str, dict]] = []

    class FakeResponse:
        ok = True
        status_code = 200
        headers: dict[str, str] = {}

        def __init__(self, body: dict) -> None:
            self.body = body

        def json(self) -> dict:
            return copy.deepcopy(self.body)

    class RecordingSession:
        def mount(self, _prefix: str, _adapter: object) -> None:
            return None

        def request(
            self,
            method: str,
            url: str,
            **kwargs,
        ) -> FakeResponse:
            calls.append((method, url, copy.deepcopy(kwargs)))
            if url.endswith(
                f"/copilotstudio/minimalBots/api/"
                f"{context['agent']['id']}/components"
            ):
                assert method == "POST"
                return FakeResponse(components)
            if url.endswith(
                "/connectivity/connectors/shared_service-now/connections"
            ):
                assert method == "GET"
                return FakeResponse(
                    {
                        "value": [
                            _connection_record(
                                connection_id=CONNECTION_ID,
                            )
                        ]
                    }
                )
            if url.endswith(
                "/connectivity/connectors/shared_service-now"
            ):
                raise AssertionError(
                    "inspect must not read connector definition"
                )
            raise AssertionError(f"Unexpected request: {method} {url}")

    session = RecordingSession()
    host = "https://example.environment.api.test.powerplatform.com"
    agent_client = snow.AgentBuilderClient(
        host,
        "token",
        ring="test",
        tenant_id=TENANT_ID,
        session=session,
    )
    connectivity_client = snow.ConnectivityClient(
        host,
        ENVIRONMENT_ID,
        "token",
        ring="test",
        session=session,
    )
    state = {
        "evidence": {},
        "phases": {},
        "componentHash": "component-hash",
        "draftSemanticHash": "draft-hash",
        "connectionBindingHash": "binding-hash",
    }
    persisted: list[dict] = []
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: agent_client,
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: connectivity_client,
    )
    monkeypatch.setattr(
        snow,
        "_state_for_components",
        lambda _context, _components_value: state,
    )
    monkeypatch.setattr(
        snow,
        "_write_lifecycle_state",
        lambda _context, value: persisted.append(copy.deepcopy(value)),
    )

    with snow.product_profile_context(profile):
        result = snow.inspect(context)

    component_calls = [
        call for call in calls if call[1].endswith("/components")
    ]
    connection_list_calls = [
        call
        for call in calls
        if call[1].endswith(
            "/connectivity/connectors/shared_service-now/connections"
        )
    ]
    connector_definition_calls = [
        call
        for call in calls
        if call[1].endswith(
            "/connectivity/connectors/shared_service-now"
        )
    ]
    assert len(component_calls) == 1
    assert component_calls[0][0] == "POST"
    assert component_calls[0][2]["params"] == {
        "api-version": "2022-03-01-preview"
    }
    assert component_calls[0][2]["json"] == {}
    assert len(connection_list_calls) == 1
    assert connection_list_calls[0][0] == "GET"
    assert connection_list_calls[0][2]["params"] == {
        "api-version": snow.CONNECTIVITY_API_VERSION,
        "$filter": f"environment eq '{ENVIRONMENT_ID}'",
    }
    assert connector_definition_calls == []
    assert len(calls) == 2
    assert "connector" not in result["connectivity"]
    assert result["connectivity"]["connections"][0]["connectionId"] == (
        CONNECTION_ID.replace("-", "")
    )
    assert result["connectivity"]["userConnectionBinding"]["mode"] == (
        "maker-ui"
    )
    assert result["progress"]["credential"]["status"] == "done"
    assert persisted[-1]["evidence"]["lastInspection"] == result
    assert "connector" not in persisted[-1]["evidence"]["lastInspection"][
        "connectivity"
    ]


def test_inspect_connection_read_failure_stops_without_partial_result(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return _components(CONNECTION_ID)

    class FakeConnectivity:
        def get_connector(self) -> dict:
            raise AssertionError("inspect must not read connector definition")

        def list_connections(self) -> list[dict]:
            raise snow.ServiceNowConnectError("connection inventory failed")

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )
    monkeypatch.setattr(
        snow,
        "_write_lifecycle_state",
        lambda _context, _state: pytest.fail(
            "failed inspection must not persist a partial result"
        ),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="connection inventory failed",
    ):
        snow.inspect(_context())


def test_inspect_offline_does_not_create_remote_clients(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    snapshot = tmp_path / "components.json"
    snapshot.write_text(
        json.dumps(_components(CONNECTION_ID)),
        encoding="utf-8",
    )
    context = _context_at(tmp_path)
    context["snapshotPath"] = snapshot
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: pytest.fail("offline inspect used AgentBuilder"),
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: pytest.fail("offline inspect used Connectivity"),
    )

    result = snow.inspect(context, offline=True)

    assert result["mode"] == "offline"
    assert "connectivity" not in result
    assert "progress" not in result


def test_parameter_sharing_has_no_active_cli_surface() -> None:
    help_text = snow.build_parser().format_help()
    phase_help = snow.build_parser()._subparsers._group_actions[0].choices[
        "record-admin-phase"
    ].format_help()

    assert "record-parameter-sharing" not in help_text
    assert not hasattr(snow, "record_parameter_sharing")
    assert "--claim" not in phase_help
    assert "--user-field" not in phase_help


def test_connectivity_scopes_are_read_only() -> None:
    scopes = snow.connectivity_scopes("test")

    assert (
        "https://api.test.powerplatform.com/Connectivity.Connections.Read"
        in scopes
    )
    assert not any(scope.endswith(".Write") for scope in scopes)


def test_set_portal_url_commits_with_fresh_readback(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(copy.deepcopy(payload))
            updated = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][1]["component"] = copy.deepcopy(
                updated
            )
            components["botComponentChanges"][1]["component"]["version"] += 1
            components["changeToken"] = "token-2"
            return {"changeToken": "token-2"}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com/hrportal",
        confirmed=True,
    )

    assert result["status"] == "committed"
    assert result["readbackVerified"] is True
    assert len(updates) == 1
    assert snow._portal_literal_value(
        components["botComponentChanges"][1]["component"]
    ) == "https://dev123.service-now.com/hrportal"
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = state["transactions"]["portal"][result["operationId"]]
    assert transaction["status"] == "committed"


def test_set_portal_url_commits_value_expression_without_shape_change(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_expression_components()

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updated = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][1]["component"] = copy.deepcopy(
                updated
            )
            components["botComponentChanges"][1]["component"]["version"] += 1
            components["changeToken"] = "token-2"
            return {"changeToken": "token-2"}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com/hrportal",
        confirmed=True,
    )

    node = snow._portal_value_node(
        components["botComponentChanges"][1]["component"]
    )
    assert result["status"] == "committed"
    assert node["value"] == {
        "$kind": "ValueExpression",
        "expressionText": (
            '"https://dev123.service-now.com/hrportal"'
        ),
    }


def test_set_portal_url_skips_update_when_already_configured(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()
    snow._portal_value_node(
        components["botComponentChanges"][1]["component"]
    )["value"]["literalValue"] = "https://dev123.service-now.com"

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, _payload: dict) -> dict:
            raise AssertionError("already configured value must not be updated")

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com",
        confirmed=True,
    )

    assert result["status"] == "already-configured"
    assert result["readbackVerified"] is True
    saved = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    receipt = saved["evidence"]["portalConfirmation"]
    assert receipt["kind"] == "maker-confirmed-portal-url"
    assert receipt["status"] == "confirmed"
    assert receipt["source"] == "already-configured"
    assert receipt["portalUrl"] == "https://dev123.service-now.com"
    assert receipt["binding"]["provider"] == "servicenow-da-hrsd"
    assert receipt["binding"]["agentId"] == AGENT_ID
    assert result["confirmation"] == receipt


def test_set_portal_url_skips_literal_value_expression_when_already_configured(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_literal_value_expression_components(
        "https://dev123.service-now.com/hrportal"
    )

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, _payload: dict) -> dict:
            raise AssertionError("already configured value must not be updated")

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com/hrportal",
        confirmed=True,
    )

    assert result["status"] == "already-configured"
    assert result["readbackVerified"] is True
    assert snow._portal_value_node(
        components["botComponentChanges"][1]["component"]
    )["value"] == {
        "$kind": "ValueExpression",
        "literalValue": "https://dev123.service-now.com/hrportal",
    }


def test_portal_confirmation_rejects_wrong_binding_or_changed_value(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, _payload: dict) -> dict:
            raise AssertionError("same-value confirmation must not update")

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com/employee",
        confirmed=True,
    )
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    topic = snow._portal_topic_component(components)

    assert snow.portal_confirmation_url(
        state,
        topic=topic,
        current_portal_url="https://dev123.service-now.com/employee",
        expected_instance_name="dev123",
    ) == "https://dev123.service-now.com/employee"

    wrong_binding = copy.deepcopy(state)
    wrong_binding["environmentId"] = str(uuid.uuid4())
    assert snow.portal_confirmation_url(
        wrong_binding,
        topic=topic,
        current_portal_url="https://dev123.service-now.com/employee",
        expected_instance_name="dev123",
    ) is None
    assert snow.portal_confirmation_url(
        state,
        topic=topic,
        current_portal_url="https://dev123.service-now.com/other",
        expected_instance_name="dev123",
    ) is None


def test_set_portal_url_reports_failed_unchanged_for_unsupported_update(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, _payload: dict) -> dict:
            raise PermissionError("unsupported")

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="did not apply the Portal BaseURI change",
    ) as raised:
        snow.set_portal_url(
            _context(),
            "https://dev123.service-now.com/hrportal",
            confirmed=True,
        )

    transaction = raised.value.details["transaction"]
    assert transaction["status"] == "failed-unchanged"
    assert transaction["rollbackVerified"] is True


def test_set_portal_url_can_verify_and_roll_back(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()
    original = snow._portal_literal_value(
        components["botComponentChanges"][1]["component"]
    )
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(copy.deepcopy(payload))
            updated = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][1]["component"] = copy.deepcopy(
                updated
            )
            components["botComponentChanges"][1]["component"]["version"] += 1
            components["changeToken"] = f"token-{len(updates) + 1}"
            return {"changeToken": components["changeToken"]}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com/hrportal",
        confirmed=True,
        rollback_after_verify=True,
    )

    assert result["status"] == "verified-and-rolled-back"
    assert result["rollbackVerified"] is True
    assert len(updates) == 2
    assert snow._portal_literal_value(
        components["botComponentChanges"][1]["component"]
    ) == original


def test_set_portal_expression_can_verify_and_roll_back(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_expression_components()
    original_value = copy.deepcopy(
        snow._portal_value_node(
            components["botComponentChanges"][1]["component"]
        )["value"]
    )
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(copy.deepcopy(payload))
            updated = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][1]["component"] = copy.deepcopy(
                updated
            )
            components["botComponentChanges"][1]["component"]["version"] += 1
            components["changeToken"] = f"token-{len(updates) + 1}"
            return {"changeToken": components["changeToken"]}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.set_portal_url(
        _context(),
        "https://dev123.service-now.com/hrportal",
        confirmed=True,
        rollback_after_verify=True,
    )

    final_value = snow._portal_value_node(
        components["botComponentChanges"][1]["component"]
    )["value"]
    assert result["status"] == "verified-and-rolled-back"
    assert result["rollbackVerified"] is True
    assert len(updates) == 2
    assert final_value == original_value


def test_set_portal_url_reports_rollback_failure(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()
    update_count = 0

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            nonlocal update_count
            update_count += 1
            if update_count == 2:
                raise PermissionError("rollback rejected")
            updated = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][1]["component"] = copy.deepcopy(
                updated
            )
            components["botComponentChanges"][1]["component"]["version"] += 1
            components["changeToken"] = "token-2"
            return {"changeToken": "token-2"}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Rollback did not restore",
    ) as raised:
        snow.set_portal_url(
            _context(),
            "https://dev123.service-now.com/hrportal",
            confirmed=True,
            rollback_after_verify=True,
        )

    transaction = raised.value.details["transaction"]
    assert transaction["status"] == "rollback-incomplete"
    assert transaction["rollbackVerified"] is False
    assert transaction["rollbackError"] == "PermissionError"


def test_set_portal_url_preserves_concurrent_authored_edit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updated = copy.deepcopy(
                payload["botComponentChanges"][0]["component"]
            )
            updated["externalConcurrentEdit"] = "must-survive"
            components["botComponentChanges"][1]["component"] = updated
            components["changeToken"] = "token-2"
            return {"changeToken": "token-2"}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="changed unexpectedly",
    ):
        snow.set_portal_url(
            _context(),
            "https://dev123.service-now.com/hrportal",
            confirmed=True,
            rollback_after_verify=True,
        )

    assert (
        components["botComponentChanges"][1]["component"][
            "externalConcurrentEdit"
        ]
        == "must-survive"
    )
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = next(iter(state["transactions"]["portal"].values()))
    assert transaction["status"] == "conflict"


def test_set_portal_url_reports_unverifiable_readback(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _portal_components()
    fetch_count = 0

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            nonlocal fetch_count
            fetch_count += 1
            if fetch_count > 1:
                raise RuntimeError("readback unavailable")
            return copy.deepcopy(components)

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updated = payload["botComponentChanges"][0]["component"]
            components["botComponentChanges"][1]["component"] = copy.deepcopy(
                updated
            )
            return {}

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="readback failed",
    ) as raised:
        snow.set_portal_url(
            _context(),
            "https://dev123.service-now.com/hrportal",
            confirmed=True,
        )

    transaction = raised.value.details["transaction"]
    assert transaction["status"] == "reconciliation-required"
    assert transaction["reconciliationError"] == "RuntimeError"
    assert transaction["mutationMayHaveOccurred"] is True


def test_inspect_portal_url_offline_reports_sanitized_configuration(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    components = _portal_components()
    _write_foundation_context(tmp_path, components=components)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.chdir(tmp_path)

    assert snow.main(["inspect-portal-url", "--offline"]) == 0
    result = json.loads(capsys.readouterr().out)

    assert result["status"] == "configured"
    assert result["mode"] == "offline"
    assert result["portal"]["origin"] == "https://dev123.service-now.com"
    assert result["portal"]["path"] == "/employee"
    assert result["input"] is None


def test_inspect_portal_url_requires_expected_maker_value(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    components = _portal_components()
    _write_foundation_context(tmp_path, components=components)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "inspect-portal-url",
            "--offline",
            "--expected-portal-url",
            "https://dev123.service-now.com/employee/",
        ]
    ) == 0
    matching = json.loads(capsys.readouterr().out)
    assert matching["status"] == "configured"
    assert matching["expectedPortalUrlMatch"] is True

    assert snow.main(
        [
            "inspect-portal-url",
            "--offline",
            "--expected-portal-url",
            "https://dev123.service-now.com/different-portal",
        ]
    ) == 0
    mismatch = json.loads(capsys.readouterr().out)
    assert mismatch["status"] == "input-required"
    assert mismatch["expectedPortalUrlMatch"] is False
    assert mismatch["portal"]["path"] == "/employee"
    assert "exact administrator-confirmed" in mismatch["input"]["description"]


def test_inspect_portal_url_requires_explicit_root_confirmation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    components = _portal_components()
    snow._portal_value_node(
        components["botComponentChanges"][1]["component"]
    )["value"]["literalValue"] = "https://dev123.service-now.com"
    _write_foundation_context(tmp_path, components=components)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.chdir(tmp_path)

    assert snow.main(["inspect-portal-url", "--offline"]) == 0
    unconfirmed = json.loads(capsys.readouterr().out)
    assert unconfirmed["status"] == "input-required"
    assert unconfirmed["expectedPortalUrlMatch"] is None
    assert "Explicit administrator confirmation" in (
        unconfirmed["input"]["description"]
    )

    assert snow.main(
        [
            "inspect-portal-url",
            "--offline",
            "--expected-portal-url",
            "https://dev123.service-now.com",
        ]
    ) == 0
    confirmed = json.loads(capsys.readouterr().out)
    assert confirmed["status"] == "configured"
    assert confirmed["expectedPortalUrlMatch"] is True


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("dev123", "dev123"),
        ("https://dev123.service-now.com", "dev123"),
        ("https://DEV123.service-now.com/", "dev123"),
    ],
)
def test_normalize_instance_name(value: str, expected: str) -> None:
    assert snow.normalize_instance_name(value) == expected


@pytest.mark.parametrize(
    "value",
    [
        "http://dev123.service-now.com",
        "https://dev123.service-now.com/path",
        "https://example.com",
        "dev123.example.com",
    ],
)
def test_normalize_instance_name_rejects_unsafe_values(value: str) -> None:
    with pytest.raises(snow.ServiceNowConnectError):
        snow.normalize_instance_name(value)


@pytest.mark.parametrize(
    ("ring", "power_automate", "power_apps", "copilot_studio"),
    [
        (
            "prod",
            "https://make.powerautomate.com",
            "https://make.powerapps.com",
            "https://copilotstudio.microsoft.com",
        ),
        (
            "preprod",
            "https://make.preprod.powerautomate.com",
            "https://make.preprod.powerapps.com",
            "https://copilotstudio.preprod.microsoft.com",
        ),
        (
            "test",
            "https://make.test.powerautomate.com",
            "https://make.test.powerapps.com",
            "https://copilotstudio.test.microsoft.com",
        ),
    ],
)
def test_maker_portal_links_are_ring_aware_and_identity_scoped(
    ring: str,
    power_automate: str,
    power_apps: str,
    copilot_studio: str,
) -> None:
    context = _context()
    context["environment"]["ring"] = ring

    links = snow.maker_portal_links(
        context,
        instance_name="dev123",
    )

    assert links["serviceNowInstance"] == {
        "label": "Open this ServiceNow instance",
        "url": "https://dev123.service-now.com/",
    }
    assert links["hrCorePlugin"] == {
        "label": "Open HR Service Delivery Core",
        "url": (
            "https://dev123.service-now.com/now/app-manager/home/plugin/"
            "id/com.sn_hr_core/details"
        ),
    }
    assert links["entra"]["url"] == "https://entra.microsoft.com/"
    assert links["powerAutomateConnections"]["url"] == (
        f"{power_automate}/environments/{ENVIRONMENT_ID}/connections"
    )
    assert links["powerApps"]["url"] == f"{power_apps}/"
    assert links["copilotStudioAgent"]["url"] == (
        f"{copilot_studio}/environments/{ENVIRONMENT_ID}/copilots/"
        f"{AGENT_ID}/details?agentBackend=cosmos"
    )
    template = "[Open Power Apps]({POWER_APPS_URL})"
    rendered = template.replace(
        "{POWER_APPS_URL}",
        links["powerApps"]["url"],
    )
    assert rendered == f"[Open Power Apps]({power_apps}/)"
    assert rendered.replace("https://", "").find("//") == -1


def test_maker_portal_links_reject_unknown_ring() -> None:
    context = _context()
    context["environment"]["ring"] = "unknown"

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Unsupported Power Platform ring",
    ):
        snow.maker_portal_links(context)


def test_plugin_requirements_are_scope_and_auth_dynamic() -> None:
    hrsd = snow.required_servicenow_prerequisites("hrsd")
    itsm = snow.required_servicenow_prerequisites("itsm")

    assert [item["id"] for item in hrsd] == ["hr-core"]
    assert itsm == []
    assert hrsd[0]["aliases"] == ["com.sn_hr_core", "sn_hr_core"]
    with pytest.raises(snow.ServiceNowConnectError):
        snow.required_servicenow_prerequisites("hrsd", "oauth2ServiceNow")


def test_preflight_requires_maker_confirmation_for_single_candidate(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return [
                {
                    "name": CONNECTION_ID.replace("-", ""),
                    "properties": {
                        "displayName": "Existing ServiceNow",
                        "statuses": [
                            {"target": "token", "status": "Connected"}
                        ],
                        "connectionParametersSet": {
                            "name": snow.AUTH_MODE,
                            "values": {
                                "token:InstanceName": {"value": "dev123"},
                                "token:ResourceUri": {
                                    "value": APP_CLIENT_ID
                                },
                            },
                        },
                    },
                }
            ]

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    discovery = snow.inspect_admin_setup(_context())
    assert discovery["instanceName"] is None
    assert discovery["instanceInputRequired"] is True
    assert discovery["instanceSelectionRequired"] is True
    assert discovery["instanceCandidates"] == [
        "https://dev123.service-now.com"
    ]
    assert discovery["instanceQuestion"]["options"] == [
        {"label": "Use https://dev123.service-now.com"},
        {"label": "Provide another ServiceNow instance URL"},
    ]
    candidate_body = discovery["instanceQuestion"]["question"]
    assert "We found ServiceNow instance candidate" in candidate_body
    assert "Employee Self-Service (HR)" in candidate_body
    assert "employee-self-service-hr" in candidate_body
    assert ENVIRONMENT_ID in candidate_body
    assert "Dev realm" in candidate_body
    assert "Which instance should this agent use" in candidate_body
    assert "**Problem**" not in candidate_body
    assert "**Category**" not in candidate_body
    assert "**Recheck and resume**" not in candidate_body
    manual = discovery["manualInstanceQuestion"]
    assert manual["allowFreeformInput"] is True
    assert "Employee Self-Service (HR)" in manual["question"]
    assert ENVIRONMENT_ID in manual["question"]
    assert "`https://dev184242.service-now.com`" in manual["question"]
    assert "ask your ServiceNow administrator" in manual["question"]
    assert "/sp" in manual["question"]
    assert "password" in manual["question"]
    assert "token" in manual["question"]
    assert discovery["instanceChoices"][
        "Use https://dev123.service-now.com"
    ] == "https://dev123.service-now.com"
    recorded = snow.record_preflight_instance(
        _context(),
        instance_url=discovery["instanceChoices"][
            "Use https://dev123.service-now.com"
        ],
    )
    refreshed = snow.inspect_admin_setup(_context())
    _record_test_identity_mapping(tmp_path)
    snow.record_admin_phase(
        _context(),
        phase="entra-registration",
        status="completed",
        client_id=APP_CLIENT_ID,
    )
    with_client_id = snow.inspect_admin_setup(_context())

    assert discovery["connectionCandidates"][0]["status"] == "Connected"
    assert "fingerprint" not in discovery
    assert "fingerprint" not in refreshed
    assert "fingerprint" not in with_client_id
    assert refreshed["instanceName"] == "dev123"
    assert refreshed["instanceInputRequired"] is False
    assert discovery["credentialResolution"]["healthyExactCount"] == 0
    assert with_client_id["credentialResolution"]["healthyExactCount"] == 1
    assert with_client_id["credentialResolution"]["unhealthyExactCount"] == 0
    assert refreshed["links"]["serviceNowInstance"]["url"] == (
        "https://dev123.service-now.com/"
    )
    assert discovery["links"]["powerAutomateConnections"]["url"].startswith(
        "https://make.test.powerautomate.com/"
    )
    assert with_client_id["entraClientId"] == APP_CLIENT_ID
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert "scenario" not in state["adminSetup"]["preflight"]
    assert "reuseDecision" not in state["adminSetup"]["preflight"]
    assert state["adminSetup"]["phaseHandoffs"]["preflight"] == {
        "status": "completed",
        "evidence": {
            "kind": "maker-confirmed-target",
            "instanceName": "dev123",
            "recordedAt": recorded["recordedAt"],
        },
    }


def test_preflight_requests_only_missing_instance_url(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return []

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    discovery = snow.inspect_admin_setup(_context())
    recorded = snow.record_preflight_instance(
        _context(),
        instance_url="https://dev123.service-now.com",
    )
    refreshed = snow.inspect_admin_setup(_context())

    assert discovery["instanceInputRequired"] is True
    assert discovery["instanceQuestion"] == (
        discovery["manualInstanceQuestion"]
    )
    manual_body = discovery["instanceQuestion"]["question"]
    assert "employee-self-service-hr" in manual_body
    assert ENVIRONMENT_ID in manual_body
    assert "Dev realm" in manual_body
    assert "`https://dev184242.service-now.com`" in manual_body
    assert "public ServiceNow HTTPS root URL" in manual_body
    assert "portal or module path" in manual_body
    assert "credential" in manual_body
    assert recorded["instanceName"] == "dev123"
    assert refreshed["instanceInputRequired"] is False
    assert "scenario" not in recorded
    assert "--scenario" not in snow.build_parser().format_help()


@pytest.mark.parametrize(
    ("profile", "slug", "display_name"),
    (
        ("hrsd", "employee-self-service-hr", "Employee Self-Service (HR)"),
        ("itsm", "employee-self-service-it", "Employee Self-Service (IT)"),
    ),
)
def test_preflight_multiple_candidates_are_independent_normal_choices(
    profile: str,
    slug: str,
    display_name: str,
) -> None:
    context = _profile_context(
        profile,
        agent_id=AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID,
        agent_slug=slug,
    )
    with snow.product_profile_context(profile):
        question, choices, manual = snow._preflight_instance_question(
            context=context,
            candidates=["dev111", "dev222"],
            discovery_error=None,
        )

    assert [option["label"] for option in question["options"]] == [
        "Use https://dev111.service-now.com",
        "Use https://dev222.service-now.com",
        "Provide another ServiceNow instance URL",
    ]
    assert not any(
        option.get("recommended") for option in question["options"]
    )
    assert choices == {
        "Use https://dev111.service-now.com": (
            "https://dev111.service-now.com"
        ),
        "Use https://dev222.service-now.com": (
            "https://dev222.service-now.com"
        ),
        "Provide another ServiceNow instance URL": "manual",
    }
    assert display_name in question["question"]
    assert slug in question["question"]
    assert "**Problem**" not in question["question"]
    assert manual["header"] == "ServiceNow instance URL"
    assert "`https://dev184242.service-now.com`" in manual["question"]


def test_preflight_keeps_confirmed_instance_when_new_connection_appears(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return [
                _connection_record(
                    instance_name="other-instance",
                    connection_id=CONNECTION_ID,
                )
            ]

    monkeypatch.chdir(tmp_path)
    snow.record_preflight_instance(
        _context(),
        instance_url="https://confirmed-instance.service-now.com",
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    discovery = snow.inspect_admin_setup(_context())

    assert discovery["instanceName"] == "confirmed-instance"
    assert discovery["instanceInputRequired"] is False
    assert discovery["instanceSelectionRequired"] is False
    assert discovery["instanceCandidates"] == [
        "https://other-instance.service-now.com"
    ]
    saved = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert saved["adminSetup"]["preflight"]["instanceName"] == (
        "confirmed-instance"
    )
    assert saved["adminSetup"]["phaseHandoffs"]["preflight"]["status"] == (
        "completed"
    )


def test_resolve_credential_automatically_records_unique_healthy_exact(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    record = _connection_record(display_name="Existing ServiceNow")

    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return [record]

        def get_connection(self, connection_id: str) -> dict:
            assert connection_id == CONNECTION_ID.replace("-", "")
            return record

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {
                "fetch_components": lambda self, _agent_id: _components(
                    CONNECTION_ID
                )
            },
        )(),
    )

    result = snow.resolve_credential_completion(_context())

    assert result["status"] == "selected"
    assert result["selection"] == {"label": "Existing ServiceNow"}
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert state["evidence"]["credential"]["connectionId"] == (
        CONNECTION_ID.replace("-", "")
    )
    assert state["adminSetup"]["phaseHandoffs"]["credential"]["status"] == (
        "completed"
    )


def test_resolve_credential_requires_bounded_choice_for_multiple_exact(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    first_id = "00000000-0000-4000-8000-000000003331"
    second_id = "00000000-0000-4000-8000-000000003332"
    records = [
        _connection_record(first_id, display_name="ServiceNow A"),
        _connection_record(second_id, display_name="ServiceNow B"),
    ]

    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return records

        def get_connection(self, connection_id: str) -> dict:
            return next(
                record
                for record in records
                if record["name"] == connection_id
            )

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {
                "fetch_components": lambda self, _agent_id: _components(
                    first_id
                )
            },
        )(),
    )

    choice = snow.resolve_credential_completion(_context())

    assert choice["status"] == "choice-required"
    assert [candidate["label"] for candidate in choice["candidates"]] == [
        "ServiceNow A (Connected, option 1)",
        "ServiceNow B (Connected, option 2)",
    ]
    serialized = json.dumps(choice)
    assert first_id not in serialized
    assert first_id.replace("-", "") not in serialized
    assert second_id not in serialized
    pending_state = json.loads(
        _lifecycle_path(tmp_path).read_text(encoding="utf-8")
    )
    assert "credential" not in pending_state["evidence"]

    selected = snow.resolve_credential_completion(
        _context(),
        selection_key=choice["candidates"][1]["selectionKey"],
    )

    assert selected["status"] == "selected"
    assert selected["selection"] == {"label": "ServiceNow B"}
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert state["evidence"]["credential"]["connectionId"] == (
        second_id.replace("-", "")
    )


@pytest.mark.parametrize(
    "record",
    [
        _connection_record(status="Disconnected"),
        _connection_record(auth_mode="oauth2"),
        _connection_record(instance_name="dev456"),
        _connection_record(
            resource_uri="00000000-0000-4000-8000-000000009999"
        ),
    ],
)
def test_resolve_credential_rejects_nonhealthy_or_nonexact_candidates(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    record: dict,
) -> None:
    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return [record]

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    result = snow.resolve_credential_completion(_context())

    assert result["status"] == "not-ready"
    assert "No exact healthy Connected" in result["remediation"]
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert "credential" not in state["evidence"]


def test_resolve_credential_read_failure_is_explicit_nonzero(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            raise snow.ServiceNowConnectError(
                "Connection inventory is unavailable."
            )

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(snow, "load_context", lambda: _context())
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )
    monkeypatch.setattr(
        sys,
        "argv",
        ["connect_servicenow_da.py", "resolve-credential"],
    )

    assert snow.main() == 1
    error = json.loads(capsys.readouterr().out)
    assert error["status"] == "error"
    assert error["message"] == "Connection inventory is unavailable."
    recovery = error["details"]["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["category"] == "Read or network error"
    assert recovery["owner"] == "Power Platform Maker/Admin"


def test_resolve_credential_rejects_stale_selection_key(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    record = _connection_record()

    class FakeConnectivity:
        def list_connections(self) -> list[dict]:
            return [record]

    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="no longer an exact healthy candidate",
    ):
        snow.resolve_credential_completion(
            _context(),
            selection_key="connection-stale",
        )

    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert "credential" not in state["evidence"]


def test_phase_local_reuse_does_not_require_global_preflight_choice(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)

    record = snow.record_admin_phase(
        _context(),
        phase="plugin-prerequisites",
        status="reused",
    )

    assert record["status"] == "reused"
    assert "record-reuse-decision" not in snow.build_parser().format_help()


def test_shared_reuse_partial_hr_to_it_requires_one_maker_acceptance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    before = target_path.read_bytes()
    connection_reads = 0

    class FakeConnectivity:
        def get_connection(self, connection_id: str) -> dict:
            nonlocal connection_reads
            connection_reads += 1
            assert uuid.UUID(connection_id).hex == CONNECTION_ID.replace("-", "")
            return _connection_record()

    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

        assert target_path.read_bytes() == before
        assert connection_reads == 0
        assert handoff["decision"] == "question-required"
        assert [item["id"] for item in handoff["candidate"]["items"]] == [
            "entra-registration",
            "servicenow-oidc",
        ]
        assert handoff["question"]["options"] == [
            {"label": "Reuse shared identity"},
            {"label": "Configure separately"},
            {"label": "Not now"},
        ]
        assert handoff["question"]["header"] == "Reuse shared identity"
        assert handoff["reuseDecisionChoices"] == [
            {
                "label": "Reuse shared identity",
                "decision": "accept",
            },
            {"label": "Configure separately", "decision": "decline"},
            {"label": "Not now", "decision": "not-now"},
        ]
        body = handoff["question"]["question"]
        assert "Employee Self-Service (IT)" in body
        assert "employee-self-service-it" in body
        assert ENVIRONMENT_ID in body
        assert "Confirmed ServiceNow instance: `dev123`" in body
        assert "HRSD agent `employee-self-service-hr`" in body
        assert "Entra administrator-confirmed application" in body
        assert "ServiceNow administrator-confirmed OIDC" in body
        assert "Physical ServiceNow" not in body
        assert "fresh read of that exact connection" not in body
        assert "each owning readiness check still runs" in body

        result = snow.apply_shared_configuration_reuse(
            context,
            source_key=handoff["candidate"]["sourceKey"],
            decision="accept",
        )
        entra_handoff = snow.identity_mapping_handoff(
            context,
            stage="entra",
        )
        oidc_handoff = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )
        resumed_entra = snow.identity_mapping_handoff(
            context,
            stage="entra",
        )
        entra_workbook_exists = snow._identity_workbook_path(
            context,
            stage="entra",
        ).exists()
        oidc_workbook_exists = snow._identity_workbook_path(
            context,
            stage="servicenow",
        ).exists()

    assert connection_reads == 0
    assert result["reusedItems"] == [
        "entra-registration",
        "servicenow-oidc",
    ]
    state = json.loads(target_path.read_text(encoding="utf-8"))
    assert state["sharedConfigurationReuse"]["decision"] == "accepted"
    assert state["adminSetup"]["phaseHandoffs"]["entra-registration"][
        "status"
    ] == "reused"
    assert state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
        "status"
    ] == "reused"
    assert state["adminSetup"]["phaseHandoffs"]["credential"] == {
        "status": "pending"
    }
    assert "credential" not in state["evidence"]
    assert state["evidence"]["identityMapping"]["binding"]["agentId"] == (
        OTHER_AGENT_ID
    )
    assert "topics" not in state["evidence"]
    assert "test" not in state["evidence"]
    assert "publish" not in state["evidence"]
    assert all(
        phase["status"] == "pending"
        for phase in state["phases"].values()
    )
    for completed_handoff in (
        entra_handoff,
        oidc_handoff,
        resumed_entra,
    ):
        assert completed_handoff["decision"] == "no-action"
        assert completed_handoff["handoffMode"] == "complete"
        assert completed_handoff["question"] is None
        assert completed_handoff["instruction"] is None
        assert completed_handoff["workbookPath"] is None
    assert entra_workbook_exists is False
    assert oidc_workbook_exists is False


def test_shared_reuse_supports_partial_it_to_hr_without_oidc_or_credential(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        include_oidc=False,
        include_credential=False,
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        instance_name="dev123",
    )

    with snow.product_profile_context("hrsd"):
        handoff = snow.shared_reuse_handoff(context)
        result = snow.apply_shared_configuration_reuse(
            context,
            source_key=handoff["candidate"]["sourceKey"],
            decision="accept",
        )
        entra_handoff = snow.identity_mapping_handoff(
            context,
            stage="entra",
        )
        oidc_handoff = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )

    assert handoff["decision"] == "question-required"
    assert [item["id"] for item in handoff["candidate"]["items"]] == [
        "entra-registration",
    ]
    assert result["reusedItems"] == ["entra-registration"]
    assert entra_handoff["decision"] == "no-action"
    assert entra_handoff["question"] is None
    assert oidc_handoff["decision"] == "question-required"
    assert oidc_handoff["question"] is not None
    assert oidc_handoff["stage"] == "servicenow"
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["adminSetup"]["phaseHandoffs"]["entra-registration"][
        "status"
    ] == "reused"
    assert saved["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] == {
        "status": "pending"
    }


def test_shared_reuse_full_it_to_hr_skips_both_admin_handoffs(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        instance_name="dev123",
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: pytest.fail(
            "Shared Identity reuse and admin handoff routing must stay local."
        ),
    )

    with snow.product_profile_context("hrsd"):
        offer = snow.shared_reuse_handoff(context)
        accepted = snow.apply_shared_configuration_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
        entra = snow.identity_mapping_handoff(context, stage="entra")
        oidc = snow.identity_mapping_handoff(context, stage="servicenow")
        resumed = snow.shared_reuse_handoff(context)

    assert accepted["reusedItems"] == [
        "entra-registration",
        "servicenow-oidc",
    ]
    assert resumed["decision"] == "already-recorded"
    assert entra["decision"] == "no-action"
    assert entra["question"] is None
    assert oidc["decision"] == "no-action"
    assert oidc["question"] is None
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["phases"]["entra-registration"]["status"] == "pending"
    assert saved["phases"]["servicenow-oidc"]["status"] == "pending"
    assert saved["phases"]["entra-registration"]["checkpointResults"] == {}
    assert saved["phases"]["servicenow-oidc"]["checkpointResults"] == {}


def test_accepted_shared_reuse_rejects_changed_source_without_target_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    with snow.product_profile_context("itsm"):
        offer = snow.shared_reuse_handoff(context)
        snow.apply_shared_configuration_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    source["updatedAt"] = "2026-10-30T00:00:00Z"
    source_path.write_text(json.dumps(source), encoding="utf-8")
    target_before = target_path.read_bytes()

    with snow.product_profile_context("itsm"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="source changed after it was accepted",
        ):
            snow.identity_mapping_handoff(context, stage="entra")

    assert target_path.read_bytes() == target_before


def test_accepted_shared_reuse_repairs_legacy_oidc_handoff_hash(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    with snow.product_profile_context("itsm"):
        offer = snow.shared_reuse_handoff(context)
        snow.apply_shared_configuration_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    target = json.loads(target_path.read_text(encoding="utf-8"))
    source_hash = source["evidence"]["identityClaimCandidates"][
        "candidateSetSha256"
    ]
    target_hash = target["evidence"]["identityClaimCandidates"][
        "candidateSetSha256"
    ]
    assert source_hash != target_hash
    target["adminSetup"]["phaseHandoffs"]["servicenow-oidc"]["evidence"][
        "candidateSetSha256"
    ] = source_hash
    target_path.write_text(json.dumps(target), encoding="utf-8")

    with snow.product_profile_context("itsm"):
        handoff = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )

    assert handoff["decision"] == "no-action"
    assert handoff["question"] is None
    repaired = json.loads(target_path.read_text(encoding="utf-8"))
    oidc_evidence = repaired["adminSetup"]["phaseHandoffs"][
        "servicenow-oidc"
    ]["evidence"]
    assert oidc_evidence["candidateSetSha256"] == target_hash
    assert oidc_evidence["mappingRelationshipSha256"] == repaired[
        "evidence"
    ]["identityMapping"]["relationshipSha256"]


def test_legacy_reuse_repair_is_atomic_when_required_flag_is_invalid(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    with snow.product_profile_context("itsm"):
        offer = snow.shared_reuse_handoff(context)
        snow.apply_shared_configuration_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    target = json.loads(target_path.read_text(encoding="utf-8"))
    oidc_evidence = target["adminSetup"]["phaseHandoffs"][
        "servicenow-oidc"
    ]["evidence"]
    oidc_evidence["candidateSetSha256"] = source["evidence"][
        "identityClaimCandidates"
    ]["candidateSetSha256"]
    oidc_evidence["runbookCompleted"] = False
    target_path.write_text(json.dumps(target), encoding="utf-8")
    state = json.loads(target_path.read_text(encoding="utf-8"))
    original_evidence = copy.deepcopy(
        state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"]["evidence"]
    )
    target_before = target_path.read_bytes()

    with snow.product_profile_context("itsm"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="incomplete or inconsistent",
        ):
            snow._identity_admin_handoff_is_ready(
                context,
                state,
                phase="servicenow-oidc",
            )

    assert target_path.read_bytes() == target_before
    assert state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
        "evidence"
    ] == original_evidence


def test_shared_reuse_can_apply_oidc_only_against_current_target_entra(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    state = json.loads(target_path.read_text(encoding="utf-8"))
    binding = {
        "tenantId": TENANT_ID,
        "environmentId": ENVIRONMENT_ID,
        "instanceName": "dev123",
        "applicationClientId": APP_CLIENT_ID,
        "provider": "servicenow-da-itsm",
        "profile": "itsm",
        "agentSlug": "employee-self-service-it",
        "agentId": OTHER_AGENT_ID,
    }
    candidates = ["employee_id"]
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "completed",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
            "tokenClaimCandidates": candidates,
            "tokenClaimEmissionConfirmed": True,
        },
    }
    state["evidence"]["identityClaimCandidates"] = {
        "kind": "administrator-confirmed-claim-candidates",
        "status": "entra-confirmed",
        "applicationClientId": APP_CLIENT_ID,
        "candidates": candidates,
        "entraAdministratorConfirmed": True,
        "tokenClaimEmissionConfirmed": True,
        "binding": binding,
        "candidateSetSha256": snow._identity_candidate_set_hash(
            application_client_id=APP_CLIENT_ID,
            candidates=candidates,
            binding=binding,
        ),
    }
    state["adminSetup"]["phaseHandoffs"]["entra-registration"][
        "evidence"
    ].update(
        {
            "tokenClaimCandidates": candidates,
            "tokenClaimEmissionConfirmed": True,
            "candidateSetSha256": state["evidence"][
                "identityClaimCandidates"
            ]["candidateSetSha256"],
            "binding": binding,
        }
    )
    target_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: pytest.fail(
            "Shared Identity reuse must not call Connectivity."
        ),
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)
        assert [item["id"] for item in handoff["candidate"]["items"]] == [
            "servicenow-oidc"
        ]
        result = snow.apply_shared_configuration_reuse(
            context,
            source_key=handoff["candidate"]["sourceKey"],
            decision="accept",
        )
        entra_handoff = snow.identity_mapping_handoff(
            context,
            stage="entra",
        )
        oidc_handoff = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )

    assert result["reusedItems"] == ["servicenow-oidc"]
    assert entra_handoff["decision"] == "no-action"
    assert entra_handoff["question"] is None
    assert oidc_handoff["decision"] == "no-action"
    assert oidc_handoff["question"] is None
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["adminSetup"]["phaseHandoffs"]["entra-registration"][
        "status"
    ] == "completed"
    assert saved["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
        "status"
    ] == "reused"
    assert saved["evidence"]["identityClaimCandidates"][
        "candidateSetSha256"
    ] == state["evidence"]["identityClaimCandidates"]["candidateSetSha256"]
    assert saved["evidence"]["identityMapping"]["binding"]["agentId"] == (
        OTHER_AGENT_ID
    )
    assert "credential" not in saved["evidence"]


def test_shared_reuse_does_not_offer_before_target_confirmation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: pytest.fail(
            "Preparing the local reuse offer must not call an API."
        ),
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    assert handoff == {
        "schemaVersion": "servicenow.shared-reuse-handoff.v1",
        "decision": "none",
        "candidates": [],
    }


@pytest.mark.parametrize(
    ("legacy_items", "expected_decision"),
    (
        (["preflight", "credential"], "question-required"),
        (
            ["preflight", "credential", "entra-registration"],
            "already-recorded",
        ),
    ),
)
def test_shared_identity_compatibility_ignores_legacy_physical_only_record(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    legacy_items: list[str],
    expected_decision: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_oidc=False,
        include_credential=True,
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    state = json.loads(target_path.read_text(encoding="utf-8"))
    with snow.product_profile_context("itsm"):
        candidate = snow._shared_reuse_candidates(context)[0]
    state["sharedConfigurationReuse"] = {
        "decision": "accepted",
        "sourceKey": candidate["sourceKey"],
        "sourceProvider": candidate["sourceProvider"],
        "sourceAgentSlug": candidate["sourceAgentSlug"],
        "applicableItems": legacy_items,
    }
    target_path.write_text(json.dumps(state), encoding="utf-8")

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    assert handoff["decision"] == expected_decision


def test_declined_entra_does_not_offer_dependent_oidc(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_oidc=False,
        include_credential=False,
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    with snow.product_profile_context("itsm"):
        first = snow.shared_reuse_handoff(context)
        assert [item["id"] for item in first["candidate"]["items"]] == [
            "entra-registration"
        ]
        snow.apply_shared_configuration_reuse(
            context,
            source_key=first["candidate"]["sourceKey"],
            decision="decline",
        )
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_oidc=True,
        include_credential=False,
    )

    with snow.product_profile_context("itsm"):
        second = snow.shared_reuse_handoff(context)

    assert second["decision"] == "already-recorded"


def test_identity_item_decision_history_preserves_prior_acceptance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_oidc=False,
        include_credential=False,
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    with snow.product_profile_context("itsm"):
        first = snow.shared_reuse_handoff(context)
        snow.apply_shared_configuration_reuse(
            context,
            source_key=first["candidate"]["sourceKey"],
            decision="accept",
        )
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_oidc=True,
        include_credential=False,
    )
    with snow.product_profile_context("itsm"):
        second = snow.shared_reuse_handoff(context)
        assert [item["id"] for item in second["candidate"]["items"]] == [
            "servicenow-oidc"
        ]
        snow.apply_shared_configuration_reuse(
            context,
            source_key=second["candidate"]["sourceKey"],
            decision="accept",
        )
        final = snow.shared_reuse_handoff(context)

    assert final["decision"] == "already-recorded"
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    decisions = saved["sharedConfigurationReuse"]["identityItemDecisions"]
    assert {
        (row["itemId"], row["decision"]) for row in decisions
    } == {
        ("entra-registration", "accepted"),
        ("servicenow-oidc", "accepted"),
    }


@pytest.mark.parametrize(
    ("mutation", "expected_items"),
    (
        ("candidate-hash", []),
        ("mapping-hash", ["entra-registration"]),
        ("mapping-client-id", ["entra-registration"]),
        ("source-binding", []),
    ),
)
def test_shared_identity_rejects_stale_hashes_and_source_bindings(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    mutation: str,
    expected_items: list[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_credential=False,
    )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    if mutation == "candidate-hash":
        source["evidence"]["identityClaimCandidates"][
            "candidateSetSha256"
        ] = "stale"
    elif mutation == "mapping-hash":
        source["evidence"]["identityMapping"]["relationshipSha256"] = (
            "stale"
        )
    elif mutation == "mapping-client-id":
        source["evidence"]["identityMapping"]["applicationClientId"] = (
            "00000000-0000-4000-8000-000000006667"
        )
        source["evidence"]["identityMapping"]["relationshipSha256"] = (
            snow._identity_mapping_relationship_hash(
                source["evidence"]["identityMapping"]
            )
        )
    else:
        source["evidence"]["identityClaimCandidates"]["binding"][
            "agentSlug"
        ] = "different-source-agent"
        source["evidence"]["identityMapping"]["binding"][
            "agentSlug"
        ] = "different-source-agent"
    source_path.write_text(json.dumps(source), encoding="utf-8")
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    if expected_items:
        assert [item["id"] for item in handoff["candidate"]["items"]] == (
            expected_items
        )
    else:
        assert handoff == {
            "schemaVersion": "servicenow.shared-reuse-handoff.v1",
            "decision": "none",
            "candidates": [],
        }


def test_shared_identity_discovers_same_profile_different_agent_slug(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr-source",
        include_credential=False,
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-hr-target",
        instance_name="dev123",
    )

    with snow.product_profile_context("hrsd"):
        handoff = snow.shared_reuse_handoff(context)

    assert handoff["decision"] == "question-required"
    assert handoff["candidate"]["sourceProfile"] == "hrsd"
    assert handoff["candidate"]["sourceAgentSlug"] == (
        "employee-self-service-hr-source"
    )
    assert [item["id"] for item in handoff["candidate"]["items"]] == [
        "entra-registration",
        "servicenow-oidc",
    ]


@pytest.mark.parametrize(
    ("source_profile", "target_profile", "source_slug", "target_slug"),
    (
        (
            "itsm",
            "hrsd",
            "employee-self-service-it",
            "employee-self-service-hr",
        ),
        (
            "hrsd",
            "itsm",
            "employee-self-service-hr",
            "employee-self-service-it",
        ),
        (
            "hrsd",
            "hrsd",
            "employee-self-service-hr-source",
            "employee-self-service-hr-target",
        ),
    ),
)
def test_saved_portal_value_persists_and_reloads_for_future_agent(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    source_profile: str,
    target_profile: str,
    source_slug: str,
    target_slug: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_agent_id = AGENT_ID
    target_agent_id = OTHER_AGENT_ID
    source_path = _write_portal_value_source_state(
        tmp_path,
        profile=source_profile,
        agent_id=source_agent_id,
        agent_slug=source_slug,
        portal_url="https://dev123.service-now.com/employee",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile=target_profile,
        agent_id=target_agent_id,
        agent_slug=target_slug,
        instance_name="dev123",
    )
    source_before = source_path.read_bytes()

    with snow.product_profile_context(target_profile):
        offer = snow.portal_value_reuse_handoff(context)
        accepted = snow.record_portal_value_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
        resumed = snow.portal_value_reuse_handoff(context)

    assert offer["decision"] == "question-required"
    assert offer["candidate"]["portalUrl"] == (
        "https://dev123.service-now.com/employee"
    )
    assert offer["question"]["options"] == [
        {"label": "Reuse saved portal URL"},
        {"label": "Enter separately"},
        {"label": "Not now"},
    ]
    assert accepted["decision"] == "accepted"
    assert resumed["decision"] == "already-recorded"
    assert resumed["record"]["portalUrl"] == accepted["portalUrl"]
    assert source_path.read_bytes() == source_before
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["portalValueReuse"] == accepted
    assert "portalConfirmation" not in saved["evidence"]
    assert saved["phases"]["portal-url"]["status"] == "pending"


@pytest.mark.parametrize(
    ("field", "value"),
    (
        ("kind", "wrong-kind"),
        ("valueType", "wrong-type"),
        ("targetProvider", "servicenow-da-itsm"),
        ("targetProfile", "itsm"),
        ("targetAgentSlug", "different-agent"),
        ("targetAgentId", "00000000-0000-4000-8000-000000002226"),
        ("tenantId", "00000000-0000-4000-8000-000000009998"),
        ("environmentId", "00000000-0000-4000-8000-000000001112"),
        ("instanceName", "different"),
    ),
)
def test_saved_portal_reuse_rejects_wrong_target_record_without_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    field: str,
    value: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_portal_value_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="it-source",
        portal_url="https://dev123.service-now.com/employee",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="hr-target",
        instance_name="dev123",
    )
    with snow.product_profile_context("hrsd"):
        offer = snow.portal_value_reuse_handoff(context)
        snow.record_portal_value_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
    target = json.loads(target_path.read_text(encoding="utf-8"))
    target["portalValueReuse"][field] = value
    target_path.write_text(json.dumps(target), encoding="utf-8")
    before = target_path.read_bytes()

    with snow.product_profile_context("hrsd"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="does not match the active target",
        ):
            snow.portal_value_reuse_handoff(context)

    assert target_path.read_bytes() == before


def test_saved_portal_reuse_survives_new_earlier_same_url_source(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_portal_value_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="zz-source",
        portal_url="https://dev123.service-now.com/employee",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="hr-target",
        instance_name="dev123",
    )
    with snow.product_profile_context("hrsd"):
        offer = snow.portal_value_reuse_handoff(context)
        accepted = snow.record_portal_value_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
    assert accepted["sourceAgentSlug"] == "zz-source"
    _write_portal_value_source_state(
        tmp_path,
        profile="hrsd",
        agent_id="00000000-0000-4000-8000-000000002226",
        agent_slug="aa-source",
        portal_url="https://dev123.service-now.com/employee",
    )

    with snow.product_profile_context("hrsd"):
        resumed = snow.portal_value_reuse_handoff(context)

    assert resumed["decision"] == "already-recorded"
    assert resumed["record"]["sourceAgentSlug"] == "zz-source"
    assert len(resumed["candidate"]["sources"]) == 2
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["portalValueReuse"] == accepted


def test_saved_portal_reuse_rejects_changed_accepted_source_receipt(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_portal_value_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="it-source",
        portal_url="https://dev123.service-now.com/employee",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="hr-target",
        instance_name="dev123",
    )
    with snow.product_profile_context("hrsd"):
        offer = snow.portal_value_reuse_handoff(context)
        snow.record_portal_value_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    source["evidence"]["portalConfirmation"]["recordedAt"] = (
        "2026-10-02T00:03:00Z"
    )
    source_path.write_text(json.dumps(source), encoding="utf-8")
    before = target_path.read_bytes()

    with snow.product_profile_context("hrsd"):
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="source changed or is no longer compatible",
        ):
            snow.portal_value_reuse_handoff(context)

    assert target_path.read_bytes() == before


def test_saved_portal_value_deduplicates_and_requires_value_selection(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_portal_value_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="hr-source-a",
        portal_url="https://dev123.service-now.com/employee",
    )
    _write_portal_value_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="it-source-a",
        portal_url="https://dev123.service-now.com/employee",
    )
    _write_portal_value_source_state(
        tmp_path,
        profile="hrsd",
        agent_id="00000000-0000-4000-8000-000000002224",
        agent_slug="hr-source-b",
        portal_url="https://dev123.service-now.com/esc",
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id="00000000-0000-4000-8000-000000002225",
        agent_slug="it-target",
        instance_name="dev123",
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.portal_value_reuse_handoff(context)

    assert handoff["decision"] == "source-selection-required"
    assert len(handoff["candidates"]) == 2
    employee = next(
        candidate
        for candidate in handoff["candidates"]
        if candidate["portalUrl"].endswith("/employee")
    )
    assert len(employee["sources"]) == 2
    assert len(employee["sourceLabels"]) == 2


@pytest.mark.parametrize(
    ("field", "value"),
    (
        ("environment", "00000000-0000-4000-8000-000000001112"),
        ("tenant", "00000000-0000-4000-8000-000000009998"),
        ("instance", "different"),
        ("url", "http://dev123.service-now.com/employee"),
        ("proof", False),
    ),
)
def test_saved_portal_value_rejects_incompatible_or_unconfirmed_sources(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    field: str,
    value: str | bool,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_portal_value_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="hr-source",
        portal_url=(
            str(value)
            if field == "url"
            else "https://dev123.service-now.com/employee"
        ),
        environment_id=(
            str(value) if field == "environment" else ENVIRONMENT_ID
        ),
        tenant_id=str(value) if field == "tenant" else TENANT_ID,
        instance_name=str(value) if field == "instance" else "dev123",
        valid_proof=bool(value) if field == "proof" else True,
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="it-target",
        instance_name="dev123",
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.portal_value_reuse_handoff(context)

    assert handoff == {
        "schemaVersion": "servicenow.portal-value-reuse.v1",
        "decision": "none",
        "candidates": [],
    }


def test_saved_portal_value_same_target_uses_fresh_readback_without_update(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_portal_value_source_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        portal_url="https://dev123.service-now.com/employee",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        instance_name="dev123",
    )
    components = _portal_components()
    topic = next(
        change["component"]
        for change in components["botComponentChanges"]
        if change["component"].get("displayName")
        == "ServiceNow HRSD Setup Configurations"
    )
    topic["dialog"]["beginDialog"]["actions"][0]["value"]["literalValue"] = (
        "https://dev123.service-now.com/employee"
    )

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return copy.deepcopy(components)

        def update_components(self, *_args, **_kwargs) -> None:
            pytest.fail("Same saved portal URL must not call update.")

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    with snow.product_profile_context("hrsd"):
        offer = snow.portal_value_reuse_handoff(context)
        accepted = snow.record_portal_value_reuse(
            context,
            source_key=offer["candidate"]["sourceKey"],
            decision="accept",
        )
        result = snow.set_portal_url(
            context,
            accepted["portalUrl"],
            confirmed=True,
        )

    assert result["status"] == "already-configured"
    assert result["readbackVerified"] is True
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    receipt = saved["evidence"]["portalConfirmation"]
    assert receipt["portalUrl"] == accepted["portalUrl"]
    assert receipt["binding"]["agentId"] == AGENT_ID
    assert receipt["binding"]["topicId"] != (
        json.loads(
            (
                tmp_path
                / ".local/connect/servicenow-da-itsm/agents/"
                "employee-self-service-it/lifecycle.json"
            ).read_text(encoding="utf-8")
        )["evidence"]["portalConfirmation"]["binding"]["topicId"]
    )


def test_local_confirmed_instance_is_preflight_candidate_only(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_portal_value_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="hr-source",
        portal_url="https://dev123.service-now.com/employee",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="it-target",
    )

    class EmptyConnectivity:
        def list_connections(self) -> list:
            return []

    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: EmptyConnectivity(),
    )
    with snow.product_profile_context("itsm"):
        discovery = snow.inspect_admin_setup(context)

    assert discovery["instanceInputRequired"] is True
    assert discovery["instanceCandidates"] == [
        "https://dev123.service-now.com"
    ]
    assert discovery["instanceChoices"][
        "Use https://dev123.service-now.com"
    ] == "https://dev123.service-now.com"
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["adminSetup"]["preflight"].get("instanceName") is None
    assert saved["adminSetup"]["phaseHandoffs"]["preflight"] == {
        "status": "pending"
    }


def test_validated_identity_mapping_rejects_application_id_mismatch(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _record_test_identity_mapping(tmp_path)
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    mapping = state["evidence"]["identityMapping"]
    mapping["applicationClientId"] = (
        "00000000-0000-4000-8000-000000006667"
    )
    mapping["relationshipSha256"] = (
        snow._identity_mapping_relationship_hash(mapping)
    )
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="valid ServiceNow-selected claim",
    ):
        snow._validated_identity_mapping(context, state)


def test_shared_reuse_does_not_fabricate_legacy_mapping_confirmation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    del source["evidence"]["identityMapping"]
    source_path.write_text(json.dumps(source), encoding="utf-8")
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    item_ids = {item["id"] for item in handoff["candidate"]["items"]}
    assert "entra-registration" in item_ids
    assert "servicenow-oidc" not in item_ids
    assert item_ids == {"entra-registration"}


def test_shared_reuse_does_not_grandfather_legacy_two_name_mapping(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    source = json.loads(source_path.read_text(encoding="utf-8"))
    source["evidence"].pop("identityClaimCandidates")
    source["evidence"]["identityMapping"] = {
        "kind": "administrator-confirmed-identity-mapping",
        "applicationClientId": APP_CLIENT_ID,
        "entraSourceAttribute": "user.employeeid",
        "tokenClaimName": "employee_id",
        "entraAdministratorConfirmed": True,
        "tokenClaimEmissionConfirmed": True,
    }
    source_path.write_text(json.dumps(source), encoding="utf-8")
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    assert handoff == {
        "schemaVersion": "servicenow.shared-reuse-handoff.v1",
        "decision": "none",
        "candidates": [],
    }


def test_shared_reuse_wait_is_zero_write_and_decline_imports_nothing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_credential=False,
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    before = target_path.read_bytes()

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)
        assert target_path.read_bytes() == before

        declined = snow.apply_shared_configuration_reuse(
            context,
            source_key=handoff["candidate"]["sourceKey"],
            decision="decline",
        )
        repeated = snow.shared_reuse_handoff(context)

    state = json.loads(target_path.read_text(encoding="utf-8"))
    assert declined["decision"] == "declined"
    assert repeated["decision"] == "already-recorded"
    assert "identityMapping" not in state["evidence"]
    assert state["adminSetup"]["phaseHandoffs"]["entra-registration"] == {
        "status": "pending"
    }


@pytest.mark.parametrize(
    ("environment_id", "instance_name"),
    [
        ("00000000-0000-4000-8000-000000009998", "dev123"),
        (ENVIRONMENT_ID, "otherinstance"),
    ],
)
def test_shared_reuse_rejects_incompatible_target_identity(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    environment_id: str,
    instance_name: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        environment_id=environment_id,
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name=instance_name,
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    assert handoff == {
        "schemaVersion": "servicenow.shared-reuse-handoff.v1",
        "decision": "none",
        "candidates": [],
    }


def test_shared_reuse_rejects_different_target_app_and_claim_relationship(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
    )
    context, target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )
    state = json.loads(target_path.read_text(encoding="utf-8"))
    target_app_id = "00000000-0000-4000-8000-000000006667"
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "completed",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": target_app_id,
            "tokenClaimCandidates": ["email"],
            "tokenClaimEmissionConfirmed": True,
        },
    }
    binding = {
        "tenantId": TENANT_ID,
        "environmentId": ENVIRONMENT_ID,
        "instanceName": "dev123",
        "applicationClientId": target_app_id,
        "provider": "servicenow-da-itsm",
        "profile": "itsm",
        "agentSlug": "employee-self-service-it",
        "agentId": OTHER_AGENT_ID,
    }
    state["evidence"]["identityClaimCandidates"] = {
        "kind": "administrator-confirmed-claim-candidates",
        "status": "entra-confirmed",
        "applicationClientId": target_app_id,
        "candidates": ["email"],
        "entraAdministratorConfirmed": True,
        "tokenClaimEmissionConfirmed": True,
        "binding": binding,
    }
    state["evidence"]["identityClaimCandidates"]["candidateSetSha256"] = (
        snow._identity_candidate_set_hash(
            application_client_id=target_app_id,
            candidates=["email"],
            binding=binding,
        )
    )
    target_path.write_text(json.dumps(state), encoding="utf-8")

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)

    assert handoff == {
        "schemaVersion": "servicenow.shared-reuse-handoff.v1",
        "decision": "none",
        "candidates": [],
    }


def test_shared_reuse_rejects_source_change_after_question(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    source_path = _write_shared_source_state(
        tmp_path,
        profile="hrsd",
        agent_id=AGENT_ID,
        agent_slug="employee-self-service-hr",
        include_credential=False,
    )
    context, _target_path = _write_shared_target_state(
        tmp_path,
        profile="itsm",
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
        instance_name="dev123",
    )

    with snow.product_profile_context("itsm"):
        handoff = snow.shared_reuse_handoff(context)
        source = json.loads(source_path.read_text(encoding="utf-8"))
        source["updatedAt"] = "changed-after-question"
        source_path.write_text(json.dumps(source), encoding="utf-8")

        with pytest.raises(
            snow.ServiceNowConnectError,
            match="stale or no longer unique",
        ):
            snow.apply_shared_configuration_reuse(
                context,
                source_key=handoff["candidate"]["sourceKey"],
                decision="accept",
            )


def test_entra_workbook_is_protected_with_only_claim_cells_unlocked(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)

    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    workbook = load_workbook(path)
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]

    assert handoff["stage"] == "entra"
    assert handoff["editableFields"] == [
        "application_client_id",
        "user_claim",
    ]
    assert workbook.sheetnames == ["Entra Handoff", "_metadata"]
    assert workbook.active.title == "Entra Handoff"
    assert sheet.max_row == 2
    assert sheet.max_column == 3
    assert tuple(sheet.cell(1, column).value for column in range(1, 4)) == (
        "Instructions",
        "Client ID",
        "User Claim",
    )
    assert sheet.protection.sheet is True
    assert sheet["A2"].protection.locked is True
    assert sheet["B2"].protection.locked is False
    assert sheet["C2"].protection.locked is False
    assert sheet["B2"].fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet["C2"].fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet["B2"].value == APP_CLIENT_ID
    assert sheet["C2"].value in {None, ""}
    assert metadata.sheet_state == "hidden"
    assert metadata.protection.sheet is True
    workbook.close()


def test_identity_workbook_short_temp_name_fits_live_windows_path_budget(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    live_parent = (
        r"C:\Users\example123\DEV\copilot-worktrees"
        r"\Employee-Self-Service-Agent-Developer-Kit"
        r"\it-validation-20261007-newenv\solutions\ess-maker-skills"
        r"\.local\connect\servicenow-da-itsm\agents"
        r"\employee-self-service-it"
    )
    final_target = (
        live_parent + r"\servicenow-identity-claims-entra.xlsx"
    )
    old_temporary = final_target + ".aucbxw_v.tmp.xlsx"
    new_temporary = live_parent + r"\.identity-aucbxw_v.tmp.xlsx"

    assert len(live_parent) == 206
    assert len(final_target) == 244
    assert len(old_temporary) == 262
    assert len(new_temporary) == 234
    assert len(new_temporary) < snow.WINDOWS_LEGACY_MAX_PATH

    target_parent_length = 206
    remaining = target_parent_length - len(str(tmp_path))
    components = []
    while remaining > 1:
        component_length = min(remaining - 1, 80)
        components.append("x" * component_length)
        remaining -= component_length + 1
    parent = tmp_path.joinpath(*components)
    context_values = {
        key: f"test-{key}"
        for key in snow.IDENTITY_WORKBOOK_CONTEXT_FIELDS
    }
    context_values.update(
        {
            "provider": snow.PROVIDER_KEY,
            "profile": snow.PROFILE_KEY,
            "agent_slug": AGENT_SLUG,
            "agent_id": AGENT_ID,
            "confirmed_user_claim_name": "email",
        }
    )
    original_mkstemp = snow.tempfile.mkstemp
    calls: list[dict[str, object]] = []

    def captured_mkstemp(*args: object, **kwargs: object) -> tuple[int, str]:
        calls.append(dict(kwargs))
        return original_mkstemp(*args, **kwargs)

    monkeypatch.setattr(snow.tempfile, "mkstemp", captured_mkstemp)
    entra_path = parent / "servicenow-identity-claims-entra.xlsx"
    servicenow_path = (
        parent / "servicenow-identity-mapping-servicenow.xlsx"
    )
    snow._write_identity_workbook(
        entra_path,
        context_values=context_values,
        stage="entra",
        candidates=[],
        steps=snow._identity_workbook_steps(
            stage="entra",
            context_values=context_values,
            missing_items=[("CLAIMS", "employee identity claim workbook")],
        ),
    )
    snow._write_identity_workbook(
        servicenow_path,
        context_values=context_values,
        stage="servicenow",
        candidates=["email"],
        steps=snow._identity_workbook_steps(
            stage="servicenow",
            context_values=context_values,
            missing_items=[],
        ),
    )

    assert len(str(parent)) == target_parent_length
    assert entra_path.is_file()
    assert servicenow_path.is_file()
    assert [call["prefix"] for call in calls] == [
        snow.IDENTITY_WORKBOOK_TEMP_PREFIX,
        snow.IDENTITY_WORKBOOK_TEMP_PREFIX,
    ]
    assert [call["suffix"] for call in calls] == [
        snow.IDENTITY_WORKBOOK_TEMP_SUFFIX,
        snow.IDENTITY_WORKBOOK_TEMP_SUFFIX,
    ]
    assert all(Path(str(call["dir"])) == parent for call in calls)


def test_identity_workbook_rejects_unsupported_windows_target_path() -> None:
    too_long = Path("x" * snow.WINDOWS_LEGACY_MAX_PATH)

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="target path is too long",
    ):
        snow._validate_identity_workbook_path_budget(
            too_long,
            platform_name="nt",
        )


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_claim_only_workbook_question_is_short_and_missing_only(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    agent_id = AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID
    slug = AGENT_SLUG if profile == "hrsd" else "employee-self-service-it"
    context, state_path = _write_shared_target_state(
        tmp_path,
        profile=profile,
        agent_id=agent_id,
        agent_slug=slug,
        instance_name="dev123",
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "reused",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
        },
    }
    prefix = "SN-DA-HRSD" if profile == "hrsd" else "SN-DA-ITSM"
    state["phases"]["entra-registration"]["checkpointResults"] = {
        f"{prefix}-ENTRA-APP-001": "Passed",
        f"{prefix}-ENTRA-CLAIMS-001": "NotConfigured",
        f"{prefix}-ENTRA-SCOPE-001": "Passed",
        f"{prefix}-ENTRA-PREAUTH-001": "Passed",
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")

    with snow.product_profile_context(profile):
        _confirm_identity_requirement(context)
        state_before = state_path.read_bytes()
        handoff = snow.identity_mapping_handoff(context, stage="entra")

    question_object = handoff["question"]
    question = handoff["instruction"]
    assert handoff["handoffMode"] == "claim-only"
    assert handoff["instructionTitle"] == (
        "Step 1 of 2 — Entra: list employee identity options"
    )
    assert question_object == {
        "header": (
            "Please enter the returned Entra admin handoff workbook URL."
        ),
        "question": (
            "Please enter the returned Entra admin handoff workbook URL."
        ),
        "allowFreeformInput": True,
    }
    assert question.startswith("**Purpose**\n")
    assert "Step 1 of 2" not in question
    assert "Complete the Entra application and identity-claim handoff." in question
    assert "Application Administrator or Cloud Application Administrator" in (
        question
    )
    assert "https://entra.microsoft.com/" in question
    assert handoff["workbookUri"] in question
    assert "options" not in question_object
    assert question_object["allowFreeformInput"] is True
    assert str(Path(handoff["workbookPath"])) in question
    assert handoff["workbookUri"] in question
    assert handoff["clientIdQuestion"] is None
    assert state_path.read_bytes() == state_before
    headings = ("**Purpose**", "**Owner**", "**Known target**", "**Action**")
    positions = [question.index(heading) for heading in headings]
    assert positions == sorted(positions)
    assert (
        "\n\n**Owner**\nApplication Administrator or Cloud Application "
        "Administrator." in question
    )
    assert "\n\n**Action**\n1. " in question
    assert "\n2. " in question
    assert "\n3. " in question
    assert "\n4. " in question
    assert "\n5. " not in question
    assert "**Expected result**" not in question
    assert "Purpose:" not in question
    assert "Owner:" not in question
    assert str(Path(handoff["workbookPath"])) in question
    assert str(Path(handoff["csvPath"])) in question
    assert "`Entra Handoff`" in question
    assert "one selected User Claim name" in question
    assert "Do not enter `aud`" in question
    assert "completed XLSX or CSV file" in question
    assert "Separate readiness checks still report" not in question
    assert "then choose Completed" not in question
    assert "Not yet" not in question
    for forbidden in (
        "no setup redo",
        "candidates",
        "token_claim_name",
        "current status",
        "signed-in employee",
    ):
        assert forbidden not in question
    assert "Application ID URI" not in question
    workbook = load_workbook(Path(handoff["workbookPath"]))
    assert workbook.sheetnames == ["Entra Handoff", "_metadata"]
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    instruction = sheet["A2"].value
    assert instruction == snow._entra_admin_instruction()
    assert "\r" not in instruction and "\n" not in instruction
    assert "**New registration.**" in instruction
    assert "*Aud* - for audience validation" in instruction
    assert "*Email* - addressable email for the user" in instruction
    assert "*UPN* - an identifier for the user" in instruction
    assert "Microsoft Graph permissions" in instruction
    assert snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL in instruction
    assert sheet["A2"].hyperlink.target == (
        snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL
    )
    assert sheet["B2"].value == APP_CLIENT_ID
    assert sheet["C2"].value in {None, ""}
    workbook.close()


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
@pytest.mark.parametrize(
    ("scenario", "expected_mode", "expected_missing"),
    (
        (
            "fresh",
            "full-setup",
            {"APP", "CLAIMS", "SCOPE", "PREAUTH"},
        ),
        (
            "partial",
            "delta",
            {"CLAIMS", "PREAUTH"},
        ),
        ("claim-only", "claim-only", {"CLAIMS"}),
    ),
)
def test_public_entra_handoff_emits_v3_workbook_for_hr_and_it(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
    scenario: str,
    expected_mode: str,
    expected_missing: set[str],
) -> None:
    scenario_root = tmp_path / f"{profile}-{scenario}"
    scenario_root.mkdir()
    monkeypatch.chdir(scenario_root)
    universal_template = _author_universal_entra_template(
        scenario_root / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    product = snow.PRODUCT_PROFILES[profile]
    slug = (
        "employee-self-service-hr"
        if profile == "hrsd"
        else "employee-self-service-it"
    )
    agent_id = AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID
    context = _profile_context(
        profile,
        agent_id=agent_id,
        agent_slug=slug,
    )
    context["root"] = scenario_root
    components = _components() if profile == "hrsd" else _itsm_components()
    prefix = "SN-DA-HRSD" if profile == "hrsd" else "SN-DA-ITSM"
    with snow.product_profile_context(profile):
        state = snow._state_base(context, components)
        state["attested"] = True
        state["acceptedContractRevision"] = 1
        state["phases"] = {
            phase: {"status": "pending", "checkpointResults": {}}
            for phase in (
                "preflight",
                "entra-registration",
                "servicenow-oidc",
                "credential",
                "topics",
                "portal-url",
                "agent-connection",
                "test",
                "publish",
            )
        }
        state["adminSetup"]["preflight"]["instanceName"] = (
            f"example-{profile}-{scenario}"
        )
        if scenario == "fresh":
            checkpoint_results = {}
        elif scenario == "partial":
            checkpoint_results = {
                f"{prefix}-ENTRA-APP-001": "Passed",
                f"{prefix}-ENTRA-CLAIMS-001": "NotConfigured",
                f"{prefix}-ENTRA-SCOPE-001": "Passed",
                f"{prefix}-ENTRA-PREAUTH-001": "NotConfigured",
            }
        else:
            checkpoint_results = {
                f"{prefix}-ENTRA-APP-001": "Passed",
                f"{prefix}-ENTRA-CLAIMS-001": "NotConfigured",
                f"{prefix}-ENTRA-SCOPE-001": "Passed",
                f"{prefix}-ENTRA-PREAUTH-001": "Passed",
            }
        state["phases"]["entra-registration"]["checkpointResults"] = (
            checkpoint_results
        )
        if scenario != "fresh":
            state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
                "status": "reused",
                "evidence": {
                    "kind": "structured-admin-attestation",
                    "recordedAt": "2026-10-01T00:00:00Z",
                    "clientId": APP_CLIENT_ID,
                },
            }
        state_path = snow._lifecycle_state_path(context)
        snow._write_json_atomic(state_path, state)
        state_before = state_path.read_bytes()
        handoff = snow.identity_mapping_handoff(context, stage="entra")

    assert state_path.read_bytes() == state_before
    assert handoff["provider"] == product.provider_key
    assert handoff["profile"] == profile
    assert handoff["handoffMode"] == expected_mode
    assert {item["checkpoint"] for item in handoff["missingItems"]} == (
        expected_missing
    )
    assert handoff["clientIdQuestion"] is None
    assert handoff["instructionTitle"] == (
        "Step 1 of 2 — Entra: list employee identity options"
    )
    assert handoff["question"] == {
        "header": (
            "Please enter the returned Entra admin handoff workbook URL."
        ),
        "question": (
            "Please enter the returned Entra admin handoff workbook URL."
        ),
        "allowFreeformInput": True,
    }
    instruction = handoff["instruction"]
    assert "No separate readiness item is included" not in instruction
    assert "Separate readiness checks still report" not in instruction
    assert "Separate checks must pass on their own evidence" not in instruction
    assert "Application ID URI and `user_impersonation` scope" not in (
        instruction
    )
    assert "ServiceNow connector pre-authorization" not in instruction
    assert "**Expected result**" not in instruction
    assert "Completed" not in instruction
    assert "Not yet" not in instruction
    assert "Complete the Entra application and identity-claim handoff." in (
        instruction
    )
    assert "one selected User Claim name" in instruction
    assert handoff["workbookUri"] in instruction
    assert handoff["csvUri"] in instruction
    assert str(Path(handoff["workbookPath"])) in instruction
    assert str(Path(handoff["csvPath"])) in instruction
    assert context["environment"]["tenant_id"] in instruction
    assert context["environment"]["id"] in instruction
    assert context["agent"]["id"] in instruction
    assert f"example-{profile}-{scenario}" in instruction
    workbook_path = Path(handoff["workbookPath"])
    assert workbook_path.is_file()
    workbook = load_workbook(
        workbook_path,
        data_only=False,
        keep_links=False,
    )
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    assert workbook.sheetnames == ["Entra Handoff", "_metadata"]
    assert sheet["A2"].value == snow._entra_admin_instruction()
    assert "\r" not in sheet["A2"].value
    assert "\n" not in sheet["A2"].value
    assert sheet["A2"].hyperlink.target == (
        snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL
    )
    assert metadata.sheet_state == "hidden"
    assert metadata.protection.sheet is True
    assert sheet.protection.sheet is True
    assert sheet["C2"].value in {None, ""}
    client_cell = sheet["B2"]
    if scenario == "fresh":
        assert handoff["applicationClientId"] is None
        assert client_cell.value in {None, ""}
        assert client_cell.protection.locked is False
    else:
        assert handoff["applicationClientId"] == APP_CLIENT_ID
        assert client_cell.value == APP_CLIENT_ID
        assert client_cell.protection.locked is False
    workbook.close()


def test_identity_requirement_not_sure_idempotency_and_scope_are_safe(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "example-instance"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    before = state_path.read_bytes()

    waiting = snow.record_identity_requirement(
        context,
        choice="not-sure",
    )
    assert waiting["status"] == "waiting"
    assert waiting["recorded"] is False
    assert state_path.read_bytes() == before
    assert not snow._identity_workbook_path(
        context,
        stage="entra",
    ).exists()

    receipt = _confirm_identity_requirement(context, choice="email")
    after_record = state_path.read_bytes()
    repeated = _confirm_identity_requirement(context, choice="email")
    assert repeated["requirementSha256"] == receipt["requirementSha256"]
    assert state_path.read_bytes() == after_record

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="immutable",
    ):
        _confirm_identity_requirement(context, choice="upn")
    assert state_path.read_bytes() == after_record

    other_context = copy.deepcopy(context)
    other_context["environment"]["id"] = (
        "99999999-9999-4999-8999-999999999999"
    )
    with pytest.raises(
        snow.ServiceNowConnectError,
        match="does not match the active agent",
    ):
        snow.identity_mapping_handoff(other_context, stage="entra")
    assert state_path.read_bytes() == after_record


def test_unfilled_legacy_entra_workbook_is_upgraded_in_place(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "example-instance"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    expected_identity = snow._required_context_identity(context, state)
    context_values = snow._identity_workbook_context_values(
        context,
        expected_identity,
        application_client_id=None,
        stage="entra",
    )
    path = snow._identity_workbook_path(context, stage="entra")
    snow._write_identity_workbook(
        path,
        context_values=context_values,
        stage="entra",
        candidates=[],
        steps=[],
        missing_items=[
            ("APP", "application"),
            ("CLAIMS", "claims"),
        ],
    )
    workbook = load_workbook(path)
    workbook.properties.identifier = None
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()

    handoff = snow.identity_mapping_handoff(context, stage="entra")

    assert handoff["handoffMode"] == "full-setup"
    assert handoff["workbookFileStatus"] == "generated"
    assert Path(handoff["workbookPath"]) == path.resolve()
    assert path.read_bytes() != before
    assert snow._identity_workbook_template_marker(path) == (
        snow._identity_workbook_template_sha256("entra")
    )
    saved = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert "identityRequirement" not in saved.get("evidence", {})


@pytest.mark.parametrize("admin_value", ("claim", "client-id"))
def test_filled_legacy_workbook_gets_versioned_current_template_copy(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    admin_value: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    template = _author_universal_entra_template(
        tmp_path / "current-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    initial = snow.identity_mapping_handoff(
        context,
        stage="entra",
    )
    canonical = Path(initial["workbookPath"])
    workbook = load_workbook(canonical)
    workbook.properties.identifier = None
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["A2"] = (
        "Legacy instructions"
    )
    if admin_value == "claim":
        workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["C2"] = "email"
    else:
        client_cell = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"]
        assert client_cell.protection.locked is False
        client_cell.value = APP_CLIENT_ID
    workbook.save(canonical)
    workbook.close()
    canonical_before = canonical.read_bytes()

    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
    )

    returned = Path(handoff["workbookPath"])
    expected_sha = snow._identity_workbook_template_sha256("entra")
    assert handoff["workbookFileStatus"] == "generated-versioned"
    assert returned != canonical.resolve()
    assert returned.name == (
        "servicenow-identity-claims-entra.template-"
        f"{expected_sha[:12].lower()}.xlsx"
    )
    assert canonical.read_bytes() == canonical_before
    assert snow._identity_workbook_template_marker(returned) == expected_sha
    current = load_workbook(returned, data_only=False, keep_links=False)
    assert current[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["A2"].value == (
        snow._entra_admin_instruction()
    )
    assert current[snow.IDENTITY_ENTRA_HANDOFF_SHEET].protection.sheet is True
    assert current["_metadata"].sheet_state == "hidden"
    assert current["_metadata"].protection.sheet is True
    current.close()

    resumed = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    assert resumed["workbookFileStatus"] == "preserved-versioned"
    assert Path(resumed["workbookPath"]) == returned


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
@pytest.mark.parametrize("needs_handoff", (True, False))
def test_legacy_candidate_receipt_requires_type_only_for_new_handoff(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
    needs_handoff: bool,
) -> None:
    scenario_root = tmp_path / f"{profile}-{needs_handoff}"
    scenario_root.mkdir()
    monkeypatch.chdir(scenario_root)
    product = snow.PRODUCT_PROFILES[profile]
    slug = (
        "employee-self-service-hr"
        if profile == "hrsd"
        else "employee-self-service-it"
    )
    agent_id = AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID
    context = _profile_context(
        profile,
        agent_id=agent_id,
        agent_slug=slug,
    )
    context["root"] = scenario_root
    components = _components() if profile == "hrsd" else _itsm_components()
    prefix = "SN-DA-HRSD" if profile == "hrsd" else "SN-DA-ITSM"
    with snow.product_profile_context(profile):
        state = snow._state_base(context, components)
        state["attested"] = True
        state["acceptedContractRevision"] = 1
        state["phases"] = {
            phase: {"status": "pending", "checkpointResults": {}}
            for phase in (
                "preflight",
                "entra-registration",
                "servicenow-oidc",
                "credential",
                "topics",
                "portal-url",
                "agent-connection",
                "test",
                "publish",
            )
        }
        state["adminSetup"]["preflight"]["instanceName"] = "example-instance"
        statuses = {
            "APP": "Passed",
            "CLAIMS": "Passed",
            "SCOPE": "Passed",
            "PREAUTH": "Failed" if needs_handoff else "Passed",
        }
        state["phases"]["entra-registration"]["checkpointResults"] = {
            f"{prefix}-ENTRA-{suffix}-001": status
            for suffix, status in statuses.items()
        }
        binding = {
            "tenantId": TENANT_ID,
            "environmentId": ENVIRONMENT_ID,
            "instanceName": "example-instance",
            "applicationClientId": APP_CLIENT_ID,
            "provider": product.provider_key,
            "profile": profile,
            "agentSlug": slug,
            "agentId": agent_id,
        }
        receipt = {
            "kind": "administrator-confirmed-claim-candidates",
            "status": "entra-confirmed",
            "applicationClientId": APP_CLIENT_ID,
            "candidates": ["email"],
            "entraAdministratorConfirmed": True,
            "tokenClaimEmissionConfirmed": True,
            "binding": binding,
            "importedWorkbookSha256": "legacy-workbook-hash",
            "sheetProtected": True,
            "recordedAt": "2026-10-01T00:00:00Z",
        }
        receipt["candidateSetSha256"] = snow._identity_candidate_set_hash(
            application_client_id=APP_CLIENT_ID,
            candidates=["email"],
            binding=binding,
        )
        state["evidence"]["identityClaimCandidates"] = receipt
        state_path = snow._lifecycle_state_path(context)
        snow._write_json_atomic(state_path, state)
        before = state_path.read_bytes()
        handoff = snow.identity_mapping_handoff(context, stage="entra")

    assert state_path.read_bytes() == before
    saved = json.loads(state_path.read_text(encoding="utf-8"))
    assert saved["evidence"]["identityClaimCandidates"][
        "candidateSetSha256"
    ] == receipt["candidateSetSha256"]
    assert "identityRequirement" not in saved["evidence"]
    if needs_handoff:
        assert handoff["handoffMode"] == "delta"
        assert Path(handoff["workbookPath"]).is_file()
        assert "options" not in handoff["question"]
        assert handoff["question"]["allowFreeformInput"] is True
        assert handoff["question"] is not None
    else:
        assert not snow._identity_workbook_path(
            context,
            stage="entra",
        ).exists()
        assert handoff["handoffMode"] == "complete"
        assert handoff["decision"] == "no-action"
        assert handoff["question"] is None
        assert handoff["legacyRequirementBypass"] == (
            "completed-candidate-receipt"
        )


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
@pytest.mark.parametrize("scenario", ("fresh", "claim-only"))
def test_custom_identity_overlay_flows_through_entra_and_servicenow(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
    scenario: str,
) -> None:
    scenario_root = tmp_path / f"{profile}-{scenario}"
    scenario_root.mkdir()
    monkeypatch.chdir(scenario_root)
    universal_template = _author_universal_entra_template(
        scenario_root / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    servicenow_template = _author_servicenow_template(
        scenario_root / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    slug = (
        "employee-self-service-hr"
        if profile == "hrsd"
        else "employee-self-service-it"
    )
    agent_id = AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID
    context = _profile_context(
        profile,
        agent_id=agent_id,
        agent_slug=slug,
    )
    context["root"] = scenario_root
    components = _components() if profile == "hrsd" else _itsm_components()
    prefix = "SN-DA-HRSD" if profile == "hrsd" else "SN-DA-ITSM"
    with snow.product_profile_context(profile):
        state = snow._state_base(context, components)
        state["attested"] = True
        state["acceptedContractRevision"] = 1
        state["phases"] = {
            phase: {"status": "pending", "checkpointResults": {}}
            for phase in (
                "preflight",
                "entra-registration",
                "servicenow-oidc",
                "credential",
                "topics",
                "portal-url",
                "agent-connection",
                "test",
                "publish",
            )
        }
        state["adminSetup"]["preflight"]["instanceName"] = (
            f"example-{profile}-{scenario}"
        )
        if scenario == "claim-only":
            state["phases"]["entra-registration"]["checkpointResults"] = {
                f"{prefix}-ENTRA-APP-001": "Passed",
                f"{prefix}-ENTRA-CLAIMS-001": "NotConfigured",
                f"{prefix}-ENTRA-SCOPE-001": "Passed",
                f"{prefix}-ENTRA-PREAUTH-001": "Passed",
                f"{prefix}-ENTRA-PERMISSIONS-001": "Passed",
                f"{prefix}-ENTRA-CONSENT-001": "Passed",
            }
            state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
                "status": "reused",
                "evidence": {
                    "kind": "structured-admin-attestation",
                    "recordedAt": "2026-10-01T00:00:00Z",
                    "clientId": APP_CLIENT_ID,
                },
            }
        state_path = snow._lifecycle_state_path(context)
        snow._write_json_atomic(state_path, state)
        requirement = _confirm_identity_requirement(
            context,
            choice="pay-employee-id",
        )
        assert requirement["category"] == "custom"
        entra = snow.identity_mapping_handoff(context, stage="entra")

    entra_path = Path(entra["workbookPath"])
    workbook = load_workbook(
        entra_path,
        data_only=False,
        keep_links=False,
    )
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    text = sheet["A2"].value
    assert text == snow._entra_admin_instruction()
    assert "\r" not in text and "\n" not in text
    assert snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL in text
    assert sheet["A2"].hyperlink.target == (
        snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL
    )
    assert "If the selected User Claim is not email" in text
    assert "one selected User Claim name" in text
    assert "*Aud* - for audience validation" in text
    assert "must not be entered as the selected User Claim" in text
    if scenario == "fresh":
        assert sheet["B2"].value in {None, ""}
        sheet["B2"] = APP_CLIENT_ID
    else:
        assert entra["handoffMode"] == "claim-only"
        assert sheet["B2"].value == APP_CLIENT_ID
    workbook.save(entra_path)
    workbook.close()
    _fill_entra_workbook(entra_path, ("approved_employee_key",))
    with snow.product_profile_context(profile):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=entra_path,
            client_id=None,
            confirmed=True,
        )
        snow.record_admin_phase(
            context,
            phase="entra-registration",
            status="completed",
        )
        servicenow = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )

    service_now_path = Path(servicenow["workbookPath"])
    workbook = load_workbook(
        service_now_path,
        data_only=False,
        keep_links=False,
    )
    steps = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
    text = "\n".join(
        str(steps.cell(row, 2).value or "")
        for row in range(2, steps.max_row + 1)
    )
    hyperlinks = [
        cell.hyperlink.target
        for row in steps.iter_rows()
        for cell in row
        if cell.hyperlink is not None
    ]
    assert hyperlinks == []
    assert [steps.cell(row, 1).value for row in range(2, 5)] == [
        "Overview",
        "Entra handoff",
        "1",
    ]
    assert steps["B3"].value == (
        f"Entra handoff - Client ID: `{APP_CLIENT_ID}`  "
        "User Claim: `approved_employee_key`"
    )
    instruction = steps["B4"].value
    assert instruction.count("`approved_employee_key`") == 3
    assert (
        "User Claim: Auto-populated from the confirmed Entra administrator "
        "result: `approved_employee_key`"
    ) in instruction
    assert "`oid`" not in instruction
    assert "User ID" not in instruction
    workbook.close()


@pytest.mark.parametrize(
    "client_id",
    [APP_CLIENT_ID, snow.IDENTITY_WORKBOOK_PENDING_CLIENT_ID],
)
def test_entra_steps_cover_known_and_unknown_app_setup(
    client_id: str,
) -> None:
    steps = snow._identity_workbook_steps(
        stage="entra",
        context_values={
            "servicenow_instance": "dev123",
            "application_client_id": client_id,
        },
        missing_items=[
            ("APP", "application"),
            ("CLAIMS", "claims"),
            ("SCOPE", "scope"),
            ("PREAUTH", "preauth"),
            ("PERMISSIONS", "permissions"),
            ("CONSENT", "consent"),
        ],
    )
    text = "\n".join(action for _step, action in steps)
    labels = [label for label, _action in steps]

    assert labels == ["1"]
    assert text == snow._entra_admin_instruction()
    assert "\r" not in text and "\n" not in text
    assert "**New registration.**" in text
    assert "**Redirect URL:** Not needed" in text
    assert "**Token type** as **Access**" in text
    assert "*Aud* - for audience validation" in text
    assert "*Email* - addressable email for the user" in text
    assert "*UPN* - an identifier for the user" in text
    assert "Microsoft Graph permissions" in text
    assert snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL in text
    assert "must not be entered as the selected User Claim" in text


def test_entra_template_copy_keeps_one_universal_branch_graph(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    base_context = {
        "purpose": "DRAFT EXAMPLE",
        "target_agent": "DRAFT EXAMPLE - Employee Self-Service (IT)",
        "tenant_id": "11111111-1111-4111-8111-111111111111",
        "environment_id": "22222222-2222-4222-8222-222222222222",
        "servicenow_instance": "example-instance",
        "stage": "entra",
        "owner": "Microsoft Entra administrator",
        "admin_instruction": "Follow Steps.",
        "privacy_note": "Do not add secrets.",
    }
    scenarios = (
        (
            "new",
            snow.IDENTITY_WORKBOOK_PENDING_CLIENT_ID,
            [
                ("APP", "application"),
                ("CLAIMS", "claims"),
                ("SCOPE", "scope"),
                ("PREAUTH", "preauth"),
                ("PERMISSIONS", "permissions"),
                ("CONSENT", "consent"),
            ],
        ),
        (
            "repair",
            "33333333-3333-4333-8333-333333333333",
            [
                ("CLAIMS", "claims"),
                ("PERMISSIONS", "permissions"),
                ("CONSENT", "consent"),
            ],
        ),
        (
            "ready",
            "44444444-4444-4444-8444-444444444444",
            [("CLAIMS", "claims")],
        ),
    )
    results = {}
    for name, client_id, missing_items in scenarios:
        context_values = {
            **base_context,
            "application_client_id": client_id,
        }
        path = tmp_path / f"{name}.xlsx"
        snow._write_identity_workbook(
            path,
            context_values=context_values,
            stage="entra",
            candidates=[],
            steps=[],
            missing_items=missing_items,
            identity_requirement={
                "label": "Email",
                "category": "standard",
            },
        )
        workbook = load_workbook(path, data_only=False, keep_links=False)
        sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
        results[name] = {
            "instruction": sheet["A2"].value,
            "client_id": sheet["B2"].value,
            "user_claim": sheet["C2"].value,
            "hyperlink": sheet["A2"].hyperlink.target,
            "shape": (sheet.max_row, sheet.max_column),
        }
        workbook.close()

    for result in results.values():
        text = result["instruction"]
        assert result["shape"] == (2, 3)
        assert text == snow._entra_admin_instruction()
        assert "\r" not in text and "\n" not in text
        assert result["hyperlink"] == snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL
        assert result["user_claim"] in {None, ""}
        assert "*Aud* - for audience validation" in text
        assert "*Email* - addressable email for the user" in text
        assert "*UPN* - an identifier for the user" in text
        assert "must not be entered as the selected User Claim" in text
    assert results["new"]["client_id"] in {None, ""}
    assert results["repair"]["client_id"] == scenarios[1][1]
    assert results["ready"]["client_id"] == scenarios[2][1]


def test_retired_graph_rows_do_not_drive_entra_handoff(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state["phases"] = {
        "entra-registration": {
            "status": "in-progress",
            "checkpointResults": {
                "SN-DA-HRSD-ENTRA-APP-001": "Passed",
                "SN-DA-HRSD-ENTRA-CLAIMS-001": "Passed",
                "SN-DA-HRSD-ENTRA-SCOPE-001": "Passed",
                "SN-DA-HRSD-ENTRA-PREAUTH-001": "NotConfigured",
                "SN-DA-HRSD-ENTRA-PERMISSIONS-001": "NotConfigured",
                "SN-DA-HRSD-ENTRA-CONSENT-001": "NotConfigured",
            },
        }
    }
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)

    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    question_object = handoff["question"]
    instruction = handoff["instruction"]

    assert handoff["handoffMode"] == "delta"
    assert handoff["instructionTitle"] == (
        "Step 1 of 2 — Entra: list employee identity options"
    )
    assert question_object == {
        "header": (
            "Please enter the returned Entra admin handoff workbook URL."
        ),
        "question": (
            "Please enter the returned Entra admin handoff workbook URL."
        ),
        "allowFreeformInput": True,
    }
    assert instruction.startswith("**Purpose**\n")
    assert str(Path(handoff["workbookPath"])) in instruction
    headings = ("**Purpose**", "**Owner**", "**Known target**", "**Action**")
    positions = [instruction.index(heading) for heading in headings]
    assert positions == sorted(positions)
    assert "**Expected result**" not in instruction
    assert "Completed" not in instruction
    assert "Not yet" not in instruction
    assert "ServiceNow connector pre-authorization" not in instruction
    assert "Graph permissions" not in instruction
    assert "tenant-wide admin consent" not in instruction
    assert "This workbook does not configure or prove those checks" not in instruction
    assert "Separate checks must pass on their own evidence" not in instruction
    assert "No separate readiness item is included" not in instruction
    assert "4. For XLSX, keep the protected `Entra Handoff`" in instruction
    assert "5. " not in instruction
    assert "token_claim_name" not in instruction
    assert handoff["workbookUri"] in instruction
    workbook = load_workbook(Path(handoff["workbookPath"]))
    entra_instruction = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET][
        "A2"
    ].value
    assert entra_instruction == snow._entra_admin_instruction()
    assert "**New registration.**" in entra_instruction
    assert "**Token configuration**" in entra_instruction
    assert "Microsoft Graph permissions" in entra_instruction
    assert snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL in entra_instruction
    assert "file://" not in entra_instruction
    workbook.close()


def test_entra_workbook_roundtrip_records_one_selected_claim(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))

    receipt = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    app = snow.record_admin_phase(
        context,
        phase="entra-registration",
        status="completed",
        client_id=APP_CLIENT_ID,
    )

    assert receipt["candidates"] == ["email"]
    assert receipt["tokenClaimEmissionConfirmed"] is True
    assert receipt["kind"] == "administrator-confirmed-claim-candidates"
    assert app["evidence"]["tokenClaimCandidates"] == receipt["candidates"]
    assert app["evidence"]["candidateSetSha256"] == receipt["candidateSetSha256"]
    assert "entraSourceAttribute" not in json.dumps(receipt)
    assert "serviceNowUserField" not in json.dumps(receipt)


@pytest.mark.parametrize(
    "returned_path",
    (
        "Completed",
        "https:/contoso.sharepoint.com/returned.xlsx",
        "missing-return.xlsx",
    ),
)
def test_entra_workbook_path_submission_must_be_local_and_existing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    returned_path: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    before = state_path.read_bytes()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="does not exist or is not a local file",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=Path(returned_path),
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )

    assert state_path.read_bytes() == before


@pytest.mark.parametrize(
    ("claim", "message"),
    [
        ("aud", "aud audience claim"),
        ("", "explicit non-secret"),
        ("not a claim", "explicit non-secret"),
    ],
)
def test_entra_workbook_rejects_invalid_candidate_sets(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    claim: str,
    message: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["C2"] = claim
    workbook.save(path)
    workbook.close()

    with pytest.raises(snow.ServiceNowConnectError, match=message):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )


def test_entra_workbook_rejects_prefilled_user_field_formula_and_marker_tamper(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)

    for mutation, message in (
        ("instruction", "instructions changed"),
        ("formula", "must not contain formulas"),
        ("ownership", "editable ownership"),
        ("metadata", "metadata keys changed"),
    ):
        handoff = snow.identity_mapping_handoff(
            context,
            stage="entra",
            client_id=APP_CLIENT_ID,
        )
        path = Path(handoff["workbookPath"])
        workbook = load_workbook(path)
        sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
        metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
        sheet["C2"] = "email"
        if mutation == "instruction":
            sheet["A2"] = "Changed instructions"
        elif mutation == "formula":
            sheet["C2"] = "=A1"
        elif mutation == "ownership":
            sheet["C2"].protection = Protection(locked=True)
        else:
            metadata["A2"] = "changed"
        workbook.save(path)
        workbook.close()
        with pytest.raises(snow.ServiceNowConnectError, match=message):
            snow.record_identity_mapping(
                context,
                stage="entra",
                workbook_path=path,
                client_id=APP_CLIENT_ID,
                confirmed=True,
            )
        path.unlink()


def test_entra_import_rejects_removed_sheet_protection(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET].protection.sheet = False
    workbook.save(path)
    workbook.close()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="protection changed",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )


def test_identity_workbook_rejects_external_links(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    before = _lifecycle_path(tmp_path).read_bytes()
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="external links or macros",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )
    assert _lifecycle_path(tmp_path).read_bytes() == before


def test_servicenow_workbook_freezes_candidates_and_selects_exactly_one_pair(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    servicenow_template = _author_servicenow_template(
        tmp_path / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("employee_id",))
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    snow.record_admin_phase(
        context,
        phase="entra-registration",
        status="completed",
        client_id=APP_CLIENT_ID,
    )

    handoff = snow.identity_mapping_handoff(context, stage="servicenow")
    path = Path(handoff["workbookPath"])
    workbook = load_workbook(path)
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    first = snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW
    assert handoff["candidates"] == ["employee_id"]
    assert workbook.sheetnames == ["Steps", "Worksheet", "_metadata"]
    assert workbook.active.title == "Steps"
    steps = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
    assert [steps.cell(row, 1).value for row in range(2, 5)] == [
        "Overview",
        "Entra handoff",
        "1",
    ]
    assert steps["B3"].value == (
        f"Entra handoff - Client ID: `{APP_CLIENT_ID}`  "
        "User Claim: `employee_id`"
    )
    assert steps["B4"].value.count("`employee_id`") == 3
    assert sheet["B6"].value == APP_CLIENT_ID
    assert sheet["B6"].protection.locked is True
    assert sheet.cell(first, 1).value == "employee_id"
    assert sheet.cell(first, 1).fill.fgColor.rgb.endswith("E7E6E6")
    assert sheet.cell(first, 2).fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet.protection.sheet is True
    assert sheet.cell(first, 1).protection.locked is True
    assert sheet.cell(first, 2).protection.locked is False
    assert metadata["C7"].value == APP_CLIENT_ID
    assert metadata["D12"].value == "employee_id"
    workbook.close()
    _fill_servicenow_workbook(
        path,
        selected_claim="employee_id",
        user_field="employee_number",
    )

    mapping = snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=path,
        client_id=None,
        confirmed=True,
    )
    oidc = snow.record_admin_phase(
        context,
        phase="servicenow-oidc",
        status="completed",
    )

    assert mapping["tokenClaimName"] == "employee_id"
    assert mapping["serviceNowUserField"] == "employee_number"
    assert oidc["evidence"]["tokenClaimName"] == "employee_id"
    assert oidc["evidence"]["serviceNowUserField"] == "employee_number"


def test_servicenow_handoff_autopopulates_bound_client_and_claim(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    servicenow_template = _author_servicenow_template(
        tmp_path / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    confirmed_claim = "custom_employee_id"
    _fill_entra_workbook(entra_path, (confirmed_claim,))
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )

    handoff = snow.identity_mapping_handoff(context, stage="servicenow")
    path = Path(handoff["workbookPath"])
    workbook = load_workbook(path, data_only=False, keep_links=False)
    steps = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
    worksheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    handoff_summary = steps["B3"].value
    instruction = steps["B4"].value

    assert steps.max_row == 4
    assert str(steps.freeze_panes) == "A4"
    assert steps["A2"].value == "Overview"
    assert steps["B2"].value == "ServiceNow administrator handoff"
    assert steps["A3"].value == "Entra handoff"
    assert handoff_summary == (
        f"Entra handoff - Client ID: `{APP_CLIENT_ID}`  "
        f"User Claim: `{confirmed_claim}`"
    )
    assert instruction.count(f"`{confirmed_claim}`") == 3
    concrete_prompt = (
        "Ensure the user claim is the same as that configured by Entra "
        f"(Entra-confirmed User Claim: `{confirmed_claim}`)."
    )
    assert instruction.count(concrete_prompt) == 2
    assert (
        "User Claim: Auto-populated from the confirmed Entra administrator "
        f"result: `{confirmed_claim}`"
    ) in instruction
    assert f"`{APP_CLIENT_ID}`" in instruction
    assert "\r" not in instruction
    assert "\n" not in instruction
    assert "<br>" not in instruction.casefold()
    for fixed_default in ("`oid`", "`upn`", "User ID", "Employee number"):
        assert fixed_default not in instruction

    assert worksheet["B6"].value == APP_CLIENT_ID
    assert worksheet["B6"].protection.locked is True
    assert worksheet["A13"].value == confirmed_claim
    assert worksheet["A13"].protection.locked is True
    assert worksheet["B13"].value in {None, ""}
    assert worksheet["B13"].protection.locked is False
    assert metadata["C7"].value == APP_CLIENT_ID
    assert metadata["D12"].value == confirmed_claim
    assert metadata["G12"].value == "yes"
    workbook.close()


def test_multiple_unselected_claims_block_new_handoff_but_legacy_import_works(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    expected_identity = snow._required_context_identity(
        context,
        snow._load_lifecycle_state(context),
    )
    entra_context = snow._identity_workbook_context_values(
        context,
        expected_identity,
        application_client_id=APP_CLIENT_ID,
        stage="entra",
    )
    entra_path = tmp_path / "legacy-entra-handoff.xlsx"
    snow._author_identity_workbook(
        entra_path,
        context_values=entra_context,
        stage="entra",
        candidates=[],
        steps=[("1", "Legacy Entra administrator instructions.")],
    )
    candidates = ("email", "custom_employee_id")
    _fill_entra_workbook(entra_path, candidates)
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="not ready.*exactly one selected User Claim",
    ):
        snow.identity_mapping_handoff(context, stage="servicenow")
    assert not snow._identity_workbook_path(
        context,
        stage="servicenow",
    ).exists()

    expected_identity = snow._required_context_identity(
        context,
        snow._load_lifecycle_state(context),
    )
    context_values = snow._identity_workbook_context_values(
        context,
        expected_identity,
        application_client_id=APP_CLIENT_ID,
        stage="servicenow",
    )
    legacy_path = tmp_path / "legacy-servicenow-handoff.xlsx"
    snow._author_identity_workbook(
        legacy_path,
        context_values=context_values,
        stage="servicenow",
        candidates=list(candidates),
        steps=snow._identity_workbook_steps(
            stage="servicenow",
            context_values=context_values,
            missing_items=[],
        ),
    )
    _fill_servicenow_workbook(
        legacy_path,
        selected_claim="custom_employee_id",
        user_field="u_employee_identity",
    )
    receipt = snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=legacy_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    assert receipt["tokenClaimName"] == "custom_employee_id"
    assert receipt["serviceNowUserField"] == "u_employee_identity"


@pytest.mark.parametrize(
    "invalid_evidence",
    ("missing-app-id", "missing-selected-claim", "wrong-binding"),
)
def test_servicenow_handoff_rejects_unready_or_wrong_bound_entra_evidence(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    invalid_evidence: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("custom_employee_id",))
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    state = snow._load_lifecycle_state(context)
    receipt = state["evidence"]["identityClaimCandidates"]
    if invalid_evidence == "missing-app-id":
        receipt["applicationClientId"] = ""
        receipt["binding"]["applicationClientId"] = ""
    elif invalid_evidence == "missing-selected-claim":
        receipt["candidates"] = []
    else:
        receipt["binding"]["tenantId"] = (
            "00000000-0000-4000-8000-000000000001"
        )
    receipt["candidateSetSha256"] = snow._identity_candidate_set_hash(
        application_client_id=receipt["applicationClientId"],
        candidates=list(receipt["candidates"]),
        binding=receipt["binding"],
    )
    snow._write_lifecycle_state(context, state)

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="not ready",
    ):
        snow.identity_mapping_handoff(context, stage="servicenow")
    assert not snow._identity_workbook_path(
        context,
        stage="servicenow",
    ).exists()


@pytest.mark.parametrize("entra_format", ("xlsx", "csv"))
@pytest.mark.parametrize("servicenow_format", ("xlsx", "csv"))
def test_identity_handoff_xlsx_csv_formats_flow_bound_values_end_to_end(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    entra_format: str,
    servicenow_format: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    claim_name = "custom_employee_id"
    user_field = "u_employee_identity"

    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(
        entra["workbookPath"]
        if entra_format == "xlsx"
        else entra["csvPath"]
    )
    if entra_format == "xlsx":
        _fill_entra_workbook(entra_path, (claim_name,))
    else:
        _fill_entra_csv(
            entra_path,
            client_id=APP_CLIENT_ID,
            user_claim=claim_name,
        )
    candidates = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    assert candidates["applicationClientId"] == APP_CLIENT_ID
    assert candidates["candidates"] == [claim_name]
    assert candidates["importedFileFormat"] == entra_format

    servicenow = snow.identity_mapping_handoff(
        context,
        stage="servicenow",
    )
    workbook = load_workbook(servicenow["workbookPath"])
    steps = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
    worksheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    summary = steps["B3"].value
    instruction = steps["B4"].value
    assert summary == (
        f"Entra handoff - Client ID: `{APP_CLIENT_ID}`  "
        f"User Claim: `{claim_name}`"
    )
    assert instruction.count(f"`{claim_name}`") == 3
    assert worksheet["B6"].value == APP_CLIENT_ID
    assert worksheet["A13"].value == claim_name
    assert worksheet["A13"].protection.locked is True
    assert worksheet["B13"].protection.locked is False
    assert metadata["C7"].value == APP_CLIENT_ID
    assert metadata["D12"].value == claim_name
    workbook.close()

    csv_headers, csv_rows = _read_csv(Path(servicenow["csvPath"]))
    assert csv_headers == list(snow.IDENTITY_SERVICENOW_CSV_HEADERS)
    assert [row["Record"] for row in csv_rows] == [
        "Entra handoff",
        "Steps",
    ]
    assert {row["Client ID"] for row in csv_rows} == {APP_CLIENT_ID}
    assert {row["User Claim"] for row in csv_rows} == {claim_name}
    assert csv_rows[0]["Content"] == summary
    assert csv_rows[1]["Content"] == instruction

    servicenow_path = Path(
        servicenow["workbookPath"]
        if servicenow_format == "xlsx"
        else servicenow["csvPath"]
    )
    if servicenow_format == "xlsx":
        _fill_servicenow_workbook(
            servicenow_path,
            selected_claim=claim_name,
            user_field=user_field,
        )
    else:
        _fill_servicenow_csv(
            servicenow_path,
            user_field=user_field,
        )
    mapping = snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=servicenow_path,
        client_id=None,
        confirmed=True,
    )
    assert mapping["tokenClaimName"] == claim_name
    assert mapping["serviceNowUserField"] == user_field
    assert mapping["importedFileFormat"] == servicenow_format


@pytest.mark.parametrize(
    "mutation",
    (
        "blank-client",
        "invalid-client",
        "blank-claim",
        "aud-claim",
        "malformed-claim",
        "duplicate-record",
        "wrong-binding",
    ),
)
def test_entra_csv_rejects_invalid_duplicate_or_wrong_bound_records(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    mutation: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["csvPath"])
    headers, rows = _read_csv(path)
    rows[0]["Client ID"] = APP_CLIENT_ID
    rows[0]["User Claim"] = "email"
    if mutation == "blank-client":
        rows[0]["Client ID"] = ""
    elif mutation == "invalid-client":
        rows[0]["Client ID"] = "not-a-guid"
    elif mutation == "blank-claim":
        rows[0]["User Claim"] = ""
    elif mutation == "aud-claim":
        rows[0]["User Claim"] = "aud"
    elif mutation == "malformed-claim":
        rows[0]["User Claim"] = "personal@example.com"
    elif mutation == "duplicate-record":
        rows.append(copy.deepcopy(rows[0]))
    elif mutation == "wrong-binding":
        rows[0]["Tenant ID"] = (
            "00000000-0000-4000-8000-000000000001"
        )
    _write_csv(path, headers, rows)
    before = state_path.read_bytes()

    with pytest.raises(snow.ServiceNowConnectError):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )
    assert state_path.read_bytes() == before


@pytest.mark.parametrize(
    "mutation",
    (
        "blank-user-field",
        "invalid-user-field",
        "conflicting-client",
        "conflicting-claim",
        "stale-candidate-hash",
        "duplicate-record",
        "changed-content",
    ),
)
def test_servicenow_csv_rejects_incomplete_conflicting_or_stale_records(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    mutation: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("email",))
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    handoff = snow.identity_mapping_handoff(
        context,
        stage="servicenow",
    )
    path = Path(handoff["csvPath"])
    headers, rows = _read_csv(path)
    rows[1]["ServiceNow User Field"] = "email"
    if mutation == "blank-user-field":
        rows[1]["ServiceNow User Field"] = ""
    elif mutation == "invalid-user-field":
        rows[1]["ServiceNow User Field"] = "User ID"
    elif mutation == "conflicting-client":
        rows[1]["Client ID"] = (
            "00000000-0000-4000-8000-000000000001"
        )
    elif mutation == "conflicting-claim":
        rows[1]["User Claim"] = "other_claim"
    elif mutation == "stale-candidate-hash":
        rows[1]["Candidate Set SHA256"] = "0" * 64
    elif mutation == "duplicate-record":
        rows.append(copy.deepcopy(rows[1]))
    elif mutation == "changed-content":
        rows[1]["Content"] = "Changed instructions"
    _write_csv(path, headers, rows)
    before = state_path.read_bytes()

    with pytest.raises(snow.ServiceNowConnectError):
        snow.record_identity_mapping(
            context,
            stage="servicenow",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )
    assert state_path.read_bytes() == before


@pytest.mark.parametrize("profile_key", ("hrsd", "itsm"))
def test_public_servicenow_handoff_emits_complete_task2_and_imports_mapping(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile_key: str,
) -> None:
    scenario_root = tmp_path / profile_key
    scenario_root.mkdir()
    monkeypatch.chdir(scenario_root)
    servicenow_template = _author_servicenow_template(
        scenario_root / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    profile = snow.PRODUCT_PROFILES[profile_key]
    slug = (
        "employee-self-service-hr"
        if profile_key == "hrsd"
        else "employee-self-service-it"
    )
    agent_id = AGENT_ID if profile_key == "hrsd" else OTHER_AGENT_ID
    context = _profile_context(
        profile_key,
        agent_id=agent_id,
        agent_slug=slug,
    )
    context["root"] = scenario_root
    components = (
        _components()
        if profile_key == "hrsd"
        else _itsm_components()
    )
    with snow.product_profile_context(profile_key):
        state = snow._state_base(context, components)
        state["adminSetup"]["preflight"]["instanceName"] = (
            f"example-{profile_key}"
        )
        state_path = snow._lifecycle_state_path(context)
        snow._write_json_atomic(state_path, state)
        _confirm_identity_requirement(
            context,
            choice="pay-employee-id",
        )
        entra = snow.identity_mapping_handoff(
            context,
            stage="entra",
            client_id=APP_CLIENT_ID,
        )
        entra_path = Path(entra["workbookPath"])
        _fill_entra_workbook(
            entra_path,
            ("custom_employee_id",),
        )
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=entra_path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )
        snow.record_admin_phase(
            context,
            phase="entra-registration",
            status="completed",
        )
        state_before_handoff = state_path.read_bytes()
        handoff = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )

    assert state_path.read_bytes() == state_before_handoff
    assert handoff["provider"] == profile.provider_key
    assert handoff["profile"] == profile_key
    assert handoff["applicationClientId"] == APP_CLIENT_ID
    assert handoff["candidates"] == ["custom_employee_id"]
    assert handoff["clientIdQuestion"] is None
    assert "options" not in handoff["question"]
    assert handoff["question"]["allowFreeformInput"] is True
    path = Path(handoff["workbookPath"])
    workbook = load_workbook(path, data_only=False, keep_links=False)
    steps = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    text = "\n".join(
        str(steps.cell(row, 2).value or "")
        for row in range(2, steps.max_row + 1)
    )
    metadata_url = (
        f"login.microsoftonline.com/{TENANT_ID}/"
        ".well-known/openid-configuration"
    )
    assert metadata_url in text
    assert APP_CLIENT_ID in text
    assert "<service-now-instance>" not in text
    assert "<customer-tenant-id>" not in text
    assert "<application-client-id>" not in text
    hyperlinks = [
        cell.hyperlink.target
        for row in steps.iter_rows()
        for cell in row
        if cell.hyperlink is not None
    ]
    assert hyperlinks == []
    first = snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW
    assert steps["B3"].value == (
        f"Entra handoff - Client ID: `{APP_CLIENT_ID}`  "
        "User Claim: `custom_employee_id`"
    )
    assert sheet.cell(first, 1).value == "custom_employee_id"
    assert sheet.cell(first, 1).protection.locked is True
    assert sheet.cell(first, 2).protection.locked is False
    assert sheet.cell(6, 2).value == APP_CLIENT_ID
    assert sheet.cell(6, 2).protection.locked is True
    workbook.close()

    _fill_servicenow_workbook(
        path,
        selected_claim="custom_employee_id",
        user_field="u_employee_number",
    )
    with snow.product_profile_context(profile_key):
        mapping = snow.record_identity_mapping(
            context,
            stage="servicenow",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )
        phase = snow.record_admin_phase(
            context,
            phase="servicenow-oidc",
            status="completed",
        )
    assert mapping["tokenClaimName"] == "custom_employee_id"
    assert mapping["serviceNowUserField"] == "u_employee_number"
    assert phase["evidence"]["tokenClaimName"] == "custom_employee_id"
    assert phase["evidence"]["serviceNowUserField"] == "u_employee_number"


@pytest.mark.parametrize(
    ("mutation", "message"),
    (
        ("confirmation", "Maker-confirmed"),
        ("binding", "active target"),
        ("hash", "hash"),
    ),
)
def test_servicenow_handoff_rejects_stale_candidate_receipt_before_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    mutation: str,
    message: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("email",))
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    persisted = json.loads(
        _lifecycle_path(tmp_path).read_text(encoding="utf-8")
    )
    receipt = persisted["evidence"]["identityClaimCandidates"]
    if mutation == "confirmation":
        receipt["entraAdministratorConfirmed"] = False
    elif mutation == "binding":
        receipt["binding"]["instanceName"] = "dev999"
    else:
        receipt["candidateSetSha256"] = "stale"
    snow._write_json_atomic(_lifecycle_path(tmp_path), persisted)
    state_before = _lifecycle_path(tmp_path).read_bytes()
    workbook_path = snow._identity_workbook_path(
        context,
        stage="servicenow",
    )

    assert not workbook_path.exists()
    with pytest.raises(snow.ServiceNowConnectError, match=message):
        snow.identity_mapping_handoff(context, stage="servicenow")

    assert not workbook_path.exists()
    assert _lifecycle_path(tmp_path).read_bytes() == state_before


def test_entra_handoff_ignores_invalid_candidate_receipt_client_id(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "reused",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
        },
    }
    state["evidence"]["identityClaimCandidates"] = {
        "kind": "administrator-confirmed-claim-candidates",
        "applicationClientId": "00000000-0000-4000-8000-000000006667",
        "candidates": ["email"],
        "candidateSetSha256": "stale",
        "entraAdministratorConfirmed": False,
        "tokenClaimEmissionConfirmed": True,
        "binding": {},
    }
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)

    handoff = snow.identity_mapping_handoff(context, stage="entra")

    assert handoff["applicationClientId"] == APP_CLIENT_ID
    assert APP_CLIENT_ID not in handoff["question"]["question"]


@pytest.mark.parametrize("profile_key", ("hrsd", "itsm"))
def test_entra_handoff_unknown_client_id_uses_same_workbook(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile_key: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    profile = snow.PRODUCT_PROFILES[profile_key]
    slug = (
        "employee-self-service-hr"
        if profile_key == "hrsd"
        else "employee-self-service-it"
    )
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": profile.schema_name,
            "workspace_slug": slug,
        },
        "active": {"slug": slug},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
        "snapshotPath": tmp_path / "components.json",
    }
    components = (
        _components()
        if profile_key == "hrsd"
        else _itsm_components()
    )
    with snow.product_profile_context(profile_key):
        state = snow._state_base(context, components)
        state["adminSetup"]["preflight"]["instanceName"] = "dev123"
        snow._write_json_atomic(
            snow._lifecycle_state_path(context),
            state,
        )
        _confirm_identity_requirement(context)
        handoff = snow.identity_mapping_handoff(context, stage="entra")

    assert handoff["applicationClientId"] is None
    assert "options" not in handoff["question"]
    assert handoff["question"]["allowFreeformInput"] is True
    assert handoff["clientIdQuestion"] is None
    workbook = load_workbook(Path(handoff["workbookPath"]))
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    client_cell = sheet["B2"]
    assert client_cell.value in {None, ""}
    assert client_cell.fill.fgColor.rgb.endswith("FFF2CC")
    assert client_cell.protection.locked is False
    assert sheet["C2"].fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet["C2"].protection.locked is False
    workbook.close()


@pytest.mark.parametrize("selected_count", (0, 2))
def test_servicenow_workbook_requires_exactly_one_selected_pair(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    selected_count: int,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    expected_identity = snow._required_context_identity(
        context,
        snow._load_lifecycle_state(context),
    )
    entra_context = snow._identity_workbook_context_values(
        context,
        expected_identity,
        application_client_id=APP_CLIENT_ID,
        stage="entra",
    )
    entra_path = tmp_path / "legacy-entra-candidates.xlsx"
    snow._author_identity_workbook(
        entra_path,
        context_values=entra_context,
        stage="entra",
        candidates=[],
        steps=[("1", "Legacy Entra administrator instructions.")],
    )
    _fill_entra_workbook(entra_path, ("email", "upn"))
    snow.record_identity_mapping(context, stage="entra", workbook_path=entra_path, client_id=APP_CLIENT_ID, confirmed=True)
    context_values = snow._identity_workbook_context_values(
        context,
        snow._required_context_identity(
            context,
            snow._load_lifecycle_state(context),
        ),
        application_client_id=APP_CLIENT_ID,
        stage="servicenow",
    )
    path = tmp_path / "legacy-servicenow-selection.xlsx"
    snow._author_identity_workbook(
        path,
        context_values=context_values,
        stage="servicenow",
        candidates=["email", "upn"],
        steps=snow._identity_workbook_steps(
            stage="servicenow",
            context_values=context_values,
            missing_items=[],
        ),
    )
    if selected_count:
        _fill_servicenow_workbook(path, selected_claim="email", user_field="email")
        if selected_count == 2:
            _fill_servicenow_workbook(path, selected_claim="upn", user_field="user_name")
    with pytest.raises(snow.ServiceNowConnectError, match="exactly one"):
        snow.record_identity_mapping(context, stage="servicenow", workbook_path=path, client_id=None, confirmed=True)


@pytest.mark.parametrize(
    "user_field",
    (
        "user.email",
        "User ID",
        "00000000-0000-4000-8000-000000001234",
    ),
)
def test_servicenow_workbook_rejects_invalid_user_field_without_state_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    user_field: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("email",))
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    handoff = snow.identity_mapping_handoff(context, stage="servicenow")
    path = Path(handoff["workbookPath"])
    _fill_servicenow_workbook(
        path,
        selected_claim="email",
        user_field=user_field,
    )
    before = state_path.read_bytes()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="explicit safe ServiceNow field name",
    ):
        snow.record_identity_mapping(
            context,
            stage="servicenow",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )

    assert state_path.read_bytes() == before


def test_servicenow_workbook_rejects_frozen_candidate_and_context_tamper(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    entra = snow.identity_mapping_handoff(context, stage="entra", client_id=APP_CLIENT_ID)
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("email",))
    snow.record_identity_mapping(context, stage="entra", workbook_path=entra_path, client_id=APP_CLIENT_ID, confirmed=True)

    for mutation, message in (
        ("claim", "candidate claims were changed"),
        ("context", "context changed"),
        ("marker", "editable ownership"),
        ("sheet-order", "sheet structure changed"),
        ("metadata-visible", "metadata sheet visibility changed"),
        ("extra-sheet", "sheet structure changed"),
    ):
        handoff = snow.identity_mapping_handoff(context, stage="servicenow")
        path = Path(handoff["workbookPath"])
        workbook = load_workbook(path)
        steps = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
        sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
        metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
        first = snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW
        metadata_first = 2 + len(snow.IDENTITY_WORKBOOK_CONTEXT_FIELDS)
        if mutation == "claim":
            sheet.cell(first, 1).value = "other"
        elif mutation == "context":
            sheet.cell(3, 2).value = "changed"
        elif mutation == "marker":
            metadata.cell(
                metadata_first,
                snow.IDENTITY_WORKBOOK_COLUMNS.index(
                    "servicenow_user_field_editable"
                )
                + 1,
            ).value = "no"
        elif mutation == "sheet-order":
            workbook._sheets = [sheet, steps, metadata]
        elif mutation == "metadata-visible":
            metadata.sheet_state = "visible"
        else:
            workbook.create_sheet("Extra")
        workbook.save(path)
        workbook.close()
        before = _lifecycle_path(tmp_path).read_bytes()
        with pytest.raises(snow.ServiceNowConnectError, match=message):
            snow.record_identity_mapping(context, stage="servicenow", workbook_path=path, client_id=None, confirmed=True)
        assert _lifecycle_path(tmp_path).read_bytes() == before
        path.unlink()


def test_workbook_handoff_preserves_maker_filled_stage_file(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    first = snow.identity_mapping_handoff(context, stage="entra", client_id=APP_CLIENT_ID)
    path = Path(first["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    before = path.read_bytes()

    resumed = snow.identity_mapping_handoff(context, stage="entra", client_id=APP_CLIENT_ID)

    assert resumed["workbookFileStatus"] == "preserved-existing"
    assert path.read_bytes() == before


def test_reimport_changed_candidates_invalidates_old_oidc_selection(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state["phases"] = {
        "entra-registration": {"status": "done", "checkpointResults": {}},
        "servicenow-oidc": {"status": "done", "checkpointResults": {}},
    }
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _record_test_identity_mapping(
        tmp_path,
        candidates=("email",),
        selected_claim="email",
        user_field="email",
    )
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    state["phases"]["entra-registration"]["checkpointResults"] = {
        "SN-DA-HRSD-ENTRA-CLAIMS-001": "NotConfigured"
    }
    state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] = {"status": "completed", "evidence": {"kind": "structured-admin-attestation", "recordedAt": "2026-10-01T00:00:00Z"}}
    _lifecycle_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    old = Path(snow.identity_mapping_handoff(context, stage="entra", client_id=APP_CLIENT_ID)["workbookPath"])
    old.unlink()
    fresh = snow.identity_mapping_handoff(context, stage="entra", client_id=APP_CLIENT_ID)
    path = Path(fresh["workbookPath"])
    _fill_entra_workbook(path, ("employee_id",))

    snow.record_identity_mapping(context, stage="entra", workbook_path=path, client_id=APP_CLIENT_ID, confirmed=True)
    changed = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))

    assert "identityMapping" not in changed["evidence"]
    assert changed["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] == {"status": "pending"}
    assert changed["phases"]["servicenow-oidc"]["status"] == "pending"


def test_entra_import_retires_legacy_mapping_without_fabricating_selection(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state["evidence"]["identityMapping"] = {
        "kind": "administrator-confirmed-identity-mapping",
        "applicationClientId": APP_CLIENT_ID,
        "entraSourceAttribute": "user.employeeid",
        "tokenClaimName": "employee_id",
    }
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("employee_id",))

    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    saved = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))

    assert "identityMapping" not in saved["evidence"]
    assert saved["migration"]["retiredIdentityMapping"]["entraSourceAttribute"] == (
        "user.employeeid"
    )
    assert saved["evidence"]["identityClaimCandidates"]["candidates"] == [
        "employee_id",
    ]


def test_admin_operation_records_only_non_secret_identity_fields(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    mapping = _record_test_identity_mapping(tmp_path)
    app = snow.record_admin_phase(_context(), phase="entra-registration", status="completed", client_id=APP_CLIENT_ID)
    oidc = snow.record_admin_phase(_context(), phase="servicenow-oidc", status="completed")
    assert app["evidence"]["tokenClaimCandidates"] == ["employee_id"]
    assert oidc["evidence"]["tokenClaimName"] == "employee_id"
    assert oidc["evidence"]["serviceNowUserField"] == "employee_number"
    assert mapping["serviceNowUserField"] == "employee_number"
    serialized = _lifecycle_path(tmp_path).read_text(encoding="utf-8").casefold()
    assert "secret" not in serialized
    assert "password" not in serialized
    assert "employee-value" not in serialized


def test_identity_workbook_cli_roundtrips_both_admin_stages(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    _write_foundation_context(tmp_path, components=_components())
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "reused",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
        },
    }
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)

    assert snow.main(
        [
            "identity-mapping-handoff",
            "--stage",
            "entra",
        ]
    ) == 0
    entra = json.loads(capsys.readouterr().out)
    assert entra["handoffMode"] == "full-setup"
    assert "options" not in entra["question"]
    assert entra["question"]["allowFreeformInput"] is True
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("employee_id",))
    assert snow.main(
        [
            "record-identity-mapping",
            "--stage",
            "entra",
            "--workbook-file",
            str(entra_path),
            "--yes",
        ]
    ) == 0
    candidates = json.loads(capsys.readouterr().out)
    assert candidates["candidates"] == ["employee_id"]

    assert snow.main(
        ["identity-mapping-handoff", "--stage", "servicenow"]
    ) == 0
    service_now = json.loads(capsys.readouterr().out)
    service_now_path = Path(service_now["workbookPath"])
    _fill_servicenow_workbook(
        service_now_path,
        selected_claim="employee_id",
        user_field="employee_number",
    )
    assert snow.main(
        [
            "record-identity-mapping",
            "--stage",
            "servicenow",
            "--workbook-file",
            str(service_now_path),
            "--yes",
        ]
    ) == 0
    mapping = json.loads(capsys.readouterr().out)

    assert mapping["tokenClaimName"] == "employee_id"
    assert mapping["serviceNowUserField"] == "employee_number"
    persisted = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    assert persisted["evidence"]["identityClaimCandidates"][
        "candidateSetSha256"
    ] == mapping["candidateSetSha256"]


def test_unknown_client_id_workbook_binds_supplied_id_through_both_stages(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    servicenow_template = _author_servicenow_template(
        tmp_path / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)

    entra = snow.identity_mapping_handoff(context, stage="entra")
    entra_path = Path(entra["workbookPath"])
    workbook = load_workbook(entra_path)
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    visible_row = 2 + [
        key for _label, key in snow.IDENTITY_WORKBOOK_VISIBLE_CONTEXT
    ].index("application_client_id")
    metadata_row = 2 + snow.IDENTITY_WORKBOOK_CONTEXT_FIELDS.index(
        "application_client_id"
    )
    entra_metadata_row = 2 + snow.IDENTITY_ENTRA_METADATA_KEYS.index(
        "application_client_id"
    )
    assert sheet["B2"].value in {None, ""}
    assert metadata.cell(entra_metadata_row, 2).value == (
        snow.IDENTITY_WORKBOOK_PENDING_CLIENT_ID
    )
    workbook.close()
    _fill_entra_workbook(entra_path, ("employee_id",))
    workbook = load_workbook(entra_path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = APP_CLIENT_ID
    workbook.active = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    workbook.save(entra_path)
    workbook.close()
    returned_bytes = entra_path.read_bytes()

    candidates = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    assert entra_path.read_bytes() == returned_bytes
    assert candidates["applicationClientId"] == APP_CLIENT_ID
    entra_phase = snow.record_admin_phase(
        context,
        phase="entra-registration",
        status="completed",
        client_id=APP_CLIENT_ID,
    )
    assert entra_phase["evidence"]["clientId"] == APP_CLIENT_ID

    service_now = snow.identity_mapping_handoff(
        context,
        stage="servicenow",
    )
    service_now_path = Path(service_now["workbookPath"])
    workbook = load_workbook(service_now_path)
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    assert sheet.cell(visible_row, 2).value == APP_CLIENT_ID
    assert metadata.cell(metadata_row, 3).value == APP_CLIENT_ID
    workbook.close()
    _fill_servicenow_workbook(
        service_now_path,
        selected_claim="employee_id",
        user_field="employee_number",
    )
    mapping = snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=service_now_path,
        client_id=None,
        confirmed=True,
    )
    assert mapping["applicationClientId"] == APP_CLIENT_ID
    assert mapping["tokenClaimName"] == "employee_id"
    assert mapping["serviceNowUserField"] == "employee_number"


@pytest.mark.parametrize("profile_key", ("hrsd", "itsm"))
def test_unknown_client_id_and_claims_import_from_same_workbook(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile_key: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    profile = snow.PRODUCT_PROFILES[profile_key]
    slug = (
        "employee-self-service-hr"
        if profile_key == "hrsd"
        else "employee-self-service-it"
    )
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": profile.schema_name,
            "workspace_slug": slug,
        },
        "active": {"slug": slug},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
        "snapshotPath": tmp_path / "components.json",
    }
    components = _components() if profile_key == "hrsd" else _itsm_components()
    with snow.product_profile_context(profile_key):
        state = snow._state_base(context, components)
        state["adminSetup"]["preflight"]["instanceName"] = "dev123"
        state_path = snow._lifecycle_state_path(context)
        snow._write_json_atomic(state_path, state)
        _confirm_identity_requirement(context)
        handoff = snow.identity_mapping_handoff(context, stage="entra")
        path = Path(handoff["workbookPath"])
        _fill_entra_workbook(path, ("employee_id",))
        workbook = load_workbook(path)
        workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = APP_CLIENT_ID
        workbook.save(path)
        workbook.close()
        returned_bytes = path.read_bytes()
        candidates = snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )
        phase = snow.record_admin_phase(
            context,
            phase="entra-registration",
            status="completed",
        )
        service_now = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )

    assert path.read_bytes() == returned_bytes
    assert candidates["applicationClientId"] == APP_CLIENT_ID
    assert candidates["candidates"] == ["employee_id"]
    assert phase["evidence"]["clientId"] == APP_CLIENT_ID
    assert service_now["clientIdQuestion"] is None
    workbook = load_workbook(Path(service_now["workbookPath"]))
    client_cell = workbook[snow.IDENTITY_WORKBOOK_SHEET].cell(6, 2)
    assert client_cell.value == APP_CLIENT_ID
    assert client_cell.fill.fgColor.rgb.endswith("E7E6E6")
    assert client_cell.protection.locked is True
    workbook.close()


@pytest.mark.parametrize("profile_key", ("hrsd", "itsm"))
def test_known_client_id_and_claims_import_without_client_id_argument(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile_key: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    profile = snow.PRODUCT_PROFILES[profile_key]
    slug = (
        "employee-self-service-hr"
        if profile_key == "hrsd"
        else "employee-self-service-it"
    )
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": profile.schema_name,
            "workspace_slug": slug,
        },
        "active": {"slug": slug},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
        "snapshotPath": tmp_path / "components.json",
    }
    components = _components() if profile_key == "hrsd" else _itsm_components()
    with snow.product_profile_context(profile_key):
        state = snow._state_base(context, components)
        state["adminSetup"]["preflight"]["instanceName"] = "dev123"
        state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
            "status": "reused",
            "evidence": {
                "kind": "structured-admin-attestation",
                "recordedAt": "2026-10-01T00:00:00Z",
                "clientId": APP_CLIENT_ID,
                "binding": {
                    "tenantId": TENANT_ID,
                    "environmentId": ENVIRONMENT_ID,
                    "instanceName": "dev123",
                    "provider": profile.provider_key,
                    "profile": profile_key,
                    "agentSlug": slug,
                    "agentId": AGENT_ID,
                    "applicationClientId": APP_CLIENT_ID,
                },
            },
        }
        state_path = snow._lifecycle_state_path(context)
        snow._write_json_atomic(state_path, state)
        _confirm_identity_requirement(context)
        handoff = snow.identity_mapping_handoff(context, stage="entra")
        path = Path(handoff["workbookPath"])
        _fill_entra_workbook(path, ("employee_id",))
        result = snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )

    assert result["applicationClientId"] == APP_CLIENT_ID
    assert result["candidates"] == ["employee_id"]
    saved = json.loads(state_path.read_text(encoding="utf-8"))
    assert saved["evidence"]["identityClaimCandidates"][
        "applicationClientId"
    ] == APP_CLIENT_ID


def test_unknown_client_id_workbook_reimports_after_admin_phase_without_argument(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(context, stage="entra")
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("employee_id",))
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = APP_CLIENT_ID
    workbook.save(path)
    workbook.close()

    first = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=path,
        client_id=None,
        confirmed=True,
    )
    phase = snow.record_admin_phase(
        context,
        phase="entra-registration",
        status="completed",
    )
    second = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=path,
        client_id=None,
        confirmed=True,
    )

    assert first["applicationClientId"] == APP_CLIENT_ID
    assert phase["evidence"]["clientId"] == APP_CLIENT_ID
    assert second["applicationClientId"] == APP_CLIENT_ID
    assert second["candidateSetSha256"] == first["candidateSetSha256"]


def test_known_client_id_workbook_mismatch_is_zero_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "reused",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
        },
    }
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(context, stage="entra")
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    other_client_id = "00000000-0000-4000-8000-000000009997"
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = other_client_id
    metadata_row = 2 + snow.IDENTITY_ENTRA_METADATA_KEYS.index(
        "application_client_id"
    )
    workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET].cell(
        metadata_row,
        2,
    ).value = other_client_id
    workbook.save(path)
    workbook.close()
    before = state_path.read_bytes()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="active lifecycle",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )

    assert state_path.read_bytes() == before


def test_generated_known_client_id_without_lifecycle_binding_needs_argument(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    before = state_path.read_bytes()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="not bound to the active lifecycle",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )

    assert state_path.read_bytes() == before
    result = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    assert result["applicationClientId"] == APP_CLIENT_ID


def test_known_lifecycle_does_not_replace_blank_pending_workbook_client_id(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    universal_template = _author_universal_entra_template(
        tmp_path / "universal-entra-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        universal_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(context, stage="entra")
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = ""
    workbook.save(path)
    workbook.close()
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "reused",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-01T00:00:00Z",
            "clientId": APP_CLIENT_ID,
        },
    }
    snow._write_json_atomic(state_path, state)
    before = state_path.read_bytes()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match=r"Enter the exact Application \(client\) ID",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )

    assert state_path.read_bytes() == before
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = APP_CLIENT_ID
    workbook.save(path)
    workbook.close()
    result = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=path,
        client_id=None,
        confirmed=True,
    )
    assert result["applicationClientId"] == APP_CLIENT_ID


@pytest.mark.parametrize(
    "workbook_client_id",
    ("", "not-a-guid"),
)
def test_same_workbook_client_id_errors_are_zero_write(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    workbook_client_id: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(context, stage="entra")
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    workbook = load_workbook(path)
    workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]["B2"] = workbook_client_id
    workbook.save(path)
    workbook.close()
    before = state_path.read_bytes()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Application",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=None,
            confirmed=True,
        )

    assert state_path.read_bytes() == before


def test_known_client_id_workbook_rejects_different_supplied_id(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="Client ID does not match",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id="00000000-0000-4000-8000-000000009997",
            confirmed=True,
        )


def test_pending_client_id_binding_still_rejects_cross_environment_workbook(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)
    handoff = snow.identity_mapping_handoff(context, stage="entra")
    path = Path(handoff["workbookPath"])
    _fill_entra_workbook(path, ("email",))
    workbook = load_workbook(path)
    metadata_row = 2 + snow.IDENTITY_ENTRA_METADATA_KEYS.index(
        "environment_id"
    )
    other_environment = "00000000-0000-4000-8000-000000009997"
    workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET].cell(
        metadata_row,
        2,
    ).value = other_environment
    workbook.save(path)
    workbook.close()

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="metadata does not match",
    ):
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )


def test_identity_workbook_import_accepts_worksheet_as_active_sheet(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    servicenow_template = _author_servicenow_template(
        tmp_path / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)

    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    _fill_entra_workbook(entra_path, ("email",))
    workbook = load_workbook(entra_path)
    workbook.active = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    workbook.save(entra_path)
    workbook.close()

    candidates = snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    assert candidates["candidates"] == ["email"]

    service_now = snow.identity_mapping_handoff(
        context,
        stage="servicenow",
    )
    service_now_path = Path(service_now["workbookPath"])
    _fill_servicenow_workbook(
        service_now_path,
        selected_claim="email",
        user_field="user_name",
    )
    workbook = load_workbook(service_now_path)
    workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]["B2"] = (
        "Updated administrator guidance."
    )
    workbook[snow.IDENTITY_WORKBOOK_SHEET]["A1"] = "Choose the matching field"
    workbook[snow.IDENTITY_WORKBOOK_SHEET]["A9"] = (
        "Complete the highlighted cells."
    )
    workbook[snow.IDENTITY_WORKBOOK_SHEET].sheet_view.zoomScale = 140
    workbook.active = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    workbook.save(service_now_path)
    workbook.close()

    mapping = snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=service_now_path,
        client_id=None,
        confirmed=True,
    )
    assert mapping["tokenClaimName"] == "email"
    assert mapping["serviceNowUserField"] == "user_name"


def test_checked_in_entra_workbook_has_expected_protection_and_editable_cells() -> None:
    root = Path(__file__).resolve().parents[2] / "solutions" / "ess-maker-skills"
    for stage, template in (
        ("entra", snow.IDENTITY_WORKBOOK_TEMPLATE_ENTRA),
    ):
        path = root / template
        workbook = load_workbook(path)
        sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
        metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
        assert workbook.sheetnames == ["Entra Handoff", "_metadata"]
        assert workbook.active.title == "Entra Handoff"
        assert tuple(
            sheet.cell(1, column).value for column in range(1, 4)
        ) == snow.IDENTITY_ENTRA_HANDOFF_HEADERS
        assert sheet.max_row == 2
        assert sheet.max_column == 3
        assert sheet["A2"].value == snow._entra_admin_instruction()
        assert "\r" not in sheet["A2"].value
        assert "\n" not in sheet["A2"].value
        assert sheet["A2"].hyperlink.target == (
            snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL
        )
        assert sheet["A2"].protection.locked is True
        assert sheet["B2"].protection.locked is False
        assert sheet["C2"].protection.locked is False
        assert sheet["B2"].fill.fgColor.rgb.endswith("FFF2CC")
        assert sheet["C2"].fill.fgColor.rgb.endswith("FFF2CC")
        assert metadata.sheet_state == "hidden"
        assert metadata.protection.sheet is True
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:C2"
        workbook.close()
        with zipfile.ZipFile(path) as archive:
            package_text = "\n".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
            )
        assert snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL in package_text
        for forbidden in (
            "file://",
            "C:\\",
            "C:/Users/",
            "/Users/",
            "/home/",
            "\\\\",
            "daphneshao",
            "EmployeeHub",
            "dev184242",
            "ac017a3f-5f05-499c-b3f2-c1c1043db7ce",
            "copilot-worktrees",
        ):
            assert forbidden not in package_text


def test_checked_in_entra_template_has_universal_branch_graph() -> None:
    path = (
        Path(__file__).resolve().parents[2]
        / "solutions"
        / "ess-maker-skills"
        / snow.IDENTITY_WORKBOOK_TEMPLATE_ENTRA
    )
    workbook = load_workbook(path, data_only=False, keep_links=False)
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
    text = sheet["A2"].value
    workbook.close()

    assert text == snow._entra_admin_instruction()
    assert "\r" not in text and "\n" not in text
    assert "**Applications** then **App registrations**" in text
    assert "**New registration.**" in text
    assert "**Name:**" in text
    assert "**Redirect URL:** Not needed" in text
    assert "**Token configuration**" in text
    assert "**Add optional claim**" in text
    assert "*Aud* - for audience validation" in text
    assert "*Email* - addressable email for the user" in text
    assert "*UPN* - an identifier for the user" in text
    assert "Microsoft Graph permissions" in text
    assert snow.SERVICENOW_NON_EMAIL_USER_MATCH_URL in text
    assert "must not be entered as the selected User Claim" in text
    assert metadata.sheet_state == "hidden"


def test_fresh_entra_generation_uses_checked_in_template_content(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["adminSetup"]["preflight"]["instanceName"] = "dev123"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    _confirm_identity_requirement(context)

    handoff = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )

    generated = Path(handoff["workbookPath"])
    assert handoff["workbookFileStatus"] == "generated"
    assert snow._identity_workbook_template_marker(generated) == (
        snow._identity_workbook_template_sha256("entra")
    )
    workbook = load_workbook(generated, data_only=False, keep_links=False)
    sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
    assert sheet["A2"].value == snow._entra_admin_instruction()
    assert sheet["B2"].value == APP_CLIENT_ID
    assert sheet["C2"].value in {None, ""}
    assert sheet.protection.sheet is True
    assert workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET].sheet_state == (
        "hidden"
    )
    workbook.close()


def test_runtime_entra_workbooks_keep_customer_context_isolated(
    tmp_path: Path,
) -> None:
    contexts = [
        {
            "purpose": "Collect claim names.",
            "target_agent": "Customer A agent",
            "tenant_id": "11111111-1111-4111-8111-111111111111",
            "environment_id": "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa",
            "servicenow_instance": "customer-a",
            "application_client_id": (
                "22222222-2222-4222-8222-222222222222"
            ),
            "stage": "entra",
            "owner": "Microsoft Entra administrator",
            "admin_instruction": "Follow Steps.",
            "privacy_note": "Claim names only.",
        },
        {
            "purpose": "Collect claim names.",
            "target_agent": "Customer B agent",
            "tenant_id": "33333333-3333-4333-8333-333333333333",
            "environment_id": "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb",
            "servicenow_instance": "customer-b",
            "application_client_id": (
                "44444444-4444-4444-8444-444444444444"
            ),
            "stage": "entra",
            "owner": "Microsoft Entra administrator",
            "admin_instruction": "Follow Steps.",
            "privacy_note": "Claim names only.",
        },
    ]
    paths = []
    for index, context_values in enumerate(contexts):
        path = tmp_path / f"customer-{index}.xlsx"
        steps = snow._identity_workbook_steps(
            stage="entra",
            context_values=context_values,
            missing_items=[("CLAIMS", "claims")],
        )
        snow._write_identity_workbook(
            path,
            context_values=context_values,
            stage="entra",
            candidates=[],
            steps=steps,
        )
        paths.append(path)

    for index, path in enumerate(paths):
        workbook = load_workbook(path, data_only=False, keep_links=False)
        sheet = workbook[snow.IDENTITY_ENTRA_HANDOFF_SHEET]
        metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
        own = contexts[index]
        other = contexts[1 - index]
        assert sheet["B2"].value == own["application_client_id"]
        assert sheet["C2"].value in {None, ""}
        metadata_text = " ".join(
            str(cell.value or "")
            for row in metadata.iter_rows()
            for cell in row
        )
        assert own["tenant_id"] in metadata_text
        assert own["servicenow_instance"] in metadata_text
        assert other["tenant_id"] not in metadata_text
        assert other["servicenow_instance"] not in metadata_text
        workbook.close()


def test_runtime_servicenow_workbooks_keep_customer_context_isolated(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    servicenow_template = _author_servicenow_template(
        tmp_path / "servicenow-template.xlsx"
    )
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_SERVICENOW",
        servicenow_template,
    )
    contexts = [
        {
            "purpose": "Choose the matching ServiceNow field.",
            "target_agent": "Customer A agent",
            "tenant_id": "11111111-1111-4111-8111-111111111111",
            "environment_id": "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa",
            "servicenow_instance": "customer-a",
            "application_client_id": (
                "22222222-2222-4222-8222-222222222222"
            ),
            "stage": "servicenow",
            "owner": "ServiceNow administrator",
            "admin_instruction": "Follow Steps.",
            "privacy_note": "Choose one field.",
        },
        {
            "purpose": "Choose the matching ServiceNow field.",
            "target_agent": "Customer B agent",
            "tenant_id": "33333333-3333-4333-8333-333333333333",
            "environment_id": "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb",
            "servicenow_instance": "customer-b",
            "application_client_id": (
                "44444444-4444-4444-8444-444444444444"
            ),
            "stage": "servicenow",
            "owner": "ServiceNow administrator",
            "admin_instruction": "Follow Steps.",
            "privacy_note": "Choose one field.",
        },
    ]
    paths = []
    for index, context_values in enumerate(contexts):
        context_values.update(
            {
                "provider": snow.PROVIDER_KEY,
                "profile": snow.PROFILE_KEY,
                "agent_slug": f"customer-{index}-agent",
                "agent_id": f"customer-{index}-agent-id",
                "confirmed_user_claim_name": (
                    f"customer_{index}_claim"
                ),
            }
        )
        path = tmp_path / f"servicenow-customer-{index}.xlsx"
        steps = snow._identity_workbook_steps(
            stage="servicenow",
            context_values=context_values,
            missing_items=[],
        )
        snow._write_identity_workbook(
            path,
            context_values=context_values,
            stage="servicenow",
            candidates=[f"customer_{index}_claim"],
            steps=steps,
        )
        paths.append(path)

    for index, path in enumerate(paths):
        workbook = load_workbook(path, data_only=False, keep_links=False)
        steps_sheet = workbook[snow.IDENTITY_WORKBOOK_STEPS_SHEET]
        sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
        metadata = workbook[snow.IDENTITY_WORKBOOK_METADATA_SHEET]
        own = contexts[index]
        other = contexts[1 - index]
        steps_text = "\n".join(
            str(steps_sheet.cell(row, 2).value or "")
            for row in range(2, steps_sheet.max_row + 1)
        )
        own_metadata_url = (
            f"login.microsoftonline.com/{own['tenant_id']}/"
            ".well-known/openid-configuration"
        )
        other_metadata_url = (
            f"login.microsoftonline.com/{other['tenant_id']}/"
            ".well-known/openid-configuration"
        )
        assert sheet.cell(2, 2).value == own["target_agent"]
        assert sheet.cell(5, 2).value == own["servicenow_instance"]
        assert sheet.cell(6, 2).value == own["application_client_id"]
        assert sheet.cell(
            snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW,
            1,
        ).value == f"customer_{index}_claim"
        metadata_text = " ".join(
            str(cell.value or "")
            for row in metadata.iter_rows()
            for cell in row
        )
        assert own["tenant_id"] in metadata_text
        assert own["servicenow_instance"] in metadata_text
        assert other["tenant_id"] not in metadata_text
        assert other["servicenow_instance"] not in metadata_text
        assert f"customer_{1 - index}_claim" not in metadata_text
        assert own_metadata_url in steps_text
        assert own["application_client_id"] in steps_text
        assert other_metadata_url not in steps_text
        assert other["application_client_id"] not in steps_text
        assert "CONFIRMED_" not in steps_text
        hyperlinks = [
            cell.hyperlink.target
            for row in steps_sheet.iter_rows()
            for cell in row
            if cell.hyperlink is not None
        ]
        assert hyperlinks == []
        workbook.close()


@pytest.mark.parametrize("mode", ("missing", "corrupt"))
def test_runtime_workbook_fails_closed_when_template_is_unavailable(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    mode: str,
) -> None:
    template = tmp_path / "template.xlsx"
    if mode == "corrupt":
        template.write_text("not an xlsx", encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "IDENTITY_WORKBOOK_TEMPLATE_ENTRA",
        template,
    )
    context_values = {
        "purpose": "Collect claims.",
        "target_agent": "Customer agent",
        "tenant_id": "11111111-1111-4111-8111-111111111111",
        "environment_id": "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa",
        "servicenow_instance": "customer-instance",
        "application_client_id": snow.IDENTITY_WORKBOOK_PENDING_CLIENT_ID,
        "stage": "entra",
        "owner": "Microsoft Entra administrator",
        "admin_instruction": "Follow Steps.",
        "privacy_note": "Claim names only.",
    }

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="template",
    ):
        snow._write_identity_workbook(
            tmp_path / "output.xlsx",
            context_values=context_values,
            stage="entra",
            candidates=[],
            steps=[],
            missing_items=[("CLAIMS", "claims")],
        )
    assert not (tmp_path / "output.xlsx").exists()


def test_v2_lifecycle_migrates_to_v8_without_obsolete_preflight_gates(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 2,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "attested": True,
                "phases": {
                    phase_id: {
                        "status": "done",
                        "actionApplied": True,
                        "checkpointResults": {"old": "Passed"},
                    }
                    for phase_id in (
                        "topics",
                        "credential",
                        "agent-connection",
                        "parameter-sharing",
                        "publish",
                        "test",
                    )
                },
                "evidence": {
                    "credential": {"connectionId": CONNECTION_ID},
                    "custom": {"mustSurvive": True},
                },
                "transactions": {
                    "topics": {"operation": {"status": "committed"}}
                },
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    state = snow._load_lifecycle_state(_context())
    snow._write_lifecycle_state(_context(), state)
    first_bytes = state_path.read_bytes()
    state_again = snow._load_lifecycle_state(_context())
    snow._write_lifecycle_state(_context(), state_again)

    assert state["schemaVersion"] == 8
    assert state["acceptedContractRevision"] == 1
    assert "reuseDecision" not in state["adminSetup"]["preflight"]
    assert set(state["adminSetup"]["phaseHandoffs"]) == {
        "preflight",
        "plugin-prerequisites",
        "entra-registration",
        "servicenow-oidc",
        "credential",
    }
    assert "operations" not in state["adminSetup"]
    assert state["phases"]["preflight"]["status"] == "in-progress"
    assert state["phases"]["credential"]["status"] == "in-progress"
    assert "actionApplied" not in state["phases"]["credential"]
    for phase_id in ("topics", "agent-connection", "test"):
        assert state["phases"][phase_id]["status"] == "pending"
    assert "parameter-sharing" not in state["phases"]
    assert state["retiredPhases"]["parameter-sharing"]["status"] == "pending"
    assert "publish" not in state["phases"]
    assert state["retiredPhases"]["publish"]["status"] == "pending"
    assert state["evidence"]["custom"]["mustSurvive"] is True
    assert state["transactions"]["topics"]["operation"]["status"] == "committed"
    assert state["migration"]["schemaV3"]["reuseApprovalInferred"] is False
    assert state["migration"]["schemaV4"][
        "removedGlobalReuseDecision"
    ] is True
    assert state["migration"]["schemaV4"][
        "retainedCompletedPreflight"
    ] is False
    assert state["migration"]["schemaV5"]["removedScenarioGate"] is True
    assert state["migration"]["schemaV7"][
        "retiredAutomaticPublishPhase"
    ] is True
    assert state["migration"]["schemaV8"][
        "retiredEntraCheckpointIds"
    ] == []
    assert first_bytes == state_path.read_bytes()


def test_v5_migration_retires_parameter_phase_without_fabricating_draft_test(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    components = _components(CONNECTION_ID)
    component_hash = snow._component_hash(components)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 5,
                "provider": snow.PROVIDER_KEY,
                "profile": snow.PROFILE_KEY,
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {
                    "parameter-sharing": {
                        "status": "done",
                        "checkpointResults": {
                            "SN-DA-HRSD-PARAMETER-SHARING-001": "Manual"
                        },
                    }
                },
                "evidence": {
                    "credential": {"connectionId": CONNECTION_ID},
                    "parameterSharing": {
                        "status": "not-exposed",
                        "makerAttested": True,
                    },
                    "publish": {
                        "status": "completed",
                        "componentHash": component_hash,
                        "publishedComponentHash": component_hash,
                    },
                    "test": {
                        "status": "completed",
                        "result": "pass",
                        "binding": {
                            "publishedComponentHash": component_hash,
                        },
                    },
                },
                "transactions": {"topics": {}},
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    migrated = snow._load_lifecycle_state(_context(), components)
    snow._write_lifecycle_state(_context(), migrated)
    first_bytes = state_path.read_bytes()
    migrated_again = snow._load_lifecycle_state(_context(), components)
    snow._write_lifecycle_state(_context(), migrated_again)

    assert "parameter-sharing" not in migrated["phases"]
    assert migrated["retiredPhases"]["parameter-sharing"]["status"] == "done"
    assert "publish" not in migrated["phases"]
    assert migrated["evidence"]["parameterSharing"]["status"] == "not-exposed"
    assert "draftSemanticHash" not in migrated["evidence"]["test"]["binding"]
    assert "testedDraftSemanticHash" not in migrated["evidence"]["publish"]
    assert first_bytes == state_path.read_bytes()


def test_v4_migration_removes_obsolete_gates_and_keeps_completed_discovery(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 4,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {
                    "preflight": {
                        "status": "done",
                        "actionApplied": True,
                        "checkpointResults": {
                            "SN-DA-HRSD-ADMIN-PREFLIGHT-001": "Passed"
                        },
                    }
                },
                "adminSetup": {
                    "schemaVersion": 2,
                    "scope": "hrsd",
                    "authMode": snow.AUTH_MODE,
                    "preflight": {
                        "scenario": "connected",
                        "instanceName": "dev123",
                        "reuseDecision": {
                            "decision": "configure-missing",
                            "discoveryFingerprint": "old",
                        },
                        "discovery": {
                            "source": "connectivity-readonly",
                            "observedAt": "2026-09-29T00:00:00Z",
                            "requiredPlugins": [],
                            "connectionCandidates": [],
                            "fingerprint": "old",
                        },
                    },
                    "phaseHandoffs": {
                        "preflight": {
                            "status": "completed",
                            "evidence": {
                                "kind": "maker-reuse-decision",
                                "decision": "configure-missing",
                                "discoveryFingerprint": "old",
                            },
                        },
                        "servicenow-oidc": {
                            "status": "completed",
                            "evidence": {
                                "kind": "structured-admin-attestation",
                                "recordedAt": "2026-09-29T00:00:00Z",
                                "claim": "upn",
                                "userField": "user_name",
                            },
                        },
                    },
                },
                "evidence": {"mustSurvive": True},
                "transactions": {"topics": {"operation": {"status": "committed"}}},
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    state = snow._load_lifecycle_state(_context())

    assert state["schemaVersion"] == 8
    assert state["adminSetup"]["schemaVersion"] == 4
    assert "scenario" not in state["adminSetup"]["preflight"]
    assert "reuseDecision" not in state["adminSetup"]["preflight"]
    assert "fingerprint" not in state["adminSetup"]["preflight"]["discovery"]
    assert state["adminSetup"]["phaseHandoffs"]["preflight"]["status"] == (
        "completed"
    )
    evidence = state["adminSetup"]["phaseHandoffs"]["preflight"]["evidence"]
    assert evidence["kind"] == "maker-confirmed-target"
    assert "decision" not in evidence
    assert "discoveryFingerprint" not in evidence
    assert state["phases"]["preflight"]["status"] == "done"
    assert state["evidence"]["mustSurvive"] is True
    assert state["transactions"]["topics"]["operation"]["status"] == "committed"
    assert state["migration"]["schemaV5"][
        "retainedCompletedPreflight"
    ] is True
    assert state["migration"]["schemaV8"][
        "retiredEntraCheckpointIds"
    ] == []
    oidc_evidence = state["adminSetup"]["phaseHandoffs"][
        "servicenow-oidc"
    ]["evidence"]
    assert oidc_evidence["oidcCapabilityConfirmed"] is True
    assert oidc_evidence["runbookCompleted"] is True
    assert oidc_evidence["mappingDetailsCollected"] is False
    assert "claim" not in oidc_evidence
    assert "userField" not in oidc_evidence
    assert state["migration"]["schemaV5"][
        "inferredOidcCapabilityFromCompletedHandoff"
    ] is True


@pytest.mark.parametrize(
    "evidence_update",
    [
        {"kind": "unrelated"},
        {"kind": None},
        {"recordedAt": ""},
        {"recordedAt": None},
        {"claim": ""},
        {"userField": ""},
    ],
)
def test_v4_migration_does_not_infer_oidc_capability_from_invalid_evidence(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    evidence_update: dict,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    evidence = {
        "kind": "structured-admin-attestation",
        "recordedAt": "2026-09-29T00:00:00Z",
        "claim": "upn",
        "userField": "user_name",
        **evidence_update,
    }
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 4,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {},
                "adminSetup": {
                    "schemaVersion": 2,
                    "scope": "hrsd",
                    "authMode": snow.AUTH_MODE,
                    "preflight": {"discovery": {}},
                    "phaseHandoffs": {
                        "servicenow-oidc": {
                            "status": "completed",
                            "evidence": evidence,
                        }
                    },
                },
                "evidence": {},
                "transactions": {"topics": {}},
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    state = snow._load_lifecycle_state(_context())
    migrated = state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"]

    assert migrated == {"status": "pending"}
    assert state["migration"]["retiredOidcMappingEvidence"][
        "qualifiedCompletion"
    ] is False
    assert state["migration"]["schemaV5"][
        "inferredOidcCapabilityFromCompletedHandoff"
    ] is False


def test_v4_migration_does_not_infer_oidc_capability_from_pending_handoff(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 4,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {},
                "adminSetup": {
                    "schemaVersion": 2,
                    "scope": "hrsd",
                    "authMode": snow.AUTH_MODE,
                    "preflight": {"discovery": {}},
                    "phaseHandoffs": {
                        "servicenow-oidc": {
                            "status": "pending",
                            "evidence": {
                                "kind": "structured-admin-attestation",
                                "recordedAt": "2026-09-29T00:00:00Z",
                                "claim": "upn",
                                "userField": "user_name",
                            },
                        }
                    },
                },
                "evidence": {},
                "transactions": {"topics": {}},
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    state = snow._load_lifecycle_state(_context())

    assert state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] == {
        "status": "pending"
    }


def test_v4_migration_removes_scenario_and_resets_incomplete_preflight_only(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 4,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {
                    "preflight": {
                        "status": "done",
                        "actionApplied": True,
                        "lastActionAt": "2026-09-29T00:00:00Z",
                        "checkpointResults": {"old": "Passed"},
                    },
                    "plugin-prerequisites": {
                        "status": "done",
                        "checkpointResults": {"keep": "Manual"},
                    },
                },
                "adminSetup": {
                    "schemaVersion": 2,
                    "scope": "hrsd",
                    "authMode": snow.AUTH_MODE,
                    "preflight": {
                        "scenario": "connected",
                        "reuseDecision": {"decision": "reuse-discovered"},
                        "discovery": {},
                    },
                    "phaseHandoffs": {},
                },
                "evidence": {},
                "transactions": {"topics": {}},
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    state = snow._load_lifecycle_state(_context())

    assert state["phases"]["preflight"]["status"] == "in-progress"
    assert state["phases"]["preflight"]["checkpointResults"] == {}
    assert "actionApplied" not in state["phases"]["preflight"]
    assert "lastActionAt" not in state["phases"]["preflight"]
    assert state["phases"]["plugin-prerequisites"]["status"] == "done"
    assert "scenario" not in state["adminSetup"]["preflight"]
    assert state["migration"]["schemaV5"][
        "retainedCompletedPreflight"
    ] is False


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("provider", "workday"),
        ("profile", "itsm"),
        ("agentSlug", "other-agent"),
        ("provider", None),
    ],
)
def test_v2_migration_rejects_canonical_identity_mismatch(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    field: str,
    value: str | None,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state = {
        "schemaVersion": 2,
        "provider": snow.PROVIDER_KEY,
        "profile": "hrsd",
        "agentSlug": AGENT_SLUG,
        "agentId": AGENT_ID,
        "environmentId": ENVIRONMENT_ID,
        "phases": {},
        "evidence": {"mustSurvive": True},
        "transactions": {"topics": {}},
        "migration": {},
    }
    if value is None:
        state.pop(field)
    else:
        state[field] = value
    original = json.dumps(state)
    state_path.write_text(original, encoding="utf-8")

    with pytest.raises(snow.ServiceNowConnectError, match=field):
        snow._load_lifecycle_state(_context())

    assert state_path.read_text(encoding="utf-8") == original


def test_future_lifecycle_schema_fails_without_rewrite(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    original = json.dumps(
        {
            "schemaVersion": snow.LIFECYCLE_SCHEMA_VERSION + 1,
            "provider": snow.PROVIDER_KEY,
            "agentSlug": AGENT_SLUG,
            "agentId": AGENT_ID,
            "environmentId": ENVIRONMENT_ID,
        }
    )
    state_path.write_text(original, encoding="utf-8")

    with pytest.raises(snow.ServiceNowConnectError, match="newer kit"):
        snow._load_lifecycle_state(_context())

    assert state_path.read_text(encoding="utf-8") == original


def test_prepare_manual_connection_returns_exact_maker_guidance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return _components()

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda *_args, **_kwargs: pytest.fail(
            "manual preparation must not call Connectivity APIs"
        ),
    )
    _seed_admin_credential_requirements(tmp_path)

    result = snow.prepare_manual_connection(
        _context(),
        instance_name=None,
        resource_uri=None,
        display_name=None,
    )

    assert result["creationMode"] == "manual"
    assert result["connection"]["instanceName"] == "dev123"
    assert result["connection"]["resourceUri"] == APP_CLIENT_ID
    assert any(
        "Connection settings" in instruction
        for instruction in result["instructions"]
    )
    instructions = result["instructions"]
    create_index = instructions.index(
        "Open the connection configuration and select Create new connection."
    )
    search_index = instructions.index(
        "Use the search box in the upper-right to search ServiceNow."
    )
    connector_index = instructions.index(
        "Select the Microsoft-published ServiceNow connector by choosing its + action."
    )
    auth_index = instructions.index("Choose Microsoft Entra ID User Login.")
    fields_index = instructions.index(
        "Enter the Instance Name and Resource URI shown in this output."
    )
    assert create_index < search_index < connector_index < auth_index < fields_index


def test_prepare_manual_connection_ignores_stale_reference_values(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    reference = components["connectionReferenceChanges"][0][
        "connectionReference"
    ]
    reference["sharedConnectionParameters"] = json.dumps(
        {
            "name": "entraIDUserLogin",
            "values": {
                "token:ResourceUri": {
                    "value": "00000000-0000-4000-8000-000000009999"
                },
                "token:InstanceName": {"value": "wrong-instance"},
            },
        }
    )

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )
    _seed_admin_credential_requirements(tmp_path)

    result = snow.prepare_manual_connection(
        _context(),
        instance_name=None,
        resource_uri=None,
        display_name=None,
    )

    assert result["connection"]["instanceName"] == "dev123"
    assert result["connection"]["resourceUri"] == APP_CLIENT_ID


def test_prepare_manual_connection_requests_missing_metadata(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    reference = components["connectionReferenceChanges"][0][
        "connectionReference"
    ]
    reference["connectionId"] = None
    reference["sharedConnectionParameters"] = json.dumps(
        {
            "name": "entraIDUserLogin",
            "values": {
                "token:ResourceUri": {"value": None},
                "token:InstanceName": {"value": None},
            },
        }
    )

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return components

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.prepare_manual_connection(
        {
            "agent": {
                "id": AGENT_ID,
                "schema_name": snow.HR_SCHEMA_NAME,
            },
            "environment": {
                "id": ENVIRONMENT_ID,
                "ring": "test",
            },
        },
        instance_name=None,
        resource_uri=None,
        display_name=None,
    )

    assert result["status"] == "input-required"
    assert result["missingFields"] == ["instanceName", "resourceUri"]
    assert not (
        tmp_path
        / ".local"
        / "connect"
        / "servicenow"
        / "agents"
        / AGENT_ID
        / "state.json"
    ).exists()
    assert not _lifecycle_path(tmp_path).exists()


def test_prepare_manual_connection_rejects_api_identifier_uri(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return _components()

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(snow.ServiceNowConnectError, match="not api://"):
        snow.prepare_manual_connection(
            _context(),
            instance_name="https://dev123.service-now.com",
            resource_uri=f"api://{APP_CLIENT_ID}",
            display_name=None,
        )


def test_phase_command_error_returns_customer_visible_recovery(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(snow, "load_context", lambda **_kwargs: _context())
    monkeypatch.setattr(
        snow,
        "prepare_manual_connection",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            snow.ServiceNowConnectError("Credential input is incomplete.")
        ),
    )

    assert snow.main(["create"]) == 1

    error = json.loads(capsys.readouterr().out)
    assert error["status"] == "error"
    recovery = error["details"]["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["owner"] == "Power Platform Maker/Admin"
    assert "dev123" in recovery["values"]
    assert APP_CLIENT_ID in recovery["values"]
    assert "/connections" in recovery["where"]
    assert "rerun only `create`" in recovery["recheck"].casefold()


def test_plugin_command_error_uses_servicenow_admin_recovery(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    monkeypatch.setattr(snow, "load_context", lambda **_kwargs: _context())
    monkeypatch.setattr(
        snow,
        "record_admin_phase",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            snow.ServiceNowConnectError("Plugin evidence is incomplete.")
        ),
    )

    assert (
        snow.main(
            [
                "record-admin-phase",
                "--phase",
                "plugin-prerequisites",
                "--status",
                "completed",
            ]
        )
        == 1
    )

    error = json.loads(capsys.readouterr().out)
    recovery = error["details"]["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["owner"] == "ServiceNow Admin"
    assert "com.sn_hr_core" in recovery["values"]
    assert "app-manager" in recovery["where"]


def test_resume_error_uses_current_credential_phase_recovery(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    _seed_admin_credential_requirements(tmp_path)
    path = _lifecycle_path(tmp_path)
    state = json.loads(path.read_text(encoding="utf-8"))
    contract = json.loads(
        (
            snow.SOLUTION_ROOT
            / "src"
            / "skills"
            / "connect"
            / "servicenow-da-hrsd"
            / "contract.json"
        ).read_text(encoding="utf-8")
    )
    state["phases"] = {
        phase["id"]: {"status": "pending", "checkpointResults": {}}
        for phase in contract["phases"]
    }
    before_credential = True
    for phase_id, phase in state["phases"].items():
        if phase_id == "credential":
            before_credential = False
            phase["status"] = "blocked"
        elif before_credential:
            phase["status"] = "done"
        else:
            phase["status"] = "pending"
    path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(snow, "load_context", lambda **_kwargs: _context())
    monkeypatch.setattr(
        snow,
        "resume_handoff",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            snow.ServiceNowConnectError("Resume state is inconsistent.")
        ),
    )

    assert snow.main(["resume-handoff"]) == 1

    error = json.loads(capsys.readouterr().out)
    recovery = error["details"]["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["owner"] == "Power Platform Maker/Admin"
    assert "/connections" in recovery["where"]
    assert "credential" in recovery["recheck"]


def test_record_test_attestation_persists_result(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context()
    components = _components(CONNECTION_ID)
    state_path = _lifecycle_path(tmp_path)
    state = _seed_passing_test(tmp_path, components)
    state["evidence"].pop("test")
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {"fetch_components": lambda self, _agent_id: components},
        )(),
    )

    result = snow.record_test_attestation(
        context,
        prompt_category="list-my-open-hr-cases",
        result="pass",
        failure_category=None,
    )

    state = json.loads(state_path.read_text(encoding="utf-8"))
    assert result["result"] == "pass"
    assert result["status"] == "completed"
    assert state["evidence"]["test"]["promptCategory"] == (
        "list-my-open-hr-cases"
    )
    assert "prompt" not in state["evidence"]["test"]
    assert "details" not in state["evidence"]["test"]
    assert state["evidence"]["test"]["binding"]["draftSemanticHash"] == (
        snow._draft_semantic_hash(components)
    )
    assert state["evidence"]["test"]["binding"]["connectionBindingHash"]
    assert "publishedComponentHash" not in state["evidence"]["test"]["binding"]


@pytest.mark.parametrize(
    ("profile", "prompt_category"),
    (
        ("hrsd", "list-my-open-hr-cases"),
        ("itsm", "list-my-open-it-tickets"),
    ),
)
def test_record_test_attestation_accepts_maker_fail_without_category(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
    prompt_category: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context()
    components = _components(CONNECTION_ID)
    if profile == "itsm":
        context = {
            **context,
            "agent": {
                **context["agent"],
                "schema_name": snow.IT_SCHEMA_NAME,
                "workspace_slug": "employee-self-service-it",
            },
            "active": {"slug": "employee-self-service-it"},
        }
        components = _itsm_components()
    with snow.product_profile_context(profile):
        state = snow._state_base(context, components)
        state["evidence"]["credential"] = {
            "connectionId": CONNECTION_ID.replace("-", "")
        }
        state["evidence"]["agentConnection"] = {
            "kind": "maker-attestation",
            "status": "completed",
            "connectionId": CONNECTION_ID.replace("-", ""),
            "physicalStatus": "Connected",
            "makerAttested": True,
            "recordedAt": "2026-10-01T00:00:00Z",
            "binding": {
                "draftSemanticHash": snow._draft_semantic_hash(components),
            },
        }
        state_path = snow._lifecycle_state_path(context)
        state_path.parent.mkdir(parents=True)
        state_path.write_text(json.dumps(state), encoding="utf-8")
        monkeypatch.setattr(
            snow,
            "_agentbuilder_client",
            lambda _context: type(
                "FakeAgentBuilder",
                (),
                {"fetch_components": lambda self, _agent_id: components},
            )(),
        )
        result = snow.record_test_attestation(
            context,
            prompt_category=prompt_category,
            result="fail",
            failure_category=None,
        )

    assert result["status"] == "failed"
    assert result["result"] == "fail"
    assert result["failureCategory"] is None


def test_record_test_attestation_keeps_legacy_failure_category_compatible(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    components = _components(CONNECTION_ID)
    state = _seed_passing_test(tmp_path, components)
    state["evidence"].pop("test")
    _lifecycle_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {"fetch_components": lambda self, _agent_id: components},
        )(),
    )

    result = snow.record_test_attestation(
        _context(),
        prompt_category="list-my-open-hr-cases",
        result="fail",
        failure_category="permission",
    )

    assert result["failureCategory"] == "permission"


def test_record_test_requires_current_agent_connection_attestation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    state = _seed_passing_test(tmp_path, components)
    state["evidence"].pop("test")
    state["evidence"].pop("agentConnection")
    _lifecycle_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {"fetch_components": lambda self, _agent_id: components},
        )(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="current Agent Connect observation",
    ):
        snow.record_test_attestation(
            _context(),
            prompt_category="list-my-open-hr-cases",
            result="pass",
            failure_category=None,
        )


def test_fresh_draft_test_bridges_trustworthy_existing_publish_receipt(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    component_hash = snow._component_hash(components)
    state = _seed_passing_test(
        tmp_path,
        components,
        publish={
            "kind": "maker-attestation",
            "status": "completed",
            "recordedAt": "2026-09-28T00:00:00Z",
            "completedAt": "2026-09-28T00:00:00Z",
            "componentHash": component_hash,
            "publishedComponentHash": component_hash,
        },
    )
    _lifecycle_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    loaded = snow._load_lifecycle_state(_context(), components)
    receipt = loaded["evidence"]["publish"]

    assert receipt["testedDraftSemanticHash"] == loaded["draftSemanticHash"]
    assert receipt["publishedSemanticHash"] == loaded["draftSemanticHash"]
    assert receipt["connectionBindingHash"] == loaded["connectionBindingHash"]


@pytest.mark.parametrize(
    ("prompt_category", "result", "failure_category", "message"),
    [
        ("free-form-prompt", "pass", None, "privacy-safe"),
        (
            "list-my-open-hr-cases",
            "pass",
            "permission",
            "cannot have a failure category",
        ),
        (
            "list-my-open-hr-cases",
            "fail",
            "not-supported",
            "failure category is unsupported",
        ),
    ],
)
def test_record_test_attestation_rejects_unbounded_or_inconsistent_evidence(
    prompt_category: str,
    result: str,
    failure_category: str | None,
    message: str,
) -> None:
    with pytest.raises(snow.ServiceNowConnectError, match=message):
        snow.record_test_attestation(
            _context(),
            prompt_category=prompt_category,
            result=result,
            failure_category=failure_category,
        )


def test_legacy_state_migration_is_idempotent_and_preserves_source(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    legacy_path = (
        tmp_path
        / ".local"
        / "connect"
        / "servicenow"
        / "agents"
        / AGENT_ID
        / "state.json"
    )
    legacy_path.parent.mkdir(parents=True)
    legacy = {
        "schemaVersion": 1,
        "agentId": AGENT_ID,
        "environmentId": ENVIRONMENT_ID,
        "topicEnablement": {
            "customerChoice": "keep-current",
            "before": {"total": 1, "active": 0, "inactive": 1},
        },
        "test": {
            "prompt": "Show my HR cases",
            "result": "pass",
        },
    }
    legacy_path.write_text(json.dumps(legacy), encoding="utf-8")
    components = _components()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: type(
            "FakeAgentBuilder",
            (),
            {"fetch_components": lambda self, _agent_id: components},
        )(),
    )

    first = snow.migrate_state(_context())
    first_bytes = _lifecycle_path(tmp_path).read_bytes()
    second = snow.migrate_state(_context())
    second_bytes = _lifecycle_path(tmp_path).read_bytes()
    state = json.loads(first_bytes)

    assert first["status"] == "migrated"
    assert second["migration"]["legacySha256"] == first["migration"][
        "legacySha256"
    ]
    assert first_bytes == second_bytes
    assert json.loads(legacy_path.read_text(encoding="utf-8")) == legacy
    assert (
        state["evidence"]["topics"]["customerChoice"]
        == "keep-current"
    )
    assert state["evidence"]["test"]["reconfirmRequired"] is True


def test_first_run_initialize_then_migrate_and_inspect_offline(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    components = _components(CONNECTION_ID)
    _write_foundation_context(tmp_path, components=components)
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "initialize-state",
            "--contract-revision",
            str(contract["contractRevision"]),
        ]
    ) == 0
    initialized = json.loads(capsys.readouterr().out)
    state_path = _lifecycle_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))

    assert initialized["status"] == "initialized"
    assert state["schemaVersion"] == snow.LIFECYCLE_SCHEMA_VERSION
    assert state["provider"] == snow.PROVIDER_KEY
    assert state["profile"] == snow.PROFILE_KEY
    assert state["agentSlug"] == AGENT_SLUG
    assert state["agentId"] == AGENT_ID
    assert state["environmentId"] == ENVIRONMENT_ID
    assert state["attested"] is True
    assert state["acceptedContractRevision"] == contract["contractRevision"]
    expected_phase_ids = [
        phase["id"] for phase in contract["phases"]
    ]
    assert initialized["phaseIds"] == expected_phase_ids
    assert set(state["phases"]) == set(expected_phase_ids)
    assert all(
        phase == {
            "status": "pending",
            "checkpointResults": {},
            "checkpointAcknowledgements": {},
        }
        for phase in state["phases"].values()
    )

    assert snow.main(["migrate-state"]) == 0
    migrated = json.loads(capsys.readouterr().out)
    assert migrated["status"] == "migrated"

    assert snow.main(["inspect", "--offline"]) == 0
    inspected = json.loads(capsys.readouterr().out)
    assert inspected["mode"] == "offline"
    assert inspected["agentId"] == AGENT_ID
    assert inspected["components"]["serviceNowTopicCount"] == 1


def test_initialize_emits_invoked_only_after_state_persists(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    _write_foundation_context(tmp_path, components=components)
    monkeypatch.chdir(tmp_path)
    context = snow.load_context()
    context["root"] = tmp_path
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    captured: list[dict] = []

    def capture(_state: dict, event: dict) -> None:
        assert snow._lifecycle_state_path(context).is_file()
        captured.append(copy.deepcopy(event))

    monkeypatch.setattr(snow, "emit_servicenow_lifecycle_event", capture)

    snow.initialize_lifecycle_state(
        context,
        contract_revision=contract["contractRevision"],
    )

    assert [event["event"] for event in captured] == ["invoked"]
    saved = json.loads(
        snow._lifecycle_state_path(context).read_text(encoding="utf-8")
    )
    assert saved["telemetry"]["backfillSuppressed"] is False
    assert saved["telemetry"]["correlationId"] == captured[0][
        "correlationId"
    ]


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_persisted_phase_transitions_emit_bounded_service_now_events(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    product = snow.PRODUCT_PROFILES[profile]
    context = _profile_context(
        profile,
        agent_id=AGENT_ID if profile == "hrsd" else OTHER_AGENT_ID,
        agent_slug=(
            "employee-self-service-hr"
            if profile == "hrsd"
            else "employee-self-service-it"
        ),
    )
    context["root"] = tmp_path
    contract = json.loads(
        (
            Path(__file__).resolve().parents[2]
            / "solutions"
            / "ess-maker-skills"
            / "src"
            / "skills"
            / "connect"
            / product.provider_key
            / "contract.json"
        ).read_text(encoding="utf-8")
    )
    contract_path = tmp_path / product.contract_path
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    captured: list[dict] = []
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, event: captured.append(copy.deepcopy(event)),
    )
    with snow.product_profile_context(profile):
        context, completed_state = _seed_complete_dev_state(
            tmp_path,
            profile=profile,
        )
        snow._lifecycle_state_path(context).unlink()
        state = copy.deepcopy(completed_state)
        for phase in state["phases"].values():
            phase["status"] = "pending"
            phase["checkpointResults"] = {}
        snow._write_lifecycle_state(context, state)
        captured.clear()
        state["phases"]["credential"]["status"] = "in-progress"
        snow._write_lifecycle_state(context, state)
        state["phases"]["credential"]["status"] = "blocked"
        state["phases"]["credential"]["checkpointResults"] = {
            f"SN-DA-{profile.upper()}-CREDENTIAL-001": "Failed"
        }
        snow._write_lifecycle_state(context, state)
        state["phases"]["credential"]["status"] = "in-progress"
        snow._write_lifecycle_state(context, state)
        before_noop = len(captured)
        snow._write_lifecycle_state(context, state)
        captured_before_success = len(captured)
        telemetry_projection = copy.deepcopy(state["telemetry"])
        state = copy.deepcopy(completed_state)
        state["telemetry"] = telemetry_projection
        snow._write_lifecycle_state(context, state)

    transition_events = captured[:captured_before_success]
    assert [event["event"] for event in transition_events] == [
        "phase-started",
        "phase-paused",
        "blocked",
        "phase-resumed",
    ]
    assert transition_events[0]["phase"] == "connections"
    assert transition_events[2]["blockerCategory"] == "connection"
    assert transition_events[2]["remediationId"] == (
        f"SN-DA-{profile.upper()}-CREDENTIAL-001"
    )
    assert transition_events[3]["retryCount"] == 1
    assert transition_events[3]["resumeCount"] == 1
    assert captured_before_success == before_noop
    success_events = captured[captured_before_success:]
    assert success_events[-1]["event"] == "completed"
    assert success_events[-1]["outcome"] == "success"
    assert sum(
        event["event"] == "phase-completed"
        for event in success_events
    ) == len(contract["phases"])


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_completed_requires_current_manual_evidence(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context, completed_state = _seed_complete_dev_state(
        tmp_path,
        profile=profile,
    )
    contract = snow._load_json(
        Path(context["root"])
        / snow.PRODUCT_PROFILES[profile].contract_path
    )
    pending = copy.deepcopy(completed_state)
    for phase in pending["phases"].values():
        phase["status"] = "pending"
        phase["checkpointResults"] = {}
    captured: list[dict] = []
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, event: captured.append(copy.deepcopy(event)),
    )
    with snow.product_profile_context(profile):
        snow._lifecycle_state_path(context).unlink()
        snow._write_lifecycle_state(context, pending)
        captured.clear()
        incomplete = copy.deepcopy(completed_state)
        incomplete["telemetry"] = copy.deepcopy(pending["telemetry"])
        incomplete["evidence"].pop("test")
        snow._write_lifecycle_state(context, incomplete)

    assert sum(
        event["event"] == "phase-completed"
        for event in captured
    ) == len(contract["phases"])
    assert all(event["event"] != "completed" for event in captured)


def test_existing_state_telemetry_baseline_does_not_backfill_history(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    state = snow._state_base(_context_at(tmp_path), _components())
    state["phases"] = {
        "preflight": {
            "status": "done",
            "checkpointResults": {
                "SN-DA-HRSD-ADMIN-PREFLIGHT-001": "Passed"
            },
        }
    }
    path = _lifecycle_path(tmp_path)
    path.parent.mkdir(parents=True)
    path.write_text(json.dumps(state), encoding="utf-8")
    captured: list[dict] = []
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, event: captured.append(copy.deepcopy(event)),
    )

    snow._write_lifecycle_state(_context_at(tmp_path), state)

    assert captured == []
    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved["telemetry"]["backfillSuppressed"] is True
    assert saved["telemetry"]["projectedPhaseStatuses"] == {
        "preflight": "done"
    }


def test_manual_evidence_write_without_phase_transition_emits_nothing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["phases"] = {
        "servicenow-oidc": {
            "status": "in-progress",
            "checkpointResults": {
                "SN-DA-HRSD-OIDC-001": "Manual"
            },
        }
    }
    snow._write_lifecycle_state(context, state)
    captured: list[dict] = []
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, event: captured.append(copy.deepcopy(event)),
    )
    state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] = {
        "status": "completed",
        "evidence": {
            "kind": "structured-admin-attestation",
            "recordedAt": "2026-10-23T23:20:00Z",
        },
    }

    snow._write_lifecycle_state(context, state)

    assert captured == []


def test_persist_failure_emits_nothing_and_telemetry_failure_is_fail_open(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    state = snow._state_base(context, _components())
    state["phases"] = {
        "preflight": {"status": "pending", "checkpointResults": {}}
    }
    snow._write_lifecycle_state(context, state)
    captured: list[dict] = []
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda _state, event: captured.append(copy.deepcopy(event)),
    )
    state["phases"]["preflight"]["status"] = "in-progress"
    original_write = snow._write_json_atomic
    monkeypatch.setattr(
        snow,
        "_write_json_atomic",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            OSError("disk unavailable")
        ),
    )

    with pytest.raises(OSError, match="disk unavailable"):
        snow._write_lifecycle_state(context, state)
    assert captured == []

    monkeypatch.setattr(snow, "_write_json_atomic", original_write)
    loaded = snow._load_json(_lifecycle_path(tmp_path))
    loaded["phases"]["preflight"]["status"] = "in-progress"
    monkeypatch.setattr(
        snow,
        "emit_servicenow_lifecycle_event",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("telemetry unavailable")
        ),
    )
    snow._write_lifecycle_state(context, loaded)
    persisted = snow._load_json(_lifecycle_path(tmp_path))
    assert persisted["phases"]["preflight"]["status"] == "in-progress"


def test_projection_prepare_failure_warns_and_persists_business_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state = snow._state_base(context, _components())
    state["phases"] = {
        phase["id"]: {"status": "pending", "checkpointResults": {}}
        for phase in contract["phases"]
    }
    snow._write_lifecycle_state(context, state)
    state["phases"]["preflight"]["status"] = "in-progress"
    monkeypatch.setattr(
        snow,
        "prepare_lifecycle_projection",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            ValueError("malformed projection")
        ),
    )

    snow._write_lifecycle_state(context, state)

    persisted = snow._load_json(_lifecycle_path(tmp_path))
    assert persisted["phases"]["preflight"]["status"] == "in-progress"
    error = capsys.readouterr().err
    assert "telemetry projection failed (ValueError)" in error
    assert "malformed projection" not in error


def test_malformed_telemetry_counts_fail_open_without_blocking_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context_at(tmp_path)
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state = snow._state_base(context, _components())
    state["phases"] = {
        phase["id"]: {"status": "pending", "checkpointResults": {}}
        for phase in contract["phases"]
    }
    snow._write_lifecycle_state(context, state)
    state["phases"]["credential"]["status"] = "blocked"
    snow._write_lifecycle_state(context, state)
    state["telemetry"]["retryCount"] = "customer-free-text"
    snow._write_json_atomic(_lifecycle_path(tmp_path), state)
    state["phases"]["credential"]["status"] = "in-progress"

    snow._write_lifecycle_state(context, state)

    persisted = snow._load_json(_lifecycle_path(tmp_path))
    assert persisted["phases"]["credential"]["status"] == "in-progress"
    assert persisted["telemetry"]["retryCount"] == "customer-free-text"
    assert "telemetry projection failed (ValueError)" in (
        capsys.readouterr().err
    )


def test_initialize_state_exact_existing_revision_is_byte_preserving_noop(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    components = _components(CONNECTION_ID)
    _write_foundation_context(tmp_path, components=components)
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state = snow._state_base(_context_at(tmp_path), components)
    phase_ids = [phase["id"] for phase in contract["phases"]]
    state.update(
        {
            "attested": True,
            "attestedAt": "2026-10-21T19:27:39.184Z",
            "acceptedContractRevision": contract["contractRevision"],
            "roleAttestations": {"Maker": {"status": "confirmed"}},
            "phases": {
                phase_id: {
                    "status": (
                        "done"
                        if phase_id
                        in {"preflight", "plugin-prerequisites"}
                        else "in-progress"
                        if phase_id == "entra-registration"
                        else "pending"
                    ),
                    "checkpointResults": (
                        {"PRE-001": "Passed"}
                        if phase_id == "preflight"
                        else {}
                    ),
                    "checkpointAcknowledgements": {},
                    **(
                        {"actionApplied": True}
                        if phase_id == "preflight"
                        else {}
                    ),
                }
                for phase_id in phase_ids
            },
            "evidence": {"preserve": {"unknown": True}},
            "transactions": {"topics": {"preserve": {"status": "committed"}}},
            "forwardCompatible": {"keep": ["all", "values"]},
        }
    )
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    before = state_path.read_bytes()
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "initialize-state",
            "--contract-revision",
            str(contract["contractRevision"]),
        ]
    ) == 0
    result = json.loads(capsys.readouterr().out)

    assert result["status"] == "existing-compatible-state"
    assert result["stateExists"] is True
    assert result["preservedExistingState"] is True
    assert state_path.read_bytes() == before

    handoff = snow.resume_handoff(
        _context_at(tmp_path),
        verified_phases=frozenset(
            {"preflight", "plugin-prerequisites"}
        ),
    )
    assert handoff["lifecycle"]["decision"] == "run-action"
    assert handoff["lifecycle"]["phase"]["id"] == "entra-registration"
    assert "guide-entra-registration.md" in handoff["action"]["actionDoc"]


@pytest.mark.parametrize(
    ("mutation", "message"),
    (
        ("revision", "requested contract revision"),
        ("provider", "active provider"),
        ("profile", "active profile"),
        ("slug", "active agentSlug"),
        ("agent", "active agentId"),
        ("tenant", "active tenantId"),
        ("environment", "active environmentId"),
        ("ring", "active ring"),
        ("attestation", "not attested"),
    ),
)
def test_initialize_state_rejects_incompatible_existing_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    mutation: str,
    message: str,
) -> None:
    components = _components(CONNECTION_ID)
    _write_foundation_context(tmp_path, components=components)
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state = snow._state_base(_context_at(tmp_path), components)
    state.update(
        {
            "attested": True,
            "attestedAt": "2026-10-21T19:27:39.184Z",
            "acceptedContractRevision": contract["contractRevision"],
            "phases": {
                phase["id"]: {
                    "status": "pending",
                    "checkpointResults": {},
                    "checkpointAcknowledgements": {},
                }
                for phase in contract["phases"]
            },
        }
    )
    if mutation == "revision":
        state["acceptedContractRevision"] -= 1
    elif mutation == "provider":
        state["provider"] = "other-provider"
    elif mutation == "profile":
        state["profile"] = "other-profile"
    elif mutation == "slug":
        state["agentSlug"] = "other-agent"
    elif mutation == "agent":
        state["agentId"] = OTHER_AGENT_ID
    elif mutation == "tenant":
        state["tenantId"] = str(uuid.uuid4())
    elif mutation == "environment":
        state["environmentId"] = str(uuid.uuid4())
    elif mutation == "ring":
        state["ring"] = "prod"
    else:
        state["attested"] = False
    state_path = _lifecycle_path(tmp_path)
    snow._write_json_atomic(state_path, state)
    before = state_path.read_bytes()
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "initialize-state",
            "--contract-revision",
            str(contract["contractRevision"]),
        ]
    ) == 1
    error = json.loads(capsys.readouterr().out)

    assert message in error["message"]
    assert state_path.read_bytes() == before


def test_initialize_state_rejects_invalid_json_without_replacement(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    _write_foundation_context(
        tmp_path,
        components=_components(CONNECTION_ID),
    )
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    before = b"{not-json\n"
    state_path.write_bytes(before)
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "initialize-state",
            "--contract-revision",
            str(contract["contractRevision"]),
        ]
    ) == 1
    error = json.loads(capsys.readouterr().out)

    assert "Could not read JSON file" in error["message"]
    assert state_path.read_bytes() == before


def test_initialize_state_never_replaces_existing_malformed_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    _write_foundation_context(
        tmp_path,
        components=_components(CONNECTION_ID),
    )
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    malformed = {
        "provider": snow.PROVIDER_KEY,
        "agentSlug": AGENT_SLUG,
    }
    state_path.write_text(json.dumps(malformed), encoding="utf-8")
    before = state_path.read_bytes()
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "initialize-state",
            "--contract-revision",
            str(contract["contractRevision"]),
        ]
    ) == 1
    error = json.loads(capsys.readouterr().out)

    assert error["status"] == "error"
    assert "does not match the active" in error["message"]
    assert state_path.read_bytes() == before


def test_initialize_state_preserves_file_created_during_atomic_publish(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    _write_foundation_context(
        tmp_path,
        components=_components(CONNECTION_ID),
    )
    contract = json.loads(_CONTRACT_PATH.read_text(encoding="utf-8"))
    contract_path = tmp_path / snow.CONTRACT_PATH
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    state_path = _lifecycle_path(tmp_path)
    sentinel = b'{"sentinel":"concurrent-state"}\n'
    original_link = snow.os.link

    def create_concurrent_state_then_link(
        source: str | bytes | Path,
        destination: str | bytes | Path,
        *args,
        **kwargs,
    ) -> None:
        Path(destination).write_bytes(sentinel)
        original_link(source, destination, *args, **kwargs)

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow.os,
        "link",
        create_concurrent_state_then_link,
    )

    assert snow.main(
        [
            "initialize-state",
            "--contract-revision",
            str(contract["contractRevision"]),
        ]
    ) == 1
    error = json.loads(capsys.readouterr().out)

    assert error["status"] == "error"
    assert "appeared during initialization" in error["message"]
    assert state_path.read_bytes() == sentinel


def test_migrate_state_cli_returns_nonzero_for_malformed_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    _write_foundation_context(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(
        json.dumps(
            {
                "provider": snow.PROVIDER_KEY,
                "profile": snow.PROFILE_KEY,
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
            }
        ),
        encoding="utf-8",
    )
    before = state_path.read_bytes()
    monkeypatch.chdir(tmp_path)

    assert snow.main(["migrate-state"]) == 1
    error = json.loads(capsys.readouterr().out)

    assert error["status"] == "error"
    assert error["message"] == (
        "ServiceNow lifecycle schemaVersion must be an integer."
    )
    _assert_recovery_handoff(error["details"]["recoveryHandoff"])
    assert state_path.read_bytes() == before


def test_migrate_state_process_exits_nonzero_for_malformed_state(
    tmp_path: Path,
) -> None:
    _write_foundation_context(tmp_path)
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(
        json.dumps(
            {
                "provider": snow.PROVIDER_KEY,
                "profile": snow.PROFILE_KEY,
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
            }
        ),
        encoding="utf-8",
    )
    before = state_path.read_bytes()

    completed = subprocess.run(
        [sys.executable, str(_SCRIPT_PATH), "migrate-state"],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
    )
    error = json.loads(completed.stdout)

    assert completed.returncode == 1
    assert completed.stderr == ""
    assert error["status"] == "error"
    assert error["message"] == (
        "ServiceNow lifecycle schemaVersion must be an integer."
    )
    _assert_recovery_handoff(error["details"]["recoveryHandoff"])
    assert state_path.read_bytes() == before


def test_enable_all_topics_rolls_back_ambiguous_partial_mutation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(json.loads(json.dumps(payload)))
            update = payload["botComponentChanges"][0]["component"]
            topic.update(update)
            topic["version"] = int(topic["version"]) + 1
            components["changeToken"] = f"token-{len(updates) + 1}"
            if len(updates) == 1:
                raise RuntimeError("ambiguous network failure")
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="rolled back safely",
    ):
        snow.enable_all_servicenow_topics(_context(), confirmed=True)

    assert len(updates) == 2
    assert topic["state"] == "Inactive"
    assert topic["status"] == "Inactive"
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = next(iter(state["transactions"]["topics"].values()))
    assert transaction["status"] == "rolled-back"
    assert list(transaction["topics"].values()) == ["restored"]


def test_enable_all_topics_accepts_server_managed_audit_changes(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"
    topic["auditInfo"] = {
        "modifiedAt": "2026-09-28T00:00:00Z",
        "modifiedBy": "before",
    }

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updated = payload["botComponentChanges"][0]["component"]
            topic.update(updated)
            topic["version"] = int(topic["version"]) + 1
            topic["auditInfo"] = {
                "modifiedAt": "2026-09-28T00:01:00Z",
                "modifiedBy": "platform",
            }
            components["changeToken"] = "token-2"
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.enable_all_servicenow_topics(
        _context(),
        confirmed=True,
    )

    assert result["transactionStatus"] == "committed"
    assert result["after"] == {"total": 1, "active": 1, "inactive": 0}
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = next(iter(state["transactions"]["topics"].values()))
    target = transaction["targets"][topic["id"]]
    assert target["ownershipStatus"] == "owned"
    assert target["changedByOperation"] is True


def test_enable_all_topics_accepts_unattributed_active_postimage_difference(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A successful write plus Active readback outranks full-content drift.

    A non-target difference may be service normalization or an external edit.
    The old implementation treated any difference as a concurrency conflict,
    attempted rollback, and failed the phase even though every requested topic
    was Active.
    """

    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(json.loads(json.dumps(payload)))
            updated = payload["botComponentChanges"][0]["component"]
            topic.update(updated)
            topic["version"] = int(topic["version"]) + 1
            topic["serverManagedNormalization"] = {
                "kind": "canonicalized",
            }
            components["changeToken"] = "token-2"
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.enable_all_servicenow_topics(
        _context(),
        confirmed=True,
    )

    assert len(updates) == 1
    assert result["transactionStatus"] == "committed"
    assert result["after"] == {"total": 1, "active": 1, "inactive": 0}
    assert result["unattributedPostimageTopics"] == [
        {
            "id": topic["id"],
            "displayName": topic["displayName"],
        }
    ]
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = next(iter(state["transactions"]["topics"].values()))
    target = transaction["targets"][topic["id"]]
    assert transaction["status"] == "committed"
    assert transaction["semanticOutcome"] == "target-state-verified"
    assert transaction["unattributedPostimageTargetCount"] == 1
    assert target["ownershipStatus"] == "content-difference"
    assert target["targetStateVerified"] is True
    assert target["changedByOperation"] is None
    assert target["fullPostimageOwned"] is False


def test_enable_all_topics_does_not_accept_content_difference_partial_success(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    first = components["botComponentChanges"][0]["component"]
    first["state"] = "Inactive"
    first["status"] = "Inactive"
    second = copy.deepcopy(first)
    second["id"] = "00000000-0000-4000-8000-000000006666"
    second["displayName"] = "ServiceNow HRSD Get Cases"
    second["schemaName"] = (
        f"{snow.HR_SCHEMA_NAME}.topic.ServiceNowHRSDGetCases"
    )
    components["botComponentChanges"].append(
        {"$kind": "BotComponentInsert", "component": second}
    )
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(json.loads(json.dumps(payload)))
            first.update(payload["botComponentChanges"][0]["component"])
            first["version"] = int(first["version"]) + 1
            first["serverManagedNormalization"] = True
            components["changeToken"] = "token-2"
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="rollback is incomplete",
    ) as raised:
        snow.enable_all_servicenow_topics(_context(), confirmed=True)

    assert len(updates) == 1
    assert raised.value.details["after"] == {
        "total": 2,
        "active": 1,
        "inactive": 1,
    }
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = state["transactions"]["topics"][
        raised.value.details["operationId"]
    ]
    assert transaction["status"] == "rollback-incomplete"
    assert "semanticOutcome" not in transaction


def test_enable_all_topics_does_not_commit_ambiguous_content_difference(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(json.loads(json.dumps(payload)))
            topic.update(payload["botComponentChanges"][0]["component"])
            topic["version"] = int(topic["version"]) + 1
            topic["serverManagedNormalization"] = True
            components["changeToken"] = "token-2"
            raise RuntimeError("ambiguous network failure")

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="rollback is incomplete",
    ) as raised:
        snow.enable_all_servicenow_topics(_context(), confirmed=True)

    assert len(updates) == 1
    assert raised.value.details["updateError"] == "RuntimeError"
    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    transaction = state["transactions"]["topics"][
        raised.value.details["operationId"]
    ]
    assert transaction["status"] == "rollback-incomplete"
    assert "semanticOutcome" not in transaction


def test_topic_rollback_does_not_overwrite_concurrent_external_edit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    first = components["botComponentChanges"][0]["component"]
    first["state"] = "Inactive"
    first["status"] = "Inactive"
    second = copy.deepcopy(first)
    second["id"] = "00000000-0000-4000-8000-000000006666"
    second["displayName"] = "ServiceNow HRSD Get Cases"
    second["schemaName"] = (
        f"{snow.HR_SCHEMA_NAME}.topic.ServiceNowHRSDGetCases"
    )
    components["botComponentChanges"].append(
        {"$kind": "BotComponentInsert", "component": second}
    )
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(json.loads(json.dumps(payload)))
            first_update = payload["botComponentChanges"][0]["component"]
            first.update(first_update)
            first["version"] = int(first["version"]) + 1
            first["externalConcurrentEdit"] = "must-survive"
            components["changeToken"] = "token-2"
            raise RuntimeError("ambiguous partial update")

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    historical_operation_id = "00000000-0000-4000-8000-000000000001"
    initial_state = snow._state_base(_context(), components)
    initial_state["transactions"]["topics"][historical_operation_id] = {
        "operationId": historical_operation_id,
        "status": "rollback-incomplete",
    }
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(json.dumps(initial_state), encoding="utf-8")

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="rollback is incomplete",
    ) as raised:
        snow.enable_all_servicenow_topics(_context(), confirmed=True)

    assert len(updates) == 1
    assert first["externalConcurrentEdit"] == "must-survive"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    current_operation_id = raised.value.details["operationId"]
    assert current_operation_id != historical_operation_id
    transaction = state["transactions"]["topics"][current_operation_id]
    assert transaction["status"] == "rollback-incomplete"
    assert transaction["topics"][first["id"]] == "conflict"
    assert raised.value.details["transactionStatus"] == "rollback-incomplete"
    assert raised.value.details["before"] == {
        "total": 2,
        "active": 0,
        "inactive": 2,
    }
    target = next(
        item
        for item in raised.value.details["targets"]
        if item["topicId"] == first["id"]
    )
    assert target["ownershipStatus"] == "conflict"
    assert target["rollbackStatus"] == "conflict"
    assert "component" not in json.dumps(raised.value.details)


def test_topic_rollback_treats_ambiguous_applied_restore_as_success(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    preimage = copy.deepcopy(topic)
    preimage["state"] = "Inactive"
    preimage["status"] = "Inactive"
    topic["state"] = "Active"
    topic["status"] = "Active"
    operation_id = "00000000-0000-4000-8000-000000007777"
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 2,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {},
                "evidence": {},
                "transactions": {
                    "topics": {
                        operation_id: {
                            "operationId": operation_id,
                            "status": "rollback-required",
                            "targets": {
                                topic["id"]: {
                                    "component": preimage,
                                    "changedByOperation": True,
                                    "ownershipStatus": "owned",
                                    "beforeContentHashExcludingVersion": (
                                        snow._component_content_hash(preimage)
                                    ),
                                    "expectedPostContentHashExcludingVersion": (
                                        snow._component_content_hash(topic)
                                    ),
                                }
                            },
                        }
                    }
                },
                "migration": {},
            }
        ),
        encoding="utf-8",
    )

    class FakeAgentBuilder:
        calls = 0

        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            self.calls += 1
            restored = payload["botComponentChanges"][0]["component"]
            topic.update(restored)
            topic["version"] = int(topic["version"]) + 1
            components["changeToken"] = f"token-{self.calls + 1}"
            if self.calls == 1:
                raise RuntimeError("ambiguous rollback response")
            return {}

    client = FakeAgentBuilder()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: client,
    )

    first = snow.rollback_topic_transaction(
        _context(),
        operation_id,
        confirmed=True,
    )
    second = snow.rollback_topic_transaction(
        _context(),
        operation_id,
        confirmed=True,
    )

    assert first["status"] == "rolled-back"
    assert second["status"] == "rolled-back"
    assert client.calls == 1
    assert topic["state"] == "Inactive"
    assert topic["status"] == "Inactive"


def test_topic_rollback_refuses_unattributed_active_content_difference(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    preimage = copy.deepcopy(topic)
    preimage["state"] = "Inactive"
    preimage["status"] = "Inactive"
    topic["state"] = "Active"
    topic["status"] = "Active"
    topic["unattributedExternalField"] = "must-survive"
    operation_id = "00000000-0000-4000-8000-000000007778"
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": 2,
                "provider": snow.PROVIDER_KEY,
                "profile": "hrsd",
                "agentSlug": AGENT_SLUG,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "phases": {},
                "evidence": {},
                "transactions": {
                    "topics": {
                        operation_id: {
                            "operationId": operation_id,
                            "status": "committed",
                            "targets": {
                                topic["id"]: {
                                    "component": preimage,
                                    "changedByOperation": None,
                                    "ownershipStatus": "content-difference",
                                    "targetStateVerified": True,
                                    "fullPostimageOwned": False,
                                    "beforeContentHashExcludingVersion": (
                                        snow._component_content_hash(preimage)
                                    ),
                                }
                            },
                        }
                    }
                },
                "migration": {},
            }
        ),
        encoding="utf-8",
    )
    updates: list[dict] = []

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            updates.append(payload)
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.rollback_topic_transaction(
        _context(),
        operation_id,
        confirmed=True,
    )

    assert result["status"] == "rollback-incomplete"
    assert result["topics"][topic["id"]] == "conflict"
    assert updates == []
    assert topic["unattributedExternalField"] == "must-survive"


def test_failed_post_update_refetch_does_not_claim_rollback_ownership(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    topic = components["botComponentChanges"][0]["component"]
    topic["state"] = "Inactive"
    topic["status"] = "Inactive"
    fail_refetch = False
    update_calls = 0

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            nonlocal fail_refetch
            if fail_refetch:
                fail_refetch = False
                raise RuntimeError("post-update fetch failed")
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            nonlocal fail_refetch, update_calls
            update_calls += 1
            update = payload["botComponentChanges"][0]["component"]
            topic.update(update)
            topic["version"] = int(topic["version"]) + 1
            components["changeToken"] = "token-2"
            fail_refetch = True
            return {}

    client = FakeAgentBuilder()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: client,
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="post-update refetch failed",
    ):
        snow.enable_all_servicenow_topics(_context(), confirmed=True)

    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    operation_id, transaction = next(
        iter(state["transactions"]["topics"].items())
    )
    assert transaction["status"] == "reconciliation-required"
    target = transaction["targets"][topic["id"]]
    assert "ownershipStatus" not in target

    topic["externalConcurrentEdit"] = "must-survive"
    result = snow.rollback_topic_transaction(
        _context(),
        operation_id,
        confirmed=True,
    )

    assert result["status"] == "rollback-incomplete"
    assert result["topics"][topic["id"]] == "conflict"
    assert topic["externalConcurrentEdit"] == "must-survive"
    assert update_calls == 1


def test_enable_all_topics_rejects_missing_transaction_target(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components()
    first = components["botComponentChanges"][0]["component"]
    first["state"] = "Inactive"
    first["status"] = "Inactive"
    second = copy.deepcopy(first)
    second["id"] = "00000000-0000-4000-8000-000000006666"
    second["displayName"] = "ServiceNow HRSD Get Cases"
    second["schemaName"] = (
        f"{snow.HR_SCHEMA_NAME}.topic.ServiceNowHRSDGetCases"
    )
    components["botComponentChanges"].append(
        {"$kind": "BotComponentInsert", "component": second}
    )

    class FakeAgentBuilder:
        def fetch_components(self, _agent_id: str) -> dict:
            return json.loads(json.dumps(components))

        def update_components(self, _agent_id: str, payload: dict) -> dict:
            second.update(payload["botComponentChanges"][1]["component"])
            second["version"] = int(second["version"]) + 1
            components["botComponentChanges"] = [
                change
                for change in components["botComponentChanges"]
                if change["component"].get("id") != first["id"]
            ]
            components["changeToken"] = "token-2"
            return {}

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="rollback is incomplete",
    ) as raised:
        snow.enable_all_servicenow_topics(_context(), confirmed=True)

    state = json.loads(_lifecycle_path(tmp_path).read_text(encoding="utf-8"))
    operation_id = raised.value.details["operationId"]
    transaction = state["transactions"]["topics"][operation_id]
    assert transaction["status"] == "rollback-incomplete"
    assert transaction["topics"][first["id"]] == "missing"
    assert transaction["status"] != "committed"
    target = next(
        item
        for item in raised.value.details["targets"]
        if item["topicId"] == first["id"]
    )
    assert target["ownershipStatus"] == "missing"
    assert target["rollbackStatus"] == "missing"


def test_main_serializes_current_topic_transaction_details(
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    operation_id = "00000000-0000-4000-8000-000000009999"
    monkeypatch.setattr(snow, "load_context", lambda: _context())
    monkeypatch.setattr(
        snow,
        "enable_all_servicenow_topics",
        lambda _context, *, confirmed: (_ for _ in ()).throw(
            snow.ServiceNowConnectError(
                "topic update failed",
                details={
                    "operationId": operation_id,
                    "transactionStatus": "rollback-incomplete",
                    "targets": [],
                },
            )
        ),
    )
    monkeypatch.setattr(
        sys,
        "argv",
        ["connect_servicenow_da.py", "enable-all-topics", "--yes"],
    )

    assert snow.main() == 1
    output = json.loads(capsys.readouterr().out)
    assert output["status"] == "error"
    assert output["message"] == "topic update failed"
    assert output["details"]["operationId"] == operation_id
    assert output["details"]["transactionStatus"] == "rollback-incomplete"
    assert output["details"]["targets"] == []
    recovery = output["details"]["recoveryHandoff"]
    _assert_recovery_handoff(recovery)
    assert recovery["owner"] == "ESS Maker / Agent Developer"
    assert recovery["where"].endswith("-> Topics")


_ITSM_IMPORTED_TOPIC_SUFFIXES = (
    "ServiceNowITSMCreateTicket",
    "ServiceNowITSMGetTicketDetails",
    "ServiceNowITSMGetTicketUpdates",
    "ServiceNowITSMGetUserTickets",
    "ServiceNowITSMSetupConfigurations",
    "ServiceNowITSMSystemCommonExecution",
    "ServiceNowITSMSystemCreateTicket",
    "ServiceNowITSMSystemGetTicketDetails",
    "ServiceNowITSMSystemGetTicketsList",
    "ServiceNowITSMSystemResponseMapper",
    "ServiceNowITSMSystemUpdateTicket",
    "ServiceNowITSMSystemUserPermissionCheck",
    "ServiceNowITSMUpdateTicket",
)


def _itsm_components() -> dict:
    components = _components(CONNECTION_ID)
    official = []
    for index, suffix in enumerate(_ITSM_IMPORTED_TOPIC_SUFFIXES):
        component = {
            "$kind": "DialogComponent",
            "id": f"00000000-0000-4000-8000-{7000 + index:012d}",
            "version": 1,
            "displayName": suffix,
            "schemaName": f"{snow.IT_SCHEMA_NAME}.topic.{suffix}",
            "state": "Active",
            "status": "Active",
            "dialog": {"$kind": "TaskDialog"},
        }
        if suffix == "ServiceNowITSMSetupConfigurations":
            component["dialog"] = {
                "$kind": "TaskDialog",
                "actions": [
                    {
                        "$kind": "SetVariable",
                        "displayName": "Set ServiceNow Portal Base URI",
                        "variable": "Global.ServiceNowITSMPortalBaseURI",
                        "value": {
                            "$kind": "PowerFx",
                            "literalValue": (
                                "https://dev123.service-now.com/esc"
                            ),
                        },
                    }
                ],
            }
        official.append(
            {"$kind": "BotComponentInsert", "component": component}
        )
    official.append(
        {
            "$kind": "BotComponentInsert",
            "component": {
                "$kind": "DialogComponent",
                "id": "00000000-0000-4000-8000-000000009999",
                "version": 1,
                "displayName": "Customer ITSM custom topic",
                "schemaName": (
                    f"{snow.IT_SCHEMA_NAME}.topic."
                    "CustomerITSMCustom"
                ),
                "state": "Inactive",
                "status": "Inactive",
                "dialog": {"$kind": "TaskDialog"},
            },
        }
    )
    components["botComponentChanges"] = official
    return components


def _seed_it_setup(
    root: Path,
    *,
    enrolled: bool = True,
    canonical_account: str | None = None,
    operational_account: str | None = None,
) -> tuple[Path, Path]:
    setup_path = root / snow.SETUP_STATE
    config_path = root / snow.ACTIVE_CONFIG
    setup_path.parent.mkdir(parents=True)
    setup_path.write_text(
        json.dumps(
            {
                "schema_version": 4,
                "intent": "DA foundation setup",
                "environment": {
                    "id": ENVIRONMENT_ID,
                    "tenant_id": "00000000-0000-4000-8000-000000009999",
                    "ring": "test",
                    "api_version": "2024-10-01",
                    "power_platform_api_endpoint": (
                        "https://example.environment.api.test.powerplatform.com"
                    ),
                    **(
                        {"account": canonical_account}
                        if canonical_account
                        else {}
                    ),
                },
                "agents": {
                    AGENT_ID: {
                        "authoring_ready": True,
                        "setup_source": "existing-dev",
                        "agent": {
                            "id": AGENT_ID,
                            "schema_name": snow.IT_SCHEMA_NAME,
                            **(
                                {"realm": "dev"}
                                if enrolled
                                else {"alm": {"isEnrolled": False}}
                            ),
                            "workspace_slug": "employee-self-service-it",
                        },
                        "workspace": {
                            "folder": "workspace/agents/employee-self-service-it",
                            "agent_path": "agent.mcs.yml",
                        },
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text(
        json.dumps(
            {
                "activeAgent": "employee-self-service-hr",
                "environmentId": ENVIRONMENT_ID,
                "tenantId": "00000000-0000-4000-8000-000000009999",
                "ring": "test",
                "agentBuilderApiVersion": "2024-10-01",
                "powerPlatformApiEndpoint": (
                    "https://example.environment.api.test.powerplatform.com"
                ),
                **(
                    {"account": operational_account}
                    if operational_account
                    else {}
                ),
                "agents": [
                    {
                        "slug": "employee-self-service-it",
                        "botId": AGENT_ID,
                        "agentBuilderChangeSetPath": (
                            "workspace/agents/employee-self-service-it/"
                            ".agentbuilder/components.json"
                        ),
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    return setup_path, config_path


def test_itsm_profile_uses_exact_official_topics_and_portal_selector() -> None:
    components = _itsm_components()
    official_topic = components["botComponentChanges"][0]["component"]
    official_topic["state"] = "Inactive"
    official_topic["status"] = "Inactive"
    custom_topic = components["botComponentChanges"][-1]["component"]

    with snow.product_profile_context("itsm"):
        summary = snow.summarize_components(components)
        portal = snow.portal_configuration_summary(
            components,
            expected_instance_name="dev123",
        )
        payload = snow.build_enable_all_topics_payload(components)

    assert summary["profile"]["provider"] == "servicenow-da-itsm"
    assert summary["serviceNowTopicCount"] == 13
    assert summary["activeServiceNowTopicCount"] == 12
    assert portal["variable"] == "Global.ServiceNowITSMPortalBaseURI"
    assert portal["valid"] is True
    assert [
        change["component"]["id"]
        for change in payload["botComponentChanges"]
    ] == [official_topic["id"]]
    assert custom_topic["id"] not in {
        change["component"]["id"]
        for change in payload["botComponentChanges"]
    }


def test_itsm_topic_discovery_accepts_new_matching_topic_and_isolates_noise() -> None:
    components = _itsm_components()
    components["botComponentChanges"].extend(
        [
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000008888",
                    "displayName": "Future ITSM topic",
                    "schemaName": (
                        f"{snow.IT_SCHEMA_NAME}.topic."
                        "ServiceNowITSMFuturePackageTopic"
                    ),
                    "state": "Active",
                    "status": "Active",
                    "dialog": {"$kind": "TaskDialog"},
                },
            },
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000008889",
                    "displayName": "Opposite provider topic",
                    "schemaName": (
                        f"{snow.IT_SCHEMA_NAME}.topic.ServiceNowHRSDCreateCase"
                    ),
                    "state": "Active",
                    "status": "Active",
                    "dialog": {"$kind": "TaskDialog"},
                },
            },
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "GlobalVariableComponent",
                    "id": "00000000-0000-4000-8000-000000008890",
                    "displayName": "Not a topic",
                    "schemaName": (
                        f"{snow.IT_SCHEMA_NAME}.topic."
                        "ServiceNowITSMNotADialog"
                    ),
                    "dialog": {"$kind": "TaskDialog"},
                },
            },
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000008891",
                    "displayName": "Missing dialog payload",
                    "schemaName": (
                        f"{snow.IT_SCHEMA_NAME}.topic."
                        "ServiceNowITSMMalformedDialog"
                    ),
                    "state": "Active",
                    "status": "Active",
                },
            },
        ]
    )

    with snow.product_profile_context("itsm"):
        summary = snow.summarize_components(components)

    schemas = {
        topic["schemaName"] for topic in summary["serviceNowTopics"]
    }
    assert summary["serviceNowTopicCount"] == 14
    assert (
        f"{snow.IT_SCHEMA_NAME}.topic.ServiceNowITSMFuturePackageTopic"
        in schemas
    )
    assert all("ServiceNowHRSD" not in schema for schema in schemas)
    assert all("NotADialog" not in schema for schema in schemas)
    assert all("MalformedDialog" not in schema for schema in schemas)


def test_topic_action_handoff_returns_complete_it_question(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    for suffix in (
        "ServiceNowITSMSystemReadExecution",
        "ServiceNowITSMSystemWriteExecution",
    ):
        components["botComponentChanges"].append(
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": str(uuid.uuid4()),
                    "displayName": suffix,
                    "schemaName": f"{snow.IT_SCHEMA_NAME}.topic.{suffix}",
                    "state": "Inactive",
                    "status": "Inactive",
                    "dialog": {"$kind": "TaskDialog"},
                },
            }
        )
    for change in components["botComponentChanges"]:
        if "ServiceNowITSM" in str(
            change["component"].get("schemaName") or ""
        ) and not str(change["component"]["schemaName"]).endswith(
            ("SystemReadExecution", "SystemWriteExecution")
        ):
            change["component"]["state"] = "Active"
            change["component"]["status"] = "Active"
    context = {
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
        "snapshotPath": tmp_path / "components.json",
    }
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
    )

    with snow.product_profile_context("itsm"):
        result = snow.topic_action_handoff(context)

    assert result["decision"] == "question-required"
    assert result["inventory"]["active"] == 13
    assert result["inventory"]["inactive"] == 2
    assert result["question"]["choices"] == [
        "Enable all topics",
        "Keep current states",
        "Not now",
    ]
    question = result["question"]["question"]
    assert "ServiceNowITSMSystemReadExecution" in question
    assert "ServiceNowITSMSystemWriteExecution" in question
    assert "13 active and 2 inactive" in question
    assert not (tmp_path / ".local/connect").exists()


def test_topic_action_handoff_all_active_is_verified_noop(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    context = {
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {"id": ENVIRONMENT_ID, "ring": "test"},
        "snapshotPath": tmp_path / "components.json",
    }
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
    )

    with snow.product_profile_context("itsm"):
        result = snow.topic_action_handoff(context)

    assert result["decision"] == "verified-noop"
    assert result["actionResult"] == "applied"
    assert result["reason"] == "all-topics-active"


def test_topic_action_handoff_has_hr_parity(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _components(CONNECTION_ID)
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"
    context = {
        **_context(),
        "snapshotPath": tmp_path / "components.json",
    }
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
    )

    result = snow.topic_action_handoff(context)

    assert result["decision"] == "question-required"
    assert result["profile"] == "hrsd"
    assert "Inactive ServiceNow HRSD topics" in result["question"]["question"]


def test_topic_action_handoff_reuses_only_current_keep_current(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"
    context = {
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {"id": ENVIRONMENT_ID, "ring": "test"},
        "snapshotPath": tmp_path / "components.json",
    }
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
    )
    with snow.product_profile_context("itsm"):
        state = snow._state_base(context, components)
        state["evidence"]["topics"] = {
            "kind": "maker-attestation",
            "status": "recorded",
            "customerChoice": "keep-current",
            "makerAttested": True,
            "boundComponentHash": snow._component_hash(components),
        }
        snow._write_json_atomic(snow._lifecycle_state_path(context), state)
        current = snow.topic_action_handoff(context)
        state["evidence"]["topics"]["boundComponentHash"] = "stale"
        snow._write_json_atomic(snow._lifecycle_state_path(context), state)
        stale = snow.topic_action_handoff(context)

    assert current["decision"] == "verified-noop"
    assert current["reason"] == "current-keep-current-attestation"
    assert stale["decision"] == "question-required"


def test_topic_action_handoff_zero_and_fetch_failures_are_explicit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    context = {
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {"id": ENVIRONMENT_ID, "ring": "test"},
        "snapshotPath": tmp_path / "components.json",
    }
    empty = _itsm_components()
    empty["botComponentChanges"] = []
    monkeypatch.chdir(tmp_path)
    with snow.product_profile_context("itsm"):
        monkeypatch.setattr(
            snow,
            "_agentbuilder_client",
            lambda _context: SimpleNamespace(
                fetch_components=lambda _agent_id: empty
            ),
        )
        with pytest.raises(snow.ServiceNowConnectError, match="no ServiceNow ITSM"):
            snow.topic_action_handoff(context)
        malformed = _itsm_components()
        malformed["botComponentChanges"] = {"invalid": True}
        monkeypatch.setattr(
            snow,
            "_agentbuilder_client",
            lambda _context: SimpleNamespace(
                fetch_components=lambda _agent_id: malformed
            ),
        )
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="invalid topic collection",
        ):
            snow.topic_action_handoff(context)
        monkeypatch.setattr(
            snow,
            "_agentbuilder_client",
            lambda _context: SimpleNamespace(
                fetch_components=lambda _agent_id: (_ for _ in ()).throw(
                    snow.AgentBuilderHTTPError("Component fetch", 403)
                )
            ),
        )
        with pytest.raises(snow.AgentBuilderHTTPError):
            snow.topic_action_handoff(context)


def test_resume_handoff_consumes_controller_and_returns_exact_action(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    for suffix in (
        "ServiceNowITSMSystemReadExecution",
        "ServiceNowITSMSystemWriteExecution",
    ):
        components["botComponentChanges"].append(
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": str(uuid.uuid4()),
                    "displayName": suffix,
                    "schemaName": f"{snow.IT_SCHEMA_NAME}.topic.{suffix}",
                    "state": "Inactive",
                    "status": "Inactive",
                    "dialog": {"$kind": "TaskDialog"},
                },
            }
        )
    for change in components["botComponentChanges"]:
        if str(change["component"].get("schemaName") or "").endswith(
            ("SystemReadExecution", "SystemWriteExecution")
        ):
            continue
        if "ServiceNowITSM" in str(
            change["component"].get("schemaName") or ""
        ):
            change["component"]["state"] = "Active"
            change["component"]["status"] = "Active"
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
        "snapshotPath": tmp_path / "components.json",
    }
    _write_resume_contract_and_state(tmp_path)
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
    )

    with snow.product_profile_context("itsm"):
        result = snow.resume_handoff(
            context,
            verified_phases=frozenset({"preflight"}),
        )

    assert result["lifecycle"]["decision"] == "run-action"
    assert result["lifecycle"]["phase"]["id"] == "topics"
    assert result["action"] == {
        "decision": "read-action-doc",
        "actionDoc": (
            "src/skills/connect/servicenow-da-itsm/"
            "actions/prepare-topics.md"
        ),
    }
    assert result["progress"]["currentPhase"]["label"] == (
        "Prepare ServiceNow topics"
    )
    assert result["progress"]["currentPhase"]["startText"] == (
        "Starting: Prepare agent topics"
    )
    assert result["progress"]["completedCount"] == 1
    assert result["progress"]["finalText"] is None

    with snow.product_profile_context("itsm"):
        checkpoint_result = snow.resume_handoff(
            context,
            executed_actions=frozenset({"topics"}),
            verified_phases=frozenset({"preflight"}),
        )

    command = checkpoint_result["checkpoints"]["commands"][0]
    assert command["target"] == "SN-DA-ITSM-TOPICS-001"
    assert command["argv"][command["argv"].index("--agent-slug") + 1] == (
        "employee-self-service-it"
    )
    assert command["expectedIdentity"] == {
        "provider": "servicenow-da-itsm",
        "profile": "itsm",
        "agentSlug": "employee-self-service-it",
        "agentId": AGENT_ID,
        "environmentId": ENVIRONMENT_ID,
    }
    assert ".local/config.json.activeAgent" not in command["command"]


@pytest.mark.parametrize(
    ("profile", "agent_id", "agent_slug"),
    (
        ("hrsd", AGENT_ID, "employee-self-service-hr"),
        ("itsm", OTHER_AGENT_ID, "employee-self-service-it"),
    ),
)
def test_resume_handoff_withholds_preflight_checkpoint_for_pending_instance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
    agent_id: str,
    agent_slug: str,
) -> None:
    product = snow.PRODUCT_PROFILES[profile]
    context = _profile_context(
        profile,
        agent_id=agent_id,
        agent_slug=agent_slug,
    )
    context["root"] = tmp_path
    source_contract = (
        snow.SOLUTION_ROOT
        / "src"
        / "skills"
        / "connect"
        / product.provider_key
        / "contract.json"
    )
    contract = json.loads(source_contract.read_text(encoding="utf-8"))
    contract_path = (
        tmp_path
        / "src"
        / "skills"
        / "connect"
        / product.provider_key
        / "contract.json"
    )
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    with snow.product_profile_context(profile):
        state = snow._state_base(context, _components())
        state["attested"] = True
        state["attestedAt"] = "2026-10-29T00:00:00Z"
        state["acceptedContractRevision"] = contract["contractRevision"]
        state["phases"] = {
            phase["id"]: {
                "status": (
                    "in-progress"
                    if phase["id"] == "preflight"
                    else "pending"
                ),
                "checkpointResults": (
                    {
                        phase["checkpoints"][0]: "NotConfigured",
                    }
                    if phase["id"] == "preflight"
                    else {}
                ),
            }
            for phase in contract["phases"]
        }
        question, choices, manual_question = (
            snow._preflight_instance_question(
                context=context,
                candidates=["dev184242"],
                discovery_error=None,
            )
        )
        state["adminSetup"]["preflight"]["discovery"] = {
            "instanceInputRequired": True,
            "instanceQuestion": question,
            "manualInstanceQuestion": manual_question,
            "instanceChoices": choices,
        }
        snow._write_lifecycle_state(context, state)
        result = snow.resume_handoff(
            context,
            executed_actions=frozenset({"preflight"}),
        )

    assert result["lifecycle"]["decision"] == "run-checkpoints"
    assert "checkpoints" not in result
    assert "action" not in result
    pending = result["pendingInput"]
    assert pending["kind"] == "instance-confirmation"
    assert pending["phase"] == "preflight"
    assert pending["question"] == question
    assert pending["manualQuestion"] == manual_question
    assert pending["choiceValues"] == choices
    body = pending["question"]["question"]
    assert "Which instance should this agent use" in body
    assert "**Problem**" not in body
    assert "**Category**" not in body
    assert "**Recheck and resume**" not in body


def test_resume_handoff_allows_preflight_checkpoint_after_instance_recorded(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    profile = "itsm"
    product = snow.PRODUCT_PROFILES[profile]
    context = _profile_context(
        profile,
        agent_id=OTHER_AGENT_ID,
        agent_slug="employee-self-service-it",
    )
    context["root"] = tmp_path
    source_contract = (
        snow.SOLUTION_ROOT
        / "src"
        / "skills"
        / "connect"
        / product.provider_key
        / "contract.json"
    )
    contract = json.loads(source_contract.read_text(encoding="utf-8"))
    contract_path = (
        tmp_path
        / "src"
        / "skills"
        / "connect"
        / product.provider_key
        / "contract.json"
    )
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    with snow.product_profile_context(profile):
        state = snow._state_base(context, _components())
        state["attested"] = True
        state["attestedAt"] = "2026-10-29T00:00:00Z"
        state["acceptedContractRevision"] = contract["contractRevision"]
        state["phases"] = {
            phase["id"]: {
                "status": (
                    "in-progress"
                    if phase["id"] == "preflight"
                    else "pending"
                ),
                "checkpointResults": {},
            }
            for phase in contract["phases"]
        }
        state["adminSetup"]["preflight"]["instanceName"] = "dev184242"
        state["adminSetup"]["preflight"]["discovery"] = {
            "instanceInputRequired": False,
        }
        snow._write_lifecycle_state(context, state)
        result = snow.resume_handoff(
            context,
            executed_actions=frozenset({"preflight"}),
        )

    assert "pendingInput" not in result
    assert result["checkpoints"]["commands"][0]["target"] == (
        "SN-DA-ITSM-ADMIN-PREFLIGHT-001"
    )


@pytest.mark.parametrize(
    ("provider", "has_plugin_phase"),
    [
        ("servicenow-da-hrsd", True),
        ("servicenow-da-itsm", False),
    ],
)
def test_phase_progress_uses_provider_contract_labels_and_order(
    provider: str,
    has_plugin_phase: bool,
) -> None:
    contract = json.loads(
        (
            snow.SOLUTION_ROOT
            / "src"
            / "skills"
            / "connect"
            / provider
            / "contract.json"
        ).read_text(encoding="utf-8")
    )
    state = {
        "phases": {
            phase["id"]: {
                "status": "pending",
                "checkpointResults": {},
            }
            for phase in contract["phases"]
        }
    }
    first = contract["phases"][0]

    progress = snow._build_phase_progress(
        contract,
        state,
        {
            "decision": "run-action",
            "phase": {"id": first["id"]},
        },
        executed_actions=frozenset(),
        verified_phases=frozenset(),
    )

    assert [item["label"] for item in progress["items"]] == [
        phase["label"] for phase in contract["phases"]
    ]
    assert (
        "Confirm HR Service Delivery Core is installed and ready"
        in [item["label"] for item in progress["items"]]
    ) is has_plugin_phase
    expected_customer_labels = [
        "Check existing setup",
        *(
            ["Confirm HR ServiceNow prerequisite"]
            if has_plugin_phase
            else []
        ),
        "Confirm Entra identity options",
        "Confirm ServiceNow user mapping",
        "Check ServiceNow connection",
        "Prepare agent topics",
        "Set employee portal URL",
        "Connect this agent",
        "Test ServiceNow",
    ]
    assert [item["customerLabel"] for item in progress["items"]] == (
        expected_customer_labels
    )
    assert progress["headline"] == (
        f"Starting {contract['displayName']} connection."
    )


def test_phase_progress_distinguishes_reused_manual_recheck_and_blocked() -> None:
    contract = {
        "displayName": "ServiceNow HRSD",
        "phases": [
            {"id": "reused", "label": "Reuse setup"},
            {"id": "manual", "label": "Confirm admin setup"},
            {"id": "cached", "label": "Recheck cached setup"},
            {"id": "later", "label": "Later work"},
            {"id": "blocked", "label": "Blocked work"},
        ],
    }
    state = {
        "adminSetup": {
            "phaseHandoffs": {
                "reused": {"status": "reused"},
            }
        },
        "phases": {
            "reused": {
                "status": "done",
                "checkpointResults": {"A": "Manual"},
            },
            "manual": {
                "status": "done",
                "checkpointResults": {"B": "Manual"},
            },
            "cached": {
                "status": "done",
                "checkpointResults": {"C": "Passed"},
            },
            "later": {
                "status": "pending",
                "checkpointResults": {},
            },
            "blocked": {
                "status": "blocked",
                "checkpointResults": {"E": "Failed"},
            },
        },
    }

    progress = snow._build_phase_progress(
        contract,
        state,
        {
            "decision": "run-checkpoints",
            "phase": {"id": "cached"},
        },
        executed_actions=frozenset(),
        verified_phases=frozenset({"reused", "manual"}),
    )

    assert [item["status"] for item in progress["items"]] == [
        "complete-reused",
        "complete-manual",
        "rechecking",
        "pending",
        "needs-attention",
    ]
    assert progress["items"][0]["text"] == "Reused — checked"
    assert progress["items"][1]["text"] == "Complete — admin confirmed"
    assert progress["currentPhase"]["startText"] == (
        "Rechecking: Recheck cached setup"
    )
    assert progress["currentPhase"]["waitingText"] == (
        "Waiting — Recheck cached setup: your input is needed."
    )


def test_phase_progress_does_not_complete_after_action_only() -> None:
    contract = {
        "displayName": "ServiceNow ITSM",
        "phases": [{"id": "topics", "label": "Prepare ServiceNow topics"}],
    }
    state = {
        "phases": {
            "topics": {
                "status": "in-progress",
                "actionApplied": True,
                "checkpointResults": {"TOPICS": "NotConfigured"},
            }
        }
    }

    progress = snow._build_phase_progress(
        contract,
        state,
        {
            "decision": "run-checkpoints",
            "phase": {"id": "topics"},
        },
        executed_actions=frozenset({"topics"}),
        verified_phases=frozenset(),
    )

    assert progress["items"][0]["status"] == "checking"
    assert progress["completedCount"] == 0
    assert progress["finalText"] is None
    assert progress["currentPhase"]["actionFinishedText"] == (
        "Action finished — checking: Prepare agent topics"
    )


def test_phase_progress_reemits_updated_board_for_transition_wait_and_error(
) -> None:
    contract = {
        "displayName": "ServiceNow ITSM",
        "phases": [
            {"id": "preflight", "label": "Check existing setup"},
            {
                "id": "entra-registration",
                "label": "Collect Entra identity options",
            },
            {"id": "topics", "label": "Prepare ServiceNow topics"},
        ],
    }
    initial_state = {
        "phases": {
            "preflight": {"status": "pending", "checkpointResults": {}},
            "entra-registration": {
                "status": "pending",
                "checkpointResults": {},
            },
            "topics": {"status": "pending", "checkpointResults": {}},
        }
    }
    transitioned_state = {
        "phases": {
            "preflight": {
                "status": "done",
                "checkpointResults": {"PRE": "Passed"},
            },
            "entra-registration": {
                "status": "in-progress",
                "checkpointResults": {},
            },
            "topics": {"status": "pending", "checkpointResults": {}},
        }
    }

    initial = snow._build_phase_progress(
        contract,
        initial_state,
        {"decision": "run-action", "phase": {"id": "preflight"}},
        executed_actions=frozenset(),
        verified_phases=frozenset(),
    )
    transitioned = snow._build_phase_progress(
        contract,
        transitioned_state,
        {
            "decision": "run-action",
            "phase": {"id": "entra-registration"},
        },
        executed_actions=frozenset(),
        verified_phases=frozenset({"preflight"}),
    )

    assert "- 🔄 Check existing setup — In progress" in initial["checklistText"]
    assert "- ⬜ Confirm Entra identity options — Pending" in (
        initial["checklistText"]
    )
    assert "- ✅ Check existing setup — Complete — checked" in (
        transitioned["checklistText"]
    )
    assert "- 🔄 Confirm Entra identity options — In progress" in (
        transitioned["checklistText"]
    )
    assert transitioned["currentPhase"]["questionText"].startswith(
        "**Progress**\n\nConnection progress:\n\n- ✅"
    )
    assert "- ✅ Check existing setup — Complete — checked" in (
        transitioned["currentPhase"]["questionText"]
    )
    assert "- ⏳ Confirm Entra identity options — Waiting" in (
        transitioned["currentPhase"]["questionText"]
    )
    assert "- ❌ Confirm Entra identity options — Needs attention" in (
        transitioned["currentPhase"]["errorChecklistText"]
    )


def test_phase_progress_final_text_requires_every_current_verified_marker() -> None:
    contract = {
        "displayName": "ServiceNow ITSM",
        "phases": [
            {"id": "test", "label": "Test the agent"},
            {"id": "publish", "label": "Publish the agent"},
        ],
    }
    state = {
        "phases": {
            "test": {
                "status": "done",
                "checkpointResults": {"TEST": "Manual"},
            },
            "publish": {
                "status": "done",
                "checkpointResults": {"PUBLISH": "Passed"},
            },
        }
    }

    incomplete = snow._build_phase_progress(
        contract,
        state,
        {"decision": "complete"},
        executed_actions=frozenset(),
        verified_phases=frozenset({"test"}),
    )
    complete = snow._build_phase_progress(
        contract,
        state,
        {"decision": "complete"},
        executed_actions=frozenset(),
        verified_phases=frozenset({"test", "publish"}),
    )

    assert incomplete["finalText"] is None
    assert incomplete["items"][1]["status"] == "rechecking"
    assert complete["finalText"] == (
        "Connection progress:\n\n"
        "- ✅ Test ServiceNow — Complete — admin confirmed\n"
        "- ✅ Publish agent — Complete — checked\n\n"
        "Complete — every ServiceNow ITSM step was checked in this run."
    )
    assert complete["items"][0]["text"] == "Complete — admin confirmed"
    assert complete["items"][1]["text"] == "Complete — checked"


@pytest.mark.parametrize(
    "provider",
    ("servicenow-da-hrsd", "servicenow-da-itsm"),
)
def test_portal_phase_reenters_action_when_checkpoint_is_incomplete(
    tmp_path: Path,
    provider: str,
) -> None:
    source = (
        snow.SOLUTION_ROOT
        / "src"
        / "skills"
        / "connect"
        / provider
        / "contract.json"
    )
    contract = json.loads(source.read_text(encoding="utf-8"))
    contract_path = (
        tmp_path
        / "src"
        / "skills"
        / "connect"
        / provider
        / "contract.json"
    )
    contract_path.parent.mkdir(parents=True)
    contract_path.write_text(json.dumps(contract), encoding="utf-8")
    agent_slug = "employee-self-service-it"
    phases = {}
    verified = set()
    before_portal = True
    for phase in contract["phases"]:
        phase_id = phase["id"]
        if phase_id == "portal-url":
            before_portal = False
            phases[phase_id] = {
                "status": "in-progress",
                "actionApplied": True,
                "checkpointResults": {
                    phase["checkpoints"][0]: "NotConfigured"
                },
            }
        elif before_portal:
            phases[phase_id] = {
                "status": "done",
                "checkpointResults": {
                    phase["checkpoints"][0]: "Passed"
                },
            }
            verified.add(phase_id)
        else:
            phases[phase_id] = {
                "status": "pending",
                "checkpointResults": {},
            }
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / provider
        / "agents"
        / agent_slug
        / "lifecycle.json"
    )
    state_path.parent.mkdir(parents=True)
    lifecycle_state = {
        "provider": provider,
        "agentSlug": agent_slug,
        "attested": True,
        "acceptedContractRevision": contract["contractRevision"],
        "phases": phases,
    }
    state_path.write_text(json.dumps(lifecycle_state), encoding="utf-8")

    decision = snow.decide_next_step(
        tmp_path,
        provider=provider,
        agent_slug=agent_slug,
        verified_phases=frozenset(verified),
    )

    assert decision["decision"] == "run-action"
    assert decision["phase"]["id"] == "portal-url"
    assert decision["phase"]["resumeActionWhenIncomplete"] is True
    progress = snow._build_phase_progress(
        contract,
        lifecycle_state,
        decision,
        executed_actions=frozenset(),
        verified_phases=frozenset(verified),
    )
    assert progress["currentPhase"]["customerLabel"] == (
        "Set employee portal URL"
    )
    assert "- ❌ Set employee portal URL — Needs attention" in (
        progress["currentPhase"]["errorChecklistText"]
    )


def _write_resume_contract(root: Path) -> None:
    contract_path = (
        root
        / "src/skills/connect/servicenow-da-itsm/contract.json"
    )
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(
        json.dumps(
            {
                "contractRevision": 1,
                "provider": "servicenow-da-itsm",
                "displayName": "ServiceNow ITSM",
                "phases": [
                    {
                        "id": "preflight",
                        "label": "Check existing setup",
                        "checkpoints": ["PRE-001"],
                    },
                    {
                        "id": "topics",
                        "label": "Prepare ServiceNow topics",
                        "checkpoints": ["SN-DA-ITSM-TOPICS-001"],
                        "actionDoc": (
                            "src/skills/connect/servicenow-da-itsm/"
                            "actions/prepare-topics.md"
                        ),
                        "actionExecution": "once",
                        "resumeActionWhenIncomplete": True,
                    },
                ],
            }
        ),
        encoding="utf-8",
    )


def test_resume_handoff_cli_returns_plan_required_without_creating_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    _seed_it_setup(tmp_path, enrolled=False)
    _write_resume_contract(tmp_path)
    connect_root = tmp_path / ".local" / "connect"
    monkeypatch.chdir(tmp_path)

    assert snow.main(
        [
            "--provider",
            "servicenow-da-itsm",
            "--agent-slug",
            "employee-self-service-it",
            "resume-handoff",
        ]
    ) == 0

    result = json.loads(capsys.readouterr().out)
    assert result["lifecycle"] == {
        "schemaVersion": "connect.lifecycle-next.v1",
        "decision": "plan-required",
        "provider": "servicenow-da-itsm",
        "agentSlug": "employee-self-service-it",
        "contractRevision": 1,
        "stateExists": False,
        "planReason": "first-run",
        "stateInitializationRequired": True,
        "stateInitializationCommand": None,
    }
    assert "action" not in result
    assert not connect_root.exists()


def _write_resume_contract_and_state(root: Path) -> None:
    _write_resume_contract(root)
    state_path = (
        root
        / ".local/connect/servicenow-da-itsm/agents/"
        "employee-self-service-it/lifecycle.json"
    )
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(
        json.dumps(
            {
                "schemaVersion": snow.LIFECYCLE_SCHEMA_VERSION,
                "provider": "servicenow-da-itsm",
                "profile": "itsm",
                "agentSlug": "employee-self-service-it",
                "agentId": AGENT_ID,
                "agentSchemaName": snow.IT_SCHEMA_NAME,
                "environmentId": ENVIRONMENT_ID,
                "ring": "test",
                "attested": True,
                "acceptedContractRevision": 1,
                "phases": {
                    "preflight": {
                        "status": "done",
                        "checkpointResults": {"PRE-001": "Passed"},
                    },
                    "topics": {
                        "status": "in-progress",
                        "actionApplied": True,
                        "lastActionAt": "2026-10-08T01:20:00Z",
                        "checkpointResults": {
                            "SN-DA-ITSM-TOPICS-001": "NotConfigured"
                        },
                    },
                },
                "transactions": {
                    "topics": {
                        "operation-1": {
                            "status": "committed",
                            "after": {"total": 13, "active": 13},
                        }
                    }
                },
                "evidence": {},
                "adminSetup": snow._empty_admin_setup(),
                "migration": {},
            }
        ),
        encoding="utf-8",
    )


def test_itsm_lifecycle_state_is_isolated_from_hrsd(
    tmp_path: Path,
) -> None:
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {"id": ENVIRONMENT_ID, "ring": "test"},
    }
    components = _itsm_components()

    with snow.product_profile_context("itsm"):
        state = snow._state_base(context, components)
        path = snow._lifecycle_state_path(context)

    assert state["provider"] == "servicenow-da-itsm"
    assert state["profile"] == "itsm"
    assert path == (
        Path(".local")
        / "connect"
        / "servicenow-da-itsm"
        / "agents"
        / "employee-self-service-it"
        / "lifecycle.json"
    )
    assert "servicenow-da-hrsd" not in str(path)


def test_agent_chooser_lists_prepared_agents_without_remote_discovery(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    setup_path, config_path = _seed_it_setup(tmp_path)
    setup_before = setup_path.read_text(encoding="utf-8")
    config_before = config_path.read_text(encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: pytest.fail("chooser must not create a remote client"),
    )

    result = snow.agent_chooser_handoff(tmp_path)

    assert result["status"] == "choice-required"
    assert result["agents"][0]["id"] == AGENT_ID
    assert result["agents"][0]["status"] == "prepared"
    assert "prepared in this workspace" in result["question"]["question"]
    assert result["question"]["allowFreeformInput"] is False
    assert result["question"]["options"][-1] == {
        "label": "Use another agent (paste link)",
        "selectionKey": "another-agent",
    }
    selected = snow.agent_chooser_handoff(
        tmp_path,
        selection_key=f"prepared:{AGENT_ID}",
    )
    assert selected["status"] == "ready"
    assert selected["selectedAgent"]["provider"] == "servicenow-da-itsm"
    assert setup_path.read_text(encoding="utf-8") == setup_before
    assert config_path.read_text(encoding="utf-8") == config_before
    assert not (tmp_path / ".local" / "connect").exists()


def test_agent_chooser_offers_link_when_no_agent_is_prepared(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    setup_path, config_path = _seed_it_setup(tmp_path)
    setup = json.loads(setup_path.read_text(encoding="utf-8"))
    setup["agents"] = {}
    setup_path.write_text(json.dumps(setup), encoding="utf-8")
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["agents"] = []
    config_path.write_text(json.dumps(config), encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: pytest.fail("chooser must not create a remote client"),
    )

    result = snow.agent_chooser_handoff(tmp_path)

    assert result["agents"] == []
    assert result["question"]["options"] == [
        {
            "label": "Use another agent (paste link)",
            "selectionKey": "another-agent",
        }
    ]
    assert "No Employee Self-Service agents are prepared" in (
        result["question"]["question"]
    )
    link_question = snow.agent_chooser_handoff(
        tmp_path,
        selection_key="another-agent",
    )
    assert link_question["question"] == {
        "header": "Use another agent",
        "question": (
            "Paste the full Copilot Studio link for the existing agent you "
            "want to use."
        ),
        "options": [],
        "allowFreeformInput": True,
    }


def test_agent_chooser_returns_typed_setup_handoff_for_new_workspace_link(
    tmp_path: Path,
) -> None:
    setup_path, config_path = _seed_it_setup(tmp_path)
    setup_before = setup_path.read_text(encoding="utf-8")
    config_before = config_path.read_text(encoding="utf-8")
    hr_id = "00000000-0000-4000-8000-000000009998"
    link = (
        "https://copilotstudio.test.microsoft.com/environments/"
        f"{ENVIRONMENT_ID}/bots/{hr_id}"
    )

    result = snow.agent_chooser_handoff(
        tmp_path,
        selection_key="another-agent",
        agent_link=link,
    )

    assert result["status"] == "setup-required"
    assert result["branch"] == "explicit-agent-link"
    assert result["setupHandoff"] == {
        "command": "/setup",
        "targetUrl": link,
        "environmentId": ENVIRONMENT_ID,
        "agentId": hr_id,
        "ring": "test",
        "requiresEnvironmentConfirmation": True,
        "requiresAccountConfirmation": True,
    }
    assert setup_path.read_text(encoding="utf-8") == setup_before
    assert config_path.read_text(encoding="utf-8") == config_before
    assert not (tmp_path / ".local" / "connect").exists()


def test_agent_chooser_resolves_prepared_agent_link_without_duplicate_setup(
    tmp_path: Path,
) -> None:
    setup_path, config_path = _seed_it_setup(tmp_path)
    config_before = config_path.read_text(encoding="utf-8")
    link = (
        "https://copilotstudio.test.microsoft.com/environments/"
        f"{ENVIRONMENT_ID}/bots/{AGENT_ID}"
    )

    result = snow.agent_chooser_handoff(
        tmp_path,
        selection_key="another-agent",
        agent_link=link,
    )

    assert result["status"] == "ready"
    assert result["selectedAgent"]["id"] == AGENT_ID
    assert result["selectedAgent"]["provider"] == "servicenow-da-itsm"
    assert config_path.read_text(encoding="utf-8") == config_before
    assert setup_path.exists()
    assert not (tmp_path / ".local" / "connect").exists()


def test_agent_chooser_rejects_invalid_or_cross_environment_link_without_write(
    tmp_path: Path,
) -> None:
    setup_path, config_path = _seed_it_setup(tmp_path)
    setup_before = setup_path.read_text(encoding="utf-8")
    config_before = config_path.read_text(encoding="utf-8")

    for link, message in (
        (
            "https://example.com/environments/"
            f"{ENVIRONMENT_ID}/bots/{AGENT_ID}",
            "full Copilot Studio link",
        ),
        (
            "https://copilotstudio.test.microsoft.com/environments/"
            "00000000-0000-4000-8000-000000009997"
            f"/bots/{AGENT_ID}",
            "different environment",
        ),
    ):
        with pytest.raises(snow.ServiceNowConnectError, match=message):
            snow.agent_chooser_handoff(
                tmp_path,
                selection_key="another-agent",
                agent_link=link,
            )

    assert setup_path.read_text(encoding="utf-8") == setup_before
    assert config_path.read_text(encoding="utf-8") == config_before
    assert not (tmp_path / ".local" / "connect").exists()


def test_list_agents_exact_fallback_selects_it_without_changing_active(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _setup_path, config_path = _seed_it_setup(tmp_path)
    original = config_path.read_text(encoding="utf-8")

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, agent_id: str) -> dict:
            assert agent_id == AGENT_ID
            return {
                "bot": {
                    "cdsBotId": AGENT_ID,
                    "displayName": "Employee Self-Service (IT)",
                    "schemaName": snow.IT_SCHEMA_NAME,
                }
            }

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"] == [
        {
            "id": AGENT_ID,
            "slug": "employee-self-service-it",
            "name": "Employee Self-Service (IT)",
            "profile": "itsm",
            "provider": "servicenow-da-itsm",
            "product": "ServiceNow ITSM",
            "status": "ready",
            "selectable": True,
        }
    ]
    assert config_path.read_text(encoding="utf-8") == original


def test_list_agents_exact_fallback_accepts_unenrolled_it_authoring(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path, enrolled=False)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, agent_id: str) -> dict:
            assert agent_id == AGENT_ID
            return {
                "bot": {
                    "cdsBotId": AGENT_ID,
                    "displayName": "Employee Self-Service (IT)",
                    "schemaName": snow.IT_SCHEMA_NAME,
                }
            }

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"][0]["id"] == AGENT_ID
    assert result["agents"][0]["profile"] == "itsm"
    assert result["agents"][0]["selectable"] is True
    assert not (tmp_path / ".local" / "connect").exists()


def test_list_agents_exact_fallback_has_hr_parity(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _write_foundation_context(tmp_path)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, agent_id: str) -> dict:
            assert agent_id == AGENT_ID
            return {
                "bot": {
                    "cdsBotId": AGENT_ID,
                    "displayName": "Employee Self-Service (HR)",
                    "schemaName": snow.HR_SCHEMA_NAME,
                }
            }

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"][0]["profile"] == "hrsd"
    assert result["agents"][0]["provider"] == "servicenow-da-hrsd"
    assert result["agents"][0]["selectable"] is True


def test_list_agents_exact_mismatch_is_not_selectable(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, _agent_id: str) -> dict:
            return {
                "bot": {
                    "cdsBotId": "00000000-0000-4000-8000-000000009999",
                    "displayName": "Employee Self-Service (IT)",
                    "schemaName": snow.IT_SCHEMA_NAME,
                }
            }

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"] == []
    assert result["discoveryErrors"][0]["status"] == "lookup-failed"
    assert "different agent identity" in result["discoveryErrors"][0]["message"]


def test_list_agents_exact_failure_is_actionable(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, _agent_id: str) -> dict:
            raise snow.AgentBuilderHTTPError(
                "Component fetch",
                403,
                request_id="request-id",
            )

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"] == []
    assert result["discoveryErrors"][0]["status"] == "not-authorized"
    assert "request request-id" in result["discoveryErrors"][0]["message"]


@pytest.mark.parametrize(
    ("status_code", "expected_status"),
    [
        (401, "authentication-required"),
        (404, "not-found"),
        (429, "throttled"),
        (503, "service-unavailable"),
    ],
)
def test_list_agents_exact_http_failures_are_explicit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    status_code: int,
    expected_status: str,
) -> None:
    _seed_it_setup(tmp_path, enrolled=False)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, _agent_id: str) -> dict:
            raise snow.AgentBuilderHTTPError(
                "Component fetch",
                status_code,
                request_id="request-id",
            )

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"] == []
    assert result["discoveryErrors"][0]["status"] == expected_status


def test_list_agents_exact_unknown_schema_is_not_selectable(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path, enrolled=False)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return []

        def fetch_components(self, _agent_id: str) -> dict:
            return {
                "bot": {
                    "cdsBotId": AGENT_ID,
                    "displayName": "Contoso Help",
                    "schemaName": "gptagent_contosohelp",
                }
            }

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert result["agents"] == []
    assert result["discoveryErrors"][0]["status"] == "lookup-failed"
    assert "supported Employee Self-Service" in (
        result["discoveryErrors"][0]["message"]
    )


def test_list_agents_rejects_environment_mismatch_before_remote_call(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _setup_path, config_path = _seed_it_setup(tmp_path, enrolled=False)
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["environmentId"] = "00000000-0000-4000-8000-000000009998"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: pytest.fail("remote client must not be created"),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="different environment environmentId values",
    ):
        snow.list_agent_choices(tmp_path)


def test_list_agents_rejects_account_mismatch_before_remote_call(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(
        tmp_path,
        enrolled=False,
        canonical_account="expected@example.test",
        operational_account="other@example.test",
    )
    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: pytest.fail("remote client must not be created"),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="different Microsoft accounts",
    ):
        snow.list_agent_choices(tmp_path)


def test_load_context_accepts_canonical_unenrolled_authoring(
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path, enrolled=False)

    with snow.product_profile_context("itsm"):
        context = snow.load_context(tmp_path, agent_id=AGENT_ID)

    assert context["agent"]["alm"] == {"isEnrolled": False}
    assert context["agentSetup"]["authoring_ready"] is True


def test_list_agents_shows_unconfigured_live_it_agent(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path)
    other_id = "00000000-0000-4000-8000-000000009999"

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return [
                {
                    "botId": AGENT_ID,
                    "fullBotName": "Employee Self-Service (IT)",
                    "realm": "dev",
                },
                {
                    "botId": other_id,
                    "fullBotName": "Employee Self-Service IT",
                    "realm": "dev",
                },
            ]

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)
    unconfigured = next(
        choice for choice in result["agents"] if choice["id"] == other_id
    )

    assert unconfigured["profile"] == "itsm"
    assert unconfigured["status"] == "setup-required"
    assert unconfigured["selectable"] is False


def test_list_agents_shows_unconfigured_live_hr_agent_without_local_setup(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path)
    hr_id = "00000000-0000-4000-8000-000000009998"

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return [
                {
                    "botId": AGENT_ID,
                    "fullBotName": "Employee Self-Service (IT)",
                    "realm": "dev",
                },
                {
                    "botId": hr_id,
                    "fullBotName": "Employee Self-Service (HR)",
                    "schemaName": snow.HR_SCHEMA_NAME,
                    "realm": "dev",
                },
            ]

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)
    unconfigured = next(
        choice for choice in result["agents"] if choice["id"] == hr_id
    )

    assert unconfigured["profile"] == "hrsd"
    assert unconfigured["status"] == "setup-required"
    assert unconfigured["selectable"] is False


def test_list_agents_deduplicates_collection_without_exact_fallback(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path, enrolled=False)

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return [
                {
                    "botId": AGENT_ID,
                    "fullBotName": "Employee Self-Service (IT)",
                    "schemaName": snow.IT_SCHEMA_NAME,
                },
                {
                    "botId": AGENT_ID.upper(),
                    "fullBotName": "Duplicate card",
                    "schemaName": snow.IT_SCHEMA_NAME,
                },
            ]

        def fetch_components(self, _agent_id: str) -> dict:
            pytest.fail("listed canonical identity must not use exact fallback")

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)

    assert [agent["id"] for agent in result["agents"]] == [AGENT_ID]
    assert result["agents"][0]["selectable"] is True
    assert not (tmp_path / ".local" / "connect").exists()


def test_list_agents_uses_live_schema_for_renamed_unconfigured_agent(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _seed_it_setup(tmp_path)
    other_id = "00000000-0000-4000-8000-000000009999"

    class FakeAgentBuilder:
        def list_agents(self) -> list[dict]:
            return [
                {
                    "botId": AGENT_ID,
                    "fullBotName": "Employee Self-Service (IT)",
                    "realm": "dev",
                },
                {
                    "botId": other_id,
                    "fullBotName": "Contoso IT Help",
                    "schemaName": snow.IT_SCHEMA_NAME,
                    "realm": "dev",
                },
            ]

    monkeypatch.setattr(
        snow,
        "_agentbuilder_client",
        lambda _context: FakeAgentBuilder(),
    )

    result = snow.list_agent_choices(tmp_path)
    renamed = next(
        choice for choice in result["agents"] if choice["id"] == other_id
    )

    assert renamed["name"] == "Contoso IT Help"
    assert renamed["profile"] == "itsm"
    assert renamed["status"] == "setup-required"
    assert renamed["selectable"] is False


def test_itsm_profile_does_not_import_legacy_hr_state(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {"id": ENVIRONMENT_ID, "ring": "test"},
    }
    legacy = tmp_path / snow._legacy_state_path(AGENT_ID)
    legacy.parent.mkdir(parents=True)
    legacy.write_text(
        json.dumps(
            {
                "schemaVersion": 1,
                "agentId": AGENT_ID,
                "environmentId": ENVIRONMENT_ID,
                "topicEnablement": {"customerChoice": "enable-all"},
            }
        ),
        encoding="utf-8",
    )

    with snow.product_profile_context("itsm"):
        state = snow._load_lifecycle_state(context, _itsm_components())

    assert state["provider"] == "servicenow-da-itsm"
    assert "legacySourcePath" not in state["migration"]
    assert "topics" not in state["evidence"]


def test_itsm_migrate_state_rejects_without_creating_state(
    tmp_path: Path,
) -> None:
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": "employee-self-service-it",
        },
        "active": {"slug": "employee-self-service-it"},
        "environment": {"id": ENVIRONMENT_ID, "ring": "test"},
    }

    with snow.product_profile_context("itsm"):
        state_path = snow._lifecycle_state_path(context)
        with pytest.raises(
            snow.ServiceNowConnectError,
            match="supported only for the HRSD provider",
        ):
            snow.migrate_state(context)

    assert not (tmp_path / state_path).exists()


def test_connectivity_discovery_timeout_is_single_attempt_and_actionable() -> None:
    calls = []

    class FakeSession:
        def request(self, *args, **kwargs):
            calls.append((args, kwargs))
            raise snow.requests.Timeout("slow read")

    client = snow.ConnectivityClient(
        "https://example.environment.api.test.powerplatform.com",
        ENVIRONMENT_ID,
        "token",
        ring="test",
        session=FakeSession(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match="60-second HTTP connect/read timeout",
    ) as raised:
        client.list_connections()

    assert len(calls) == 1
    assert calls[0][1]["timeout"] == snow.CONNECTIVITY_READ_TIMEOUT_SECONDS
    assert raised.value.details == {
        "category": "timeout",
        "timeoutSeconds": 60,
        "mutationMayHaveOccurred": False,
        "retryable": True,
    }
    assert "No remote write was attempted" in str(raised.value)
    assert "not a whole-invocation deadline" in str(raised.value)


def test_connectivity_throttle_stops_without_retry_and_preserves_retry_after() -> None:
    calls = []

    class FakeResponse:
        ok = False
        status_code = 429
        headers = {
            "Retry-After": "45",
            "x-ms-request-id": "request-123",
        }

        @staticmethod
        def json():
            return {"error": {"code": "TooManyRequests"}}

    class FakeSession:
        def request(self, *args, **kwargs):
            calls.append((args, kwargs))
            return FakeResponse()

    client = snow.ConnectivityClient(
        "https://example.environment.api.test.powerplatform.com",
        ENVIRONMENT_ID,
        "token",
        ring="test",
        session=FakeSession(),
    )

    with pytest.raises(
        snow.ServiceNowConnectError,
        match=r"HTTP 429.*retry-after 45",
    ) as raised:
        client.list_connections()

    assert len(calls) == 1
    assert raised.value.details == {
        "category": "throttled",
        "statusCode": 429,
        "errorCode": "TooManyRequests",
        "retryAfter": "45",
        "requestId": "request-123",
        "mutationMayHaveOccurred": False,
        "retryable": True,
    }


@pytest.mark.parametrize("profile", ("hrsd", "itsm"))
def test_inspect_admin_timeout_preserves_error_and_allows_manual_target(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    profile: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    context = _context()
    components = _components()
    if profile == "itsm":
        context = {
            **context,
            "agent": {
                **context["agent"],
                "schema_name": snow.IT_SCHEMA_NAME,
                "workspace_slug": "employee-self-service-it",
            },
            "active": {"slug": "employee-self-service-it"},
        }
        components = _itsm_components()
    with snow.product_profile_context(profile):
        state = snow._state_base(context, components)
        state_path = snow._lifecycle_state_path(context)
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(json.dumps(state), encoding="utf-8")
        class FakeConnectivity:
            def list_connections(self):
                raise snow.ServiceNowConnectError(
                    "ServiceNow connection discovery reached the configured "
                    "60-second HTTP connect/read timeout.",
                    details={
                        "category": "timeout",
                        "mutationMayHaveOccurred": False,
                    },
                )

        monkeypatch.setattr(
            snow,
            "_connectivity_client",
            lambda _context: FakeConnectivity(),
        )

        discovery = snow.inspect_admin_setup(context)

    assert discovery["connectionCandidates"] == []
    assert discovery["instanceCandidates"] == []
    assert discovery["instanceInputRequired"] is True
    assert discovery["instanceQuestion"]["allowFreeformInput"] is True
    assert discovery["discoveryError"]["details"] == {
        "category": "timeout",
        "mutationMayHaveOccurred": False,
    }
    assert "not an empty inventory" in (
        discovery["instanceQuestion"]["question"]
    )
    assert "not an empty inventory" in (
        discovery["manualInstanceQuestion"]["question"]
    )
    assert "**Problem**" not in discovery["instanceQuestion"]["question"]
    assert "ask your ServiceNow administrator" in (
        discovery["instanceQuestion"]["question"]
    )
    saved = json.loads(state_path.read_text(encoding="utf-8"))
    assert saved["adminSetup"]["preflight"]["discovery"][
        "discoveryError"
    ] == discovery["discoveryError"]
    assert saved["adminSetup"]["phaseHandoffs"]["preflight"] == {
        "status": "pending"
    }


def test_inspect_admin_inventory_marker_describes_transport_not_wall_clock(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    state = snow._state_base(_context(), _components())
    state_path = _lifecycle_path(tmp_path)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps(state), encoding="utf-8")

    class FakeConnectivity:
        def list_connections(self):
            return []

    monkeypatch.setattr(
        snow,
        "_connectivity_client",
        lambda _context: FakeConnectivity(),
    )

    snow.inspect_admin_setup(_context())

    marker = next(
        line
        for line in capsys.readouterr().out.splitlines()
        if line.startswith("DA_ADMIN_PREFLIGHT_PROGRESS_JSON:")
    )
    progress = json.loads(marker.split(":", 1)[1])
    assert progress == {
        "stage": "connection-inventory",
        "state": "starting",
        "mode": "read-only",
        "httpConnectReadTimeoutSeconds": 60,
        "wholeInvocationDeadline": None,
    }


def test_main_emits_authentication_marker_before_admin_discovery(
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr(snow, "load_context", lambda: _context())
    monkeypatch.setattr(
        snow,
        "inspect_admin_setup",
        lambda _context: {"status": "ok"},
    )

    assert snow.main(["inspect-admin-setup"]) == 0

    lines = capsys.readouterr().out.splitlines()
    marker = next(
        line
        for line in lines
        if line.startswith("DA_ADMIN_PREFLIGHT_PROGRESS_JSON:")
    )
    assert json.loads(marker.split(":", 1)[1]) == {
        "stage": "authentication",
        "state": "starting",
        "mode": "read-only-preparation",
    }
