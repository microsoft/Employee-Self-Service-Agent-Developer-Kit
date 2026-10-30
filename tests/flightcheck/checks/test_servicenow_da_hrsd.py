# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.

from __future__ import annotations

from dataclasses import asdict
import json
from pathlib import Path
import threading
import time
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from agentbuilder import AgentBuilderHTTPError
import flightcheck.checks.servicenow_da_hrsd as hrsd_checks
from flightcheck.checks.servicenow_da_hrsd import (
    _acquire_servicenow_reads,
    invalidate_servicenow_read_snapshot,
    run_servicenow_da_hrsd_checks,
)
from flightcheck.runner import FlightCheckRunner, Status

import connect_servicenow_da as snow
from tests.conftest import require_validated_mock
from tests.mocks import graph as graph_mock


ENVIRONMENT_ID = "00000000-0000-4000-8000-000000001111"
TENANT_ID = "00000000-0000-4000-8000-000000009999"
AGENT_ID = "00000000-0000-4000-8000-000000002222"
CONNECTION_ID = "00000000000040008000000000003333"
APP_CLIENT_ID = "00000000-0000-4000-8000-000000006666"
AGENT_SLUG = "employee-self-service-hr"
_RECOVERY_HEADINGS = (
    "**Problem**",
    "**Category**",
    "**Owner**",
    "**Where**",
    "**Values to use**",
    "**Action**",
    "**Success**",
    "**Recheck and resume**",
)


def _components() -> dict:
    return {
        "changeToken": "token",
        "botComponentChanges": [
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000005555",
                    "version": 1,
                    "displayName": "ServiceNow HRSD Get Cases",
                    "schemaName": (
                        f"{snow.HR_SCHEMA_NAME}.topic.ServiceNowHRSDGetCases"
                    ),
                    "state": "Active",
                    "status": "Active",
                    "dialog": {"$kind": "TaskDialog"},
                },
            },
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000005556",
                    "version": 2,
                    "displayName": "ServiceNow HRSD Setup Configurations",
                    "schemaName": (
                        f"{snow.HR_SCHEMA_NAME}.topic."
                        "ServiceNowHRSDSetupConfigurations"
                    ),
                    "state": "Active",
                    "status": "Active",
                    "dialog": {
                        "$kind": "TaskDialog",
                        "beginDialog": {
                            "actions": [
                                {
                                    "$kind": "SetVariable",
                                    "displayName": (
                                        "Set ServiceNow Portal BaseURI"
                                    ),
                                    "variable": (
                                        "Global.ServiceNowHRSDPortalBaseURI"
                                    ),
                                    "value": {
                                        "$kind": "Literal",
                                        "literalValue": (
                                            "https://dev123.service-now.com/"
                                            "employee"
                                        ),
                                    },
                                }
                            ]
                        },
                    },
                },
            },
        ],
        "connectionReferenceChanges": [
            {
                "connectionReference": {
                    "id": "00000000-0000-4000-8000-000000004444",
                    "connectorId": snow.CONNECTOR_ID,
                    "connectionReferenceLogicalName": "servicenow-ref",
                    "displayName": "ServiceNow",
                    "connectionId": CONNECTION_ID,
                    "sharedConnectionParameters": json.dumps(
                        {
                            "name": "entraIDUserLogin",
                            "values": {
                                "token:InstanceName": {"value": "dev123"},
                                "token:ResourceUri": {
                                    "value": APP_CLIENT_ID
                                },
                            },
                        }
                    ),
                }
            }
        ],
        "connectorDefinitionChanges": [],
    }


def test_all_hrsd_blocking_phase_outputs_include_actionable_recovery() -> None:
    results = run_servicenow_da_hrsd_checks(
        SimpleNamespace(config={}, checkpoint_target=None)
    )

    expected_suffixes = {
        "PKG",
        "ADMIN-PREFLIGHT",
        "PLUGIN",
        "ENTRA-APP",
        "ENTRA-CLAIMS",
        "ENTRA-SCOPE",
        "ENTRA-PREAUTH",
        "OIDC",
        "TOPICS",
        "PORTAL",
        "CREDENTIAL",
        "AGENT-CONNECTION",
        "TEST",
        "PUBLISH",
    }
    assert {
        result.checkpoint_id.removeprefix("SN-DA-HRSD-").removesuffix("-001")
        for result in results
    } == expected_suffixes
    for result in results:
        assert result.status == Status.FAILED.value
        assert result.result in result.remediation
        assert all(
            heading in result.remediation for heading in _RECOVERY_HEADINGS
        )
        assert "{" not in result.remediation
        assert "dependent later steps remain pending" in result.remediation


@pytest.mark.parametrize(
    ("status", "problem", "category", "retry"),
    (
        (
            Status.ERROR.value,
            "HTTP 429 throttled",
            "Read throttled",
            "Retry-After",
        ),
        (
            Status.ERROR.value,
            "HTTP 401 authentication required",
            "Read permission",
            "minimum read permission",
        ),
        (
            Status.ERROR.value,
            "HTTP 403 permission denied",
            "Read permission",
            "minimum read permission",
        ),
        (
            Status.ERROR.value,
            "Socket timeout",
            "Read or network error",
            "Do not treat the error as missing configuration",
        ),
        (
            Status.NOT_CONFIGURED.value,
            "The value is absent",
            "Missing or repairable configuration",
            "Repair the current value",
        ),
    ),
)
def test_recovery_output_preserves_error_category(
    status: str,
    problem: str,
    category: str,
    retry: str,
) -> None:
    result = hrsd_checks._result(
        "SN-DA-HRSD-PORTAL-001",
        status,
        "Portal",
        problem,
        "Repair the current value",
    )

    assert f"**Category**\n{category}" in result.remediation
    assert retry in result.remediation


def _connection() -> dict:
    return {
        "name": CONNECTION_ID,
        "properties": {
            "apiId": snow.CONNECTOR_ID,
            "displayName": "ServiceNow",
            "statuses": [{"target": "token", "status": "Connected"}],
            "connectionParametersSet": {
                "name": "entraIDUserLogin",
                "values": {
                    "token:InstanceName": {"value": "dev123"},
                    "token:ResourceUri": {"value": APP_CLIENT_ID},
                },
            },
        },
    }


def _runner(components: dict):
    return SimpleNamespace(
        config={
            "activeAgent": AGENT_SLUG,
            "agents": [
                {
                    "slug": AGENT_SLUG,
                    "botId": AGENT_ID,
                    "schemaName": snow.HR_SCHEMA_NAME,
                }
            ],
        },
        agentbuilder=SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
        connectivity=SimpleNamespace(
            list_connections=lambda _environment_id: pytest.fail(
                "ServiceNow HRSD must not use environment-wide inventory."
            ),
            list_connector_connections=lambda environment_id, connector: (
                [_connection()]
                if (
                    environment_id == ENVIRONMENT_ID
                    and connector == snow.CONNECTOR_NAME
                )
                else pytest.fail("Unexpected connector inventory target.")
            ),
        ),
        env_id=ENVIRONMENT_ID,
    )


def test_servicenow_read_acquisition_overlaps_two_independent_sources() -> None:
    lock = threading.Lock()
    barrier = threading.Barrier(2)
    active = 0
    max_active = 0
    calls = {"components": 0, "connections": 0}

    def read(name: str, value):
        nonlocal active, max_active
        with lock:
            calls[name] += 1
            active += 1
            max_active = max(max_active, active)
        barrier.wait(timeout=1)
        time.sleep(0.05)
        with lock:
            active -= 1
        return value

    runner = SimpleNamespace(
        agentbuilder=SimpleNamespace(
            fetch_components=lambda _agent_id: read(
                "components",
                _components(),
            )
        ),
        connectivity=SimpleNamespace(
            list_connector_connections=lambda _env_id, _connector: read(
                "connections",
                [_connection()],
            )
        ),
    )

    started = time.perf_counter()
    components, component_error, connections, connectivity_error = (
        _acquire_servicenow_reads(
            runner,
            agent_id=AGENT_ID,
            environment_id=ENVIRONMENT_ID,
            needs_inventory=True,
        )
    )
    elapsed = time.perf_counter() - started

    assert component_error is None
    assert connectivity_error is None
    assert components == _components()
    assert connections == [_connection()]
    assert calls == {"components": 1, "connections": 1}
    assert max_active == 2
    assert elapsed < 0.09


def test_servicenow_read_snapshot_reuses_one_exact_components_read() -> None:
    calls = {"components": 0, "connections": 0}
    runner = SimpleNamespace(
        servicenow_parallel_reads=False,
        agentbuilder=SimpleNamespace(
            fetch_components=lambda _agent_id: (
                calls.__setitem__("components", calls["components"] + 1)
                or _components()
            )
        ),
        connectivity=SimpleNamespace(
            list_connector_connections=lambda _env_id, _connector: (
                calls.__setitem__("connections", calls["connections"] + 1)
                or [_connection()]
            )
        ),
    )

    for _ in range(2):
        _acquire_servicenow_reads(
            runner,
            agent_id=AGENT_ID,
            environment_id=ENVIRONMENT_ID,
            needs_inventory=True,
        )

    assert calls == {"components": 1, "connections": 1}


def test_servicenow_parallel_read_output_matches_serial_evaluation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.chdir(tmp_path)
    serial = _runner(_components())
    serial.checkpoint_target = "SN-DA-HRSD-TEST-001"
    serial.servicenow_parallel_reads = False
    parallel = _runner(_components())
    parallel.checkpoint_target = "SN-DA-HRSD-TEST-001"

    serial_rows = run_servicenow_da_hrsd_checks(serial)
    parallel_rows = run_servicenow_da_hrsd_checks(parallel)

    assert [asdict(row) for row in parallel_rows] == [
        asdict(row) for row in serial_rows
    ]
    assert not (tmp_path / ".local").exists()


def test_servicenow_throttle_quiesces_new_read_dispatch() -> None:
    calls = {"components": 0, "connections": 0}

    def throttled(_agent_id: str):
        calls["components"] += 1
        raise AgentBuilderHTTPError("Component fetch", 429)

    runner = SimpleNamespace(
        agentbuilder=SimpleNamespace(fetch_components=throttled),
        connectivity=SimpleNamespace(
            list_connector_connections=lambda _env_id, _connector: (
                calls.__setitem__("connections", calls["connections"] + 1)
                or [_connection()]
            )
        ),
    )

    first = _acquire_servicenow_reads(
        runner,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        needs_inventory=True,
    )
    first_calls = dict(calls)
    second = _acquire_servicenow_reads(
        runner,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        needs_inventory=True,
    )
    invalidate_servicenow_read_snapshot(runner)
    third = _acquire_servicenow_reads(
        runner,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        needs_inventory=True,
    )

    assert isinstance(first[1], AgentBuilderHTTPError)
    assert runner._servicenow_read_quiesced is True
    assert calls == first_calls
    assert second[0] is None
    assert isinstance(second[1], AgentBuilderHTTPError)
    assert second[1].status_code == 429
    assert isinstance(third[1], AgentBuilderHTTPError)
    assert third[1].status_code == 429


def test_servicenow_serial_throttle_stops_second_prepared_dispatch() -> None:
    calls = {"components": 0, "connections": 0}

    def throttled(_agent_id: str):
        calls["components"] += 1
        raise AgentBuilderHTTPError("Component fetch", 429)

    runner = SimpleNamespace(
        servicenow_parallel_reads=False,
        agentbuilder=SimpleNamespace(fetch_components=throttled),
        connectivity=SimpleNamespace(
            list_connector_connections=lambda _env_id, _connector: (
                calls.__setitem__("connections", calls["connections"] + 1)
                or [_connection()]
            )
        ),
    )

    result = _acquire_servicenow_reads(
        runner,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        needs_inventory=True,
    )

    assert calls == {"components": 1, "connections": 0}
    assert isinstance(result[1], AgentBuilderHTTPError)
    assert isinstance(result[3], AgentBuilderHTTPError)
    assert result[1].status_code == result[3].status_code == 429


def test_servicenow_read_snapshot_invalidates_at_external_boundary() -> None:
    calls = {"components": 0, "connections": 0}
    runner = SimpleNamespace(
        servicenow_parallel_reads=False,
        agentbuilder=SimpleNamespace(
            fetch_components=lambda _agent_id: (
                calls.__setitem__("components", calls["components"] + 1)
                or _components()
            )
        ),
        connectivity=SimpleNamespace(
            list_connector_connections=lambda _env_id, _connector: (
                calls.__setitem__("connections", calls["connections"] + 1)
                or [_connection()]
            )
        ),
    )

    _acquire_servicenow_reads(
        runner,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        needs_inventory=True,
    )
    invalidate_servicenow_read_snapshot(runner)
    _acquire_servicenow_reads(
        runner,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        needs_inventory=True,
    )

    assert calls == {"components": 2, "connections": 2}


def _registered_runner(
    components: dict,
    *,
    target_test_only: bool = False,
) -> FlightCheckRunner:
    source = _runner(components)
    runner = FlightCheckRunner(
        scope="servicenow-da-hrsd",
        target_matcher=(
            lambda checkpoint_id: checkpoint_id == "SN-DA-HRSD-TEST-001"
        )
        if target_test_only
        else None,
        checkpoint_target=(
            "SN-DA-HRSD-TEST-001" if target_test_only else None
        ),
    )
    for name in ("config", "agentbuilder", "connectivity", "env_id"):
        setattr(runner, name, getattr(source, name))
    runner.register("ServiceNow DA HRSD", run_servicenow_da_hrsd_checks)
    return runner


def _checkpoint(results, checkpoint_id: str):
    return next(
        row for row in results if row.checkpoint_id == checkpoint_id
    )


def _graph_client(
    *,
    missing_claims: bool = False,
    optional_claim_names: tuple[str, ...] | None = None,
    foreign_consent: bool = False,
):
    require_validated_mock(graph_mock)
    application = graph_mock.application(
        app_id=APP_CLIENT_ID,
        connector_app_id=snow.SERVICENOW_CONNECTOR_APP_ID,
        optional_claim_names=(
            optional_claim_names
            if optional_claim_names is not None
            else ("email",)
            if missing_claims
            else ("email", "upn")
        ),
    )
    service_principal = graph_mock.service_principal(
        app_id=APP_CLIENT_ID,
        sp_id="00000000-0000-4000-8000-000000007777",
        application_template_id=None,
    )
    graph_service_principal = graph_mock.service_principal(
        app_id=graph_mock.MS_GRAPH_RESOURCE_APP_ID,
        sp_id="00000000-0000-4000-8000-000000008888",
        application_template_id=None,
    )
    grant = graph_mock.oauth2_permission_grant(
        client_id=service_principal["id"],
        resource_id=(
            "00000000-0000-4000-8000-000000009999"
            if foreign_consent
            else graph_service_principal["id"]
        ),
    )

    class FakeGraph:
        def get_all(
            self,
            path: str,
            params: dict | None = None,
            *,
            raise_on_permission_error: bool = False,
        ) -> list[dict]:
            del params, raise_on_permission_error
            if path == "/applications":
                return [application]
            if path == "/oauth2PermissionGrants":
                return [grant]
            raise AssertionError(f"Unexpected Graph path: {path}")

        def get_service_principals(
            self,
            *,
            filter_expr: str = "",
            **_kwargs,
        ) -> list[dict]:
            if graph_mock.MS_GRAPH_RESOURCE_APP_ID in filter_expr:
                return [graph_service_principal]
            return [service_principal]

    return FakeGraph()


def _state_path(root: Path) -> Path:
    return (
        root
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )


def _write_state(
    root: Path,
    components: dict,
    *,
    candidates: tuple[str, ...] = ("upn",),
    selected_claim: str = "upn",
    user_field: str = "user_name",
) -> None:
    path = _state_path(root)
    path.parent.mkdir(parents=True)
    component_hash = snow._component_hash(components)
    discovery = {
        "source": "connectivity-readonly",
        "observedAt": "2026-09-29T00:00:00Z",
        "connectionCandidates": [],
        "requiredPlugins": [],
    }
    handoffs = {
        phase: {
            "status": "completed",
            "evidence": {
                "kind": "structured-admin-attestation",
                "recordedAt": "2026-09-29T00:00:00Z",
            },
        }
        for phase in snow.ADMIN_PHASES
    }
    handoffs["entra-registration"]["evidence"].update(
        {
            "clientId": APP_CLIENT_ID,
            "tokenClaimCandidates": list(candidates),
            "tokenClaimEmissionConfirmed": True,
        }
    )
    handoffs["plugin-prerequisites"]["evidence"].update(
        {
            "pluginId": "com.sn_hr_core",
            "applicationTitle": "Human Resources Scoped App: Core",
            "versionDetails": "Installed",
            "installedReadyMessageConfirmed": True,
        }
    )
    handoffs["servicenow-oidc"]["evidence"].update(
        {
            "oidcCapabilityConfirmed": True,
            "runbookCompleted": True,
            "mappingDetailsCollected": True,
            "claimMappingConfirmed": True,
            "tokenClaimName": selected_claim,
            "serviceNowUserField": user_field,
        }
    )
    draft_semantic_hash = snow._draft_semantic_hash(components)
    agent_connection = {
        "kind": "maker-attestation",
        "status": "completed",
        "makerAttested": True,
        "physicalStatus": "Connected",
        "connectionId": CONNECTION_ID,
        "recordedAt": "2026-09-29T00:00:00Z",
        "binding": {
            "draftSemanticHash": draft_semantic_hash,
            "connectionId": CONNECTION_ID,
        },
    }
    state = {
        "schemaVersion": 6,
        "provider": snow.PROVIDER_KEY,
        "profile": snow.PROFILE_KEY,
        "agentSlug": AGENT_SLUG,
        "agentId": AGENT_ID,
        "tenantId": TENANT_ID,
        "environmentId": ENVIRONMENT_ID,
        "draftSemanticHash": draft_semantic_hash,
        "adminSetup": {
            "schemaVersion": 4,
            "scope": "hrsd",
            "authMode": "entraIDUserLogin",
            "preflight": {
                "instanceName": "dev123",
                "discovery": discovery,
            },
            "phaseHandoffs": handoffs,
        },
        "evidence": {
            "identityClaimCandidates": {
                "kind": "administrator-confirmed-claim-candidates",
                "status": "entra-confirmed",
                "applicationClientId": APP_CLIENT_ID,
                "candidates": list(candidates),
                "entraAdministratorConfirmed": True,
                "tokenClaimEmissionConfirmed": True,
                "binding": {
                    "tenantId": TENANT_ID,
                    "environmentId": ENVIRONMENT_ID,
                    "instanceName": "dev123",
                    "applicationClientId": APP_CLIENT_ID,
                    "provider": snow.PROVIDER_KEY,
                    "profile": snow.PROFILE_KEY,
                    "agentSlug": AGENT_SLUG,
                    "agentId": AGENT_ID,
                },
            },
            "identityMapping": {
                "kind": "administrator-confirmed-identity-mapping",
                "status": "servicenow-confirmed",
                "applicationClientId": APP_CLIENT_ID,
                "tokenClaimName": selected_claim,
                "serviceNowUserField": user_field,
                "entraAdministratorConfirmed": True,
                "tokenClaimEmissionConfirmed": True,
                "serviceNowAdministratorConfirmed": True,
                "binding": {
                    "tenantId": TENANT_ID,
                    "environmentId": ENVIRONMENT_ID,
                    "instanceName": "dev123",
                    "applicationClientId": APP_CLIENT_ID,
                    "provider": snow.PROVIDER_KEY,
                    "profile": snow.PROFILE_KEY,
                    "agentSlug": AGENT_SLUG,
                    "agentId": AGENT_ID,
                },
            },
            "credential": {"connectionId": CONNECTION_ID},
            "agentConnection": agent_connection,
        },
    }
    candidate_hash = snow._identity_candidate_set_hash(
        application_client_id=APP_CLIENT_ID,
        candidates=list(candidates),
        binding=state["evidence"]["identityClaimCandidates"]["binding"],
    )
    state["evidence"]["identityClaimCandidates"]["candidateSetSha256"] = (
        candidate_hash
    )
    state["evidence"]["identityMapping"]["candidateSetSha256"] = candidate_hash
    relationship_hash = snow._identity_mapping_relationship_hash(
        state["evidence"]["identityMapping"]
    )
    state["evidence"]["identityMapping"]["relationshipSha256"] = (
        relationship_hash
    )
    state["adminSetup"]["phaseHandoffs"]["entra-registration"]["evidence"][
        "candidateSetSha256"
    ] = candidate_hash
    state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"]["evidence"][
        "mappingRelationshipSha256"
    ] = relationship_hash
    for phase in ("entra-registration", "servicenow-oidc"):
        state["adminSetup"]["phaseHandoffs"][phase]["evidence"]["binding"] = {
            "tenantId": TENANT_ID,
            "environmentId": ENVIRONMENT_ID,
            "instanceName": "dev123",
            "applicationClientId": APP_CLIENT_ID,
            "provider": snow.PROVIDER_KEY,
            "profile": snow.PROFILE_KEY,
            "agentSlug": AGENT_SLUG,
            "agentId": AGENT_ID,
        }
    connection_binding_hash = snow._connection_binding_hash(
        {
            "agent": {
                "id": AGENT_ID,
                "workspace_slug": AGENT_SLUG,
            },
            "environment": {"id": ENVIRONMENT_ID},
        },
        state,
        components,
    )
    state["connectionBindingHash"] = connection_binding_hash
    state["evidence"].update(
        {
            "publish": {
                "status": "completed",
                "componentHash": component_hash,
                "publishedComponentHash": component_hash,
                "testedDraftSemanticHash": draft_semantic_hash,
                "publishedSemanticHash": draft_semantic_hash,
                "connectionBindingHash": connection_binding_hash,
                "completedAt": "2026-09-28T00:00:00Z",
            },
            "test": {
                "kind": "maker-attestation",
                "status": "completed",
                "promptCategory": "list-my-open-hr-cases",
                "result": "pass",
                "failureCategory": None,
                "binding": {
                    "draftSemanticHash": draft_semantic_hash,
                    "connectionBindingHash": connection_binding_hash,
                },
            },
        }
    )
    path.write_text(
        json.dumps(state),
        encoding="utf-8",
    )


def _record_new_mapping_workflow(
    root: Path,
    *,
    candidates: tuple[str, ...],
    selected_claim: str,
    user_field: str,
) -> None:
    context = {
        "root": root,
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.HR_SCHEMA_NAME,
            "workspace_slug": AGENT_SLUG,
        },
        "active": {"slug": AGENT_SLUG},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": TENANT_ID,
            "ring": "test",
        },
    }
    state_path = _state_path(root)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state.setdefault("phases", {})["entra-registration"] = {
        "status": "in-progress",
        "checkpointResults": {
            "SN-DA-HRSD-ENTRA-CLAIMS-001": Status.NOT_CONFIGURED.value
        },
    }
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "pending"
    }
    state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] = {
        "status": "pending"
    }
    state.setdefault("evidence", {}).pop("identityClaimCandidates", None)
    state["evidence"].pop("identityMapping", None)
    state_path.write_text(json.dumps(state), encoding="utf-8")
    snow.record_identity_requirement(context, choice="email")
    entra = snow.identity_mapping_handoff(
        context,
        stage="entra",
        client_id=APP_CLIENT_ID,
    )
    entra_path = Path(entra["workbookPath"])
    workbook = load_workbook(entra_path)
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    for offset, claim in enumerate(candidates):
        sheet.cell(
            snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW + offset,
            1,
        ).value = claim
    workbook.save(entra_path)
    workbook.close()
    snow.record_identity_mapping(
        context,
        stage="entra",
        workbook_path=entra_path,
        client_id=APP_CLIENT_ID,
        confirmed=True,
    )
    snow.record_admin_phase(
        context,
        phase="entra-registration",
        status="completed",
        client_id=APP_CLIENT_ID,
    )
    service_now = snow.identity_mapping_handoff(
        context,
        stage="servicenow",
    )
    service_now_path = Path(service_now["workbookPath"])
    workbook = load_workbook(service_now_path)
    sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
    for row in range(
        snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW,
        sheet.max_row + 1,
    ):
        if sheet.cell(row, 1).value == selected_claim:
            sheet.cell(row, 2).value = user_field
    workbook.save(service_now_path)
    workbook.close()
    snow.record_identity_mapping(
        context,
        stage="servicenow",
        workbook_path=service_now_path,
        client_id=None,
        confirmed=True,
    )
    snow.record_admin_phase(
        context,
        phase="servicenow-oidc",
        status="completed",
    )


def test_hrsd_checks_preserve_manual_maker_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    statuses = {result.checkpoint_id: result.status for result in results}

    assert statuses["SN-DA-HRSD-PKG-001"] == Status.PASSED.value
    assert statuses["SN-DA-HRSD-TOPICS-001"] == Status.PASSED.value
    assert statuses["SN-DA-HRSD-PORTAL-001"] == Status.PASSED.value
    assert statuses["SN-DA-HRSD-CREDENTIAL-001"] == Status.PASSED.value
    assert (
        statuses["SN-DA-HRSD-AGENT-CONNECTION-001"]
        == Status.MANUAL.value
    )
    assert "SN-DA-HRSD-PARAMETER-SHARING-001" not in statuses
    assert statuses["SN-DA-HRSD-PUBLISH-001"] == Status.PASSED.value
    assert statuses["SN-DA-HRSD-TEST-001"] == Status.MANUAL.value


def test_hrsd_package_and_topics_remain_dynamic_with_one_matching_topic() -> None:
    components = _components()
    components["botComponentChanges"] = components["botComponentChanges"][:1]

    package = next(
        result
        for result in run_servicenow_da_hrsd_checks(_runner(components))
        if result.checkpoint_id == "SN-DA-HRSD-PKG-001"
    )
    topics = next(
        result
        for result in run_servicenow_da_hrsd_checks(_runner(components))
        if result.checkpoint_id == "SN-DA-HRSD-TOPICS-001"
    )

    assert package.status == Status.PASSED.value
    assert package.result.endswith("contains 1 ServiceNow HRSD topic(s).")
    assert topics.status == Status.PASSED.value
    assert topics.result == "1/1 ServiceNow HRSD topics are Active."


@pytest.mark.parametrize(
    ("case", "expected_phrase"),
    [
        ("disconnected", "not a Connected Microsoft Entra ID User Login"),
        ("wrong-auth", "not a Connected Microsoft Entra ID User Login"),
        ("wrong-instance", "different Instance Name or Resource URI"),
        ("wrong-resource", "different Instance Name or Resource URI"),
    ],
)
def test_test_checkpoint_propagates_remediable_credential_state(
    monkeypatch,
    tmp_path: Path,
    case: str,
    expected_phrase: str,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    connection = _connection()
    if case == "disconnected":
        connection["properties"]["statuses"] = [
            {"target": "token", "status": "Disconnected"}
        ]
    elif case == "wrong-auth":
        connection["properties"]["connectionParametersSet"]["name"] = (
            "keyBasedAuth"
        )
    elif case == "wrong-instance":
        connection["properties"]["connectionParametersSet"]["values"][
            "token:InstanceName"
        ]["value"] = "dev456"
    else:
        connection["properties"]["connectionParametersSet"]["values"][
            "token:ResourceUri"
        ]["value"] = "00000000-0000-4000-8000-000000009999"
    runner = _runner(components)
    runner.connectivity = SimpleNamespace(
        list_connector_connections=lambda _environment_id, _connector: [
            connection
        ]
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    credential = _checkpoint(results, "SN-DA-HRSD-CREDENTIAL-001")
    test = _checkpoint(results, "SN-DA-HRSD-TEST-001")

    assert credential.status == Status.NOT_CONFIGURED.value
    assert test.status == Status.NOT_CONFIGURED.value
    assert "Current credential prerequisite is NotConfigured" in test.result
    assert expected_phrase in test.result
    assert test.remediation == credential.remediation


def test_test_checkpoint_propagates_connectivity_error(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    runner = _runner(components)
    runner.connectivity = SimpleNamespace(
        list_connector_connections=lambda *_args: (_ for _ in ()).throw(
            RuntimeError("synthetic inventory failure")
        )
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    credential = _checkpoint(results, "SN-DA-HRSD-CREDENTIAL-001")
    test = _checkpoint(results, "SN-DA-HRSD-TEST-001")

    assert credential.status == Status.ERROR.value
    assert test.status == Status.ERROR.value
    assert "Current credential prerequisite is Error" in test.result
    assert "synthetic inventory failure" in test.result
    assert test.remediation == credential.remediation


@pytest.mark.parametrize(
    ("case", "prerequisite_label"),
    [
        ("credential-pending", "credential"),
        ("topics-inactive", "topics"),
        ("agent-binding-stale", "Agent Connect"),
    ],
)
def test_test_checkpoint_blocks_noncompleted_prerequisite(
    monkeypatch,
    tmp_path: Path,
    case: str,
    prerequisite_label: str,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    if case == "credential-pending":
        state["adminSetup"]["phaseHandoffs"]["credential"] = {
            "status": "pending"
        }
    elif case == "topics-inactive":
        components["botComponentChanges"][0]["component"]["state"] = "Inactive"
        components["botComponentChanges"][0]["component"]["status"] = "Inactive"
    else:
        state["evidence"]["agentConnection"]["binding"][
            "draftSemanticHash"
        ] = "stale-draft"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    test = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-TEST-001",
    )

    assert test.status == Status.NOT_CONFIGURED.value
    assert (
        f"Current {prerequisite_label} prerequisite is NotConfigured"
        in test.result
    )


def test_test_checkpoint_keeps_healthy_metadata_only_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    components["changeToken"] = "server-metadata-only"
    monkeypatch.chdir(tmp_path)

    test = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-TEST-001",
    )

    assert test.status == Status.MANUAL.value
    assert test.result == "Maker recorded a passing HRSD Test pane result."


def test_test_checkpoint_rejects_stale_test_connection_binding(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["evidence"]["test"]["binding"]["connectionBindingHash"] = (
        "stale-connection-binding"
    )
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    test = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-TEST-001",
    )

    assert test.status == Status.NOT_CONFIGURED.value
    assert "selected connection" in test.result


@pytest.mark.parametrize("disconnected", [False, True])
def test_test_checkpoint_full_and_fixed_runs_match(
    monkeypatch,
    tmp_path: Path,
    disconnected: bool,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    full_runner = _registered_runner(components)
    fixed_runner = _registered_runner(components, target_test_only=True)
    if disconnected:
        connection = _connection()
        connection["properties"]["statuses"] = [
            {"target": "token", "status": "Disconnected"}
        ]
        inventory = SimpleNamespace(
            list_connector_connections=lambda *_args: [connection]
        )
        full_runner.connectivity = inventory
        fixed_runner.connectivity = inventory
    monkeypatch.chdir(tmp_path)

    full = full_runner.run()
    fixed = fixed_runner.run()
    full_test = _checkpoint(full.results, "SN-DA-HRSD-TEST-001")
    fixed_test = _checkpoint(fixed.results, "SN-DA-HRSD-TEST-001")

    assert vars(fixed_test) == vars(full_test)
    assert fixed.results == [fixed_test]
    assert fixed.overall == full.overall


def test_admin_prerequisites_pass_with_graph_and_structured_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    runner = _runner(components)
    runner.graph = _graph_client()
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    statuses = {result.checkpoint_id: result.status for result in results}

    assert statuses["SN-DA-HRSD-ADMIN-PREFLIGHT-001"] == Status.PASSED.value
    assert statuses["SN-DA-HRSD-PLUGIN-001"] == Status.MANUAL.value
    for suffix in (
        "APP",
        "SCOPE",
        "PREAUTH",
    ):
        assert (
            statuses[f"SN-DA-HRSD-ENTRA-{suffix}-001"]
            == Status.PASSED.value
        )
    assert statuses["SN-DA-HRSD-ENTRA-CLAIMS-001"] == Status.MANUAL.value
    assert statuses["SN-DA-HRSD-OIDC-001"] == Status.MANUAL.value


@pytest.mark.parametrize(
    ("candidates", "optional_claim_names"),
    [
        (("upn",), ("upn",)),
        (("email",), ("email",)),
        (("employee_id",), ()),
        (("email", "upn", "employee_id"), ("email", "upn")),
    ],
)
def test_claim_checkpoint_uses_confirmed_candidate_set(
    monkeypatch,
    tmp_path: Path,
    candidates: tuple[str, ...],
    optional_claim_names: tuple[str, ...],
) -> None:
    components = _components()
    _write_state(
        tmp_path,
        components,
        candidates=candidates,
        selected_claim=candidates[0],
    )
    runner = _runner(components)
    runner.graph = _graph_client(
        optional_claim_names=optional_claim_names,
    )
    monkeypatch.chdir(tmp_path)

    result = _checkpoint(
        run_servicenow_da_hrsd_checks(runner),
        "SN-DA-HRSD-ENTRA-CLAIMS-001",
    )

    assert result.status == Status.MANUAL.value
    assert all(candidate in result.result for candidate in candidates)
    assert "administrator confirmed actual claim emission" in result.result
    assert result.remediation == ""


def test_claim_checkpoint_rejects_missing_or_mismatched_confirmation(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["phaseHandoffs"]["entra-registration"]["evidence"][
        "tokenClaimEmissionConfirmed"
    ] = False
    state_path.write_text(json.dumps(state), encoding="utf-8")
    runner = _runner(components)
    runner.graph = _graph_client(optional_claim_names=("upn",))
    monkeypatch.chdir(tmp_path)

    missing = _checkpoint(
        run_servicenow_da_hrsd_checks(runner),
        "SN-DA-HRSD-ENTRA-CLAIMS-001",
    )
    assert missing.status == Status.NOT_CONFIGURED.value
    assert "claim candidates" in missing.result
    assert "record the emitted employee identity claim candidates" in (
        missing.remediation
    )

    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["phaseHandoffs"]["entra-registration"]["evidence"][
        "tokenClaimEmissionConfirmed"
    ] = True
    state["adminSetup"]["phaseHandoffs"]["entra-registration"]["evidence"][
        "tokenClaimCandidates"
    ] = ["email"]
    state_path.write_text(json.dumps(state), encoding="utf-8")

    mismatched = _checkpoint(
        run_servicenow_da_hrsd_checks(runner),
        "SN-DA-HRSD-ENTRA-CLAIMS-001",
    )
    assert mismatched.status == Status.NOT_CONFIGURED.value
    assert "does not match the Entra handoff" in (
        mismatched.result
    )


def test_oidc_checkpoint_requires_capability_owned_by_oidc_phase(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"]["evidence"].pop(
        "oidcCapabilityConfirmed"
    )
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = next(
        row
        for row in run_servicenow_da_hrsd_checks(_runner(components))
        if row.checkpoint_id == "SN-DA-HRSD-OIDC-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "OIDC capability confirmation is missing" in result.result


def test_oidc_checkpoint_rejects_legacy_mapping_details(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    evidence = state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
        "evidence"
    ]
    evidence["claim"] = "upn"
    evidence["userField"] = "user_name"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = next(
        row
        for row in run_servicenow_da_hrsd_checks(_runner(components))
        if row.checkpoint_id == "SN-DA-HRSD-OIDC-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "privacy-minimized" in result.result


def test_oidc_checkpoint_rejects_legacy_bounded_shape_without_two_names(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    evidence = state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
        "evidence"
    ]
    evidence["mappingDetailsCollected"] = False
    evidence.pop("claimMappingConfirmed")
    evidence.pop("tokenClaimName")
    evidence.pop("serviceNowUserField")
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-OIDC-001",
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "legacy OIDC attestation" in result.result
    assert "two-name mapping handoff" in result.remediation


def test_recorded_mapping_workflow_reaches_real_hr_checkpoints(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    monkeypatch.chdir(tmp_path)
    _record_new_mapping_workflow(
        tmp_path,
        candidates=("email", "upn", "employee_id"),
        selected_claim="employee_id",
        user_field="employee_number",
    )
    runner = _runner(components)
    runner.graph = _graph_client(optional_claim_names=())

    results = run_servicenow_da_hrsd_checks(runner)
    claims = _checkpoint(results, "SN-DA-HRSD-ENTRA-CLAIMS-001")
    oidc = _checkpoint(results, "SN-DA-HRSD-OIDC-001")

    assert claims.status == Status.MANUAL.value
    assert "employee_id" in claims.result
    assert oidc.status == Status.MANUAL.value
    assert "employee_id" in oidc.result
    assert "employee_number" in oidc.result
    assert oidc.remediation == ""


@pytest.mark.parametrize(
    ("portal_value", "expected_status", "expected_phrase"),
    [
        ("", Status.NOT_CONFIGURED.value, "does not contain"),
        (
            "https://other.service-now.com/sp",
            Status.FAILED.value,
            "different ServiceNow instance",
        ),
        (
            "http://dev123.service-now.com/sp",
            Status.FAILED.value,
            "not a safe plain HTTPS URL",
        ),
    ],
)
def test_portal_checkpoint_fails_closed_on_invalid_configuration(
    monkeypatch,
    tmp_path: Path,
    portal_value: str,
    expected_status: str,
    expected_phrase: str,
) -> None:
    components = _components()
    node = components["botComponentChanges"][1]["component"]["dialog"][
        "beginDialog"
    ]["actions"][0]
    node["value"]["literalValue"] = portal_value
    _write_state(tmp_path, components)
    monkeypatch.chdir(tmp_path)

    result = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )

    assert result.status == expected_status
    assert expected_phrase in result.result
    assert result.remediation


def test_portal_checkpoint_requires_explicit_root_url_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    components["botComponentChanges"][1]["component"]["dialog"][
        "beginDialog"
    ]["actions"][0]["value"]["literalValue"] = (
        "https://dev123.service-now.com"
    )
    _write_state(tmp_path, components)
    monkeypatch.chdir(tmp_path)

    unconfirmed = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )
    assert unconfirmed.status == Status.NOT_CONFIGURED.value
    assert "no evidence" in unconfirmed.result

    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["transactions"] = {
        "portal": {
            "operation-root": {
                "preparedAt": "2026-10-05T18:00:00Z",
                "status": "rolled-back",
                "requestedPortalUrl": "https://dev123.service-now.com",
            }
        }
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")
    rolled_back = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )
    assert rolled_back.status == Status.NOT_CONFIGURED.value

    state["transactions"]["portal"]["operation-root"]["status"] = "committed"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    legacy_transaction = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )
    assert legacy_transaction.status == Status.NOT_CONFIGURED.value

    topic = snow._portal_topic_component(components)
    snow._record_portal_confirmation(
        state,
        topic=topic,
        portal_url="https://dev123.service-now.com",
        instance_name="dev123",
        source="already-configured",
    )
    state_path.write_text(json.dumps(state), encoding="utf-8")
    confirmed = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )
    assert confirmed.status == Status.PASSED.value
    assert "confirmed ServiceNow instance" in confirmed.result


def test_test_checkpoint_requires_current_portal_configuration(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    components["botComponentChanges"][1]["component"]["dialog"][
        "beginDialog"
    ]["actions"][0]["value"]["literalValue"] = ""
    _write_state(tmp_path, components)
    monkeypatch.chdir(tmp_path)

    test = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-TEST-001",
    )

    assert test.status == Status.NOT_CONFIGURED.value
    assert "portal Base URI prerequisite" in test.result


def test_portal_checkpoint_requires_latest_requested_url(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["transactions"] = {
        "portal": {
            "operation-1": {
                "preparedAt": "2026-10-05T18:00:00Z",
                "status": "failed-unchanged",
                "requestedPortalUrl": (
                    "https://dev123.service-now.com/requested-portal"
                ),
            }
        }
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    mismatch = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )
    assert mismatch.status == Status.NOT_CONFIGURED.value
    assert "does not match" in mismatch.result
    assert "exact requested" in mismatch.remediation

    components["botComponentChanges"][1]["component"]["dialog"][
        "beginDialog"
    ]["actions"][0]["value"]["literalValue"] = (
        "https://dev123.service-now.com/requested-portal"
    )
    matching = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PORTAL-001",
    )
    assert matching.status == Status.PASSED.value


def test_publish_checkpoint_ignores_metadata_only_drift_and_reopens_on_semantic_drift(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    current_hash = snow._component_hash(components)
    semantic_hash = snow._draft_semantic_hash(components)
    state["evidence"]["publish"] = {
        "status": "completed",
        "requestedComponentHash": "pre-publish-hash",
        "publishedComponentHash": current_hash,
        "componentHash": current_hash,
        "testedDraftSemanticHash": semantic_hash,
        "publishedSemanticHash": semantic_hash,
        "connectionBindingHash": state["connectionBindingHash"],
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    current = run_servicenow_da_hrsd_checks(_runner(components))
    publish = next(
        row
        for row in current
        if row.checkpoint_id == "SN-DA-HRSD-PUBLISH-001"
    )
    assert publish.status == Status.PASSED.value
    target_runner = _runner(components)
    target_runner.checkpoint_target = "SN-DA-HRSD-PUBLISH-001"
    targeted = run_servicenow_da_hrsd_checks(target_runner)
    assert targeted[0].status == Status.PASSED.value

    components["changeToken"] = "later-edit"
    metadata_only = run_servicenow_da_hrsd_checks(_runner(components))
    publish = next(
        row
        for row in metadata_only
        if row.checkpoint_id == "SN-DA-HRSD-PUBLISH-001"
    )
    assert publish.status == Status.NOT_CONFIGURED.value

    state["evidence"]["publish"]["publishedComponentHash"] = snow._component_hash(
        components
    )
    state["evidence"]["publish"]["componentHash"] = snow._component_hash(
        components
    )
    state_path.write_text(json.dumps(state), encoding="utf-8")
    metadata_only = run_servicenow_da_hrsd_checks(_runner(components))
    publish = next(
        row
        for row in metadata_only
        if row.checkpoint_id == "SN-DA-HRSD-PUBLISH-001"
    )
    assert publish.status == Status.PASSED.value

    components["botComponentChanges"][0]["component"]["displayName"] = (
        "Edited ServiceNow topic"
    )
    semantic_drift = run_servicenow_da_hrsd_checks(_runner(components))
    publish = next(
        row
        for row in semantic_drift
        if row.checkpoint_id == "SN-DA-HRSD-PUBLISH-001"
    )
    assert publish.status == Status.NOT_CONFIGURED.value


def test_publish_checkpoint_surfaces_current_semantic_drift_as_failed(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["evidence"]["publish"] = {
        "kind": "maker-attestation",
        "status": "needs_remediation",
        "publishedComponentHash": snow._component_hash(components),
        "componentHash": snow._component_hash(components),
        "testedDraftSemanticHash": "pre-publish-semantic-hash",
        "publishedSemanticHash": snow._draft_semantic_hash(components),
        "remediation": "Run Test again before any later publish.",
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = next(
        row
        for row in run_servicenow_da_hrsd_checks(_runner(components))
        if row.checkpoint_id == "SN-DA-HRSD-PUBLISH-001"
    )

    assert result.status == Status.FAILED.value
    assert "needs remediation" in result.result


@pytest.mark.parametrize(
    "test_record",
    [
        {
            "prompt": "Show case ABC-123 for Jane Doe",
            "details": "Sensitive response",
            "result": "pass",
        },
        {
            "promptCategory": "unknown",
            "result": "pass",
            "failureCategory": None,
        },
        {
            "promptCategory": "list-my-open-hr-cases",
            "result": "pass",
            "failureCategory": "permission",
        },
    ],
)
def test_test_checkpoint_rejects_legacy_or_unbounded_evidence(
    monkeypatch,
    tmp_path: Path,
    test_record: dict,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    current_binding = state["evidence"]["test"]["binding"]
    test_record["binding"] = {
        **current_binding,
    }
    state["evidence"]["test"] = test_record
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = next(
        row
        for row in run_servicenow_da_hrsd_checks(_runner(components))
        if row.checkpoint_id == "SN-DA-HRSD-TEST-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "privacy-safe bounded evidence contract" in result.result


def test_test_checkpoint_accepts_bounded_failure_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    current_binding = state["evidence"]["test"]["binding"]
    state["evidence"]["test"] = {
        "kind": "maker-attestation",
        "status": "failed",
        "promptCategory": "list-my-open-hr-cases",
        "result": "fail",
        "failureCategory": "permission",
        "binding": current_binding,
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = next(
        row
        for row in run_servicenow_da_hrsd_checks(_runner(components))
        if row.checkpoint_id == "SN-DA-HRSD-TEST-001"
    )

    assert result.status == Status.FAILED.value


def test_test_checkpoint_accepts_maker_failure_without_category(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    current_binding = state["evidence"]["test"]["binding"]
    state["evidence"]["test"] = {
        "kind": "maker-attestation",
        "status": "failed",
        "promptCategory": "list-my-open-hr-cases",
        "result": "fail",
        "failureCategory": None,
        "binding": current_binding,
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    test_result = _checkpoint(results, "SN-DA-HRSD-TEST-001")
    publish = _checkpoint(results, "SN-DA-HRSD-PUBLISH-001")

    assert test_result.status == Status.FAILED.value
    assert publish.status == Status.FAILED.value
    assert "maker recorded Not pass" in publish.result

    target_runner = _runner(components)
    target_runner.checkpoint_target = "SN-DA-HRSD-PUBLISH-001"
    targeted = run_servicenow_da_hrsd_checks(target_runner)
    assert targeted[0].status == Status.FAILED.value


@pytest.mark.parametrize("result", ["pass", "fail"])
def test_publish_rejects_test_bound_to_previous_connection(
    monkeypatch,
    tmp_path: Path,
    result: str,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = _state_path(tmp_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["evidence"]["test"]["result"] = result
    state["evidence"]["test"]["status"] = (
        "completed" if result == "pass" else "failed"
    )
    state["evidence"]["test"]["failureCategory"] = None
    state["evidence"]["credential"]["connectionId"] = (
        "00000000000040008000000000009999"
    )
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    full_publish = _checkpoint(
        run_servicenow_da_hrsd_checks(_runner(components)),
        "SN-DA-HRSD-PUBLISH-001",
    )
    target_runner = _runner(components)
    target_runner.checkpoint_target = "SN-DA-HRSD-PUBLISH-001"
    targeted = run_servicenow_da_hrsd_checks(target_runner)

    assert full_publish.status == Status.NOT_CONFIGURED.value
    assert targeted[0].status == Status.NOT_CONFIGURED.value
    assert "current passing Test pane result" in targeted[0].result


def test_entra_phase_evaluation_uses_one_logical_read_set(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    counters = {
        "agentbuilder": 0,
        "connectivity": 0,
        "applications": 0,
        "app_service_principal": 0,
        "graph_service_principal": 0,
        "permission_grants": 0,
    }
    runner = _runner(components)

    def fetch_components(_agent_id: str) -> dict:
        counters["agentbuilder"] += 1
        return components

    def list_connector_connections(
        environment_id: str,
        connector: str,
    ) -> list[dict]:
        assert environment_id == ENVIRONMENT_ID
        assert connector == snow.CONNECTOR_NAME
        counters["connectivity"] += 1
        return [_connection()]

    base_graph = _graph_client()

    class CountingGraph:
        def get_all(
            self,
            path: str,
            params: dict | None = None,
            *,
            raise_on_permission_error: bool = False,
        ) -> list[dict]:
            if path == "/applications":
                counters["applications"] += 1
            elif path == "/oauth2PermissionGrants":
                counters["permission_grants"] += 1
            return base_graph.get_all(
                path,
                params,
                raise_on_permission_error=raise_on_permission_error,
            )

        def get_service_principals(
            self,
            *,
            filter_expr: str = "",
            **kwargs,
        ) -> list[dict]:
            if graph_mock.MS_GRAPH_RESOURCE_APP_ID in filter_expr:
                counters["graph_service_principal"] += 1
            else:
                counters["app_service_principal"] += 1
            return base_graph.get_service_principals(
                filter_expr=filter_expr,
                **kwargs,
            )

    runner.agentbuilder = SimpleNamespace(fetch_components=fetch_components)
    runner.connectivity = SimpleNamespace(
        list_connections=lambda _environment_id: pytest.fail(
            "ServiceNow HRSD must not use environment-wide inventory."
        ),
        list_connector_connections=list_connector_connections,
    )
    runner.graph = CountingGraph()
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    entra_results = [
        result
        for result in results
        if result.checkpoint_id.startswith("SN-DA-HRSD-ENTRA-")
    ]

    assert len(entra_results) == 4
    assert counters == {
        "agentbuilder": 1,
        "connectivity": 1,
        "applications": 1,
        "app_service_principal": 0,
        "graph_service_principal": 0,
        "permission_grants": 0,
    }


@pytest.mark.parametrize(
    ("target", "expected_rows", "inventory_reads", "graph_reads"),
    [
        ("SN-DA-HRSD-PKG-001", 1, 0, 0),
        ("SN-DA-HRSD-ADMIN-PREFLIGHT-001", 1, 0, 0),
        ("SN-DA-HRSD-PLUGIN-001", 1, 0, 0),
        ("SN-DA-HRSD-ENTRA-*", 4, 0, 1),
        ("SN-DA-HRSD-ENTRA-CLAIMS-001", 1, 0, 1),
        ("SN-DA-HRSD-OIDC-001", 1, 0, 0),
        ("SN-DA-HRSD-CREDENTIAL-001", 1, 1, 0),
        ("SN-DA-HRSD-TOPICS-001", 1, 0, 0),
        ("SN-DA-HRSD-AGENT-CONNECTION-001", 1, 1, 0),
        ("SN-DA-HRSD-TEST-001", 1, 1, 0),
        ("SN-DA-HRSD-PUBLISH-001", 1, 0, 0),
    ],
)
def test_target_aware_evaluation_uses_only_required_clients(
    monkeypatch,
    tmp_path: Path,
    target: str,
    expected_rows: int,
    inventory_reads: int,
    graph_reads: int,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    counters = {"agentbuilder": 0, "connectivity": 0, "graph": 0}
    runner = _runner(components)
    runner.checkpoint_target = target

    def fetch_components(_agent_id: str) -> dict:
        counters["agentbuilder"] += 1
        return components

    def list_connector_connections(
        environment_id: str,
        connector: str,
    ) -> list[dict]:
        assert environment_id == ENVIRONMENT_ID
        assert connector == snow.CONNECTOR_NAME
        counters["connectivity"] += 1
        return [_connection()]

    graph = _graph_client()

    class CountingGraph:
        def get_all(self, *args, **kwargs):
            counters["graph"] += 1
            return graph.get_all(*args, **kwargs)

        def get_service_principals(self, *args, **kwargs):
            counters["graph"] += 1
            return graph.get_service_principals(*args, **kwargs)

    runner.agentbuilder = SimpleNamespace(fetch_components=fetch_components)
    runner.connectivity = SimpleNamespace(
        list_connector_connections=list_connector_connections,
    )
    runner.graph = CountingGraph()
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)

    assert len(results) == expected_rows
    assert all(
        result.checkpoint_id == target
        or (
            target == "SN-DA-HRSD-ENTRA-*"
            and result.checkpoint_id.startswith("SN-DA-HRSD-ENTRA-")
        )
        for result in results
    )
    assert counters == {
        "agentbuilder": 1,
        "connectivity": inventory_reads,
        "graph": graph_reads,
    }


def test_preflight_discovery_change_does_not_require_global_reapproval(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["preflight"]["discovery"]["connectionCandidates"] = [
        {"connectionId": "new-candidate", "status": "Connected"}
    ]
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = next(
        item
        for item in run_servicenow_da_hrsd_checks(_runner(components))
        if item.checkpoint_id == "SN-DA-HRSD-ADMIN-PREFLIGHT-001"
    )

    assert result.status == Status.PASSED.value
    assert "exact supported agent target" in result.result


def test_connected_connection_does_not_complete_other_admin_phases(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    for phase in (
        "plugin-prerequisites",
        "entra-registration",
        "servicenow-oidc",
    ):
        state["adminSetup"]["phaseHandoffs"][phase] = {"status": "pending"}
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    statuses = {item.checkpoint_id: item.status for item in results}

    assert statuses["SN-DA-HRSD-ADMIN-PREFLIGHT-001"] == Status.PASSED.value
    assert statuses["SN-DA-HRSD-PLUGIN-001"] == Status.NOT_CONFIGURED.value
    assert statuses["SN-DA-HRSD-ENTRA-APP-001"] == Status.NOT_CONFIGURED.value
    assert statuses["SN-DA-HRSD-OIDC-001"] == Status.NOT_CONFIGURED.value


def test_optional_claim_absence_does_not_override_confirmed_emission(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    runner = _runner(components)
    runner.graph = _graph_client(missing_claims=True)
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    claim_result = next(
        result
        for result in results
        if result.checkpoint_id == "SN-DA-HRSD-ENTRA-CLAIMS-001"
    )

    assert claim_result.status == Status.MANUAL.value
    assert "upn: not listed in optionalClaims" in claim_result.result
    assert "administrator confirmed actual claim emission" in (
        claim_result.result
    )
    assert "does not prove emitted token contents" in claim_result.result
    assert claim_result.remediation == ""


def test_graph_unavailable_uses_manual_fallback_only_with_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    runner = _runner(components)
    runner.graph = None
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    app = next(
        result
        for result in results
        if result.checkpoint_id == "SN-DA-HRSD-ENTRA-APP-001"
    )
    assert app.status == Status.MANUAL.value
    assert "Structured admin evidence is present" in app.result

    state_path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
        "status": "pending"
    }
    state_path.write_text(json.dumps(state), encoding="utf-8")

    results = run_servicenow_da_hrsd_checks(runner)
    app = next(
        result
        for result in results
        if result.checkpoint_id == "SN-DA-HRSD-ENTRA-APP-001"
    )
    assert app.status == Status.NOT_CONFIGURED.value
    assert "No Application client ID" in app.result


def test_credential_rejects_resource_uri_mismatch(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    connection = _connection()
    connection["properties"]["connectionParametersSet"]["values"][
        "token:ResourceUri"
    ]["value"] = "00000000-0000-4000-8000-000000009999"
    runner = _runner(components)
    runner.connectivity = SimpleNamespace(
        list_connector_connections=lambda _environment_id, _connector: [
            connection
        ],
    )
    monkeypatch.chdir(tmp_path)

    result = next(
        item
        for item in run_servicenow_da_hrsd_checks(runner)
        if item.checkpoint_id == "SN-DA-HRSD-CREDENTIAL-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "different Instance Name or Resource URI" in result.result
    assert "exact Application client ID" in result.result


def test_hrsd_checks_use_connector_scoped_connections_only(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    runner = _runner(components)

    calls: list[tuple[str, str]] = []
    runner.connectivity = SimpleNamespace(
        list_connections=lambda _environment_id: pytest.fail(
            "ServiceNow HRSD must not use environment-wide inventory."
        ),
        list_connector_connections=lambda environment_id, connector: (
            calls.append((environment_id, connector)) or [_connection()]
        ),
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    statuses = {result.checkpoint_id: result.status for result in results}

    assert calls == [(ENVIRONMENT_ID, snow.CONNECTOR_NAME)]
    assert statuses["SN-DA-HRSD-CREDENTIAL-001"] == Status.PASSED.value
    assert (
        statuses["SN-DA-HRSD-AGENT-CONNECTION-001"]
        == Status.MANUAL.value
    )


@pytest.mark.parametrize("status_code", [429, 401, 403, 500])
def test_hrsd_connector_inventory_error_has_no_fallback(
    monkeypatch,
    tmp_path: Path,
    status_code: int,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    calls = 0
    runner = _runner(components)

    def fail_connector_inventory(
        _environment_id: str,
        _connector: str,
    ) -> list[dict]:
        nonlocal calls
        calls += 1
        raise AgentBuilderHTTPError(
            "Connector-scoped connection listing",
            status_code,
            request_id="request-123",
        )

    runner.connectivity = SimpleNamespace(
        list_connector_connections=fail_connector_inventory,
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    credential = next(
        result
        for result in results
        if result.checkpoint_id == "SN-DA-HRSD-CREDENTIAL-001"
    )

    assert calls == 1
    assert credential.status == Status.ERROR.value
    assert "request-123" in credential.result


def test_hrsd_checks_reopen_stale_maker_evidence(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    components["botComponentChanges"][0]["component"]["displayName"] = (
        "Edited ServiceNow topic"
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    statuses = {result.checkpoint_id: result.status for result in results}

    assert (
        statuses["SN-DA-HRSD-AGENT-CONNECTION-001"]
        == Status.NOT_CONFIGURED.value
    )
    assert "SN-DA-HRSD-PARAMETER-SHARING-001" not in statuses
    assert statuses["SN-DA-HRSD-PUBLISH-001"] == Status.NOT_CONFIGURED.value
    assert statuses["SN-DA-HRSD-TEST-001"] == Status.NOT_CONFIGURED.value


def test_agent_connection_accepts_legacy_root_only_connection_id(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(path.read_text(encoding="utf-8"))
    del state["evidence"]["agentConnection"]["binding"]["connectionId"]
    path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    result = next(
        item
        for item in results
        if item.checkpoint_id == "SN-DA-HRSD-AGENT-CONNECTION-001"
    )

    assert result.status == Status.MANUAL.value


def test_agent_connection_rejects_conflicting_nested_and_root_ids(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(path.read_text(encoding="utf-8"))
    state["evidence"]["agentConnection"]["connectionId"] = (
        "00000000000040008000000000009999"
    )
    path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    result = next(
        item
        for item in results
        if item.checkpoint_id == "SN-DA-HRSD-AGENT-CONNECTION-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "conflicting connection identities" in result.result


def test_agent_connection_rejects_wrong_selected_connection_id(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    other_id = "00000000000040008000000000009999"
    path = (
        tmp_path
        / ".local"
        / "connect"
        / snow.PROVIDER_KEY
        / "agents"
        / AGENT_SLUG
        / "lifecycle.json"
    )
    state = json.loads(path.read_text(encoding="utf-8"))
    state["evidence"]["credential"]["connectionId"] = other_id
    path.write_text(json.dumps(state), encoding="utf-8")
    runner = _runner(components)
    other_connection = _connection()
    other_connection["name"] = other_id
    runner.connectivity = SimpleNamespace(
        list_connections=lambda _environment_id: pytest.fail(
            "ServiceNow HRSD must not use environment-wide inventory."
        ),
        list_connector_connections=(
            lambda _environment_id, _connector: [other_connection]
        ),
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(runner)
    result = next(
        item
        for item in results
        if item.checkpoint_id == "SN-DA-HRSD-AGENT-CONNECTION-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "different selected credential" in result.result


def test_agent_connection_reports_component_revision_stale(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    _write_state(tmp_path, components)
    components["botComponentChanges"][0]["component"]["displayName"] = (
        "Edited ServiceNow topic"
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))
    result = next(
        item
        for item in results
        if item.checkpoint_id == "SN-DA-HRSD-AGENT-CONNECTION-001"
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "stale for this saved draft" in result.result


def test_hrsd_package_absence_is_not_configured(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    components["botComponentChanges"] = []
    components["connectionReferenceChanges"] = []
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_hrsd_checks(_runner(components))

    assert {
        result.status for result in results
    } == {Status.NOT_CONFIGURED.value}


def test_hrsd_checks_prefer_explicit_agent_slug(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _components()
    requested_id = "00000000-0000-4000-8000-000000008888"
    fetched: list[str] = []
    runner = _runner(components)
    runner.agent_slug = "requested-hr"
    runner.config = {
        "activeAgent": AGENT_SLUG,
        "agents": [
            {
                "slug": AGENT_SLUG,
                "botId": AGENT_ID,
                "schemaName": snow.HR_SCHEMA_NAME,
            },
            {
                "slug": "requested-hr",
                "botId": requested_id,
                "schemaName": snow.HR_SCHEMA_NAME,
            },
        ],
    }
    runner.agentbuilder = SimpleNamespace(
        fetch_components=lambda agent_id: (
            fetched.append(agent_id) or components
        )
    )
    monkeypatch.chdir(tmp_path)

    run_servicenow_da_hrsd_checks(runner)

    assert fetched == [requested_id]
