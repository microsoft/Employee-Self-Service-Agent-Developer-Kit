# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.

from __future__ import annotations

import copy
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

import connect_servicenow_da as snow
from flightcheck.checks.servicenow_da_itsm import (
    run_servicenow_da_itsm_checks,
)
from flightcheck.registry import resolve
from flightcheck.runner import (
    CheckResult,
    Priority,
    Role,
    RunResult,
    Status,
    compact_checkpoint_result_to_dict,
)
from tests.flightcheck.checks.test_servicenow_da_hrsd import (
    AGENT_ID,
    APP_CLIENT_ID,
    CONNECTION_ID,
    ENVIRONMENT_ID,
    _RECOVERY_HEADINGS,
    _components as _hrsd_components,
    _connection,
    _graph_client,
)


IT_SLUG = "employee-self-service-it"
ITSM_IMPORTED_TOPIC_SUFFIXES = (
    "ServiceNowITSMCreateTicket",
    "ServiceNowITSMGetTicketDetails",
    "ServiceNowITSMGetTicketUpdates",
    "ServiceNowITSMGetUserTickets",
    "ServiceNowITSMSetupConfigurations",
    "ServiceNowITSMSystemCommonExecution",
    "ServiceNowITSMSystemCreateTicket",
    "ServiceNowITSMSystemGetTicketDetails",
    "ServiceNowITSMSystemGetTicketsList",
    "ServiceNowITSMSystemResponseMapper",
    "ServiceNowITSMSystemUpdateTicket",
    "ServiceNowITSMSystemUserPermissionCheck",
    "ServiceNowITSMUpdateTicket",
)


def _itsm_components() -> dict:
    components = copy.deepcopy(_hrsd_components())
    changes = []
    for index, suffix in enumerate(ITSM_IMPORTED_TOPIC_SUFFIXES):
        dialog = {"$kind": "TaskDialog"}
        if suffix == "ServiceNowITSMSetupConfigurations":
            dialog = {
                "$kind": "TaskDialog",
                "beginDialog": {
                    "actions": [
                        {
                            "$kind": "SetVariable",
                            "displayName": "Set ServiceNow Portal Base URI",
                            "variable": "Global.ServiceNowITSMPortalBaseURI",
                            "value": {
                                "$kind": "Literal",
                                "literalValue": (
                                    "https://dev123.service-now.com/employee"
                                ),
                            },
                        }
                    ]
                },
            }
        changes.append(
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": (
                        "00000000-0000-4000-8000-"
                        f"{8000 + index:012d}"
                    ),
                    "version": 1,
                    "displayName": suffix,
                    "schemaName": f"{snow.IT_SCHEMA_NAME}.topic.{suffix}",
                    "state": "Active",
                    "status": "Active",
                    "dialog": dialog,
                },
            }
        )
    changes.append(
        {
            "$kind": "BotComponentInsert",
            "component": {
                "$kind": "DialogComponent",
                "id": "00000000-0000-4000-8000-000000009999",
                "version": 1,
                "displayName": "Customer ITSM custom topic",
                "schemaName": (
                    f"{snow.IT_SCHEMA_NAME}.topic."
                    "CustomerITSMCustom"
                ),
                "state": "Inactive",
                "status": "Inactive",
                "dialog": {"$kind": "TaskDialog"},
            },
        }
    )
    components["botComponentChanges"] = changes
    return components


def _runner(components: dict, target: str | None = None):
    return SimpleNamespace(
        config={
            "activeAgent": "employee-self-service-hr",
            "agents": [
                {
                    "slug": IT_SLUG,
                    "botId": AGENT_ID,
                    "schemaName": snow.IT_SCHEMA_NAME,
                }
            ],
        },
        agent_slug=IT_SLUG,
        checkpoint_target=target,
        agentbuilder=SimpleNamespace(
            fetch_components=lambda _agent_id: components
        ),
        connectivity=SimpleNamespace(
            list_connector_connections=lambda environment_id, connector: (
                [_connection()]
                if (
                    environment_id == ENVIRONMENT_ID
                    and connector == snow.CONNECTOR_NAME
                )
                else []
            )
        ),
        env_id=ENVIRONMENT_ID,
    )


def _checkpoint(results, checkpoint_id: str):
    return next(
        result for result in results if result.checkpoint_id == checkpoint_id
    )


def test_all_itsm_blocking_phase_outputs_include_actionable_recovery() -> None:
    results = run_servicenow_da_itsm_checks(
        SimpleNamespace(config={}, checkpoint_target=None)
    )

    expected_suffixes = {
        "PKG",
        "ADMIN-PREFLIGHT",
        "ENTRA-APP",
        "ENTRA-CLAIMS",
        "ENTRA-SCOPE",
        "ENTRA-PREAUTH",
        "OIDC",
        "TOPICS",
        "PORTAL",
        "CREDENTIAL",
        "AGENT-CONNECTION",
        "TEST",
        "PUBLISH",
    }
    assert {
        result.checkpoint_id.removeprefix("SN-DA-ITSM-").removesuffix("-001")
        for result in results
    } == expected_suffixes
    for result in results:
        assert result.status == Status.FAILED.value
        assert result.result in result.remediation
        assert all(
            heading in result.remediation for heading in _RECOVERY_HEADINGS
        )
        assert "{" not in result.remediation


def _state_path(root: Path) -> Path:
    return (
        root
        / ".local"
        / "connect"
        / "servicenow-da-itsm"
        / "agents"
        / IT_SLUG
        / "lifecycle.json"
    )


def _write_it_state(root: Path, components: dict) -> dict:
    with snow.product_profile_context("itsm"):
        context = {
            "agent": {
                "id": AGENT_ID,
                "schema_name": snow.IT_SCHEMA_NAME,
                "workspace_slug": IT_SLUG,
            },
            "environment": {
                "id": ENVIRONMENT_ID,
                "tenant_id": "00000000-0000-4000-8000-000000009999",
                "ring": "test",
            },
        }
        state = snow._state_base(context, components)
        component_hash = snow._component_hash(components)
        draft_hash = snow._draft_semantic_hash(components)
        state["adminSetup"]["preflight"] = {
            "instanceName": "dev123",
            "discovery": {
                "source": "connectivity-readonly",
                "observedAt": "2026-10-01T00:00:00Z",
            },
        }
        for phase in snow.ADMIN_PHASES:
            state["adminSetup"]["phaseHandoffs"][phase] = {
                "status": "completed",
                "evidence": {
                    "kind": "structured-admin-attestation",
                    "recordedAt": "2026-10-01T00:00:00Z",
                },
            }
        state["adminSetup"]["phaseHandoffs"]["entra-registration"][
            "evidence"
        ]["clientId"] = APP_CLIENT_ID
        state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
            "evidence"
        ].update(
            {
                "oidcCapabilityConfirmed": True,
                "runbookCompleted": True,
                "mappingDetailsCollected": True,
                "claimMappingConfirmed": True,
                "tokenClaimName": "upn",
                "serviceNowUserField": "user_name",
            }
        )
        state["adminSetup"]["phaseHandoffs"]["entra-registration"][
            "evidence"
        ].update(
            {
                "tokenClaimCandidates": ["email", "upn"],
                "tokenClaimEmissionConfirmed": True,
            }
        )
        binding = {
            "tenantId": "00000000-0000-4000-8000-000000009999",
            "environmentId": ENVIRONMENT_ID,
            "instanceName": "dev123",
            "applicationClientId": APP_CLIENT_ID,
            "provider": snow.ITSM_PROFILE.provider_key,
            "profile": "itsm",
            "agentSlug": IT_SLUG,
            "agentId": AGENT_ID,
        }
        state["evidence"]["identityClaimCandidates"] = {
            "kind": "administrator-confirmed-claim-candidates",
            "status": "entra-confirmed",
            "applicationClientId": APP_CLIENT_ID,
            "candidates": ["email", "upn"],
            "entraAdministratorConfirmed": True,
            "tokenClaimEmissionConfirmed": True,
            "binding": binding,
        }
        candidate_hash = snow._identity_candidate_set_hash(
            application_client_id=APP_CLIENT_ID,
            candidates=["email", "upn"],
            binding=binding,
        )
        state["evidence"]["identityClaimCandidates"][
            "candidateSetSha256"
        ] = candidate_hash
        state["evidence"]["identityMapping"] = {
            "kind": "administrator-confirmed-identity-mapping",
            "status": "servicenow-confirmed",
            "applicationClientId": APP_CLIENT_ID,
            "tokenClaimName": "upn",
            "serviceNowUserField": "user_name",
            "candidateSetSha256": candidate_hash,
            "entraAdministratorConfirmed": True,
            "tokenClaimEmissionConfirmed": True,
            "serviceNowAdministratorConfirmed": True,
            "binding": binding,
        }
        relationship_hash = snow._identity_mapping_relationship_hash(
            state["evidence"]["identityMapping"]
        )
        state["evidence"]["identityMapping"]["relationshipSha256"] = (
            relationship_hash
        )
        state["adminSetup"]["phaseHandoffs"]["entra-registration"][
            "evidence"
        ]["candidateSetSha256"] = candidate_hash
        state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"][
            "evidence"
        ]["mappingRelationshipSha256"] = relationship_hash
        for phase in ("entra-registration", "servicenow-oidc"):
            state["adminSetup"]["phaseHandoffs"][phase]["evidence"][
                "binding"
            ] = copy.deepcopy(binding)
        state["evidence"]["credential"] = {"connectionId": CONNECTION_ID}
        state["evidence"]["agentConnection"] = {
            "kind": "maker-attestation",
            "status": "completed",
            "makerAttested": True,
            "physicalStatus": "Connected",
            "connectionId": CONNECTION_ID,
            "recordedAt": "2026-10-01T00:00:00Z",
            "binding": {"draftSemanticHash": draft_hash},
        }
        binding_hash = snow._connection_binding_hash(
            context,
            state,
            components,
        )
        state["connectionBindingHash"] = binding_hash
        state["evidence"]["test"] = {
            "kind": "maker-attestation",
            "status": "completed",
            "promptCategory": "list-my-open-it-tickets",
            "result": "pass",
            "failureCategory": None,
            "binding": {
                "draftSemanticHash": draft_hash,
                "connectionBindingHash": binding_hash,
            },
        }
        state["evidence"]["publish"] = {
            "status": "completed",
            "componentHash": component_hash,
            "publishedComponentHash": component_hash,
            "testedDraftSemanticHash": draft_hash,
            "publishedSemanticHash": draft_hash,
            "connectionBindingHash": binding_hash,
        }
    path = _state_path(root)
    path.parent.mkdir(parents=True)
    path.write_text(json.dumps(state), encoding="utf-8")
    return state


def test_itsm_credential_drift_is_repairable_not_configured(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    _write_it_state(tmp_path, components)
    connection = _connection()
    connection["properties"]["connectionParametersSet"]["values"][
        "token:ResourceUri"
    ]["value"] = "00000000-0000-4000-8000-000000009999"
    runner = _runner(components)
    runner.connectivity = SimpleNamespace(
        list_connector_connections=lambda _environment_id, _connector: [
            connection
        ],
    )
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_itsm_checks(runner)

    credential = _checkpoint(results, "SN-DA-ITSM-CREDENTIAL-001")
    test = _checkpoint(results, "SN-DA-ITSM-TEST-001")
    assert credential.status == Status.NOT_CONFIGURED.value
    assert test.status == Status.NOT_CONFIGURED.value
    assert "different Instance Name or Resource URI" in credential.result


def test_itsm_credential_inventory_error_remains_error(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    _write_it_state(tmp_path, components)
    runner = _runner(components)

    def fail_inventory(_environment_id: str, _connector: str) -> list[dict]:
        raise RuntimeError("inventory unavailable")

    runner.connectivity = SimpleNamespace(
        list_connector_connections=fail_inventory,
    )
    monkeypatch.chdir(tmp_path)

    credential = _checkpoint(
        run_servicenow_da_itsm_checks(runner),
        "SN-DA-ITSM-CREDENTIAL-001",
    )

    assert credential.status == Status.ERROR.value
    assert "inventory unavailable" in credential.result


def test_itsm_portal_checkpoint_accepts_bound_root_confirmation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    with snow.product_profile_context("itsm"):
        topic = snow._portal_topic_component(components)
        snow._portal_value_node(topic)["value"]["literalValue"] = (
            "https://dev123.service-now.com"
        )
        state = _write_it_state(tmp_path, components)
        snow._record_portal_confirmation(
            state,
            topic=topic,
            portal_url="https://dev123.service-now.com",
            instance_name="dev123",
            source="already-configured",
        )
    _state_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    result = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-PORTAL-001")
        ),
        "SN-DA-ITSM-PORTAL-001",
    )

    assert result.status == Status.PASSED.value
    assert "confirmed ServiceNow instance" in result.result


def test_recorded_mapping_workflow_reaches_real_it_checkpoints(
    monkeypatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    _write_it_state(tmp_path, components)
    monkeypatch.chdir(tmp_path)
    context = {
        "root": tmp_path,
        "agent": {
            "id": AGENT_ID,
            "schema_name": snow.IT_SCHEMA_NAME,
            "workspace_slug": IT_SLUG,
        },
        "active": {"slug": IT_SLUG},
        "environment": {
            "id": ENVIRONMENT_ID,
            "tenant_id": "00000000-0000-4000-8000-000000009999",
            "ring": "test",
        },
    }
    with snow.product_profile_context("itsm"):
        state_path = _state_path(tmp_path)
        state = json.loads(state_path.read_text(encoding="utf-8"))
        state.setdefault("phases", {})["entra-registration"] = {
            "status": "in-progress",
            "checkpointResults": {
                "SN-DA-ITSM-ENTRA-CLAIMS-001": (
                    Status.NOT_CONFIGURED.value
                )
            },
        }
        state["adminSetup"]["phaseHandoffs"]["entra-registration"] = {
            "status": "pending"
        }
        state["adminSetup"]["phaseHandoffs"]["servicenow-oidc"] = {
            "status": "pending"
        }
        state.setdefault("evidence", {}).pop(
            "identityClaimCandidates",
            None,
        )
        state["evidence"].pop("identityMapping", None)
        state_path.write_text(json.dumps(state), encoding="utf-8")
        snow.record_identity_requirement(context, choice="email")
        entra = snow.identity_mapping_handoff(
            context,
            stage="entra",
            client_id=APP_CLIENT_ID,
        )
        entra_path = Path(entra["workbookPath"])
        workbook = load_workbook(entra_path)
        sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
        for offset, claim in enumerate(("email", "upn", "employee_id")):
            sheet.cell(
                snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW + offset,
                1,
            ).value = claim
        workbook.save(entra_path)
        workbook.close()
        snow.record_identity_mapping(
            context,
            stage="entra",
            workbook_path=entra_path,
            client_id=APP_CLIENT_ID,
            confirmed=True,
        )
        snow.record_admin_phase(
            context,
            phase="entra-registration",
            status="completed",
            client_id=APP_CLIENT_ID,
        )
        service_now = snow.identity_mapping_handoff(
            context,
            stage="servicenow",
        )
        service_now_path = Path(service_now["workbookPath"])
        workbook = load_workbook(service_now_path)
        sheet = workbook[snow.IDENTITY_WORKBOOK_SHEET]
        for row in range(
            snow.IDENTITY_WORKBOOK_FIRST_DATA_ROW,
            sheet.max_row + 1,
        ):
            if sheet.cell(row, 1).value == "employee_id":
                sheet.cell(row, 2).value = "employee_number"
        workbook.save(service_now_path)
        workbook.close()
        snow.record_identity_mapping(
            context,
            stage="servicenow",
            workbook_path=service_now_path,
            client_id=None,
            confirmed=True,
        )
        snow.record_admin_phase(
            context,
            phase="servicenow-oidc",
            status="completed",
        )
    runner = _runner(components)
    runner.graph = _graph_client(optional_claim_names=())

    results = run_servicenow_da_itsm_checks(runner)
    claims = _checkpoint(results, "SN-DA-ITSM-ENTRA-CLAIMS-001")
    oidc = _checkpoint(results, "SN-DA-ITSM-OIDC-001")

    assert claims.status == Status.MANUAL.value
    assert "email" in claims.result
    assert "upn" in claims.result
    assert "employee_id" in claims.result
    assert oidc.status == Status.MANUAL.value
    assert "employee_id" in oidc.result


def test_itsm_package_and_topics_checks_use_selected_it_agent() -> None:
    components = _itsm_components()

    package = run_servicenow_da_itsm_checks(
        _runner(components, "SN-DA-ITSM-PKG-001")
    )
    topics = run_servicenow_da_itsm_checks(
        _runner(components, "SN-DA-ITSM-TOPICS-001")
    )

    assert _checkpoint(
        package,
        "SN-DA-ITSM-PKG-001",
    ).status == Status.PASSED.value
    assert _checkpoint(
        topics,
        "SN-DA-ITSM-TOPICS-001",
    ).status == Status.PASSED.value
    assert all("HRSD" not in result.checkpoint_id for result in package + topics)


def test_itsm_topics_compact_checkpoint_contract_accepts_actual_inventory() -> None:
    target = "SN-DA-ITSM-TOPICS-001"
    topic_rows = run_servicenow_da_itsm_checks(
        _runner(_itsm_components(), target)
    )
    client_row = CheckResult(
        checkpoint_id="CHECKPOINT-CLIENT-AGENTBUILDER",
        category="Client availability",
        priority=Priority.HIGH.value,
        status=Status.PASSED.value,
        description="AgentBuilder client",
        result="Available.",
        remediation="",
        roles=[Role.ESS_MAKER.value],
    )
    run_result = RunResult(
        scope="single-checkpoint",
        started="2026-10-08T00:00:00Z",
        results=[client_row, *topic_rows],
        total=2,
        passed=2,
        overall="READY",
    )

    payload = compact_checkpoint_result_to_dict(
        run_result,
        target=target,
        provider="servicenow-da-itsm",
        profile="itsm",
        agent_slug=IT_SLUG,
        agent_id=AGENT_ID,
        environment_id=ENVIRONMENT_ID,
        invocation_id="dynamic-topic-test",
        invocation_source="lifecycle",
        exit_code=0,
    )

    assert payload["overall"] == "READY"
    assert len(payload["results"]) == 1
    row = payload["results"][0]
    assert row["checkpointId"] == target
    assert row["status"] == Status.PASSED.value
    assert row["description"] == "ServiceNow ITSM topic readiness"
    assert row["result"] == "13/13 ServiceNow ITSM topics are Active."
    assert row["remediation"] == ""


def test_itsm_registry_has_no_hr_plugin_checkpoint() -> None:
    assert resolve("SN-DA-ITSM-PKG-001").provider == "servicenow-da-itsm"
    assert resolve("SN-DA-ITSM-ENTRA-*").profile == "itsm"
    assert resolve("SN-DA-ITSM-PLUGIN-001") is None


@pytest.mark.parametrize("retained_count", (1, 5, 13))
def test_itsm_package_and_topics_accept_any_positive_imported_inventory(
    retained_count: int,
) -> None:
    components = _itsm_components()
    components["botComponentChanges"] = [
        change
        for change in components["botComponentChanges"][:retained_count]
    ]

    package = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-PKG-001")
        ),
        "SN-DA-ITSM-PKG-001",
    )
    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert package.status == Status.PASSED.value
    assert topics.status == Status.PASSED.value
    assert f"contains {retained_count} ServiceNow ITSM topic(s)" in (
        package.result
    )
    assert topics.result == (
        f"{retained_count}/{retained_count} ServiceNow ITSM topics are Active."
    )


def test_itsm_package_and_topics_accept_new_matching_topic() -> None:
    components = _itsm_components()
    components["botComponentChanges"].append(
        {
            "$kind": "BotComponentInsert",
            "component": {
                "$kind": "DialogComponent",
                "id": "00000000-0000-4000-8000-000000008888",
                "displayName": "Future package topic",
                "schemaName": (
                    f"{snow.IT_SCHEMA_NAME}.topic."
                    "ServiceNowITSMFuturePackageTopic"
                ),
                "state": "Active",
                "status": "Active",
                "dialog": {"$kind": "TaskDialog"},
            },
        }
    )

    package = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-PKG-001")
        ),
        "SN-DA-ITSM-PKG-001",
    )
    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert package.status == Status.PASSED.value
    assert topics.status == Status.PASSED.value
    assert "contains 14 ServiceNow ITSM topic(s)" in package.result


def test_itsm_topic_inventory_isolates_opposite_and_non_topic_components() -> None:
    components = _itsm_components()
    components["botComponentChanges"].extend(
        [
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": "00000000-0000-4000-8000-000000008889",
                    "displayName": "HRSD topic",
                    "schemaName": (
                        f"{snow.IT_SCHEMA_NAME}.topic.ServiceNowHRSDCreateCase"
                    ),
                    "state": "Active",
                    "status": "Active",
                    "dialog": {"$kind": "TaskDialog"},
                },
            },
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "GlobalVariableComponent",
                    "id": "00000000-0000-4000-8000-000000008890",
                    "displayName": "Not a dialog",
                    "schemaName": (
                        f"{snow.IT_SCHEMA_NAME}.topic."
                        "ServiceNowITSMNotADialog"
                    ),
                    "dialog": {"$kind": "TaskDialog"},
                },
            },
        ]
    )

    with snow.product_profile_context("itsm"):
        summary = snow.summarize_components(components)

    assert summary["serviceNowTopicCount"] == 13
    assert all(
        "ServiceNowHRSD" not in str(topic["schemaName"])
        and "NotADialog" not in str(topic["schemaName"])
        for topic in summary["serviceNowTopics"]
    )


def test_itsm_zero_recognized_topics_is_not_configured() -> None:
    components = _itsm_components()
    components["botComponentChanges"] = [
        change
        for change in components["botComponentChanges"]
        if "ServiceNowITSM"
        not in str(change["component"].get("schemaName") or "")
    ]

    package = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-PKG-001")
        ),
        "SN-DA-ITSM-PKG-001",
    )
    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert package.status == Status.NOT_CONFIGURED.value
    assert topics.status == Status.NOT_CONFIGURED.value


def test_itsm_inactive_inventory_requires_enablement() -> None:
    components = _itsm_components()
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"

    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert topics.status == Status.NOT_CONFIGURED.value
    assert topics.result == "12/13 ServiceNow ITSM topics are Active."


def test_itsm_newly_discovered_inactive_topics_block_until_enabled() -> None:
    components = _itsm_components()
    for index, suffix in enumerate(
        (
            "ServiceNowITSMSystemReadExecution",
            "ServiceNowITSMSystemWriteExecution",
        )
    ):
        components["botComponentChanges"].append(
            {
                "$kind": "BotComponentInsert",
                "component": {
                    "$kind": "DialogComponent",
                    "id": (
                        "00000000-0000-4000-8000-"
                        f"{8880 + index:012d}"
                    ),
                    "displayName": suffix,
                    "schemaName": f"{snow.IT_SCHEMA_NAME}.topic.{suffix}",
                    "state": "Inactive",
                    "status": "Inactive",
                    "dialog": {"$kind": "TaskDialog"},
                },
            }
        )

    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert topics.status == Status.NOT_CONFIGURED.value
    assert topics.result == "13/15 ServiceNow ITSM topics are Active."
    assert "Enable the inactive ITSM topics" in topics.remediation


def test_itsm_keep_current_accepts_dynamic_inactive_inventory(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"
    state = _write_it_state(tmp_path, components)
    state["evidence"]["topics"] = {
        "kind": "maker-attestation",
        "status": "recorded",
        "customerChoice": "keep-current",
        "boundComponentHash": snow._component_hash(components),
    }
    _state_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert topics.status == Status.MANUAL.value
    assert topics.result == "12/13 ServiceNow ITSM topics are Active."


def test_itsm_stale_keep_current_hash_requires_fresh_choice(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    components["botComponentChanges"][0]["component"]["state"] = "Inactive"
    components["botComponentChanges"][0]["component"]["status"] = "Inactive"
    state = _write_it_state(tmp_path, components)
    state["evidence"]["topics"] = {
        "kind": "maker-attestation",
        "status": "recorded",
        "customerChoice": "keep-current",
        "boundComponentHash": "stale-component-hash",
    }
    _state_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    topics = _checkpoint(
        run_servicenow_da_itsm_checks(
            _runner(components, "SN-DA-ITSM-TOPICS-001")
        ),
        "SN-DA-ITSM-TOPICS-001",
    )

    assert topics.status == Status.NOT_CONFIGURED.value
    assert "Enable the inactive ITSM topics" in topics.remediation


def test_itsm_checkpoint_rejects_hr_agent_schema() -> None:
    runner = _runner(_itsm_components(), "SN-DA-ITSM-PKG-001")
    runner.config["agents"][0]["schemaName"] = snow.HR_SCHEMA_NAME

    result = _checkpoint(
        run_servicenow_da_itsm_checks(runner),
        "SN-DA-ITSM-PKG-001",
    )

    assert result.status == Status.NOT_CONFIGURED.value
    assert "IT agent" in result.result


def test_itsm_not_pass_blocks_test_and_publish(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    components = _itsm_components()
    state = _write_it_state(tmp_path, components)
    state["evidence"]["test"]["status"] = "failed"
    state["evidence"]["test"]["result"] = "fail"
    state["evidence"]["test"]["failureCategory"] = None
    _state_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    results = run_servicenow_da_itsm_checks(_runner(components))

    assert _checkpoint(
        results,
        "SN-DA-ITSM-TEST-001",
    ).status == Status.FAILED.value
    assert _checkpoint(
        results,
        "SN-DA-ITSM-PUBLISH-001",
    ).status == Status.FAILED.value


@pytest.mark.parametrize("result", ("pass", "fail"))
def test_itsm_publish_rejects_test_bound_to_previous_connection(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    result: str,
) -> None:
    components = _itsm_components()
    state = _write_it_state(tmp_path, components)
    state["evidence"]["test"]["result"] = result
    state["evidence"]["test"]["status"] = (
        "completed" if result == "pass" else "failed"
    )
    state["evidence"]["credential"]["connectionId"] = (
        "00000000000040008000000000009999"
    )
    _state_path(tmp_path).write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.chdir(tmp_path)

    full = _checkpoint(
        run_servicenow_da_itsm_checks(_runner(components)),
        "SN-DA-ITSM-PUBLISH-001",
    )
    targeted_runner = _runner(components, "SN-DA-ITSM-PUBLISH-001")
    targeted_runner.connectivity = SimpleNamespace(
        list_connector_connections=lambda *_args: pytest.fail(
            "Publish-only checkpoint must not read connection inventory."
        )
    )
    targeted = run_servicenow_da_itsm_checks(targeted_runner)

    assert full.status == Status.NOT_CONFIGURED.value
    assert targeted[0].status == Status.NOT_CONFIGURED.value
